"""Render a document with built-in styles, validated headings and linked contents."""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from folio.document.figures import DocumentFigures
from folio.document.lists import ProseListLayout
from folio.document.pagination import paginate
from folio.document.presentation import color_proposal, prose_paragraphs
from folio.document.reader import heading_level, visible_text, xpath
from folio.document.style_layout import apply_style_layout, style_generated_table
from folio.document.template_layout import (
    TemplateLayout,
    template_heading_matches,
    validate_template_output,
)
from folio.document.toc import rebuild_toc
from folio.document.typography import (
    BODY,
    DOCUMENT_STATUS,
    TABLE_BODY,
    TABLE_HEADER,
    TOC_HEADING,
    clean_heading,
    disable_table_numbering,
    install_styles,
    prop,
    repair_table_pagination,
)
from folio.models import DocumentPlan, SectionState
from folio.runtime.language import localize_text


def render(
    target: Path,
    plan: DocumentPlan,
    states: list[SectionState],
    *,
    title: str,
    draft: bool,
    demo: bool = False,
    cover: dict | None = None,
    comments: list | None = None,
    preview: bool = False,
    include_evidence_appendix: bool = False,
    pdf: Path | None = None,
    require_pagination: bool = False,
    figures: dict[str, list[dict]] | None = None,
    defer_pagination: bool = False,
    writing_profile: dict | None = None,
    language: str = "zh-CN",
) -> dict:
    if plan.template and include_evidence_appendix:
        raise ValueError('Template mode cannot add a top-level references appendix automatically. Reserve a section in the template.')
    figure_layout = DocumentFigures(plan, states, figures, language=language)
    template = TemplateLayout(plan, title=title, cover=cover, language=language) if plan.template else None
    doc = template.document if template else Document()
    typography = template.typography if template else install_styles(doc, language=language)
    style_report = apply_style_layout(doc, writing_profile, template=bool(template))
    # The anchor is local rendering machinery, never an input document contract.
    marker = doc.add_paragraph()._p if template is None else None

    def place(element):
        marker.addprevious(element)

    title = (cover or {}).get("document_title") or plan.document_title or title
    english = language == "en"
    if template is None:
        place(doc.add_paragraph(title, "Title")._p)
        if preview or draft:
            status = (
                "OUTLINE PREVIEW" if preview else "DRAFT"
            ) if english else (
                "目录预览" if preview else "草稿"
            )
            place(doc.add_paragraph(status, DOCUMENT_STATUS)._p)
    if demo and template is None:
        place(doc.add_paragraph(
            "Demo output: produced by the deterministic example runtime without real model semantic review."
            if english else "演示产物：使用确定性示例运行器，未经真实模型语义审查。"
        )._p)
    if plan.sections and template is None:
        place(doc.add_paragraph("Contents" if english else "目录", TOC_HEADING)._p)
        toc = OxmlElement("w:sdt")
        properties = OxmlElement("w:sdtPr")
        gallery = OxmlElement("w:docPartObj")
        name = OxmlElement("w:docPartGallery")
        name.set(qn("w:val"), "Table of Contents")
        gallery.append(name)
        properties.append(gallery)
        toc.append(properties)
        content = OxmlElement("w:sdtContent")
        paragraph = OxmlElement("w:p")
        from folio.document.toc import field_run

        paragraph.extend(
            [
                field_run("begin"),
                field_run("instruction", ' TOC \\o "1-3" \\h \\z \\u '),
                field_run("end"),
            ]
        )
        content.append(paragraph)
        toc.append(content)
        place(toc)
        place(doc.add_page_break()._p)
    entries = []
    by_id = {s.section_id: s for s in states}
    parents = {s.parent_id for s in plan.sections if s.parent_id}
    numbers = doc.part.numbering_part.element
    abstract_id = max(map(int, xpath(numbers, "./w:abstractNum/@w:abstractNumId")), default=-1) + 1
    definition = OxmlElement("w:abstractNum")
    definition.set(qn("w:abstractNumId"), str(abstract_id))
    prop(definition, "multiLevelType", val="multilevel")
    for level in range(9):
        if xpath(definition, f'./w:lvl[@w:ilvl="{level}"]'):
            continue
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        for tag, value in [
            ("start", "1"),
            ("numFmt", "decimal"),
            ("lvlText", ".".join(f"%{i + 1}" for i in range(level + 1))),
            ("suff", "space"),
            ("lvlJc", "left"),
        ]:
            prop(lvl, tag, val=value)
        definition.append(lvl)
    numbers.insert(0, definition)
    next_num_id = max(map(int, xpath(numbers, "./w:num/@w:numId")), default=0) + 1
    bookmark_ids = xpath(doc.element, ".//w:bookmarkStart/@w:id")
    bookmark_id = max((int(v) for v in bookmark_ids if v.isdigit()), default=0) + 1
    entries = []

    def numbering(prefix):
        # A section's children must not change the fixed numbering of later DocPlan slots.
        nonlocal next_num_id
        current = next_num_id
        next_num_id += 1
        number = OxmlElement("w:num")
        number.set(qn("w:numId"), str(current))
        prop(number, "abstractNumId", val=abstract_id)
        for level, value in enumerate(prefix):
            override = OxmlElement("w:lvlOverride")
            override.set(qn("w:ilvl"), str(level))
            prop(override, "startOverride", val=value)
            number.append(override)
        numbers.append(number)
        return current

    def heading(text, level, number, num_id, bookmark, existing=None):
        nonlocal bookmark_id
        text = clean_heading(text)
        if existing is not None:
            from docx.text.paragraph import Paragraph

            p = Paragraph(existing, doc._body)
        else:
            p = doc.add_paragraph(text, f"Folio Heading {level}")
        ppr = p._p.get_or_add_pPr()
        if existing is None:
            ppr.get_or_add_numPr().get_or_add_numId().val = num_id
            ppr.get_or_add_numPr().get_or_add_ilvl().val = level - 1
        prop(ppr, "outlineLvl", val=level - 1)
        begin, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        begin.set(qn("w:id"), str(bookmark_id))
        begin.set(qn("w:name"), bookmark)
        end.set(qn("w:id"), str(bookmark_id))
        p._p.insert(1, begin)
        p._p.append(end)
        bookmark_id += 1
        entries.append((bookmark, text, level, number))
        if existing is None:
            place(p._p)
        return p

    for section in plan.sections:
        if template:
            marker = template.start_section(section)
        prefix = list(map(int, section.display_number.split(".")))
        num_id = numbering(prefix)
        counts = prefix + [0] * (9 - len(prefix))
        heading(
            section.title,
            section.level,
            section.display_number,
            num_id,
            section.section_id,
            existing=template.heading(section) if template else None,
        )
        state = by_id.get(section.section_id)
        if state is None:
            if section.content_mode != "container":
                paragraph = (
                    template.paragraph(section, "Content pending." if english else "待生成。")
                    if template
                    else doc.add_paragraph("Content pending." if english else "待生成。", BODY)
                )
                place(paragraph._p)
            if template:
                for required in template.missing_tables(section):
                    table = template.placeholder_table(section, required)
                    place(table._tbl)
            continue
        list_layout = ProseListLayout()
        stack = [(section.level, section.level)]
        for index, block in enumerate(state.content.blocks):
            if block.kind == "heading":
                list_layout.reset()
                if (
                    section.section_id in parents
                    or section.level >= 6
                    or (template and not template.allows_headings())
                ):
                    p = (
                        template.paragraph(section, clean_heading(block.text))
                        if template
                        else doc.add_paragraph(clean_heading(block.text), BODY)
                    )
                    place(p._p)
                else:
                    requested = max(section.level + 1, block.heading_level)
                    while len(stack) > 1 and stack[-1][0] >= requested:
                        stack.pop()
                    level = min(stack[-1][1] + 1, 6)
                    stack.append((requested, level))
                    counts[level - 1] += 1
                    counts[level:] = [0] * (9 - level)
                    number = ".".join(map(str, counts[:level]))
                    p = heading(block.text, level, number, num_id, f"{section.section_id}_b{index}")
                color_proposal(p._p, block)
            elif block.kind == "paragraph":
                for text in prose_paragraphs(block.text):
                    p = (
                        template.paragraph(section, text)
                        if template
                        else doc.add_paragraph(text, BODY)
                    )
                    if template is None:
                        list_layout.apply(p, text, block.block_id)
                    color_proposal(p._p, block)
                    place(p._p)
            else:
                list_layout.reset()
                rows = [list(row) for row in block.rows]
                headers = section.expected_table_headers
                if headers and len(rows[0]) == len(headers) and rows[0] != headers:
                    rows.insert(0, list(headers))
                table = template.table(section, rows) if template else None
                if table is None:
                    table = doc.add_table(rows=0, cols=len(rows[0]))
                    table.style = "Table Grid"
                    for row_index, row in enumerate(rows):
                        for cell, value in zip(table.add_row().cells, row, strict=True):
                            cell.text = "\n".join(prose_paragraphs(value))
                            cell.paragraphs[0].style = (
                                TABLE_HEADER if row_index == 0 else TABLE_BODY
                            )
                    disable_table_numbering(table)
                    repair_table_pagination(table)
                    if template is None:
                        style_generated_table(table, writing_profile)
                if hasattr(table, "template_header_rows"):
                    for row in table.rows[table.template_header_rows :]:
                        color_proposal(row._tr, block)
                else:
                    color_proposal(table._tbl, block)
                place(table._tbl)
            figure_layout.emit(doc, section.section_id, block.block_id, place=place, anchor=marker)
        figure_layout.emit(doc, section.section_id, place=place, anchor=marker)
        if template:
            for required in template.missing_tables(section):
                if not (preview or draft):
                    raise ValueError(f'A required template table is incomplete: {section.title} / {required.table_id}')
                table = template.placeholder_table(section, required)
                place(table._tbl)
        for pending in state.content.pending_items:
            pending = localize_text(pending, language)
            paragraph = (
                template.paragraph(section, f"Pending confirmation: {pending}" if english else f"待确认：{pending}")
                if template
                else doc.add_paragraph(f"Pending confirmation: {pending}" if english else f"待确认：{pending}")
            )
            place(paragraph._p)
    if include_evidence_appendix:
        place(doc.add_heading("References" if english else "参考资料", 1)._p)
        evidence = {e.evidence_id: e for s in states for e in s.evidence}
        for e in evidence.values():
            name = e.metadata.get("document_name") or e.source_id
            paragraph = (
                template.paragraph(plan.sections[-1], f"{name} · {e.locator}")
                if template
                else doc.add_paragraph(f"{name} · {e.locator}", BODY)
            )
            place(paragraph._p)
    if template:
        template.finish()
    else:
        marker.getparent().remove(marker)
    figure_layout.validate_complete()
    rebuild_toc(doc, entries, typography)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    validate_render(target, plan)
    validate_template_output(target, plan)
    layout = (
        {"status": "deferred"}
        if defer_pagination
        else paginate(target, pdf=pdf, required=require_pagination)
    )
    if not defer_pagination:
        validate_template_output(target, plan)
    if writing_profile:
        layout["style_application"] = style_report
    return layout


def validate_render(path: Path, plan: DocumentPlan) -> None:
    doc = Document(path)
    bookmarks = doc.element.xpath(".//w:bookmarkStart")
    names = [node.get(qn("w:name")) for node in bookmarks]
    ids = [node.get(qn("w:id")) for node in bookmarks]
    if len(names) != len(set(names)) or len(ids) != len(set(ids)):
        raise ValueError("Duplicate document bookmarks")
    by_name = {node.get(qn("w:name")): node for node in bookmarks}
    for section in plan.sections:
        bookmark = by_name.get(section.section_id)
        if bookmark is None or not (
            template_heading_matches(bookmark.getparent(), section.title)
            if plan.template
            else visible_text(bookmark.getparent()).strip() == clean_heading(section.title)
        ):
            raise ValueError(f"Missing or incorrect rendered heading: {section.title}")
        if heading_level(bookmark.getparent(), doc) != section.level:
            raise ValueError("Rendered heading level differs from plan")
    if any(anchor not in by_name for anchor in doc.element.xpath(".//w:hyperlink/@w:anchor")):
        raise ValueError("Broken table-of-contents link")
