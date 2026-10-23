"""Fill a frozen user DOCX while retaining its authored structure and formatting."""

import re
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from lxml import etree

from folio.document.reader import heading_level, table_rows, visible_text, xpath
from folio.document.template_contract import _template_mapping
from folio.document.typography import (
    BODY,
    TABLE_BODY,
    clean_heading,
    install_template_styles,
    resolved_format,
)


def _replace_text(paragraph, text):
    """Keep a paragraph's properties and representative run properties, never its sample text."""
    runs = xpath(paragraph, "./w:r[w:t]")
    prototype = max(runs, key=lambda run: len(visible_text(run)), default=None)
    properties = prototype.find(qn("w:rPr")) if prototype is not None else None
    for child in list(paragraph):
        if child.tag != qn("w:pPr"):
            paragraph.remove(child)
    run = OxmlElement("w:r")
    if properties is not None:
        run.append(deepcopy(properties))
    for i, line in enumerate(text.split("\n")):
        if i:
            run.append(OxmlElement("w:br"))
        node = OxmlElement("w:t")
        node.set(qn("xml:space"), "preserve")
        node.text = line
        run.append(node)
    paragraph.append(run)


def _empty_data_row(table):
    """Header-only templates need ordinary cells, even when their header is merged."""
    row = OxmlElement("w:tr")
    for column in xpath(table, "./w:tblGrid/w:gridCol"):
        cell = OxmlElement("w:tc")
        properties = OxmlElement("w:tcPr")
        width = OxmlElement("w:tcW")
        width.set(qn("w:type"), "dxa")
        width.set(qn("w:w"), column.get(qn("w:w"), "0"))
        properties.append(width)
        cell.append(properties)
        paragraph = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        style = OxmlElement("w:pStyle")
        style.set(qn("w:val"), TABLE_BODY.replace(" ", ""))
        ppr.append(style)
        paragraph.append(ppr)
        cell.append(paragraph)
        row.append(cell)
    return row


def _replace_placeholders(paragraph, values):
    """Replace even split-run placeholders without removing cover art or fields."""
    nodes = xpath(paragraph, ".//w:t")
    for key, value in values.items():
        for token in ("{{" + key + "}}", "{" + key + "}"):
            joined = "".join(node.text or "" for node in nodes)
            for match in reversed(list(re.finditer(re.escape(token), joined))):
                start, end = match.span()
                offset, inserted = 0, False
                for node in nodes:
                    text = node.text or ""
                    stop = offset + len(text)
                    if offset < end and stop > start:
                        before = text[: max(0, start - offset)]
                        after = text[max(0, end - offset) :]
                        node.text = before + (str(value) if not inserted else "") + after
                        node.set(qn("xml:space"), "preserve")
                        inserted = True
                    offset = stop


class TemplateLayout:
    """One mutable document; all relationships stay attached to the original package."""

    def __init__(self, plan, *, title, cover=None, language="zh-CN"):
        self.language = language
        contract = plan.template
        path = Path(contract.path)
        if sha256(path.read_bytes()).hexdigest() != contract.source_hash:
            raise ValueError('The document template has changed. Regenerate and confirm the outline.')
        self.document = Document(path)
        self.contract = contract
        self.markers = {}
        self.headings = {}
        self.prototypes = {}
        self.tables = {}
        self.consumed_tables = set()
        body = self.document.element.body
        original = list(body)
        styles = {}
        first_body = None
        for index, section in enumerate(contract.sections):
            heading = original[section.body_index]
            if (
                heading_level(heading, self.document) != section.level
                or visible_text(heading).strip() != section.title
            ):
                raise ValueError('Template section locations do not match the confirmed outline.')
            self.headings[section.section_id] = heading
            styles.setdefault(section.level, heading)
            stop = (
                contract.sections[index + 1].body_index
                if index + 1 < len(contract.sections)
                else len(original)
            )
            elements = original[section.body_index + 1 : stop]
            prototype = next(
                (
                    element
                    for element in elements
                    if element.tag == qn("w:p") and visible_text(element).strip()
                ),
                None,
            )
            self.prototypes[section.section_id] = prototype
            if first_body is None and prototype is not None:
                first_body = prototype
            self.tables[section.section_id] = [
                (table, deepcopy(original[table.body_index])) for table in section.tables
            ]
            # All prose/data cells are instructions or examples, not project facts.
            # Section properties are independent of that text and must survive.
            for element in elements:
                if element.tag == qn("w:sectPr"):
                    continue
                for properties in xpath(element, ".//w:sectPr"):
                    paragraph = OxmlElement("w:p")
                    ppr = OxmlElement("w:pPr")
                    ppr.append(deepcopy(properties))
                    paragraph.append(ppr)
                    element.addprevious(paragraph)
                body.remove(element)
            marker = OxmlElement("w:p")
            heading.addnext(marker)
            self.markers[section.section_id] = marker
        self.typography = install_template_styles(self.document, body=first_body, headings=styles)
        self.body_styles = {}
        for section_id, prototype in self.prototypes.items():
            name = "Folio Template Body " + section_id
            if name not in self.document.styles:
                style = self.document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
                ppr, rpr = resolved_format(
                    self.document, prototype if prototype is not None else first_body
                )
                for item in list(ppr):
                    if item.tag in {qn("w:" + key) for key in ("pStyle", "numPr", "sectPr", "rPr")}:
                        ppr.remove(item)
                style.element.extend([ppr, rpr])
            self.body_styles[section_id] = self.document.styles[name]
        # Replace explicit cover placeholders only; retain logos and fixed labels.
        values = {"document_title": title} | (cover or {})
        for element in original[: contract.sections[0].body_index]:
            for paragraph in [element] if element.tag == qn("w:p") else xpath(element, ".//w:p"):
                _replace_placeholders(paragraph, values)
        self.current_section = None

    def start_section(self, section):
        if section.template_section_id:
            self.current_section = section.template_section_id
        if self.current_section is None:
            raise ValueError('New sections must be placed inside template sections that allow expansion.')
        return self.markers[self.current_section]

    def heading(self, section):
        return self.headings.get(section.template_section_id)

    def allows_headings(self):
        return next(
            section.allow_children
            for section in self.contract.sections
            if section.section_id == self.current_section
        )

    def paragraph(self, section, text):
        style = self.body_styles[section.template_section_id or self.current_section]
        return self.document.add_paragraph(text, style)

    def table(self, section, rows):
        for definition, prototype in self.tables.get(section.template_section_id, []):
            if definition.table_id in self.consumed_tables:
                continue
            if rows[0] != definition.headers:
                continue
            if (
                sum(any(value.strip() for value in row) for row in rows[1:])
                < definition.min_data_rows
            ):
                raise ValueError(f'A required template table has insufficient valid data rows: {definition.table_id}')
            self.consumed_tables.add(definition.table_id)
            original_rows = xpath(prototype, "./w:tr")
            headers = original_rows[: definition.header_rows]
            data_prototype = (
                original_rows[definition.header_rows]
                if len(original_rows) > definition.header_rows
                else _empty_data_row(prototype)
            )
            for original_row in original_rows[definition.header_rows :]:
                prototype.remove(original_row)
            for row in rows[1:]:
                element = deepcopy(data_prototype)
                for marker in xpath(element, ".//w:tblHeader"):
                    marker.getparent().remove(marker)
                # A repeating data row must have one physical cell per column.
                # Horizontally/vertically merged data is rejected at contract extraction.
                cells = xpath(element, "./w:tc")
                if len(cells) != len(row):
                    raise ValueError('Template data rows contain merged cells that cannot be safely populated.')
                for cell, value in zip(cells, row, strict=True):
                    paragraphs = xpath(cell, "./w:p")
                    paragraph = deepcopy(paragraphs[0]) if paragraphs else OxmlElement("w:p")
                    for child in list(cell):
                        if child.tag != qn("w:tcPr"):
                            cell.remove(child)
                    _replace_text(paragraph, value)
                    cell.append(paragraph)
                prototype.append(element)
            if len(headers) != definition.header_rows:
                raise ValueError('The template table is missing its declared header rows.')
            self.document.element.body.insert(len(self.document.element.body) - 1, prototype)
            table = Table(prototype, self.document._body)
            table.template_header_rows = definition.header_rows
            return table
        return None

    def missing_tables(self, section):
        return [
            definition
            for definition, _ in self.tables.get(section.template_section_id, [])
            if definition.table_id not in self.consumed_tables
        ]

    def placeholder_table(self, section, required):
        rows = [required.headers] + [
            ["To be completed" if self.language == "en" else "待填写"] * len(required.headers)
            for _ in range(required.min_data_rows)
        ]
        return self.table(section, rows)

    def finish(self):
        for marker in self.markers.values():
            marker.getparent().remove(marker)


def template_heading_matches(rendered, expected):
    return clean_heading(visible_text(rendered).strip()) == clean_heading(expected)


def _canonical(element):
    return etree.tostring(element, method="c14n", exclusive=True) if element is not None else None


def _check_owned_format(baseline, actual, *, label):
    paragraph_keys = {
        "jc",
        "spacing",
        "ind",
        "tabs",
        "textAlignment",
        "keepNext",
        "keepLines",
        "widowControl",
        "pageBreakBefore",
    }
    run_keys = {"rFonts", "sz", "szCs", "b", "bCs", "i", "iCs", "caps", "smallCaps", "vertAlign"}
    for expected, rendered, names in zip(baseline, actual, (paragraph_keys, run_keys), strict=True):
        for name in names:
            value = expected.find(qn("w:" + name))
            if value is not None and _canonical(value) != _canonical(
                rendered.find(qn("w:" + name))
            ):
                raise ValueError(f'Template body or cell formatting was modified: {label} / {name}')


def validate_template_output(path, plan):
    """Guard template-owned formatting after rendering and after cosmetic reviews."""
    if plan.template is None:
        return
    contract = plan.template
    original_path = Path(contract.path)
    if sha256(original_path.read_bytes()).hexdigest() != contract.source_hash:
        raise ValueError('The document template has changed. Regenerate and confirm the outline.')
    original, output = Document(original_path), Document(path)
    original_styles = {style.style_id: style.element for style in original.styles}
    output_styles = {style.style_id: style.element for style in output.styles}
    for key, style in original_styles.items():
        if _canonical(style) != _canonical(output_styles.get(key)):
            raise ValueError(f'A template style was modified: {key}')
    if _canonical(original.styles.element.find(qn("w:docDefaults"))) != _canonical(
        output.styles.element.find(qn("w:docDefaults"))
    ):
        raise ValueError("The template's default font or paragraph formatting was modified.")
    for element in original.part.numbering_part.element:
        match = next(
            (
                node
                for node in output.part.numbering_part.element
                if node.tag == element.tag and node.attrib == element.attrib
            ),
            None,
        )
        if _canonical(element) != _canonical(match):
            raise ValueError("The template's existing numbering definitions were modified.")
    if [_canonical(section._sectPr) for section in original.sections] != [
        _canonical(section._sectPr) for section in output.sections
    ]:
        raise ValueError('Template page or section settings were modified.')
    for relationship in original.part.rels.values():
        if relationship.reltype.rsplit("/", 1)[-1] not in {"header", "footer"}:
            continue
        match = output.part.rels.get(relationship.rId)
        if match is None or relationship.target_part.blob != match.target_part.blob:
            raise ValueError('A template header or footer was modified.')
    original_body = list(original.element.body)
    output_body = list(output.element.body)
    by_name = {
        bookmark.get(qn("w:name")): bookmark.getparent()
        for bookmark in output.element.xpath(".//w:bookmarkStart")
    }
    sections = {section.template_section_id: section for section in plan.sections}
    fixed_section_ids = {value.section_id for key, value in sections.items() if key}
    _template_mapping(
        [
            SimpleNamespace(
                title=visible_text(element).strip(), level=level, content_mode="generate"
            )
            for element in output_body
            if (level := heading_level(element, output))
        ],
        contract,
    )
    body_prototypes = {}
    for index, expected in enumerate(contract.sections):
        stop = (
            contract.sections[index + 1].body_index
            if index + 1 < len(contract.sections)
            else len(original_body)
        )
        body_prototypes[expected.section_id] = next(
            (
                element
                for element in original_body[expected.body_index + 1 : stop]
                if element.tag == qn("w:p") and visible_text(element).strip()
            ),
            None,
        )
    default_body = next((value for value in body_prototypes.values() if value is not None), None)
    previous_position = -1
    for expected in contract.sections:
        section = sections.get(expected.section_id)
        heading = by_name.get(section.section_id) if section else None
        if heading is None or not template_heading_matches(heading, expected.title):
            raise ValueError(f'A fixed template section is missing or modified: {expected.title}')
        position = output_body.index(heading)
        if position <= previous_position or heading_level(heading, output) != expected.level:
            raise ValueError(f'The order or level of a fixed template section was modified: {expected.title}')
        previous_position = position
        baseline_heading = original_body[expected.body_index]
        for name in ("pPr", "rPr"):
            # Compare effective properties so the explicit outline level used for
            # bookmarking is equivalent to the template's inherited heading level.
            baseline = resolved_format(original, baseline_heading)[name == "rPr"]
            actual = resolved_format(output, heading)[name == "rPr"]
            for value in (baseline, actual):
                for node in xpath(value, "./w:outlineLvl"):
                    value.remove(node)
            if _canonical(baseline) != _canonical(actual):
                raise ValueError(f'Template section styling or numbering was modified: {expected.title}')
        start = output_body.index(heading) + 1
        elements = []
        for element in output_body[start:]:
            # Added child sections inherit their owning fixed section's body
            # prototype. Stop only at the next fixed template heading.
            if any(
                name in fixed_section_ids for name in xpath(element, "./w:bookmarkStart/@w:name")
            ):
                break
            elements.append(element)
        prototype = body_prototypes[expected.section_id]
        if prototype is None:
            prototype = default_body
        baseline_body = resolved_format(original, prototype)
        for element in elements:
            if (
                element.tag == qn("w:p")
                and visible_text(element).strip()
                and not heading_level(element, output)
                # Figure captions use the separate, built-in figure layout;
                # authored business prose has a section-specific template style.
                and xpath(element, "./w:pPr/w:pStyle/@w:val") != [BODY.replace(" ", "")]
            ):
                _check_owned_format(
                    baseline_body, resolved_format(output, element), label=expected.title
                )
        tables = [element for element in elements if element.tag == qn("w:tbl")]
        for required in expected.tables:
            prototype = original_body[required.body_index]
            header_rows = table_rows(prototype)[: required.header_rows]
            table = next(
                (
                    item
                    for item in tables
                    if table_rows(item)[: required.header_rows] == header_rows
                ),
                None,
            )
            if table is None:
                raise ValueError(f'A required template table is missing: {required.table_id}')
            tables.remove(table)
            if (
                sum(
                    any(value.strip() for value in row)
                    for row in table_rows(table)[required.header_rows :]
                )
                < required.min_data_rows
            ):
                raise ValueError(f'A required template table has insufficient valid data rows: {required.table_id}')
            for tag in ("tblPr", "tblGrid"):
                if _canonical(prototype.find(qn("w:" + tag))) != _canonical(
                    table.find(qn("w:" + tag))
                ):
                    raise ValueError(f'Template table layout was modified: {required.table_id}')
            baseline_rows = xpath(prototype, "./w:tr")
            if len(baseline_rows) > required.header_rows:
                baseline_cells = xpath(baseline_rows[required.header_rows], "./w:tc")
            else:
                baseline_cells = []
            for row in xpath(table, "./w:tr")[required.header_rows :]:
                cells = xpath(row, "./w:tc")
                if baseline_cells and len(cells) != len(baseline_cells):
                    raise ValueError(f'The template table data column count was modified: {required.table_id}')
                for index, cell in enumerate(cells):
                    expected_paragraphs = (
                        xpath(baseline_cells[index], "./w:p") if baseline_cells else []
                    )
                    baseline_formats = (
                        resolved_format(original, expected_paragraphs[0])
                        if expected_paragraphs
                        else resolved_format(original, default_body)
                    )
                    for paragraph in xpath(cell, "./w:p"):
                        _check_owned_format(
                            baseline_formats,
                            resolved_format(output, paragraph),
                            label=required.table_id,
                        )
            for baseline, actual in zip(
                xpath(prototype, "./w:tr")[: required.header_rows],
                xpath(table, "./w:tr")[: required.header_rows],
                strict=True,
            ):
                baseline, actual = deepcopy(baseline), deepcopy(actual)
                for tag in ("tblHeader", "cantSplit"):
                    if not xpath(baseline, "./w:trPr/w:" + tag):
                        for node in xpath(actual, "./w:trPr/w:" + tag):
                            node.getparent().remove(node)
                for element in (baseline, actual):
                    for node in xpath(element, "./w:trPr"):
                        if not len(node) and not node.attrib:
                            node.getparent().remove(node)
                if _canonical(baseline) != _canonical(actual):
                    raise ValueError(f'Template table header styling was modified: {required.table_id}')
