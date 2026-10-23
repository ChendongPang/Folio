"""Insert validated, locally rendered figures without changing business paragraphs."""

import math
from collections import defaultdict
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.section import Section as WordSection
from docx.shared import Emu, Inches, Pt, RGBColor

from folio.document.figure_layout import MIN_FLOW_TEXT_PT, detail_views, word_crop
from folio.document.typography import BODY, prop
from folio.models import DocumentPlan, SectionState


class DocumentFigures:
    """Validate the whole manifest first, then emit in actual document order."""

    def __init__(
        self,
        plan: DocumentPlan,
        states: list[SectionState],
        figures: dict[str, list[dict]] | None,
        *,
        language: str = "zh-CN",
    ):
        self._english = language == "en"
        self._by_position = defaultdict(list)
        self._emitted = set()
        self._count = 0
        self._total = 0
        sections = {section.section_id for section in plan.sections}
        by_id = {state.section_id: state for state in states}
        ids = set()
        for section_id, items in (figures or {}).items():
            if section_id not in sections:
                raise ValueError(f"Figure refers to unknown section: {section_id}")
            if not items:
                continue
            if section_id not in by_id:
                raise ValueError(f"Figure section has no body state: {section_id}")
            block_ids = {block.block_id for block in by_id[section_id].content.blocks}
            for item in items:
                figure_id = item.get("figure_id")
                if (
                    not isinstance(figure_id, str)
                    or not figure_id.strip()
                    or (section_id, figure_id) in ids
                ):
                    raise ValueError(f"Missing or duplicate figure ID: {figure_id}")
                ids.add((section_id, figure_id))
                caption = item.get("caption")
                if not isinstance(caption, str) or not caption.strip():
                    raise ValueError(f"Figure caption is required: {figure_id}")
                if item.get("origin") not in {"evidence", "model_proposal"}:
                    raise ValueError(f"Invalid figure origin: {figure_id}")
                after = item.get("after_block_id") or None
                if after is not None and after not in block_ids:
                    raise ValueError(f"Unknown figure insertion block: {section_id}/{after}")
                if item.get("status") == "unresolved":
                    if not isinstance(item.get("reason"), str) or not item["reason"].strip():
                        raise ValueError(f"Unresolved figure requires a reason: {figure_id}")
                    sources = item.get("source_block_ids", [])
                    if not sources or not set(sources).issubset(block_ids):
                        raise ValueError(
                            f"Unresolved figure has invalid source blocks: {figure_id}"
                        )
                    self._by_position[(section_id, after)].append(dict(item))
                    self._total += 1
                    continue
                try:
                    image_path = Path(item["image_path"])
                    if not image_path.is_file():
                        raise ValueError("image file does not exist")
                    picture = Image.from_file(str(image_path))
                    if picture.content_type != "image/png":
                        raise ValueError("only rendered PNG files are accepted")
                    width, height = item["width"], item["height"]
                    if (
                        type(width) is not int
                        or type(height) is not int
                        or width <= 0
                        or height <= 0
                        or (width, height) != (picture.px_width, picture.px_height)
                    ):
                        raise ValueError("manifest dimensions do not match PNG dimensions")
                except Exception as exc:
                    raise ValueError(f"Invalid figure image {figure_id}: {exc}") from exc
                self._by_position[(section_id, after)].append(dict(item))
                self._total += 1

    def emit(self, document, section_id: str, after_block_id=None, *, place=None, anchor=None):
        """Place a figure after its complete block, or at the end of the section body."""
        key = (section_id, after_block_id)
        items = self._by_position.get(key, [])
        if not items:
            return
        if key in self._emitted:
            raise ValueError(f"Figure insertion position emitted twice: {key}")
        self._emitted.add(key)
        section = _word_section(document, anchor)
        for item in items:
            self._count += 1
            caption = f"{'Figure' if self._english else '图'} {self._count}  {item['caption'].strip()}"
            if item["origin"] == "model_proposal":
                caption += " (proposed design)" if self._english else "（拟议设计）"
            if item.get("status") == "unresolved":
                paragraph = document.add_paragraph(
                    caption + (" (pending confirmation: validation failed; figure not inserted)" if self._english else "（配图待确认：未通过校验，暂未插入）"), BODY
                )
                _figure_paragraph(paragraph, keep_next=False)
                if item["origin"] == "model_proposal":
                    paragraph.runs[0].font.color.rgb = RGBColor.from_string("FF0000")
                if place is not None:
                    place(paragraph._p)
                continue
            # Keep the complete figure. Only figures that would become unreadable
            # receive overlapping, labeled detail pages using the identical PNG.
            detail_caption = caption + self._detail_label(36, 36, 36, 36)
            max_width, max_height = _image_bounds(document, section, detail_caption)
            details = detail_views(item, max_width, max_height)
            views = [None, *details]
            for view_index, view in enumerate(views):
                label = caption
                if view is None:
                    if details:
                        label += " (overview; enlarged views follow)" if self._english else "（总览，放大图见随后各页）"
                    width, height = _fit_image(document, section, item, label)
                else:
                    label += self._detail_label(view['row'], view['rows'], view['column'], view['columns'])
                    left, top, right, bottom = view["box"]
                    pixel_width, pixel_height = right - left, bottom - top
                    scale = min(
                        max(1, MIN_FLOW_TEXT_PT / (item["text_font_px"] * 0.75)),
                        max_width / (pixel_width * 9525),
                        max_height / (pixel_height * 9525),
                    )
                    width, height = (
                        int(pixel_width * 9525 * scale),
                        int(pixel_height * 9525 * scale),
                    )
                image_paragraph = document.add_paragraph(style=BODY)
                _figure_paragraph(image_paragraph, keep_next=True)
                image_paragraph.paragraph_format.page_break_before = view is not None
                picture = image_paragraph.add_run().add_picture(
                    str(item["image_path"]), width=Emu(width), height=Emu(height)
                )
                if view is not None:
                    crop = OxmlElement("a:srcRect")
                    for attribute, value in word_crop(
                        view["box"], item["width"], item["height"]
                    ).items():
                        crop.set(attribute, str(value))
                    fill = picture._inline.graphic.graphicData.pic.blipFill
                    fill.insert(1, crop)
                picture._inline.docPr.set("descr", label)
                picture._inline.docPr.set("title", item["caption"].strip())
                picture._inline.docPr.set(
                    "name", f"Folio Figure {self._count} View {view_index}"
                )
                if place is not None:
                    place(image_paragraph._p)
                caption_paragraph = document.add_paragraph(label, BODY)
                _figure_paragraph(caption_paragraph, keep_next=False)
                caption_paragraph.paragraph_format.space_after = Pt(6)
                if item["origin"] == "model_proposal":
                    caption_paragraph.runs[0].font.color.rgb = RGBColor.from_string("FF0000")
                if place is not None:
                    place(caption_paragraph._p)

    def validate_complete(self):
        if self._count != self._total:
            raise ValueError(f"Some figures were not inserted: {self._count}/{self._total}")

    def _detail_label(self, row, rows, column, columns):
        if self._english:
            return f" (enlarged view: row {row}/{rows}, column {column}/{columns}; adjacent views overlap)"
        return f"（放大分图：第 {row}/{rows} 行、第 {column}/{columns} 列，相邻图有重叠）"


def _word_section(document, anchor):
    if anchor is not None:
        # A paragraph's section ends at the next sectPr, not at the previous break.
        for element in [anchor, *anchor.itersiblings()]:
            properties = (
                element
                if element.tag == qn("w:sectPr")
                else element.find("./" + qn("w:pPr") + "/" + qn("w:sectPr"))
            )
            if properties is not None:
                return WordSection(properties, document.part)
    return document.sections[-1]


def _image_bounds(document, section, caption):
    page_width = section.page_width or Inches(8.27)
    page_height = section.page_height or Inches(11.69)
    max_width = page_width - (section.left_margin or 0) - (section.right_margin or 0)
    max_width -= section.gutter or 0
    columns = section._sectPr.find(qn("w:cols"))
    if columns is not None:
        widths = [int(col.get(qn("w:w"))) * 635 for col in columns if col.get(qn("w:w"))]
        count = int(columns.get(qn("w:num"), "1"))
        if widths:
            max_width = min(max_width, min(widths))
        elif count > 1:
            gap = int(columns.get(qn("w:space"), "720")) * 635
            max_width = (max_width - gap * (count - 1)) / count
    font_size = document.styles[BODY].font.size or Pt(11)
    caption_lines = max(1, math.ceil(len(caption) * font_size / max(1, max_width)))
    max_height = page_height - (section.top_margin or 0) - (section.bottom_margin or 0)
    # Reserve room for the caption, baseline and modest paragraph spacing on the same page.
    max_height -= font_size * 1.5 * (caption_lines + 1) + Pt(12)
    if max_width <= 0 or max_height <= 0:
        raise ValueError("Page layout has no room for figure")
    return max_width, max_height


def _fit_image(document, section, item, caption):
    max_width, max_height = _image_bounds(document, section, caption)
    width, height = item["width"] * Inches(1) / 96, item["height"] * Inches(1) / 96
    scale = min(1, max_width / width, max_height / height)
    return max(1, int(width * scale)), max(1, int(height * scale))


def _figure_paragraph(paragraph, *, keep_next):
    formatting = paragraph.paragraph_format
    formatting.alignment = WD_ALIGN_PARAGRAPH.CENTER
    formatting.keep_with_next = keep_next
    formatting.keep_together = True
    formatting.page_break_before = False
    formatting.space_before = Pt(0)
    formatting.space_after = Pt(0)
    # Exact line heights inherited from body text otherwise crop inline drawings.
    formatting.line_spacing = 1.0 if keep_next else 1.15
    ppr = paragraph._p.get_or_add_pPr()
    prop(
        ppr,
        "ind",
        left=0,
        right=0,
        firstLine=0,
        hanging=0,
        leftChars=0,
        rightChars=0,
        firstLineChars=0,
        hangingChars=0,
    )
    prop(ppr, "outlineLvl", val=9)
    ppr.get_or_add_numPr().get_or_add_numId().val = 0
