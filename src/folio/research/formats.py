"""Bounded, versioned source readers shared by planning, retrieval and research.

Readers never execute Office macros, formulas, shell commands or OCR on the host.
Scanned pages and images remain explicitly unread until an isolated extraction is
registered. Source locations are generated here, never accepted from a model.
"""

from __future__ import annotations

import hashlib
import io
import math
import posixpath
import re
from bisect import insort
from dataclasses import asdict, dataclass, field
from pathlib import Path
from zipfile import BadZipFile, ZipFile

SUPPORTED_SUFFIXES = frozenset(
    {
        ".docx",
        ".md",
        ".txt",
        ".xlsx",
        ".pdf",
        ".pptx",
        ".png",
        ".jpg",
        ".jpeg",
        ".tif",
        ".tiff",
        ".webp",
    }
)
IMAGE_SUFFIXES = frozenset({".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp"})
MAX_TEXT_CHARS = 8_000_000
MAX_BLOCKS = 100_000
MAX_PAGES = 2000
MAX_CELLS = 500_000
MAX_BOOKMARKS = 10_000
MAX_MERGE_RANGES = 500
READER_VERSION = "source-reader-v2.3"


class SourceFormatError(ValueError):
    """A malformed, unsupported or over-capacity source, with a useful reason."""


@dataclass
class SourceBlock:
    locator: str
    text: str
    kind: str = "paragraph"
    metadata: dict = field(default_factory=dict)


@dataclass
class SourceDocument:
    source_version: str
    metadata: dict
    blocks: list[SourceBlock] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def model_dump(self):
        return asdict(self)

    @classmethod
    def from_dict(cls, value):
        return cls(
            source_version=value["source_version"],
            metadata=value["metadata"],
            blocks=[SourceBlock(**block) for block in value["blocks"]],
            warnings=value.get("warnings", []),
        )


def source_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _package(path: Path):
    expected = {
        ".docx": "word/document.xml",
        ".xlsx": "xl/workbook.xml",
        ".pptx": "ppt/presentation.xml",
    }
    try:
        with ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > 5000:
                raise SourceFormatError('The Office file contains too many entries (maximum 5,000)')
            if expected[path.suffix.lower()] not in archive.namelist():
                raise SourceFormatError('The Office file structure does not match its extension')
            if any(item.flag_bits & 1 for item in members):
                raise SourceFormatError('Encrypted Office files are not supported')
    except BadZipFile as exc:
        raise SourceFormatError('The Office file is not a valid ZIP document') from exc


def _append(document: SourceDocument, block: SourceBlock):
    if not block.text.strip():
        return
    if len(document.blocks) >= MAX_BLOCKS:
        raise SourceFormatError(f'The source exceeds {MAX_BLOCKS} text blocks. Split the file')
    count = document.metadata.get("characters", 0) + len(block.text)
    if count > MAX_TEXT_CHARS:
        raise SourceFormatError(f'The source text exceeds {MAX_TEXT_CHARS} characters. Split the file')
    document.metadata["characters"] = count
    document.blocks.append(block)


def _text(path: Path, document: SourceDocument):
    text = path.read_text(encoding="utf-8-sig")
    fence = None
    for number, line in enumerate(text.splitlines(), 1):
        markdown = path.suffix.lower() == ".md"
        marker = re.match(r"^ {0,3}(`{3,}|~{3,})", line) if markdown else None
        in_code = fence is not None
        if marker:
            if fence is None:
                fence = marker[1]
            elif marker[1][0] == fence[0] and len(marker[1]) >= len(fence) and not line[marker.end():].strip():
                fence = None
        match = re.match(r"^ {0,3}(#{1,6})[ \t]+(.+)", line) if markdown and not in_code and not marker else None
        metadata = {"line": number}
        if match:
            metadata["level"] = len(match[1])
        if markdown and not in_code and not marker and re.fullmatch(r" {0,3}(?:=+|-+)[ \t]*", line):
            previous = document.blocks[-1] if document.blocks else None
            if (
                previous and previous.kind == "paragraph"
                and previous.metadata.get("line") == number - 1
                and not re.match(r"^(?: {4}|\t| {0,3}(?:[-+*>]|\d+[.)])[ \t])", previous.text)
                and not re.match(r"^ {0,3}(`{3,}|~{3,})", previous.text)
            ):
                previous.kind = "heading"
                previous.metadata["level"] = 1 if line.lstrip().startswith("=") else 2
        _append(
            document,
            SourceBlock(f"line {number}", line, "heading" if match else "paragraph", metadata),
        )


def _word(path: Path, document: SourceDocument):
    from folio.document.reader import read_docx

    parsed = read_docx(path)
    document.warnings.extend(parsed.warnings)
    document.metadata["comments"] = parsed.comments
    document.metadata["revisions"] = parsed.revisions
    document.metadata["unread_visuals"] = any(
        "图片或文本框" in warning or "images or text boxes" in warning.lower()
        for warning in parsed.warnings
    )
    document.metadata["needs_conversion"] = document.metadata["unread_visuals"]
    for block in parsed.blocks:
        metadata = {"heading_path": block.path, "comment_ids": block.comment_ids}
        if block.level:
            metadata["level"] = block.level
        if block.kind == "table":
            for index, row in enumerate(block.rows, 1):
                _append(
                    document,
                    SourceBlock(
                        f"{block.locator}/row[{index}]",
                        " | ".join(row),
                        "table_row",
                        metadata | {"row": index, "cells": row},
                    ),
                )
        else:
            _append(document, SourceBlock(block.locator, block.text, block.kind, metadata))


def _excel(path: Path, document: SourceDocument):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    sheets, formulas, total_cells = [], 0, 0
    sheet_metadata = _excel_sheet_metadata(path)
    unread_sheets = []
    try:
        for sheet in book.worksheets:
            metadata = sheet_metadata.get(sheet.title, {
                "rows": 0, "columns": 0, "range_cells": 0, "data_cells": 0,
                "merged_cells": [], "merged_cells_total": 0,
                "merged_cells_truncated": False, "unread_visual_count": 0,
            })
            rows, columns = metadata["rows"], metadata["columns"]
            # Worksheet dimensions are untrusted metadata. They can be inflated by
            # deleted content or formatting while the actual cells stay small.
            if metadata["range_cells"] + total_cells > MAX_CELLS:
                raise SourceFormatError(
                    f'The actual Excel data range exceeds {MAX_CELLS} cells. Reduce the used range or split the file'
                )
            total_cells += metadata["range_cells"]
            summary = {
                "name": sheet.title,
                "rows": rows,
                "columns": columns,
                "data_cells": metadata["data_cells"],
                "range_cells": metadata["range_cells"],
                "hidden": sheet.sheet_state != "visible",
                "headers": [],
                "header_basis": "first_nonempty_row",
                "merged_cells": metadata["merged_cells"],
                "merged_cells_total": metadata["merged_cells_total"],
                "merged_cells_truncated": metadata["merged_cells_truncated"],
            }
            visual_count = metadata["unread_visual_count"]
            summary["unread_visual_count"] = visual_count
            if visual_count:
                unread_sheets.append(sheet.title)
            sheets.append(summary)
            selected_range = {
                "min_row": 1,
                "max_row": rows,
                "min_col": 1,
                "max_col": columns,
            }
            cached_rows = cached[sheet.title].iter_rows(**selected_range)
            for row_index, (row, cached_row) in enumerate(
                zip(sheet.iter_rows(**selected_range), cached_rows, strict=True), 1
            ):
                values, cells, formula_cells = [], [], []
                for index, cell in enumerate(row):
                    value = cell.value
                    coordinate = f"{get_column_letter(index + 1)}{row_index}"
                    if cell.data_type == "f":
                        formulas += 1
                        cached_value = cached_row[index].value
                        formula_cells.append(
                            {
                                "cell": coordinate,
                                "formula": str(value),
                                "cached_value": None if cached_value is None else str(cached_value),
                                "cache_verified": False,
                            }
                        )
                        rendered = f"{value} [cached={cached_value if cached_value is not None else 'unavailable'}; unverified]"
                    else:
                        rendered = "" if value is None else str(value)
                    values.append(rendered)
                    cells.append({"cell": coordinate, "value": rendered})
                while values and not values[-1]:
                    values.pop()
                    cells.pop()
                if not any(values):
                    continue
                if not summary["headers"]:
                    summary["headers"] = [value[:120] for value in values[:20]]
                    summary["headers_truncated"] = len(values) > 20 or any(
                        len(value) > 120 for value in values[:20]
                    )
                    summary["header_row"] = row_index
                quoted = sheet.title.replace("'", "''")
                locator = (
                    f"sheet '{quoted}'!A{row_index}:{get_column_letter(len(values))}{row_index}"
                )
                _append(
                    document,
                    SourceBlock(
                        locator,
                        " | ".join(values),
                        "table_row",
                        {
                            "sheet": sheet.title,
                            "row": row_index,
                            "cells": cells,
                            "formulas": formula_cells,
                        },
                    ),
                )
        document.metadata["sheets"] = sheets
        document.metadata["formula_count"] = formulas
        document.metadata.update(unread_sheets=unread_sheets, unread_visuals=bool(unread_sheets),
                                 needs_conversion=bool(unread_sheets))
        if unread_sheets:
            document.warnings.append('Excel contains unread images, charts, or drawings. Convert and verify the visual content of the relevant sheets.')
        if formulas:
            document.warnings.append(
                'Excel formulas were not calculated while reading. Cached values are unverified; recalculate in an isolated Office environment and verify the method before using formula results.'
            )
    finally:
        book.close()
        cached.close()


def _excel_sheet_metadata(path: Path) -> dict:
    """Read actual data bounds and merge metadata without trusting dimensions."""
    from lxml import etree
    from openpyxl.utils import range_boundaries

    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    with ZipFile(path) as archive:
        workbook = etree.fromstring(archive.read("xl/workbook.xml"), parser)
        relationships = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"), parser)
        targets = {
            item.get("Id"): item.get("Target")
            for item in relationships
            if item.get("TargetMode") != "External" and item.get("Type", "").endswith("/worksheet")
        }
        result = {}
        for sheet in workbook.iter():
            if not isinstance(sheet.tag, str) or sheet.tag.rsplit("}", 1)[-1] != "sheet":
                continue
            relationship_id = next((value for key, value in sheet.attrib.items() if key.endswith("}id")), None)
            target = targets.get(relationship_id)
            if not target:
                continue
            member = posixpath.normpath(target.lstrip("/") if target.startswith("/") else "xl/" + target)
            if not member.startswith("xl/") or member not in archive.namelist():
                continue
            unread_visual_count = _excel_visual_count(archive, member, parser)
            ranges, count = [], 0
            min_row = min_column = max_row = max_column = None
            data_cells = 0
            for _, element in etree.iterparse(io.BytesIO(archive.read(member)), events=("end",), resolve_entities=False, no_network=True):
                tag = element.tag.rsplit("}", 1)[-1] if isinstance(element.tag, str) else ""
                if tag == "c" and _excel_cell_has_data(element):
                    coordinate = element.get("r", "")
                    match = re.fullmatch(r"([A-Z]{1,3})([1-9]\d*)", coordinate)
                    if match:
                        column = _excel_column_index(match.group(1))
                        row = int(match.group(2))
                        data_cells += 1
                        min_row = row if min_row is None else min(min_row, row)
                        max_row = row if max_row is None else max(max_row, row)
                        min_column = column if min_column is None else min(min_column, column)
                        max_column = column if max_column is None else max(max_column, column)
                elif tag == "mergeCell":
                    reference = element.get("ref", "")
                    if re.fullmatch(r"[A-Z]{1,3}[1-9]\d*:[A-Z]{1,3}[1-9]\d*", reference):
                        count += 1
                        column, row, max_cell_column, max_cell_row = range_boundaries(reference)
                        min_row = row if min_row is None else min(min_row, row)
                        max_row = max_cell_row if max_row is None else max(max_row, max_cell_row)
                        min_column = column if min_column is None else min(min_column, column)
                        max_column = (
                            max_cell_column if max_column is None
                            else max(max_column, max_cell_column)
                        )
                        insort(ranges, (row, column, reference))
                        if len(ranges) > MAX_MERGE_RANGES:
                            ranges.pop()
                if tag in {"c", "mergeCell"}:
                    # Values live below the cell element, so clear only after the
                    # complete cell has been inspected.
                    element.clear()
                    while element.getprevious() is not None:
                        del element.getparent()[0]
            rows = max_row or 0
            columns = max_column or 0
            result[sheet.get("name", "")] = {
                "rows": rows,
                "columns": columns,
                "data_cells": data_cells,
                "range_cells": rows * columns,
                "min_row": min_row,
                "max_row": max_row,
                "min_column": min_column,
                "max_column": max_column,
                "merged_cells": [item[2] for item in ranges],
                "merged_cells_total": count,
                "merged_cells_truncated": count > len(ranges),
                "unread_visual_count": unread_visual_count,
            }
        return result


def _excel_cell_has_data(cell) -> bool:
    """A present cell object can be style-only; count actual values and formulas."""
    for child in cell:
        tag = child.tag.rsplit("}", 1)[-1] if isinstance(child.tag, str) else ""
        if tag == "f":
            return True
        if tag == "v" and (child.text or "").strip():
            return True
        if tag == "is" and any((node.text or "").strip() for node in child.iter() if isinstance(node.tag, str)):
            return True
    return False


def _excel_column_index(letters: str) -> int:
    value = 0
    for letter in letters:
        value = value * 26 + ord(letter) - ord("A") + 1
    return value


def _excel_visual_count(archive, member, parser):
    """Inspect drawings; only explicit author decorative markers exempt shapes."""
    from lxml import etree

    directory, filename = posixpath.split(member)
    relname = directory + "/_rels/" + filename + ".rels"
    if relname not in archive.namelist():
        return 0
    relationships = etree.fromstring(archive.read(relname), parser)
    count = 0
    for relationship in relationships:
        if not relationship.get("Type", "").endswith(("/drawing", "/vmlDrawing")):
            continue
        if relationship.get("TargetMode") == "External":
            count += 1
            continue
        target = relationship.get("Target", "")
        drawing = posixpath.normpath(target.lstrip("/") if target.startswith("/") else directory + "/" + target)
        if not drawing.startswith("xl/") or drawing not in archive.namelist():
            count += 1
            continue
        root = etree.fromstring(archive.read(drawing), parser)
        for element in root.iter():
            if not isinstance(element.tag, str) or element.tag.rsplit("}", 1)[-1] not in {
                "pic", "graphicFrame", "sp", "cxnSp", "shape",
            }:
                continue
            decorative = any(
                isinstance(child.tag, str) and child.tag.rsplit("}", 1)[-1] == "decorative"
                and child.get("val") in {"1", "true"} for child in element.iter()
            )
            if not decorative:
                count += 1
    return count


def _pdf_has_visuals(page):
    """Find images and vector diagrams without decoding pixels.

    Explicit PDF /Artifact marks exempt decorative content. Position and size
    cannot establish that an ordinary image contains no business requirement.
    """
    from pypdf.generic import ContentStream

    seen = set()

    def graphics(contents):
        used, decorative, visible = set(), [], False
        for operands, operator in contents.operations if contents is not None else []:
            if operator in {b"BMC", b"BDC"}:
                decorative.append(bool(decorative and decorative[-1]) or bool(operands and operands[0] == "/Artifact"))
            elif operator == b"EMC":
                if decorative:
                    decorative.pop()
            elif not (decorative and decorative[-1]):
                if operator == b"Do" and operands:
                    used.add(str(operands[0]))
                elif operator in {b"INLINE IMAGE", b"S", b"s", b"f", b"F", b"f*", b"B", b"B*", b"b", b"b*"}:
                    # Text extraction loses arrows, connections, and chart
                    # geometry even when every node label is native text.
                    visible = True
        return used, visible

    used, visible = graphics(page.get_contents())
    if visible:
        return True

    def images(resources, depth=0, names=None):
        if not resources or depth > 30:
            return depth > 30
        resources = resources.get_object()
        xobjects = resources.get("/XObject", {})
        xobjects = xobjects.get_object() if hasattr(xobjects, "get_object") else xobjects
        for name, reference in xobjects.items():
            if names is not None and str(name) not in names:
                continue
            item = reference.get_object()
            identity = id(item)
            if identity in seen:
                continue
            seen.add(identity)
            if item.get("/Subtype") == "/Image":
                # Even one pixel can be stretched into a meaningful status
                # colour or chart bar. Only explicit /Artifact marks exempt it.
                return True
            elif item.get("/Subtype") == "/Form":
                nested_names, nested_visuals = graphics(ContentStream(item, getattr(page, "pdf", None)))
                if nested_visuals or images(item.get("/Resources", resources), depth + 1, nested_names):
                    return True
        return False

    return images(page.get("/Resources"), names=used)


def _pdf_outline(reader, document):
    """Keep real bookmark destinations and hierarchy, never infer PDF headings."""
    outlines, by_path, skipped, limited = [], {}, 0, False
    try:
        source = reader.outline
    except Exception:
        document.metadata.update(outline=[], outline_status="unreadable")
        document.warnings.append('PDF bookmarks could not be parsed. Navigation uses actual page numbers without inferring sections.')
        return
    stack = [(iter(enumerate(source, 1)), (), None, 1)]
    visited = 0
    while stack:
        iterator, prefix, parent_path, level = stack[-1]
        try:
            ordinal, item = next(iterator)
        except StopIteration:
            stack.pop()
            continue
        visited += 1
        if visited > MAX_BOOKMARKS or len(stack) > 100:
            limited = True
            break
        bookmark_path = prefix + (ordinal,)
        if isinstance(item, list):
            # pypdf represents child lists immediately after their parent.
            previous_path = prefix + (ordinal - 1,)
            parent = by_path.get(previous_path)
            stack.append((iter(enumerate(item, 1)), bookmark_path,
                          parent["bookmark_path"] if parent else parent_path,
                          parent["level"] + 1 if parent else level))
            continue
        try:
            page_index = reader.get_destination_page_number(item)
            title = str(getattr(item, "title", "") or item.get("/Title", "")).strip()
            if not title or isinstance(page_index, bool) or not isinstance(page_index, int) or not 0 <= page_index < len(reader.pages):
                raise ValueError("invalid bookmark destination")
            entry = {
                "title": title[:1000], "title_truncated": len(title) > 1000,
                "level": level, "bookmark_path": list(bookmark_path),
                "parent_bookmark_path": parent_path,
                "page": page_index + 1,
                "destination_type": str(getattr(item, "typ", "") or ""),
            }
            try:
                top = getattr(item, "top", None)
                if top is not None and math.isfinite(float(top)):
                    entry["top"] = float(top)
            except (TypeError, ValueError):
                pass
            outlines.append(entry)
            by_path[bookmark_path] = entry
        except (AttributeError, KeyError, TypeError, ValueError):
            skipped += 1
    ordered = all(left["page"] <= right["page"] for left, right in zip(outlines, outlines[1:], strict=False))
    document.metadata.update(
        outline=outlines, outline_status="partial" if skipped or limited else "available" if outlines else "absent",
        outline_ordered=ordered,
    )
    if skipped:
        document.warnings.append(f'Skipped {skipped} PDF bookmarks lacking valid page numbers in this file. Their section ranges were not inferred.')
    if limited:
        document.warnings.append('The PDF bookmark outline exceeds the structure-reading limit. Bookmarks are navigation aids; continue reading body text by page number.')
    if not ordered:
        document.warnings.append('PDF bookmark order differs from actual page order. Only bookmark target pages are provided; complete section boundaries are not inferred.')


def _pdf(path: Path, document: SourceDocument):
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted:
        raise SourceFormatError('Encrypted PDFs are not supported. Provide a readable copy first')
    if len(reader.pages) > MAX_PAGES:
        raise SourceFormatError(f'The PDF exceeds {MAX_PAGES} pages. Split the file')
    _pdf_outline(reader, document)
    unread, visual_pages = [], []
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        has_visuals = _pdf_has_visuals(page)
        if has_visuals:
            visual_pages.append(number)
        if not text.strip() or has_visuals:
            unread.append(number)
        for index, paragraph in enumerate(re.split(r"\n\s*\n", text), 1):
            _append(
                document,
                SourceBlock(
                    f"page {number}/block {index}",
                    paragraph,
                    "paragraph",
                    {"page": number, "block": index},
                ),
            )
    document.metadata.update(
        page_count=len(reader.pages), needs_ocr=bool(unread), unread_pages=unread,
        unread_visual_pages=visual_pages, unread_visuals=bool(visual_pages)
    )
    if unread:
        document.warnings.append(
            f'{len(unread)} pages were not fully read: no text was extracted, or images or drawing regions remain. Inspect and verify original pages; titles or page numbers alone cannot establish completeness.'
        )


def _slides(path: Path, document: SourceDocument):
    from pptx import Presentation

    presentation = Presentation(path)
    if len(presentation.slides) > MAX_PAGES:
        raise SourceFormatError(f'The PPTX exceeds {MAX_PAGES} pages. Split the file')
    slides, unread_visuals = [], []

    def walk(shapes, prefix, number):
        for index, shape in enumerate(shapes, 1):
            locator = f"{prefix}/shape {index}"
            if shape.shape_type in {3, 7, 10, 11, 12, 13, 16, 21, 24, 26} or hasattr(
                shape, "image"
            ):
                unread_visuals.append(locator)
            if hasattr(shape, "shapes"):
                walk(shape.shapes, locator, number)
            if shape.has_text_frame:
                _append(
                    document,
                    SourceBlock(
                        locator,
                        shape.text_frame.text,
                        "paragraph",
                        {"slide": number, "shape": index},
                    ),
                )
            if shape.has_table:
                for row_index, row in enumerate(shape.table.rows, 1):
                    cells = [cell.text for cell in row.cells]
                    _append(
                        document,
                        SourceBlock(
                            f"{locator}/row {row_index}",
                            " | ".join(cells),
                            "table_row",
                            {"slide": number, "row": row_index, "cells": cells},
                        ),
                    )

    for number, slide in enumerate(presentation.slides, 1):
        slides.append(
            {"slide": number, "title": slide.shapes.title.text[:300] if slide.shapes.title else ""}
        )
        walk(slide.shapes, f"slide {number}", number)
        if slide.has_notes_slide:
            _append(
                document,
                SourceBlock(
                    f"slide {number}/notes",
                    slide.notes_slide.notes_text_frame.text,
                    "notes",
                    {"slide": number},
                ),
            )
    document.metadata.update(
        slide_count=len(slides),
        slides=slides,
        unread_visuals=bool(unread_visuals),
        unread_slides=sorted({int(re.match(r"slide (\d+)", item)[1]) for item in unread_visuals}),
        needs_conversion=bool(unread_visuals),
    )
    document.warnings.append(
        'PPTX text, tables, and notes were read. Embedded images, charts, and complex drawings require verification after isolated conversion and OCR.'
    )


def _image(path: Path, document: SourceDocument):
    from PIL import Image

    with Image.open(path) as image:
        if image.width * image.height > 80_000_000:
            raise SourceFormatError('The image exceeds 80 million pixels. Resize it before uploading')
        if getattr(image, "n_frames", 1) > MAX_PAGES:
            raise SourceFormatError(f'The multipage image exceeds {MAX_PAGES} pages. Split the file')
        document.metadata.update(
            width=image.width,
            height=image.height,
            page_count=getattr(image, "n_frames", 1),
            needs_ocr=True,
        )
        image.verify()
    document.warnings.append(
        'The image has not undergone OCR. Do not treat the filename or dimensions as business content. Use the isolated OCR tool.'
    )


def read_source(path: Path | str) -> SourceDocument:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise SourceFormatError(f"Unsupported source format: {suffix or '(no extension)'}")
    if not path.is_file():
        raise SourceFormatError('The source file does not exist')
    document = SourceDocument(
        source_hash(path),
        {
            "format": suffix[1:],
            "size": path.stat().st_size,
            "reader_version": READER_VERSION,
            "characters": 0,
        },
    )
    try:
        if suffix in {".docx", ".xlsx", ".pptx"}:
            _package(path)
        if suffix in {".txt", ".md"}:
            _text(path, document)
        elif suffix == ".docx":
            _word(path, document)
        elif suffix == ".xlsx":
            _excel(path, document)
        elif suffix == ".pdf":
            _pdf(path, document)
        elif suffix == ".pptx":
            _slides(path, document)
        else:
            _image(path, document)
    except SourceFormatError:
        raise
    except ImportError as exc:
        raise SourceFormatError(f"The dependency for reading {suffix} is missing. Install the project's complete dependencies") from exc
    except Exception as exc:
        raise SourceFormatError(f'Cannot read the {suffix} source: {type(exc).__name__}') from exc
    document.metadata["block_count"] = len(document.blocks)
    document.metadata["complete"] = not (
        document.metadata.get("needs_ocr", False) or document.metadata.get("unread_visuals", False)
    )
    return document
