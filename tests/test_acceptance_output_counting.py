"""The offline DOCX precheck must not count TOC caches as authored body text."""

import importlib
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1] / "examples" / "acceptance_cases"


@pytest.fixture
def checker(monkeypatch):
    monkeypatch.syspath_prepend(str(ROOT))
    return importlib.import_module("check_output")


def inspect(doc, tmp_path, checker):
    path = tmp_path / "result.docx"
    doc.save(path)
    return checker.inspect_document(path)


def field_char(paragraph, kind):
    node = OxmlElement("w:fldChar")
    node.set(qn("w:fldCharType"), kind)
    paragraph.add_run()._r.append(node)


def instruction(paragraph, text):
    node = OxmlElement("w:instrText")
    node.text = text
    paragraph.add_run()._r.append(node)


def test_body_heading_can_inherit_through_custom_styles(tmp_path, checker):
    doc = Document()
    base = doc.styles.add_style("业务章标题", WD_STYLE_TYPE.PARAGRAPH)
    base.base_style = doc.styles["Heading 1"]
    custom = doc.styles.add_style("业务章标题变体", WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = base
    doc.add_paragraph("封面不计入")
    doc.add_paragraph("第一章", style=custom)
    doc.add_paragraph("正文甲乙")
    result = inspect(doc, tmp_path, checker)
    assert result["effective_body_chars"] == 4
    assert result["headings"] == ["第一章"]
    assert result["heading_detection"] == "recognized"
    assert result["counting_warnings"] == []


def test_custom_style_outline_level_is_respected(tmp_path, checker):
    doc = Document()
    custom = doc.styles.add_style("Section Label", WD_STYLE_TYPE.PARAGRAPH)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "2")
    custom.element.get_or_add_pPr().append(outline)
    doc.add_paragraph("第一节", style=custom)
    doc.add_paragraph("正文甲乙")
    overridden = doc.add_paragraph("普通段落", style=custom)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "9")
    overridden._p.get_or_add_pPr().append(outline)
    result = inspect(doc, tmp_path, checker)
    assert result["effective_body_chars"] == 8
    assert result["headings"] == ["第一节"]


def test_simple_toc_field_cache_is_excluded(tmp_path, checker):
    doc = Document()
    doc.add_heading("目录", 1)
    paragraph = doc.add_paragraph()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), ' TOC \\o "1-3" ')
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "目录缓存ABCDE123"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)
    doc.add_heading("第一章", 1)
    doc.add_paragraph("正文甲乙")
    result = inspect(doc, tmp_path, checker)
    assert result["effective_body_chars"] == 4


def test_complex_toc_spans_paragraphs_and_nested_pageref(tmp_path, checker):
    doc = Document()
    doc.add_heading("目录", 1)
    first = doc.add_paragraph()
    field_char(first, "begin")
    instruction(first, " TO")
    second = doc.add_paragraph()
    instruction(second, 'C \\o "1-3" ')
    field_char(second, "separate")
    second.add_run("第一条目录缓存")
    third = doc.add_paragraph("第二条目录缓存")
    field_char(third, "begin")
    instruction(third, " PAGEREF _Toc123 \\h ")
    field_char(third, "separate")
    third.add_run("12")
    field_char(third, "end")
    third.add_run("页码后的目录文字")
    doc.add_paragraph("下一段仍然属于目录")
    field_char(doc.add_paragraph(), "end")
    doc.add_heading("第一章", 1)
    doc.add_paragraph("正文甲乙")
    result = inspect(doc, tmp_path, checker)
    assert result["effective_body_chars"] == 4
    assert result["headings"] == ["目录", "第一章"]


def test_other_field_results_remain_visible_body_text(tmp_path, checker):
    doc = Document()
    doc.add_heading("第一章", 1)
    paragraph = doc.add_paragraph("页")
    field_char(paragraph, "begin")
    instruction(paragraph, " PAGE ")
    field_char(paragraph, "separate")
    paragraph.add_run("12")
    field_char(paragraph, "end")
    result = inspect(doc, tmp_path, checker)
    assert result["effective_body_chars"] == 3


def test_empty_document_keeps_zero_and_reports_missing_heading(tmp_path, checker):
    result = inspect(Document(), tmp_path, checker)
    assert result["effective_body_chars"] == 0
    assert result["headings"] == []
    assert result["heading_detection"] == "no_recognized_heading"
    assert result["counting_warnings"]
