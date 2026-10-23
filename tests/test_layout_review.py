import asyncio
import base64
import hashlib
import json
import shutil
import threading
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from docx import Document

from folio.document.formatting import (
    FORMAT_POLICY,
    LayoutPatch,
    collect_targets,
    document_signature,
    protected_signature,
    replay_formatting,
    replay_page_review,
)
from folio.document.page_preview import (
    materialize_page_images,
    prepare_page_preview,
    render_page_preview,
)
from folio.document.pagination import office_runtime
from folio.document.reader import visible_text
from folio.document.word import render
from folio.model_settings import ModelSettingsInput, ModelSettingsStore
from folio.models import (
    Block,
    DocumentPlan,
    Extraction,
    JobInput,
    Section,
    SectionContent,
    SectionState,
)
from folio.orchestration.layout_review import (
    MAX_REVIEW_CALLS,
    MAX_REVIEW_PAGES,
    MAX_REVIEW_SECONDS,
    PageLayoutIssue,
    PageLayoutReview,
    review_rendered_document,
)
from folio.runtime.llm import LLMRuntime

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jX1sAAAAASUVORK5CYII="
)


@pytest.fixture
def source(tmp_path):
    plan = DocumentPlan(
        sections=[
            Section(
                section_id="s1",
                node_id="o1",
                level=1,
                display_number="1",
                title="审批规则",
                instruction="写审批规则",
            )
        ]
    )
    states = [
        SectionState(
            section_id="s1",
            content=SectionContent(
                blocks=[
                    Block(
                        block_id="b1", text="金额 100 万元以上需要主管审批。", evidence_ids=["e1"]
                    ),
                    Block(block_id="b2", text="建议保留审批日志。", origin="model_proposal"),
                ]
            ),
            extraction=Extraction(summary="审批"),
            evidence=[],
        )
    ]
    target = tmp_path / "result.docx"
    render(target, plan, states, title="审批需求", draft=True, defer_pagination=True)
    return target, plan, states


def preview(source, directory, *, max_pages):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    document = directory / "paginated.docx"
    pdf = directory / "document.pdf"
    image = directory / "page.png"
    shutil.copyfile(source, document)
    pdf.write_bytes(b"PDF")
    image.write_bytes(PNG)
    return {
        "status": "rendered",
        "total_pages": 1,
        "pages": [{"page": 1, "path": str(image), "text": visible_text(Document(source).element)}],
        "document": str(document),
        "pdf": str(pdf),
        "layout": {"status": "paginated", "total_pages": 1, "pages": {"s1": 1}},
    }


async def run_review(source, tmp_path, call, **kwargs):
    path, plan, states = source
    return await review_rendered_document(
        path,
        plan,
        states,
        tmp_path,
        call,
        lambda _: None,
        enabled=True,
        supports_images=True,
        renderer=preview,
        **kwargs,
    )


async def test_page_patches_commit_only_after_rendering_and_clean_recheck(source, tmp_path):
    path, plan, states = source
    before = Document(path)
    signature = protected_signature(before, collect_targets(before, plan, states))
    contexts = []

    async def call(task, context, model):
        assert task == "review_layout"
        contexts.append(context)
        if len(contexts) == 1:
            target = context["targets"][0]["target_id"]
            return PageLayoutReview(
                issues=[PageLayoutIssue(page=1, target_ids=[target], description="段落间距不足")],
                patches=[LayoutPatch(target_id=target, space_after_pt=9)],
            )
        assert context["targets"][0]["current_layout"]["space_after_pt"] == 9
        return PageLayoutReview()

    report = await run_review(source, tmp_path, call)
    assert report["status"] == "passed" and report["changed"] and report["visual_checked"]
    after = Document(path)
    assert protected_signature(after, collect_targets(after, plan, states)) == signature
    assert report["layout"]["pages"] == {"s1": 1}
    assert Path(report["pdf"]).read_bytes() == b"PDF"
    assert contexts[0]["page_images"] == contexts[1]["page_images"]  # content-addressed inputs
    assert "base64" not in (tmp_path / "layout-review.json").read_text()


@pytest.mark.parametrize("outcome", ["issue", "timeout", "bad_patch", "calls"])
async def test_incomplete_or_unsafe_correction_leaves_input_identical(source, tmp_path, outcome):
    path, _, _ = source
    original = path.read_bytes()
    count = 0

    async def call(task, context, model):
        nonlocal count
        count += 1
        target = context["targets"][0]["target_id"]
        if outcome == "bad_patch":
            return PageLayoutReview(patches=[LayoutPatch(target_id=target, font_delta_pt=-1)])
        if count == 1:
            return PageLayoutReview(patches=[LayoutPatch(target_id=target, space_after_pt=9)])
        if outcome == "timeout":
            raise TimeoutError()
        return PageLayoutReview(issues=[PageLayoutIssue(page=1, description="图中文字不可读")])

    result = await run_review(source, tmp_path, call, max_calls=1 if outcome == "calls" else 12)
    assert result["status"] in {"needs_attention", "timeout", "failed", "partial"}
    assert not result["changed"] and path.read_bytes() == original


async def test_missing_capabilities_never_call_model_or_claim_visual_review(source, tmp_path):
    async def forbidden(*args):
        pytest.fail("missing capabilities must not call a model")

    path, plan, states = source
    for enabled, supported, reason in [
        (False, True, "disabled"),
        (True, False, "image_input_not_enabled"),
    ]:
        result = await review_rendered_document(
            path,
            plan,
            states,
            tmp_path,
            forbidden,
            lambda _: None,
            enabled=enabled,
            supports_images=supported,
        )
        assert result["status"] == "skipped" and result["reason"] == reason
        assert not result["visual_checked"]
    result = await review_rendered_document(
        path,
        plan,
        states,
        tmp_path,
        forbidden,
        lambda _: None,
        enabled=True,
        supports_images=True,
        renderer=lambda *a, **k: {"status": "skipped", "reason": "requires_office"},
    )
    assert result["reason"] == "requires_office" and not result["visual_checked"]


async def test_partial_page_coverage_is_recorded_without_claiming_full_check(source, tmp_path):
    def many_pages(*args, **kwargs):
        result = preview(*args, **kwargs)
        result["total_pages"] = 3
        result["pages"] = [result["pages"][0] | {"page": number} for number in range(1, 4)]
        return result

    async def call(*args):
        return PageLayoutReview()

    path, plan, states = source
    result = await review_rendered_document(
        path,
        plan,
        states,
        tmp_path,
        call,
        lambda _: None,
        enabled=True,
        supports_images=True,
        renderer=many_pages,
        max_calls=1,
    )
    assert result["status"] == "partial" and not result["visual_checked"]
    assert result["rounds"][0]["reviewed_pages"] == [1, 2]


@pytest.mark.parametrize("page_counts", [(30,), (16, 16), (16, 18, 18)])
async def test_default_budget_covers_long_documents_and_every_correction(
    source, tmp_path, page_counts
):
    rendered = 0
    inspected = {}

    def renderer(*args, **kwargs):
        nonlocal rendered
        result = preview(*args, **kwargs)
        count = page_counts[rendered]
        rendered += 1
        result["total_pages"] = count
        result["pages"] = [
            result["pages"][0] | {"page": number}
            for number in range(1, min(count, kwargs["max_pages"]) + 1)
        ]
        return result

    async def call(task, context, model):
        inspected.setdefault(rendered, []).extend(page["page"] for page in context["page_images"])
        if rendered < len(page_counts) and context["page_images"][0]["page"] == 1:
            target = context["targets"][0]["target_id"]
            return PageLayoutReview(
                issues=[PageLayoutIssue(page=1, target_ids=[target], description="段落间距不足")],
                patches=[LayoutPatch(target_id=target, space_after_pt=6 + rendered * 3)],
            )
        return PageLayoutReview()

    path, plan, states = source
    result = await review_rendered_document(
        path,
        plan,
        states,
        tmp_path,
        call,
        lambda _: None,
        enabled=True,
        supports_images=True,
        renderer=renderer,
    )
    assert result["status"] == "passed" and result["visual_checked"]
    assert result["changed"] == (len(page_counts) > 1)
    assert result["calls"] == sum((count + 1) // 2 for count in page_counts)
    for round_number, count in enumerate(page_counts, 1):
        assert inspected[round_number] == list(range(1, count + 1))
        assert result["rounds"][round_number - 1]["complete"]
        assert result["rounds"][round_number - 1]["remaining_pages"] == 0
    assert result["budget"]["max_calls"] >= result["calls"]
    assert 180 < result["budget"]["stage_seconds"] <= MAX_REVIEW_SECONDS


async def test_page_hard_limit_reports_unreviewed_pages_and_retains_input(source, tmp_path):
    path, plan, states = source
    original = path.read_bytes()

    def renderer(*args, **kwargs):
        result = preview(*args, **kwargs)
        result["total_pages"] = MAX_REVIEW_PAGES + 3
        result["pages"] = [
            result["pages"][0] | {"page": number} for number in range(1, kwargs["max_pages"] + 1)
        ]
        return result

    async def call(*args):
        return PageLayoutReview()

    result = await review_rendered_document(
        path,
        plan,
        states,
        tmp_path,
        call,
        lambda _: None,
        enabled=True,
        supports_images=True,
        renderer=renderer,
    )
    assert result["status"] == "partial" and result["reason"] == "page_limit"
    assert not result["visual_checked"] and not result["changed"]
    assert result["budget"]["max_calls"] == MAX_REVIEW_CALLS
    assert result["budget"]["stage_seconds"] == MAX_REVIEW_SECONDS
    assert result["rounds"][0]["remaining_pages"] == 3
    assert path.read_bytes() == original
    assert not (tmp_path / "layout-preview.pdf").exists()


async def test_explicit_deadline_still_stops_page_review(source, tmp_path):
    path, _, _ = source
    original = path.read_bytes()

    async def call(*args):
        await asyncio.sleep(5)
        return PageLayoutReview()

    report = await run_review(source, tmp_path, call, stage_seconds=0.1)
    assert report["status"] == "timeout"
    assert report["budget"]["stage_seconds"] == 0.1
    assert not report["changed"] and path.read_bytes() == original


def test_page_count_is_known_before_rasterizing_only_requested_batch(source, tmp_path, monkeypatch):
    from folio.document import page_preview

    commands = []

    def paginate(path, *, pdf):
        pdf.write_bytes(b"PDF")
        return {"status": "paginated"}

    def run(command, **kwargs):
        commands.append(command)
        if Path(command[0]).name == "pdfinfo":
            return SimpleNamespace(stdout="Pages: 30\n")
        if command[0] == "pdftotext":
            return SimpleNamespace(stdout="第1页\f第2页\f")
        assert command[0] == "pdftoppm"
        first = int(command[command.index("-f") + 1])
        last = int(command[command.index("-l") + 1])
        for number in range(first, last + 1):
            Path(f"{command[-1]}-{number:02}.png").write_bytes(PNG)
        return SimpleNamespace(stdout=b"")

    monkeypatch.setattr(page_preview.shutil, "which", lambda name: f"/usr/bin/{name}")
    monkeypatch.setattr(page_preview, "paginate", paginate)
    monkeypatch.setattr(page_preview.subprocess, "run", run)
    result = prepare_page_preview(source[0], tmp_path / "lazy-pages")
    assert result["total_pages"] == 30
    assert len(commands) == 1 and not list((tmp_path / "lazy-pages").glob("*.png"))
    assert all("path" not in page for page in result["pages"])
    batch = materialize_page_images(result, result["pages"][:2])
    assert [page["text"] for page in batch] == ["第1页", "第2页"]
    assert len(list((tmp_path / "lazy-pages").glob("*.png"))) == 2
    assert all("path" not in page for page in result["pages"][2:])


async def test_cancellation_waits_for_local_renderer_before_releasing_files(source, tmp_path):
    started, release, ended = threading.Event(), threading.Event(), threading.Event()

    def renderer(*args, **kwargs):
        started.set()
        release.wait(5)
        try:
            return preview(*args, **kwargs)
        finally:
            ended.set()

    async def forbidden(*args):
        pytest.fail("no model work after cancellation")

    path, plan, states = source
    original = path.read_bytes()
    task = asyncio.create_task(
        review_rendered_document(
            path,
            plan,
            states,
            tmp_path,
            forbidden,
            lambda _: None,
            enabled=True,
            supports_images=True,
            renderer=renderer,
        )
    )
    assert await asyncio.to_thread(started.wait, 5)
    task.cancel()
    await asyncio.sleep(0)
    assert not task.done()
    release.set()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert ended.is_set() and path.read_bytes() == original
    assert json.loads((tmp_path / "layout-review.json").read_text())["status"] == "cancelled"


async def test_explicit_vision_input_uses_images_and_bounded_capacity_without_probe(tmp_path, monkeypatch):
    settings = ModelSettingsStore(tmp_path)
    settings.save(
        ModelSettingsInput(
            base_url="http://example.test/v1",
            model="vision",
            no_auth=True,
            supports_image_input=True,
        )
    )
    assert settings.public()["supports_image_input"]
    runtime = LLMRuntime(settings.path)

    async def forbidden_probe(*args, **kwargs):
        pytest.fail("Image capacity must be estimated locally without tokenizer probes")

    monkeypatch.setattr(runtime.capacity, "request", forbidden_probe)
    image = tmp_path / "page.png"
    image.write_bytes(PNG)
    context = {
        "page_images": [{"page": 1, "path": str(image), "sha256": hashlib.sha256(PNG).hexdigest()}]
    }
    try:
        messages, wire_schema, *_ = runtime.request_parts(
            "review_layout", context, PageLayoutReview.model_json_schema()
        )
        assert messages[1]["content"][-1]["image_url"]["url"].startswith("data:image/png;base64,")
        assert str(image) not in messages[1]["content"][0]["text"]
        measured = await runtime.measure(
            "review_layout", context, PageLayoutReview.model_json_schema(), 20000
        )
        assert measured["image_token_reserve"] == 4096
        assert measured["measurement"] == "default_tokenizer_estimate"
        from folio.runtime.token_estimate import estimate_input

        text_messages = [message | {"content": [part for part in message["content"]
                                                if part["type"] != "image_url"]}
                         if isinstance(message.get("content"), list) else message for message in messages]
        strict_schema = wire_schema if runtime.config["structured_output"] == "json_schema" else None
        assert measured["estimated_input_tokens"] == estimate_input(text_messages, schema=strict_schema) + 4096
        assert measured["fits"] and measured["output_tokens"] == 128000
        image.write_bytes(PNG + b"changed")
        with pytest.raises(ValueError, match="changed"):
            runtime.request_parts("review_layout", context, PageLayoutReview.model_json_schema())
    finally:
        await runtime.aclose()


def test_image_capability_is_explicit_and_settings_roundtrip(tmp_path):
    settings = ModelSettingsStore(tmp_path)
    saved = settings.save(
        ModelSettingsInput(base_url="http://example.test/v1", model="text", no_auth=True)
    )
    assert not saved["supports_image_input"]
    assert "layout_review_enabled" not in saved


def test_actual_office_pdf_and_page_preview_preserve_source(source, tmp_path):
    if not office_runtime() or not all(
        shutil.which(name) for name in ("pdftoppm", "pdftotext", "pdfinfo")
    ):
        pytest.skip("LibreOffice/UNO and Poppler required")
    path, _, _ = source
    original = path.read_bytes()
    directory = tmp_path / "real-preview"
    directory.mkdir()
    (directory / "page-999.png").write_bytes(PNG)
    result = render_page_preview(path, directory, max_pages=2)
    assert result["status"] == "rendered" and result["total_pages"] >= 1
    assert path.read_bytes() == original
    assert Path(result["pdf"]).read_bytes().startswith(b"%PDF")
    assert Path(result["pages"][0]["path"]).read_bytes().startswith(b"\x89PNG")
    assert "100" in "".join(page["text"] for page in result["pages"])
    assert not (directory / "page-999.png").exists()


async def approved_layout(source, folder):
    calls = 0

    async def call(task, context, model):
        nonlocal calls
        calls += 1
        if calls == 1:
            return PageLayoutReview(
                patches=[
                    LayoutPatch(
                        target_id=context["targets"][0]["target_id"],
                        space_after_pt=9,
                    )
                ]
            )
        return PageLayoutReview()

    report = await run_review(source, folder, call)
    assert report["status"] == "passed" and report["changed"]
    return report


@pytest.mark.parametrize("change", [None, "body", "baseline", "media", "rejected", "result"])
async def test_only_exact_approved_page_layout_can_replay(source, tmp_path, change):
    path, plan, states = source
    original = path.read_bytes()
    report = await approved_layout(source, tmp_path)
    target_signature = report["target_document_signature"]
    path.write_bytes(original)
    if change in {"body", "baseline"}:
        doc = Document(path)
        if change == "body":
            doc.add_paragraph("变更后的正文")
        else:
            # A property allowed for editing still belongs to the immutable source baseline.
            collect_targets(doc, plan, states)[0].element.get_or_add_pPr().get_or_add_spacing().set(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}after", "180"
            )
        doc.save(path)
    if change == "media":
        with ZipFile(path, "a") as archive:
            archive.writestr("word/media/changed.png", PNG)
    if change in {"rejected", "result"}:
        if change == "rejected":
            report["status"] = "needs_attention"
        else:
            report["target_document_signature"] = "not-the-approved-output"
        (tmp_path / "layout-review.json").write_text(json.dumps(report))
    before = path.read_bytes()
    result = replay_page_review(path, plan, states, tmp_path / "layout-review.json")
    assert not result["visual_checked"]
    if change is None:
        assert result["status"] == "replayed" and result["changed"]
        assert document_signature(path) == target_signature
    else:
        assert not result["changed"] and path.read_bytes() == before


@pytest.mark.parametrize("pagination_failure", [False, True])
async def test_reformat_reuses_approved_patches_without_visual_check_or_model(
    source, tmp_path, monkeypatch, pagination_failure
):
    from folio.document.reformat import reformat_job

    path, plan, states = source
    baseline_signature = document_signature(path)
    folder = tmp_path / "jobs" / "job1"
    folder.mkdir(parents=True)
    await approved_layout(source, folder)
    (folder / "formatting.json").write_text(
        json.dumps(
            {
                "policy": FORMAT_POLICY,
                "status": "skipped",
                "changed": False,
                "reason": "page_review_enabled",
            }
        )
    )
    request = JobInput(
        title="审批需求",
        user_request="说明审批业务",
        knowledge_path=str(tmp_path / "knowledge"),
        runtime="llm",
    )
    job = {
        "status": "draft",
        "plan": plan.model_dump(),
        "input": request.model_dump(),
        "artifact": str(path),
    }
    store = SimpleNamespace(root=tmp_path, get=lambda _: job, states=lambda _: states)
    calls = []

    def paginate(path, *, pdf, required):
        calls.append(path)
        if pagination_failure and len(calls) == 2:
            raise RuntimeError("candidate pagination failed")
        pdf.write_bytes(b"PDF")
        return {"status": "paginated"}

    monkeypatch.setattr("folio.document.reformat.paginate", paginate)
    result = reformat_job(store, "job1", tmp_path / "reformatted")
    assert result["model_calls"] == 0 and not result["business_state_changed"]
    assert result["layout"]["formatting"]["status"] == "skipped"
    assert result["layout"]["page_review"]["status"] == (
        "discarded" if pagination_failure else "replayed"
    )
    assert not result["layout"]["page_review"]["visual_checked"]
    # Paginate the built-in contents baseline first, then the approved layout candidate.
    assert len(calls) == (3 if pagination_failure else 2)
    if pagination_failure:
        assert document_signature(result["artifact"]) == baseline_signature


def test_skipped_legacy_formatting_manifest_does_not_require_baselines(source, tmp_path):
    path, plan, states = source
    (tmp_path / "formatting.json").write_text(
        json.dumps({"status": "skipped", "policy": FORMAT_POLICY})
    )
    assert replay_formatting(path, plan, states, tmp_path / "formatting.json") == {
        "status": "skipped",
        "changed": False,
    }
