"""Directory research grounds conclusions in originals actually delivered to the model."""

import asyncio
import json
import threading

import pytest
from docx import Document

from folio.research.planning import PlanningResearchLimits, research_for_directory
from folio.runtime.base import RuntimeResult


def action(tool, **args):
    return {"tool": tool, "arguments": args,
            "summary": tool, "notes": "保留资料真实范围与未解决问题"}


class Runtime:
    def __init__(self, policy):
        self.policy = policy
        self.calls = []

    @staticmethod
    def prompt_chars(task, context, schema):
        return len(json.dumps(context, ensure_ascii=False)) + len(json.dumps(schema)) + 3000

    async def complete(self, task, context, schema, session_key):
        assert task == "directory_research"
        self.calls.append(context)
        value = self.policy(context)
        if asyncio.iscoroutine(value):
            value = await value
        if isinstance(value, RuntimeResult):
            return value
        if "tool" not in value:
            return RuntimeResult(json.dumps(value, ensure_ascii=False))
        parameters = dict(value["arguments"])
        if value["tool"] != "finish":
            parameters["summary"] = value.get("summary", "")
        if value["tool"] != "research.save":
            parameters["notes"] = value.get("notes", "")
        for field in ("research_notes", "review_reason"):
            if field in value:
                parameters[field] = value[field]
        return RuntimeResult("", usage={"input_tokens": 10, "output_tokens": 5}, tool_calls=[{
            "id": f"call_{len(self.calls)}", "type": "function", "function": {
                "name": value["tool"].replace(".", "_"),
                "arguments": json.dumps(parameters, ensure_ascii=False),
            },
        }])


@pytest.fixture
def sources(tmp_path):
    root = tmp_path / "sources"
    root.mkdir()
    (root / "rules.md").write_text("# 技术要求\n必须支持撤单后释放库存。\n必须记录审批日志。\n", encoding="utf-8")
    return root


class ReadPolicy:
    """Adaptive fake uses returned node IDs, pagination and original evidence."""

    def __init__(self, path="rules.md", title=None, excluded=None, incomplete=False, fake=False):
        self.path, self.title = path, title
        self.excluded = excluded or []
        self.node = None
        self.evidence = []
        self.incomplete, self.fake = incomplete, fake

    def __call__(self, context):
        observations = context["observations"]
        if not observations:
            return action("fs.inspect", path=self.path, limit=8)
        latest = observations[-1]
        result = latest["result"]
        if latest["tool"] == "fs.inspect":
            if self.title:
                match = next((node for node in result["structure"] if node["title"] == self.title), None)
                if match:
                    self.node = match["node_id"]
            if result["next_offset"] is not None:
                return action("fs.inspect", path=self.path, limit=8, offset=result["next_offset"])
            kwargs = {"node_id": self.node} if self.node else {}
            return action("fs.read", path=self.path, limit=1, **kwargs)
        if latest["tool"] == "fs.read":
            self.evidence.extend(item["evidence_id"] for item in result["evidence"])
            if result["next_offset"] is not None and not self.incomplete:
                kwargs = {"node_id": self.node} if self.node else {}
                return action("fs.read", path=self.path, limit=1, offset=result["next_offset"], **kwargs)
        return action("finish", summary="已研究技术要求，投标人案例在正文阶段补充。",
                      requirements=[{"requirement_id": "technical", "text": "覆盖技术要求", "evidence_ids":
                                     ["ev_fabricated"] if self.fake else list(dict.fromkeys(self.evidence))}],
                      excluded_files=self.excluded, unresolved=["投标人案例待补充"], blocking_gaps=[])


@pytest.mark.asyncio
async def test_notes_survive_tool_failure_and_are_saved_as_markdown(sources, tmp_path):
    following = ReadPolicy()

    def policy(context):
        if context["observations"] and context["observations"][-1]["tool"] == "fs.read":
            observation = context["observations"][-1]
            return action("fs.inspect", path="missing.md") | {"research_notes": [{
                "topic": "取消与审计", "text": "撤单后释放库存，审批需要留日志。", "kind": "source_statement",
                "evidence_ids": [raw["evidence_id"] for raw in observation["result"]["evidence"]],
                "request_ids": [observation["request_id"]],
            }]}
        return following(context)

    output = tmp_path / "research"
    report = await research_for_directory(Runtime(policy), sources, output, "写目录",
                                          limits=PlanningResearchLimits(max_calls=3))
    assert report["status"] == "partial"
    assert report["research_memory"]["notes"][0]["text"] == "撤单后释放库存，审批需要留日志。"
    assert "撤单后释放库存" in (output / "research.md").read_text()
    saved = json.loads((output / "report.json").read_text())
    assert saved["research_memory"] == report["research_memory"]
    assert (output / "events/research_003-error.json").exists()


@pytest.mark.asyncio
async def test_repeated_navigation_stops_without_an_impossible_source_note_checkpoint(sources, tmp_path):
    runtime = Runtime(lambda context: action("fs.inspect", path="rules.md"))
    report = await research_for_directory(runtime, sources, tmp_path / "research", "写目录")
    assert report["calls"] == 4  # first navigation progresses, then three no-op attempts
    assert any('Three consecutive' in gap for gap in report["blocking_gaps"])
    assert not any(context.get("memory_checkpoint_required") for context in runtime.calls)


@pytest.mark.asyncio
async def test_identical_source_reads_require_a_reason_and_stop_without_refetching(sources, tmp_path, monkeypatch):
    import folio.research.planning as planning

    original = planning.WorkspaceManager.read
    fetched = []

    def read(workspace, **args):
        fetched.append(args)
        return original(workspace, **args)

    monkeypatch.setattr(planning.WorkspaceManager, "read", read)
    count = 0

    def policy(context):
        nonlocal count
        count += 1
        return action("fs.inspect", path="rules.md") if count == 1 else action("fs.read", path="rules.md")

    runtime = Runtime(policy)
    output = tmp_path / "research"
    report = await research_for_directory(runtime, sources, output, "写目录")
    assert len(fetched) == 1
    assert report["calls"] == 6  # five research turns, then one bounded finish-only handoff
    assert runtime.calls[-1]["research_finalizing"]
    assert report["finalization_attempted"]
    assert not runtime.calls[-1].get("memory_checkpoint_required")
    error = json.loads((output / "events/research_003-error.json").read_text())
    assert "research_002" in error["message"] and "review_reason" in error["message"]
    assert report["research_memory"]["batches"]["research_002"]["delivered"]
    assert not report["research_memory"]["batches"]["research_002"]["reviewed"]


@pytest.mark.asyncio
async def test_late_technical_chapter_is_read_after_full_structure_paging(tmp_path):
    sources = tmp_path / "sources"
    sources.mkdir()
    document = Document()
    for number in range(1, 13):
        document.add_heading(f"第{number}章 公告与须知", 1)
        document.add_paragraph("须知说明" * 400)
    document.add_heading("第十三章 技术需求", 1)
    document.add_paragraph("VaR和ES计量、返回检验、压力测试必须逐项响应。" * 120)
    document.save(sources / "tender.docx")
    runtime = Runtime(ReadPolicy("tender.docx", "第十三章 技术需求"))
    report = await research_for_directory(runtime, sources, tmp_path / "research", "编制技术标目录")
    assert report["status"] == "ready"
    assert report["unresolved"] == ["投标人案例待补充"]
    assert report["blocking_gaps"] == []
    assert report["files"][0]["structure_complete"]
    assert report["files"][0]["readings"][0]["label"] == "第十三章 技术需求"
    assert all("须知说明" not in item["text"] for item in report["evidence"])
    assert "VaR" in "".join(item["text"] for item in report["evidence"])
    observations = [item for call in runtime.calls for item in call["observations"]]
    assert any(item["tool"] == "fs.inspect" and item["arguments"].get("offset") == 8 for item in observations)
    assert any(item["tool"] == "fs.read" and item["arguments"].get("offset", 0) > 0 for item in observations)
    saved = json.loads((tmp_path / "research/report.json").read_text())
    assert saved["evidence"] == report["evidence"]
    assert len(list((tmp_path / "research/events").glob("*-response.json"))) == report["calls"]


@pytest.mark.asyncio
async def test_inspect_without_actual_read_never_ready(sources, tmp_path):
    count = 0

    def policy(context):
        nonlocal count
        count += 1
        return action("fs.inspect", path="rules.md") if count == 1 else action("finish", summary="目录结构已看过")

    report = await research_for_directory(Runtime(policy), sources, tmp_path / "research", "编写要求")
    assert report["status"] == "partial"
    assert not report["evidence"]
    assert report["rounds"] == 3
    assert any('key source text' in gap for gap in report["blocking_gaps"])


@pytest.mark.asyncio
async def test_fabricated_evidence_is_rejected_and_not_forwarded(sources, tmp_path):
    report = await research_for_directory(Runtime(ReadPolicy(fake=True)), sources, tmp_path / "research", "编写要求")
    assert report["status"] == "partial"
    assert report["requirements"] == []
    assert any("ev_fabricated" in item for item in report["validation_issues"])
    assert report["evidence"]


@pytest.mark.asyncio
async def test_citing_delivered_fragments_does_not_claim_the_entire_read_range_is_complete(sources, tmp_path):
    report = await research_for_directory(Runtime(ReadPolicy(incomplete=True)), sources, tmp_path / "research", "编写要求")
    assert report["status"] == "ready"
    reading = report["files"][0]["read_ranges"][0]
    assert not reading["complete"] and not reading["delivered_complete"]
    assert report["selected_ranges"][0]["range_id"] == reading["range_id"]
    assert report["requirements"][0]["evidence_ids"] == reading["delivered_evidence_ids"]
    assert report["validation_issues"] == []


@pytest.mark.asyncio
async def test_uncited_read_evidence_stays_preserved_without_blocking_finish(sources, tmp_path):
    def policy(context):
        if not context["observations"]:
            return action("fs.read", file="F1")
        return action("finish", summary="目录需要覆盖撤单库存。",
                      requirements=[{"requirement_id": "inventory", "text": "安排撤单库存要求",
                                     "evidence_ids": ["E2"]}])

    report = await research_for_directory(Runtime(policy), sources, tmp_path / "research", "编写要求")
    assert report["status"] == "ready" and report["calls"] == 2
    assert len(report["evidence"]) > 1
    batch = report["research_memory"]["batches"]["research_001"]
    assert set(batch["delivered_evidence_ids"]) == {item["evidence_id"] for item in report["evidence"]}
    assert set(batch["reviewed_evidence_ids"]) == set(report["requirements"][0]["evidence_ids"])
    assert batch["reviewed"] is False
    assert not any("归纳" in issue for issue in report["validation_issues"])


@pytest.mark.asyncio
@pytest.mark.parametrize("legacy_ranges", [None, "obsolete", [{"range_id": "F1/N037", "reason": "旧格式"}]])
async def test_direct_tpu_page_read_can_finish_with_only_evidence_references(tmp_path, legacy_ranges):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    sources = tmp_path / "sources"
    sources.mkdir()
    writer = PdfWriter()
    for text in ("Introduction.", "TPU ICI bandwidth is 19.2 Tb/s.", "Appendix."):
        page = writer.add_blank_page(width=600, height=800)
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({
                NameObject("/F1"): DictionaryObject({
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }),
            }),
        })
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 50 750 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(sources / "TPU.pdf")

    def policy(context):
        if not context["observations"]:
            return action("fs.read", file="F1", page=2)
        reading = context["observations"][-1]
        assert reading["tool"] == "fs.read"
        assert "TPU ICI bandwidth is 19.2 Tb/s." in reading["model_text"]
        result = action("finish", summary="已读取 TPU 互连带宽资料。",
                        requirements=[{"requirement_id": "tpu", "text": "分析 TPU ICI 互连带宽",
                                       "evidence_ids": ["E1"]}])
        if legacy_ranges is not None:
            result["arguments"]["selected_ranges"] = legacy_ranges
        return result

    runtime = Runtime(policy)
    report = await research_for_directory(runtime, sources, tmp_path / "research", "分析 TPU 通信瓶颈")
    assert report["status"] == "ready"
    assert report["calls"] == 2
    assert not report["files"][0]["inspected"]
    assert not report["files"][0]["structure_complete"]
    reading = report["files"][0]["read_ranges"][0]
    assert reading["selector"] == {"path": "TPU.pdf", "page": 2}
    assert report["selected_ranges"][0]["range_id"] == reading["range_id"]
    assert report["requirements"][0]["evidence_ids"] == reading["delivered_evidence_ids"]
    assert report["research_memory"]["notes"][0]["request_ids"] == ["research_001"]


@pytest.mark.asyncio
@pytest.mark.parametrize("field,replacement", [
    ("source_id", "another-file.md"), ("source_version", "another-version"), ("locator", "another-position"),
])
async def test_finish_rejects_evidence_with_changed_source_provenance(
    sources, tmp_path, monkeypatch, field, replacement,
):
    from folio.research.planning import _DirectoryResearch

    original = _DirectoryResearch.accept_conclusion

    def accept(self, arguments):
        original(self, arguments)
        for key in self.delivered_evidence:
            self.evidence[key][field] = replacement

    monkeypatch.setattr(_DirectoryResearch, "accept_conclusion", accept)
    report = await research_for_directory(Runtime(ReadPolicy()), sources, tmp_path / "research", "编写要求")
    assert report["status"] == "partial"
    assert report["requirements"] == []
    assert report["selected_ranges"] == []
    assert any('source version or location' in issue for issue in report["validation_issues"])


@pytest.mark.asyncio
async def test_search_hits_are_not_read_evidence(sources, tmp_path):
    count = 0

    def policy(context):
        nonlocal count
        count += 1
        if count == 1:
            return action("fs.search", query="库存")
        evidence = context["observations"][-1]["result"]["evidence"][0]
        return action("finish", requirements=[{"requirement_id": "r1", "text": "库存", "evidence_ids": [evidence["evidence_id"]]}])

    report = await research_for_directory(Runtime(policy), sources, tmp_path / "research", "库存要求")
    assert report["status"] == "partial"
    assert report["requirements"] == report["evidence"] == []
    assert not report["files"][0]["read"]


@pytest.mark.asyncio
async def test_path_escape_is_a_correctable_error_and_never_reads_secret(sources, tmp_path):
    secret = tmp_path / "private.txt"
    secret.write_text("must-not-be-read")
    following = ReadPolicy()
    calls = 0

    def policy(context):
        nonlocal calls
        calls += 1
        if calls == 1:
            return action("fs.read", path="../private.txt")
        if calls == 2:
            assert 'snapshot root' in context["feedback"][0]
            return action("fs.inspect", path="rules.md")
        return following(context)

    report = await research_for_directory(Runtime(policy), sources, tmp_path / "research", "编写要求")
    assert report["status"] == "ready"
    assert "must-not-be-read" not in json.dumps(report)
    errors = list((tmp_path / "research/events").glob("*-error.json"))
    assert len(errors) == 1


@pytest.mark.asyncio
@pytest.mark.parametrize("excluded,expected", [([], "partial"), ([{"path": "poster.txt", "reason": "营销海报与库存需求无关"}], "ready")])
async def test_every_unused_file_requires_an_explicit_reason(sources, tmp_path, excluded, expected):
    (sources / "poster.txt").write_text("营销海报说明")
    report = await research_for_directory(Runtime(ReadPolicy(excluded=excluded)), sources, tmp_path / "research", "库存要求")
    assert report["status"] == expected
    assert not next(file for file in report["files"] if file["path"] == "poster.txt")["read"]


@pytest.mark.asyncio
async def test_call_budget_stops_and_preserves_auditable_reads(sources, tmp_path):
    report = await research_for_directory(Runtime(ReadPolicy()), sources, tmp_path / "research", "编写要求",
                                          limits=PlanningResearchLimits(max_calls=2))
    assert report["status"] == "partial"
    assert report["calls"] == 2
    assert report["files"][0]["read"]
    assert (tmp_path / "research/events/research_002-tool.json").exists()
    assert any('budget' in gap for gap in report["blocking_gaps"])


@pytest.mark.asyncio
async def test_empty_sources_does_not_call_model(tmp_path):
    sources = tmp_path / "sources"
    sources.mkdir()
    runtime = Runtime(lambda context: pytest.fail("must not call model"))
    report = await research_for_directory(runtime, sources, tmp_path / "research", "写通用文档")
    assert report["status"] == "no_sources"
    assert report["calls"] == 0
    assert report["blocking_gaps"] == []


@pytest.mark.asyncio
async def test_cancellation_preserves_report_and_completed_tools(sources, tmp_path):
    entered = asyncio.Event()
    calls = 0

    async def policy(context):
        nonlocal calls
        calls += 1
        if calls == 1:
            return action("fs.inspect", path="rules.md")
        entered.set()
        await asyncio.Future()

    task = asyncio.create_task(research_for_directory(Runtime(policy), sources, tmp_path / "research", "编写要求"))
    await entered.wait()
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    report = json.loads((tmp_path / "research/report.json").read_text())
    assert report["status"] == "partial"
    assert report["files"][0]["inspected"]
    assert report["calls"] == 2
    assert any('interrupted' in gap for gap in report["blocking_gaps"])


@pytest.mark.asyncio
async def test_timeout_finishes_partial_without_retrying_forever(sources, tmp_path):
    async def policy(context):
        await asyncio.sleep(10)

    report = await research_for_directory(Runtime(policy), sources, tmp_path / "research", "编写要求",
                                          limits=PlanningResearchLimits(call_timeout_seconds=0.01))
    assert report["status"] == "partial"
    assert report["calls"] == 1
    assert any('timed out' in gap for gap in report["blocking_gaps"])


@pytest.mark.asyncio
async def test_manifest_pages_are_available_and_exclusions_must_be_seen(sources, tmp_path):
    excluded = []
    for number in range(20):
        path = f"z{number:03}.txt"
        (sources / path).write_text("与库存无关的宣传参考")
        excluded.append({"path": path, "reason": "宣传材料不在本次库存文档范围内"})
    unread = await research_for_directory(Runtime(ReadPolicy(excluded=excluded)), sources,
                                         tmp_path / "unread", "库存要求")
    assert unread["status"] == "partial"
    assert any("Manifest" in gap for gap in unread["blocking_gaps"])
    following = ReadPolicy(excluded=excluded)
    count = 0

    def policy(context):
        nonlocal count
        count += 1
        if count == 1:
            assert context["manifest"]["next_offset"] == 20
            return action("fs.list", offset=20)
        if count == 2:
            assert context["observations"][-1]["result"]["complete"]
            return action("fs.inspect", path="rules.md")
        return following(context)

    report = await research_for_directory(Runtime(policy), sources, tmp_path / "read", "库存要求")
    assert report["status"] == "ready"
    assert len(report["excluded_files"]) == 20


@pytest.mark.asyncio
async def test_long_user_requirements_remain_inline_when_request_fits(sources, tmp_path):
    requirements = "具体编制要求。" * 2000 + "最终范围：只包含包件二。"
    following = ReadPolicy()

    def policy(context):
        assert context["requirements"] == requirements
        return following(context)

    report = await research_for_directory(Runtime(policy), sources, tmp_path / "read", requirements)
    assert report["status"] == "ready"


@pytest.mark.asyncio
async def test_final_structure_page_does_not_prove_skipped_pages_were_seen(sources, tmp_path):
    following = ReadPolicy()
    count = 0

    def policy(context):
        nonlocal count
        count += 1
        if count == 1:
            return action("fs.inspect", path="rules.md", offset=999)
        return following(context)

    report = await research_for_directory(Runtime(policy), sources, tmp_path / "research", "库存要求")
    assert report["status"] == "ready"
    assert not report["files"][0]["structure_complete"]
    assert not report["files"][0]["structure_delivered"]
    assert report["requirements"]
    assert not any("内部结构" in gap for gap in report["blocking_gaps"])


@pytest.mark.asyncio
async def test_evidence_budget_is_hard_and_keeps_real_already_read_chunks(sources, tmp_path):
    (sources / "rules.md").write_text("# 技术要求\n" + "库存规则" * 2000)
    calls = 0

    def policy(context):
        nonlocal calls
        calls += 1
        if calls == 1:
            return action("fs.inspect", path="rules.md")
        return action("fs.read", path="rules.md", offset=max(0, calls - 2), max_chars=2000)

    report = await research_for_directory(Runtime(policy), sources, tmp_path / "research", "库存要求",
                                          limits=PlanningResearchLimits(max_calls=4, max_evidence_chars=2000))
    assert report["status"] == "partial"
    assert sum(len(item["text"]) for item in report["evidence"]) <= 2000
    errors = [json.loads(path.read_text()) for path in (tmp_path / "research/events").glob("*-error.json")]
    assert any('source-text budget' in item["message"] for item in errors)


@pytest.mark.asyncio
@pytest.mark.parametrize("blocks_directory,expected", [(None, "partial"), (True, "partial"), (False, "ready")])
async def test_unread_visual_content_requires_visible_scope_assessment(tmp_path, blocks_directory, expected):
    from PIL import Image

    sources = tmp_path / "sources"
    sources.mkdir()
    image_path = tmp_path / "diagram.png"
    Image.new("RGB", (10, 10), "white").save(image_path)
    document = Document()
    document.add_picture(str(image_path))
    document.add_heading("技术需求", 1)
    document.add_paragraph("系统必须支持撤单后释放库存，技术响应必须逐项说明。")
    document.save(sources / "visual.docx")
    following = ReadPolicy("visual.docx")

    def policy(context):
        result = following(context)
        if result["tool"] == "finish" and blocks_directory is not None:
            args = result["arguments"]
            args["unread_assessments"] = [{
                "path": "visual.docx", "blocks_directory": blocks_directory,
                "reason": "已读技术要求足以安排目录，图片内容尚未核验，保留到正文阶段补查；不能判断其中是否还有业务细节。",
            }]
        return result

    report = await research_for_directory(Runtime(policy), sources, tmp_path / "research", "编制技术响应目录")
    assert report["status"] == expected
    assert report["files"][0]["source_complete"] is False
    assert any("图片或文本框" in warning for warning in report["files"][0]["warnings"])
    if blocks_directory is None:
        assert any('completeness warning on the outline has not been assessed' in gap for gap in report["blocking_gaps"])
    elif blocks_directory:
        assert any('Unread content affects the outline' in gap for gap in report["blocking_gaps"])
    else:
        assert any("图片或文本框" in gap and 'does not verify' in gap for gap in report["unresolved"])
        assert report["blocking_gaps"] == []


@pytest.mark.asyncio
@pytest.mark.parametrize("operation", ["initialize", "read"])
async def test_cancellation_joins_workspace_threads_before_returning(sources, tmp_path, monkeypatch, operation):
    import folio.research.planning as planning

    started, release, finished = threading.Event(), threading.Event(), threading.Event()
    original = planning.WorkspaceManager

    def pause(function, *args, **kwargs):
        started.set()
        assert release.wait(5), "test must release the local worker"
        try:
            return function(*args, **kwargs)
        finally:
            finished.set()

    class DelayedWorkspace(original):
        def __init__(self, *args, **kwargs):
            if operation == "initialize":
                pause(super().__init__, *args, **kwargs)
            else:
                super().__init__(*args, **kwargs)

        def read(self, *args, **kwargs):
            return pause(super().read, *args, **kwargs)

    monkeypatch.setattr(planning, "WorkspaceManager", DelayedWorkspace)
    task = asyncio.create_task(research_for_directory(Runtime(ReadPolicy()), sources, tmp_path / "research", "库存要求"))
    try:
        assert await asyncio.to_thread(started.wait, 3)
        task.cancel()
        # These event-loop turns let cancellation reach the shielded worker.
        await asyncio.sleep(0)
        await asyncio.sleep(0)
        assert not task.done()
        assert not finished.is_set()
    finally:
        release.set()
        with pytest.raises(asyncio.CancelledError):
            await task
    assert finished.is_set()
    if operation == "read":
        report = json.loads((tmp_path / "research/report.json").read_text())
        assert report["status"] == "partial"
        assert report["files"][0]["inspected"]
        # Completed workers are saved, but extraction is never confused with model delivery.
        assert report["evidence"]
        assert report["files"][0]["delivered_characters"] == 0
        assert report["research_memory"]["batches"]["research_002"]["delivered"] is False


@pytest.mark.asyncio
async def test_specific_omission_warning_requires_assessment_without_complete_flag(sources, tmp_path, monkeypatch):
    import folio.research.planning as planning

    original = planning.WorkspaceManager.read

    def read_with_warning(self, *args, **kwargs):
        result = original(self, *args, **kwargs)
        result["warnings"] = ["文件包含尚未读取的视觉内容。"]
        return result

    monkeypatch.setattr(planning.WorkspaceManager, "read", read_with_warning)
    report = await research_for_directory(Runtime(ReadPolicy()), sources, tmp_path / "research", "库存要求")
    assert report["status"] == "partial"
    assert any('source-reading completeness warning' in gap for gap in report["blocking_gaps"])


@pytest.mark.asyncio
async def test_progress_describes_actual_files_chapters_and_waiting_work(tmp_path):
    sources = tmp_path / "sources"
    sources.mkdir()
    name = "a" * 32 + "__招标资料.md"
    (sources / name).write_text("# 技术要求\n必须支持撤单后释放库存。\n必须记录审批日志。", encoding="utf-8")
    messages = []
    read = ReadPolicy(name, "技术要求")

    def policy(context):
        # The relevant description must be visible while waiting for the model,
        # not first appear after that potentially slow request returns.
        if not context["observations"]:
            assert messages[-1] == "正在分析 1 份资料的文件清单，选择与本次文档有关的资料。"
        elif context["observations"][-1]["tool"] == "fs.inspect":
            assert "《招标资料.md》的 1 个章节标题" in messages[-1]
            assert "正在判断需要阅读的内容" in messages[-1]
        else:
            assert "《招标资料.md》的“技术要求”原文" in messages[-1]
            assert "正在归纳本批原文，提取目录需要覆盖的要求" in messages[-1]
        value = read(context)
        value["notes"] = "这个内部研究笔记不能显示给用户"
        return value

    report = await research_for_directory(Runtime(policy), sources, tmp_path / "research", "编制技术标",
                                          on_progress=messages.append)
    assert report["status"] == "ready"
    rendered = "\n".join(messages)
    assert "正在查看《招标资料.md》的内部结构" in rendered
    assert "正在读取《招标资料.md》的“技术要求”" in rendered
    assert "正在继续读取《招标资料.md》的“技术要求”" in rendered
    assert "正在核对资料研究结果" in rendered
    assert "实际读取 1 份资料" in messages[-1]
    for hidden in ("fs.inspect", "fs.read", "内部研究笔记", "a" * 32, str(tmp_path)):
        assert hidden not in rendered


@pytest.mark.asyncio
async def test_progress_describes_correctable_invalid_model_reply_and_partial(sources, tmp_path):
    messages = []
    calls = 0

    def policy(context):
        nonlocal calls
        calls += 1
        if calls == 1:
            return {"private_json": "不可展示的格式错误内容"}
        if calls == 2:
            assert "AI 返回的操作格式不完整" in messages[-1]
        return action("finish", summary="还没有读取相关原文")

    report = await research_for_directory(Runtime(policy), sources, tmp_path / "research", "编制技术标",
                                          on_progress=messages.append,
                                          limits=PlanningResearchLimits(max_rounds=1))
    assert report["status"] == "partial"
    text = "\n".join(messages)
    assert "关键原文仍需补查" in text
    assert "轮次上限" in messages[-1]
    assert "private_json" not in text
    assert "不可展示的格式错误内容" not in text


@pytest.mark.asyncio
async def test_progress_reports_budget_and_timeout_without_claiming_read(sources, tmp_path):
    messages = []
    report = await research_for_directory(Runtime(ReadPolicy()), sources, tmp_path / "budget", "编制技术标",
                                          on_progress=messages.append,
                                          limits=PlanningResearchLimits(max_calls=1))
    assert report["calls"] == 1
    assert "调用或时间上限" in messages[-1]
    assert not any("已取得" in message and "原文" in message for message in messages)

    async def delayed(context):
        await asyncio.sleep(10)

    messages.clear()
    report = await research_for_directory(Runtime(delayed), sources, tmp_path / "timeout", "编制技术标",
                                          on_progress=messages.append,
                                          limits=PlanningResearchLimits(call_timeout_seconds=0.01))
    assert report["status"] == "partial"
    assert "等待超时" in messages[-1]


def test_public_progress_formats_supported_ranges_and_rejects_internal_summaries():
    from folio.research.planning_progress import action_message, result_message

    files = {"资料/需求.xlsx": {}, "附件.pdf": {}}
    args = {"path": "资料/需求.xlsx", "sheet": "功能清单", "cell_range": "A2:E20"}
    message = action_message("fs.read", args, files, {}, "核对各模块的必填功能")
    assert "《需求.xlsx》的工作表“功能清单”（A2:E20）" in message
    assert "本步目的：核对各模块的必填功能" in message
    for bad_summary in ("fs.read", "fs_read", "fs_inspect", '{"path":"/home/internal"}',
                        "打开 sources/internal.xlsx", "node_id=secret", "tool_call_id=private"):
        assert bad_summary not in action_message("fs.read", args, files, {}, bad_summary)
    assert "第 12 页" in action_message("fs.read", {"path": "附件.pdf", "page": 12}, files, {})
    assert "第 2 张幻灯片" in action_message("fs.read", {"path": "附件.pdf", "slide": 2}, files, {})
    assert "完整编制要求" in action_message("fs.read", {"scope": "context", "field": "requirements"}, files, {})
    assert "此前的资料阅读记录" in action_message("fs.read", {"scope": "history"}, files, {})
    assert "已经读取的原文依据" in action_message("fs.read", {"scope": "evidence"}, files, {})
    result = result_message("fs.search", {}, {"evidence": [{"text": "命中片段"}]}, files, {})
    assert "已找到 1 条相关片段的位置" in result
    assert "已读" not in result


def test_progress_distinguishes_whole_file_scope_from_partial_page_batch():
    from folio.research.planning_progress import action_message, result_message

    files = {"report.pdf": {}}
    titles = {("report.pdf", "root"): "全文"}
    args = {"path": "report.pdf", "node_id": "root"}
    result = {"offset": 10, "total": 40, "next_offset": 12, "evidence": [
        {"text": "实际原文。", "metadata": {"page": 7}},
        {"text": "下一页原文。", "metadata": {"page": 8}},
    ]}
    message = result_message("fs.read", args, result, files, titles)
    assert "的正文" in message and "全文" not in message
    assert "本批 11 字符" in message
    assert "第 7—8 页" in message and "第 11—12 段 / 共 40 段" in message
    assert "还有后续内容可继续读取" in message
    assert "正在归纳本批原文" in message
    assert "全文" not in action_message("fs.read", args, files, titles)


def test_progress_does_not_imply_skipped_pdf_pages_were_read():
    from folio.research.planning_progress import result_message

    result = {"offset": 0, "total": 2, "next_offset": None, "evidence": [
        {"text": "原文A", "metadata": {"page": 2}},
        {"text": "原文B", "metadata": {"page": 5}},
    ]}
    message = result_message("fs.read", {"path": "report.pdf"}, result, {"report.pdf": {}}, {})
    assert "第 2、5 页" in message
    assert "第 2—5 页" not in message
    assert "还有后续" not in message


@pytest.mark.asyncio
async def test_native_call_and_real_result_are_paired_without_rewriting_file_alias(sources, tmp_path):
    from folio.research.transport import research_messages

    def policy(context):
        if not context["observations"]:
            return action("fs.inspect", file="F1")
        observed = context["observations"][-1]
        assert observed["arguments"]["path"] == "rules.md"
        raw_call = observed["tool_call"]
        assert json.loads(raw_call["function"]["arguments"])["file"] == "F1"
        assert "path" not in json.loads(raw_call["function"]["arguments"])
        messages = research_messages(context, {})
        assistant = next(message for message in messages if message.get("tool_calls"))
        result = next(message for message in messages if message["role"] == "tool")
        assert assistant["tool_calls"] == [raw_call]
        assert result["tool_call_id"] == raw_call["id"]
        assert "技术要求" in result["content"]
        return action("fs.read", file="F1")

    output = tmp_path / "research"
    report = await research_for_directory(Runtime(policy), sources, output, "编写要求",
                                          limits=PlanningResearchLimits(max_calls=2))
    assert report["calls"] == 2
    response = json.loads((output / "events/research_001-response.json").read_text())
    request = json.loads((output / "events/research_002-request.json").read_text())
    tool = json.loads((output / "events/research_001-tool.json").read_text())
    assert request["tools"]
    assert tool["tool_call"] == response["tool_calls"][0]
    assert any(message.get("tool_calls") == response["tool_calls"] for message in request["messages"])


@pytest.mark.asyncio
async def test_text_actions_never_execute_and_do_not_override_native_call(sources, tmp_path):
    content = json.dumps({"tool": "finish", "arguments": {"summary": "未经研究就完成"}}, ensure_ascii=False)
    native = {"id": "call_real", "type": "function", "function": {
        "name": "fs_inspect", "arguments": '{"path":"rules.md","summary":"核对结构"}',
    }}

    def policy(context):
        if not context["observations"]:
            return RuntimeResult(content + content, tool_calls=[native])
        assert context["observations"][-1]["tool"] == "fs.inspect"
        return RuntimeResult(content)

    output = tmp_path / "research"
    report = await research_for_directory(Runtime(policy), sources, output, "编写要求",
                                          limits=PlanningResearchLimits(max_calls=2))
    assert report["status"] == "partial"
    assert report["files"][0]["inspected"] is True
    assert report["summary"] == ""
    assert not (output / "events/research_002-assessment.json").exists()
    failure = json.loads((output / "events/research_002-error.json").read_text())
    assert failure["tool_call"] == {}


@pytest.mark.asyncio
async def test_omitted_or_empty_legacy_notes_do_not_erase_previous_notes(sources, tmp_path):
    calls = 0

    def policy(context):
        nonlocal calls
        calls += 1
        if calls == 1:
            result = action("fs.inspect", path="rules.md")
            result["notes"] = "已确定需要核对库存和审批范围。"
            return result
        assert context["notes"] == "已确定需要核对库存和审批范围。"
        result = action("fs.read", path="rules.md")
        if calls == 2:
            result["notes"] = ""
        else:
            result.pop("notes")
        return result

    report = await research_for_directory(Runtime(policy), sources, tmp_path / "research", "编写要求",
                                          limits=PlanningResearchLimits(max_calls=3))
    assert report["calls"] == 3


@pytest.mark.asyncio
async def test_rejected_finish_is_returned_to_its_actual_call(sources, tmp_path):
    from folio.research.transport import research_messages

    def policy(context):
        if not context["observations"]:
            return action("finish", summary="还没读原文")
        observed = context["observations"][-1]
        assert observed["tool"] == "finish"
        assert observed["result"]["status"] == "partial"
        assert observed["result"]["blocking_gaps"]
        messages = research_messages(context, {})
        result = next(message for message in messages if message["role"] == "tool")
        assert result["tool_call_id"] == observed["tool_call"]["id"]
        assert 'key source text' in result["content"]
        return action("fs.inspect", path="rules.md")

    report = await research_for_directory(Runtime(policy), sources, tmp_path / "research", "编写要求",
                                          limits=PlanningResearchLimits(max_calls=2))
    assert report["rounds"] == 2
    assert report["files"][0]["inspected"]


@pytest.mark.asyncio
@pytest.mark.parametrize("arguments", [
    '{"path":"rules.md"} {"path":"other.md"}',
    '{"path":"rules.md","summary":"核对资料","page":"1-15"}',
])
async def test_invalid_native_arguments_return_error_with_original_call_id(sources, tmp_path, arguments):
    from folio.research.transport import research_messages

    bad = {"id": "call_bad_arguments", "type": "function", "function": {
        "name": "fs_read", "arguments": arguments,
    }}

    def policy(context):
        if not context["observations"]:
            return RuntimeResult("", tool_calls=[bad])
        observed = context["observations"][-1]
        assert observed["tool_call"] == bad
        assert "error" in observed["result"]
        messages = research_messages(context, {})
        result = next(message for message in messages if message["role"] == "tool")
        assert result["tool_call_id"] == bad["id"]
        return action("fs.inspect", path="rules.md")

    report = await research_for_directory(Runtime(policy), sources, tmp_path / "research", "编写要求",
                                          limits=PlanningResearchLimits(max_calls=2))
    assert report["files"][0]["inspected"]
    assert not report["files"][0]["read"]
