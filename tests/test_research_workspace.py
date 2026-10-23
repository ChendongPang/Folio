"""Real file formats, exact source provenance, paging and workspace isolation."""

from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pypdf import PdfWriter

from folio.api.app import create_app
from folio.knowledge.local import LocalKnowledge
from folio.research.formats import SourceFormatError, read_source, source_hash
from folio.research.source_index import SourceIndex
from folio.research.workspace import WorkspaceError, WorkspaceManager
from folio.source_documents import _prepare_inputs


def excel(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "业务字段"
    sheet.append(["字段", "必填", "说明"])
    sheet.append(["amount", "是", "交易金额"])
    sheet.append(["comment", "否", "交易备注"])
    sheet.append(["calculation", "=1+2", None])
    workbook.save(path)
    return path


@pytest.fixture
def workspace(tmp_path):
    source = tmp_path / "sources"
    source.mkdir()
    (source / "rules.md").write_text(
        "# 交易要求\n金额必须大于零\n审批需要双人复核\n", encoding="utf-8"
    )
    excel(source / "fields.xlsx")
    return WorkspaceManager(source, tmp_path / "task")


def test_formats_excel_has_exact_rows_formulas_and_version(workspace):
    parsed = read_source(workspace.source_root / "fields.xlsx")
    assert parsed.source_version == source_hash(workspace.source_root / "fields.xlsx")
    assert parsed.metadata["sheets"][0]["name"] == "业务字段"
    assert parsed.blocks[1].locator == "sheet '业务字段'!A2:C2"
    assert parsed.blocks[1].metadata["cells"][0] == {"cell": "A2", "value": "amount"}
    assert parsed.blocks[-1].metadata["formulas"][0]["cache_verified"] is False
    assert parsed.metadata["formula_count"] == 1
    assert 'unverified' in parsed.warnings[0]


def test_read_exact_excel_cells_retains_canonical_source_slice(workspace):
    result = workspace.read("fields.xlsx", sheet="业务字段", cell_range="B2:C3")
    assert [item["text"] for item in result["evidence"]] == ["是 | 交易金额", "否 | 交易备注"]
    evidence = result["evidence"][0]
    assert evidence["metadata"]["source_locator"] == "sheet '业务字段'!A2:C2"
    assert evidence["locator"].endswith("/characters 9:17")
    assert workspace.resolve_evidence(evidence).text == "是 | 交易金额"
    reread = workspace.read("fields.xlsx", locator=evidence["locator"])
    assert reread["evidence"] == [evidence]
    assert result["complete"] and not result["truncated"]


def test_workspace_paging_and_search_are_explicit(workspace):
    files = workspace.list(limit=1)
    assert files["total"] == 2 and files["next_offset"] == 1 and not files["complete"]
    next_page = workspace.list(offset=1, limit=1)
    assert next_page["complete"] and next_page["files"] != files["files"]
    read = workspace.read("rules.md", limit=1)
    assert read["total"] == 3 and read["next_offset"] == 1
    found = workspace.search("金额", limit=1)
    assert found["total"] == 2 and found["next_offset"] == 1
    assert 'do not establish complete source coverage' in found["coverage_note"]
    assert "金额" in found["evidence"][0]["text"]


def test_long_blocks_are_sliced_without_losing_characters(tmp_path):
    source = tmp_path / "sources"
    source.mkdir()
    text = "甲乙" * 3000
    (source / "large.txt").write_text(text)
    workspace = WorkspaceManager(source, tmp_path / "task")
    whole = workspace.read("large.txt")
    assert whole["total"] == 1 and whole["complete"]
    assert whole["evidence"][0]["text"] == text
    assert whole["evidence"][0]["locator"] == "line 1"
    page = workspace.read("large.txt", max_chars=1753)
    assert page["truncated"] and page["next_offset"] == 0
    assert page["next_character_offset"] == 1753
    texts = [page["evidence"][0]["text"]]
    segments = page["read_segments"]
    while page["next_offset"] is not None:
        page = workspace.read(
            "large.txt", offset=page["next_offset"],
            character_offset=page["next_character_offset"], max_chars=937,
        )
        assert page["total"] == 1
        assert page["characters"] <= 937
        for evidence in page["evidence"]:
            assert workspace.resolve_evidence(evidence).model_dump() == evidence
        texts.extend(item["text"] for item in page["evidence"])
        segments.extend(page["read_segments"])
    assert "".join(texts) == text
    assert [segment["character_offset"] for segment in segments[1:]] == [
        segment["character_end"] for segment in segments[:-1]
    ]
    assert all(segment["offset"] == 0 and segment["total_characters"] == len(text) for segment in segments)
    assert segments[-1]["character_end"] == len(text)
    assert page["next_character_offset"] is None


def test_selected_section_reads_all_blocks_without_old_page_limits(tmp_path):
    source = tmp_path / "sources"
    source.mkdir()
    paragraphs = [f"需求{i}：完整保留这一条。" for i in range(130)]
    (source / "sections.md").write_text(
        "# 选用章节\n" + "\n".join(paragraphs) + "\n# 未选章节\n不可混入\n",
        encoding="utf-8",
    )
    workspace = WorkspaceManager(source, tmp_path / "task")
    result = workspace.read("sections.md", node_id="N001")
    assert result["total"] == 131 and result["complete"]
    assert [item["text"] for item in result["evidence"]] == ["# 选用章节", *paragraphs]
    explicit_limit = workspace.read("sections.md", node_id="N001", limit=125)
    assert len(explicit_limit["evidence"]) == 125
    assert (explicit_limit["next_offset"], explicit_limit["next_character_offset"]) == (125, 0)
    final = workspace.read("sections.md", node_id="N001", offset=125)
    assert final["complete"]
    assert explicit_limit["evidence"] + final["evidence"] == result["evidence"]


def test_whole_read_is_not_capped_at_six_hundred_thousand_characters(tmp_path):
    source = tmp_path / "sources"
    source.mkdir()
    text = "长" * 600123
    (source / "large.txt").write_text(text)
    workspace = WorkspaceManager(source, tmp_path / "task")
    result = workspace.read("large.txt")
    assert result["complete"] and result["characters"] == len(text)
    assert len(result["evidence"]) == result["total"] == 1
    assert workspace.read("large.txt", max_chars=len(text)) == result


def test_character_budget_continues_across_short_and_long_blocks(tmp_path):
    source = tmp_path / "sources"
    source.mkdir()
    paragraphs = ["短段", "第二段很长" * 13, "结束"]
    (source / "mixed.txt").write_text("\n".join(paragraphs))
    workspace = WorkspaceManager(source, tmp_path / "task")
    first = workspace.read("mixed.txt", max_chars=9)
    assert first["total"] == 3
    assert [item["text"] for item in first["evidence"]] == [paragraphs[0], paragraphs[1][:7]]
    assert (first["next_offset"], first["next_character_offset"]) == (1, 7)
    remaining = workspace.read("mixed.txt", offset=1, character_offset=7)
    assert remaining["complete"] and remaining["total"] == 3
    assert "".join(item["text"] for item in first["evidence"] + remaining["evidence"]) == "".join(paragraphs)
    assert remaining["read_segments"] == [
        {"offset": 1, "character_offset": 7, "character_end": len(paragraphs[1]), "total_characters": len(paragraphs[1])},
        {"offset": 2, "character_offset": 0, "character_end": 2, "total_characters": 2},
    ]
    one_character = workspace.read("mixed.txt", max_chars=1)
    assert one_character["evidence"][0]["text"] == "短"
    assert (one_character["next_offset"], one_character["next_character_offset"]) == (0, 1)


def test_explicit_character_range_has_stable_relative_paging(tmp_path):
    source = tmp_path / "sources"
    source.mkdir()
    text = "甲乙丙丁戊己庚辛壬癸" * 600
    (source / "large.txt").write_text(text)
    workspace = WorkspaceManager(source, tmp_path / "task")
    locator = "line 1/characters 137:5251"
    whole = workspace.read("large.txt", locator=locator)
    assert whole["total"] == 1 and whole["complete"]
    page = workspace.read("large.txt", locator=locator, max_chars=701)
    evidence = list(page["evidence"])
    assert page["read_segments"][0]["total_characters"] == 5251 - 137
    assert evidence[0]["metadata"]["characters_start"] == 137
    while page["next_offset"] is not None:
        page = workspace.read(
            "large.txt", locator=locator, offset=page["next_offset"],
            character_offset=page["next_character_offset"], max_chars=911,
        )
        assert page["total"] == 1
        evidence.extend(page["evidence"])
    assert "".join(item["text"] for item in evidence) == text[137:5251]
    for item in evidence:
        assert workspace.resolve_evidence(item).model_dump() == item
        assert workspace.read("large.txt", locator=item["locator"])["evidence"] == [item]


def test_excel_cell_range_pages_without_fixed_slices_or_losing_provenance(tmp_path):
    source = tmp_path / "sources"
    source.mkdir()
    workbook = Workbook()
    selected = "选定内容" * 900
    workbook.active.append(["前置列", selected, "不选列"])
    workbook.active.append(["前置列", "第二行选定", "不选列"])
    workbook.save(source / "range.xlsx")
    workspace = WorkspaceManager(source, tmp_path / "task")
    selection = {"sheet": "Sheet", "cell_range": "B1:B2"}
    whole = workspace.read("range.xlsx", **selection)
    assert whole["total"] == 2
    assert [item["text"] for item in whole["evidence"]] == [selected, "第二行选定"]
    page = workspace.read("range.xlsx", **selection, max_chars=873)
    evidence = list(page["evidence"])
    while page["next_offset"] is not None:
        page = workspace.read(
            "range.xlsx", **selection, offset=page["next_offset"],
            character_offset=page["next_character_offset"], max_chars=617,
        )
        assert page["total"] == 2
        evidence.extend(page["evidence"])
    assert "".join(item["text"] for item in evidence) == selected + "第二行选定"
    for item in evidence:
        assert workspace.resolve_evidence(item).model_dump() == item


def test_search_keeps_long_matching_block_and_can_continue_budgeted_excerpt(tmp_path):
    source = tmp_path / "sources"
    source.mkdir()
    text = "独特检索词 " + "正文" * 1800
    (source / "search.txt").write_text(text)
    workspace = WorkspaceManager(source, tmp_path / "task")
    whole = workspace.search("独特检索词")
    assert whole["complete"] and whole["total"] == 1
    assert whole["evidence"][0]["text"] == text
    first = workspace.search("独特检索词", max_chars=37)
    assert (first["next_offset"], first["next_character_offset"]) == (0, 37)
    rest = workspace.search(
        "独特检索词", offset=first["next_offset"],
        character_offset=first["next_character_offset"],
    )
    assert rest["complete"] and rest["total"] == 1
    assert first["evidence"][0]["text"] + rest["evidence"][0]["text"] == text
    assert workspace.resolve_evidence(rest["evidence"][0]).model_dump() == rest["evidence"][0]


@pytest.mark.parametrize(
    "path", ["../outside.txt", "/etc/passwd", "nested/../../rules.md", "..\\rules.md", "\x00"]
)
def test_path_escape_rejected(workspace, path):
    with pytest.raises(WorkspaceError):
        workspace.read(path)


def test_symlink_and_changed_sources_rejected(workspace, tmp_path):
    outside = tmp_path / "outside.md"
    outside.write_text("private")
    (workspace.source_root / "link.md").symlink_to(outside)
    with pytest.raises(WorkspaceError, match='symbolic link'):
        workspace.read("link.md")
    with pytest.raises(WorkspaceError, match='symbolic link'):
        WorkspaceManager(workspace.source_root, tmp_path / "second")
    (workspace.source_root / "link.md").unlink()
    workspace.read("rules.md")
    (workspace.source_root / "rules.md").write_text("changed")
    with pytest.raises(WorkspaceError, match='changed'):
        workspace.read("rules.md")


@pytest.mark.parametrize(
    "kwargs", [
        {"limit": 0}, {"offset": -1}, {"limit": True}, {"max_chars": 0},
        {"max_chars": True}, {"character_offset": -1}, {"character_offset": True},
        {"character_offset": 999}, {"offset": 999, "character_offset": 1},
    ]
)
def test_invalid_output_budgets_rejected(workspace, kwargs):
    with pytest.raises(WorkspaceError):
        workspace.read("rules.md", **kwargs)


def test_source_evidence_tampering_rejected(workspace):
    evidence = workspace.read("rules.md")["evidence"][0]
    evidence["text"] = "invented"
    with pytest.raises(WorkspaceError, match='does not match'):
        workspace.resolve_evidence(evidence)


def test_persistent_read_cache_and_task_isolation(workspace, monkeypatch):
    first = workspace.read("rules.md")
    monkeypatch.setattr(
        "folio.research.source_index.read_source",
        lambda _: pytest.fail("should reuse versioned cache"),
    )
    restored = WorkspaceManager(workspace.source_root, workspace.task_root)
    assert restored.read("rules.md") == first
    with pytest.raises(WorkspaceError, match='cannot be inside'):
        WorkspaceManager(workspace.source_root, workspace.source_root / "nested")
    assert all(
        (workspace.task_root / folder).is_dir()
        for folder in ("context", "scripts", "scratch", "artifacts")
    )


def test_word_preserves_heading_table_and_source_locations(tmp_path):
    path = tmp_path / "source.docx"
    document = Document()
    document.add_heading("业务约束", 1)
    document.add_paragraph("必须双人复核")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "字段"
    table.rows[0].cells[1].text = "金额"
    document.save(path)
    parsed = read_source(path)
    assert parsed.blocks[0].kind == "heading"
    assert parsed.blocks[-1].kind == "table_row"
    assert parsed.blocks[-1].locator.endswith("/row[1]")


def test_source_larger_than_previous_byte_limit_still_reads_actual_image(tmp_path):
    path = tmp_path / "large.png"
    Image.new("RGB", (20, 30), "white").save(path)
    # Trailing sparse data raises the original's real byte size without
    # allocating a huge decoded image or bypassing the production reader.
    size = 64 * 1024 * 1024 + 1
    with path.open("r+b") as stream:
        stream.truncate(size)
    parsed = read_source(path)
    assert path.stat().st_size == parsed.metadata["size"] == size
    assert parsed.metadata["width"] == 20 and parsed.metadata["height"] == 30
    assert parsed.metadata["needs_ocr"] is True
    assert parsed.source_version == source_hash(path)


@pytest.mark.parametrize("suffix", [".docx", ".xlsx", ".pptx"])
def test_office_above_previous_unpacked_limit_preserves_document_content(tmp_path, suffix):
    path = tmp_path / f"large{suffix}"
    expected = "实际业务内容"
    if suffix == ".docx":
        document = Document()
        document.add_paragraph(expected)
    elif suffix == ".xlsx":
        document = Workbook()
        document.active.append([expected])
    else:
        document = Presentation()
        document.slides.add_slide(document.slide_layouts[0]).shapes.title.text = expected
    document.save(path)
    # The unused ZIP member is valid compressed data, not fabricated ZIP
    # metadata, and does not force the Office reader to build enormous text.
    with ZipFile(path, "a", compression=ZIP_DEFLATED) as archive:
        with archive.open("unused/padding.bin", "w") as stream:
            chunk = b"\0" * (1024 * 1024)
            for _ in range(81):
                stream.write(chunk)
    with ZipFile(path) as archive:
        assert sum(item.file_size for item in archive.infolist()) > 80 * 1024 * 1024
    assert path.stat().st_size < 64 * 1024 * 1024
    parsed = read_source(path)
    assert any(expected in block.text for block in parsed.blocks)
    assert parsed.metadata["format"] == suffix[1:]


def test_office_entry_count_limit_remains_independent_of_byte_size(tmp_path):
    path = tmp_path / "many-parts.docx"
    Document().save(path)
    with ZipFile(path, "a", compression=ZIP_DEFLATED) as archive:
        for index in range(5001 - len(archive.infolist())):
            archive.writestr(f"unused/part-{index}", b"")
    with pytest.raises(SourceFormatError, match='maximum 5,000'):
        read_source(path)


def test_previous_unreadable_index_is_retried_without_invalidating_ready_cache(tmp_path, monkeypatch):
    source = tmp_path / "sources"
    source.mkdir()
    path = source / "previously-rejected.txt"
    path.write_text("原先被文件大小规则拒绝的资料。", encoding="utf-8")
    index = SourceIndex(source, cache_root=tmp_path / "cache")
    original_reader = read_source

    def previous_reader(_):
        raise SourceFormatError("资料不存在或超过单文件 64 MB 读取上限")

    with monkeypatch.context() as patch:
        patch.setattr("folio.research.source_index.read_source", previous_reader)
        first = index.prepare()
    assert first["files"][path.name]["status"] == "unreadable"
    assert not list((tmp_path / "cache" / "objects").glob("*.json"))
    calls = []

    def current_reader(path):
        calls.append(path)
        return original_reader(path)

    monkeypatch.setattr("folio.research.source_index.read_source", current_reader)
    recovered = SourceIndex(source, cache_root=tmp_path / "cache")
    assert recovered.prepare()["files"][path.name]["status"] == "ready"
    assert recovered.load(path.name).blocks[0].text == path.read_text()
    assert len(calls) == 1
    again = SourceIndex(source, cache_root=tmp_path / "cache")
    assert again.prepare()["files"][path.name]["status"] == "ready"
    assert len(calls) == 1


def test_slides_extract_visible_text_and_tables(tmp_path):
    path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "交易接口"
    slide.placeholders[1].text = "amount 为必填字段"
    presentation.save(path)
    parsed = read_source(path)
    assert parsed.metadata["slide_count"] == 1
    assert parsed.metadata["slides"] == [{"slide": 1, "title": "交易接口"}]
    assert any("amount" in block.text and block.metadata["slide"] == 1 for block in parsed.blocks)


def test_scanned_pdf_and_images_require_explicit_ocr(tmp_path):
    source = tmp_path / "sources"
    source.mkdir()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.write(source / "scan.pdf")
    Image.new("RGB", (20, 20), "white").save(source / "scan.png")
    workspace = WorkspaceManager(source, tmp_path / "task")
    for path in ("scan.pdf", "scan.png"):
        result = workspace.inspect(path)
        assert result["metadata"]["needs_ocr"] and not result["complete"]
        assert workspace.read(path)["evidence"] == []
        assert not workspace.read(path)["source_complete"]
    image_warnings = "\n".join(workspace.inspect("scan.png")["warnings"])
    assert "Use the isolated OCR tool." in image_warnings
    assert "请调用隔离 OCR 工具" not in image_warnings


def test_ocr_registration_preserves_unverified_provenance_and_restart(tmp_path):
    source = tmp_path / "sources"
    source.mkdir()
    Image.new("RGB", (20, 20), "white").save(source / "scan.png")
    workspace = WorkspaceManager(source, tmp_path / "task")
    assert workspace.search("金额")["total"] == 0
    workspace.register_extraction(
        "scan.png",
        {"blocks": [{"locator": "page 1", "text": "金额上限 100"}]},
        execution_id="run_1",
        image_digest="sha256:image",
    )
    evidence = workspace.read("scan.png")["evidence"][0]
    assert evidence["evidence_id"].startswith("extracted:")
    assert evidence["metadata"]["requires_extraction_review"]
    assert evidence["metadata"]["source_kind"] == "ocr"
    assert evidence["metadata"]["extraction"]["execution_id"] == "run_1"
    assert workspace.search("金额")["total"] == 1
    assert not any(
        "图片尚未进行 OCR" in warning for warning in workspace.inspect("scan.png")["warnings"]
    )
    restored = WorkspaceManager(source, workspace.task_root)
    assert restored.resolve_evidence(evidence).text == "金额上限 100"
    restored.register_extraction(
        "scan.png",
        {"blocks": [{"locator": "page:1", "text": "金额上限 100"}]},
        execution_id="run_1",
        image_digest="sha256:image",
    )
    assert restored.read("scan.png")["total"] == 1
    evidence["metadata"]["requires_extraction_review"] = False
    with pytest.raises(WorkspaceError, match='does not match'):
        restored.resolve_evidence(evidence)


def test_office_ocr_requires_trusted_page_count_and_keeps_original_hash(tmp_path):
    source = tmp_path / "sources"
    source.mkdir()
    document = Document()
    document.add_paragraph("原始文字")
    document.save(source / "office.docx")
    workspace = WorkspaceManager(source, tmp_path / "task")
    extraction = {
        "operation": "office_ocr",
        "page_count": 2,
        "blocks": [
            {"text": "第一页识别结果", "metadata": {"page": 1}},
            {"text": "第二页识别结果", "metadata": {"page": 2}},
        ],
    }
    result = workspace.register_extraction(
        "office.docx", extraction, execution_id="office_run", image_digest="sha256:office"
    )
    assert result["complete"]
    evidence = workspace.read("office.docx", page=2)["evidence"][0]
    assert evidence["source_version"] == source_hash(source / "office.docx")
    assert evidence["metadata"]["extraction"]["operation"] == "office_ocr"
    assert evidence["metadata"]["requires_extraction_review"]
    del extraction["page_count"]
    with pytest.raises(WorkspaceError, match='actual page count produced'):
        workspace.register_extraction(
            "office.docx", extraction, execution_id="bad_run", image_digest="sha256:office"
        )


def test_large_cell_metadata_does_not_bypass_output_budget(tmp_path):
    source = tmp_path / "sources"
    source.mkdir()
    workbook = Workbook()
    workbook.active.append(["X" * 30000, "第二列"])
    workbook.save(source / "large.xlsx")
    workspace = WorkspaceManager(source, tmp_path / "task")
    result = workspace.read("large.xlsx", max_chars=2000)
    assert len(result["evidence"][0]["text"]) == 2000
    assert "cells" not in result["evidence"][0]["metadata"]
    assert len(str(workspace.inspect("large.xlsx"))) < 4000


def test_malformed_oversized_and_encrypted_files_are_explicit(tmp_path, monkeypatch):
    path = tmp_path / "bad.xlsx"
    path.write_text("not an office document")
    with pytest.raises(SourceFormatError, match="ZIP"):
        read_source(path)
    path = tmp_path / "encrypted.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("password")
    writer.write(path)
    with pytest.raises(SourceFormatError, match='[Ee]ncrypted'):
        read_source(path)
    path = excel(tmp_path / "large.xlsx")
    monkeypatch.setattr("folio.research.formats.MAX_CELLS", 2)
    with pytest.raises(SourceFormatError, match='cells'):
        read_source(path)


async def test_local_retrieval_reuses_exact_source_evidence(workspace):
    knowledge = LocalKnowledge(workspace.source_root)
    results = await knowledge.search(["金额"], limit=10, max_chars=12000)
    assert results
    for item in results:
        assert workspace.resolve_evidence(item).evidence_id == item.evidence_id


def test_planning_catalog_is_metadata_only_and_defers_excel_reading_to_research(tmp_path):
    path = excel(tmp_path / "fields.xlsx")
    catalog = _prepare_inputs(
        tmp_path / "planning",
        [{"id": "source1", "name": "字段.xlsx", "path": str(path)}],
    )
    assert catalog[0]["source_version"] == source_hash(path)
    assert catalog[0]["format"] == "xlsx"
    assert catalog[0]["name"] == "字段.xlsx"
    assert catalog[0]["path"].endswith("字段.xlsx")
    assert "headings" not in catalog[0] and "excerpt" not in catalog[0]
    workspace = WorkspaceManager(tmp_path / "planning" / "sources", tmp_path / "research")
    navigation = workspace.inspect_navigation(catalog[0]["path"])
    assert "业务字段" in navigation["text"]
    assert "amount" not in navigation["text"]
    original = workspace.read(catalog[0]["path"], node_id="N001")
    assert any("amount" in item["text"] for item in original["evidence"])


@pytest.mark.parametrize("suffix", ["xlsx", "pdf", "pptx", "png"])
def test_upload_accepts_valid_extended_formats(tmp_path, suffix):
    path = tmp_path / ("source." + suffix)
    if suffix == "xlsx":
        excel(path)
    elif suffix == "pdf":
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.write(path)
    elif suffix == "pptx":
        Presentation().save(path)
    else:
        Image.new("RGB", (10, 10), "white").save(path)
    with TestClient(create_app(tmp_path / "app")) as client:
        response = client.post(f"/uploads?filename=source.{suffix}", content=path.read_bytes())
        assert response.status_code == 201, response.text
        bad = client.post(f"/uploads?filename=bad.{suffix}", content=b"not valid")
        assert bad.status_code == 422
