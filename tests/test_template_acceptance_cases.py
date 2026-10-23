"""Validate uploadable template fixtures, not real-model long-document quality."""

import hashlib
import importlib
import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from folio.document.reader import xpath
from folio.document.template_contract import read_template

ROOT = Path(__file__).resolve().parents[1] / "examples" / "acceptance_cases"
SLUGS = (
    "01_expense_urd", "02_order_srs", "05_inventory_report",
    "06_operations_manual", "08_data_platform_srs",
)


def json_file(path):
    return json.loads(path.read_text(encoding="utf-8"))


@pytest.mark.parametrize("slug", SLUGS)
def test_uploaded_template_variants_reuse_sources_without_leaking_expected_answers(slug):
    folder = ROOT / "template_variants" / slug
    expected = json_file(folder / "expected.json")
    manifest = json_file(folder / "manifest.json")
    recipe = json_file(ROOT / "recipes" / f"{slug}.json")
    template = folder / "文档模板.docx"
    contract = read_template(template)
    assert [s.title for s in contract.sections] == [c["title"] for c in recipe["expected"]["chapters"]]
    assert [s.model_dump(exclude={"body_index"}) for s in contract.sections] == expected["fixed_sections"]
    assert all(s.level == 1 for s in contract.sections)
    assert expected["target_chars"] == recipe["target_chars"]
    assert expected["full_long_document_generation_tested"] is False
    assert expected["requires_manual_review"] is True
    assert (folder / expected["base_expected_path"]).resolve() == ROOT / "cases" / slug / "expected" / "expected.json"
    original_prompt = (ROOT / "cases" / slug / "prompt.txt").read_text(encoding="utf-8").strip()
    assert (folder / "prompt.txt").read_text(encoding="utf-8").startswith(original_prompt + "\n\n")
    templates = [item for item in manifest["files"] if item["role"] == "template"]
    assert len(templates) == 1 and templates[0]["path"] == "文档模板.docx"
    for item in manifest["files"]:
        path = folder / item["path"]
        assert path.stat().st_size == item["bytes"]
        assert hashlib.sha256(path.read_bytes()).hexdigest() == item["sha256"]
        if item["role"] == "source":
            assert path.resolve().is_relative_to(ROOT / "cases" / slug / "sources")
    assert len(manifest["files"]) == 1 + len(recipe["sources"])

    doc = Document(template)
    text = "\n".join(p.text for p in doc.paragraphs)
    planner_context = json.dumps(contract.planning_context(), ensure_ascii=False)
    for example in expected["forbidden_template_examples"]:
        assert example in text
        assert example not in planner_context
    assert all(fact["statement"] not in text for fact in recipe["expected"]["facts"])
    assert "{{document_title}}" in text
    assert doc.styles["Normal"].font.name == "Arial"
    assert doc.styles["Normal"].font.size.pt == 10.5
    assert doc.styles["Normal"].element.get_or_add_rPr().rFonts.get(qn("w:eastAsia")) == "仿宋"
    for key, value in expected["format_expectations"]["margins_cm"].items():
        assert getattr(doc.sections[0], key + "_margin").cm == pytest.approx(value, abs=0.001)
    assert expected["variant_id"] in doc.sections[0].header.paragraphs[0].text
    assert xpath(doc.sections[0].footer._element, ".//w:fldSimple[@w:instr='PAGE']")
    assert [i for i, s in enumerate(contract.sections) if s.allow_children] == (
        list(range(2, 12)) if expected["variant_id"] == "TV08" else []
    )


def test_order_variant_has_two_independent_required_tables_with_identical_headers():
    contract = read_template(ROOT / "template_variants" / "02_order_srs" / "文档模板.docx")
    section = next(s for s in contract.sections if s.title == "支付窗口与并发竞争")
    assert len(section.tables) == 2
    assert section.tables[0].headers == section.tables[1].headers
    assert section.tables[0].table_id != section.tables[1].table_id
    assert section.tables[0].body_index != section.tables[1].body_index


def test_inventory_variant_has_two_physical_headers_and_unmerged_empty_data():
    path = ROOT / "template_variants" / "05_inventory_report" / "文档模板.docx"
    contract = read_template(path)
    required = [table for section in contract.sections for table in section.tables]
    assert len(required) == 1
    assert required[0].header_rows == 2
    assert required[0].headers == ["商品资料 / SKU", "商品资料 / 计量单位", "补货结论 / 建议数量", "补货结论 / 计算依据"]
    table = Document(path).tables[0]
    assert all(xpath(row._tr, "./w:trPr/w:tblHeader") for row in table.rows[:2])
    assert len(table.rows) == 3
    assert not xpath(table.rows[2]._tr, ".//w:gridSpan | .//w:vMerge")
    assert all(not cell.text for cell in table.rows[2].cells)


def test_variant_builder_is_offline_and_leaves_original_sources_untouched(tmp_path, monkeypatch):
    import httpx

    def forbidden(*args, **kwargs):
        raise AssertionError("Template fixture generation cannot call a network service")

    monkeypatch.setattr(httpx.Client, "send", forbidden)
    monkeypatch.setattr(httpx.AsyncClient, "send", forbidden)
    monkeypatch.syspath_prepend(str(ROOT))
    build = importlib.import_module("build_template_variants")
    paths = [path for slug in SLUGS for path in (ROOT / "cases" / slug).rglob("*") if path.is_file()]
    before = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in paths}
    folders = build.build_all(tmp_path)
    assert len(folders) == 5
    for folder in folders:
        contract = read_template(folder / build.TEMPLATE_NAME)
        manifest = json_file(folder / "manifest.json")
        assert contract.source_hash == manifest["files"][0]["sha256"]
        assert (folder / "prompt.txt").is_file()
        assert json_file(folder / "expected.json")["validation_status"] == "fixture_structure_checked_only"
    assert before == {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in paths}
