"""Style examples survive preview/confirmation without entering factual evidence."""

import json
import re

from docx import Document
from docx.shared import Pt

from folio.models import JobInput
from folio.orchestration.engine import Engine
from folio.previews import PreviewLibrary
from folio.runtime.base import RuntimeResult
from folio.runtime.demo import DemoRuntime
from folio.source_documents import SourceDirectoryDrafts
from folio.store import Store


def example(path):
    document = Document()
    document.styles["Normal"].font.size = Pt(15)
    for index in range(12):
        document.add_heading(f"范文专用栏目{index}", 1)
        document.add_paragraph("秘密旧公司预算为八百万元。" + "本段用于观察表达形式和段落习惯。" * 10)
    document.save(path)
    return path


async def test_style_is_analyzed_before_planning_and_frozen_for_writing(tmp_path, monkeypatch):
    calls = []

    class Planner(DemoRuntime):
        def __init__(self, config=None):
            pass

        async def complete(self, task, context, schema, session_key):
            calls.append((task, context))
            if task == "style_analysis":
                ids = re.findall(r"### (D\d+-S\d+) /[^\n]*\n\nType: paragraph", context["style_analysis_text"])
                assert len(ids) >= 2
                return RuntimeResult(json.dumps({
                    "rules": [{
                        "dimension": dimension, "scope": "body", "description": description,
                        "sample_ids": ids[:2], "confidence": "observed",
                    } for dimension, description in [
                        ("organization", "先交代背景，再分项解释，最后给出小结。"),
                        ("writing", "使用客观陈述和简明的分项说明。"),
                        ("structure", "全文通过多级小节组织信息。"),
                        ("detail", "正文用完整句子具体展开。"),
                        ("headings", "小节采用简短主题词命名。"),
                        ("expression", "同类概念保持一致的称谓。"),
                    ]],
                    "dimension_checks": [{"dimension": "figures_tables", "status": "not_present", "reason": "已读文字未出现图表说明。", "sample_ids": ids[:2]}],
                    "request_sample_ids": [], "uncertainties": [],
                }, ensure_ascii=False))
            assert "秘密旧公司" not in json.dumps(context, ensure_ascii=False)
            assert "范文专用栏目" not in json.dumps(context, ensure_ascii=False)
            return await super().complete(task, context, schema, session_key)

    monkeypatch.setattr("folio.directory_drafts.LLMRuntime", Planner)
    monkeypatch.setattr("folio.document.word.paginate", lambda *a, **k: {"status": "not_required"})
    monkeypatch.setattr("folio.document.pagination.paginate", lambda *a, **k: {"status": "not_required"})
    root = tmp_path / "state"
    factual = tmp_path / "rules.md"
    factual.write_text("# 订单规则\n订单付款前允许撤销。", encoding="utf-8")
    sample = example(tmp_path / "example.docx")
    attachments = [
        {"id": "a" * 32, "name": factual.name, "path": str(factual), "role": "source"},
        {"id": "b" * 32, "name": sample.name, "path": str(sample), "role": "style_reference"},
    ]
    drafts = SourceDirectoryDrafts(root, PreviewLibrary(root))
    result = await drafts.create_from_sources(None, "conversation", attachments, "# 订单\n说明撤销规则")
    loaded = drafts.load(result["revision"], "conversation")
    assert [task for task, _ in calls].index("style_analysis") < [task for task, _ in calls].index("directory_research")
    assert "先交代背景" in next(context for task, context in calls if task == "plan")["writing_profile"]
    assert len(list(loaded["knowledge_path"].iterdir())) == 1
    assert len(loaded["sample_paths"]) == 1
    assert (loaded["folder"] / "style" / "style.md").is_file()
    assert "style_url" in result
    profile = loaded["writing_profile"]
    assert profile["coverage"]["full_document_review"] is False
    assert profile["coverage"]["text_complete"] is True
    assert profile["layout"]["body"]["font_size_pt"] == 15

    store = Store(root)
    job_id = store.create(JobInput(
        title="订单说明", user_request="# 订单\n说明撤销规则", runtime="demo",
        knowledge_path=str(loaded["knowledge_path"]), sample_paths=loaded["sample_paths"],
    ), plan=loaded["plan"], writing_profile=profile)
    frozen = root / "jobs" / job_id / "inputs" / "style" / "profile.json"
    assert json.loads(frozen.read_text()) == profile
    sample.unlink()
    calls.clear()
    job = await Engine(store, Planner()).run(job_id)
    assert job["status"] in {"draft", "completed", "published"}
    assert not any(task in {"style_analysis", "plan"} for task, _ in calls)
    writing = [context for task, context in calls if task in {"generate", "repair"}]
    assert writing and all("客观陈述" in context["writing_profile"] for context in writing)
    for state in store.states(job_id):
        assert all("example" not in evidence.source_id for evidence in state.evidence)
    exported = list((root / "jobs" / job_id).rglob("draft.docx"))
    assert exported
    assert Document(exported[0]).styles["Folio Body"].font.size.pt == 15


def test_style_analysis_request_has_no_project_context_and_uses_readable_wire(tmp_path):
    from types import SimpleNamespace

    from folio.document.style_profile import StyleAnalysisStep
    from folio.orchestration.requests import RequestWorkflow
    from folio.runtime.base import prompt_chars
    from folio.runtime.llm import LLMRuntime

    owner = SimpleNamespace(request=SimpleNamespace(user_request="项目需求" * 10000))
    context = {"style_analysis_text": "# 风格分析\n\n## 样本\n原始片段。"}
    assert RequestWorkflow.request_context(owner, "style_analysis", context) == context
    runtime = LLMRuntime.__new__(LLMRuntime)
    runtime.config = {"structured_output": "json_schema", "provider_mode": "custom"}
    runtime.model = "mock"
    schema = StyleAnalysisStep.model_json_schema()
    messages, _, *_ = runtime.request_parts("style_analysis", context, schema)
    assert messages[1]["content"] == context["style_analysis_text"]
    assert "\\n" not in messages[1]["content"]
    assert '"properties"' not in messages[0]["content"]
    assert prompt_chars(runtime, "style_analysis", context, schema) > sum(len(item["content"]) for item in messages)


async def test_job_style_calls_are_accounted_and_recovery_uses_frozen_result(store, inputs, tmp_path):
    from folio.orchestration.style_processing import prepare_job_style

    sample = example(tmp_path / "sample.docx")
    job_id = store.create(JobInput(**inputs, sample_paths=[str(sample)], runtime="llm"))
    job = store.start_execution(job_id)
    engine = Engine(store, DemoRuntime())
    engine.job_id = job_id
    engine.request = JobInput.model_validate(job["input"])
    engine.limits = engine.request.limits
    profile = await prepare_job_style(engine)
    calls = profile["coverage"]["model_calls"]
    assert calls >= 1
    assert store.call_budget_usage(job_id, "style_analysis") == calls
    assert store.get(job_id)["calls"] == calls
    assert profile["source_hashes"]
    assert await prepare_job_style(engine) == profile
    assert store.get(job_id)["calls"] == calls


async def test_style_image_transport_keeps_markdown_and_reserves_image_capacity(tmp_path):
    import hashlib

    from PIL import Image

    from folio.document.style_profile import StyleAnalysisStep
    from folio.runtime.capacity import Capacity
    from folio.runtime.image_inputs import IMAGE_TOKEN_RESERVE
    from folio.runtime.llm import LLMRuntime

    path = tmp_path / "page.png"
    Image.new("RGB", (10, 10), "white").save(path)
    runtime = LLMRuntime.__new__(LLMRuntime)
    runtime.model = "mock"
    runtime.config = {
        "structured_output": "json_schema", "provider_mode": "custom",
        "supports_image_input": True, "max_tokens": 4096,
    }
    runtime.capacity = Capacity(runtime)
    context = {"style_analysis_text": "# 风格分析\n\n查看页面留白。", "page_images": [{
        "page": 1, "path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    }]}
    schema = StyleAnalysisStep.model_json_schema()
    messages, _, *_ = runtime.request_parts("style_analysis", context, schema)
    parts = messages[1]["content"]
    assert parts[0]["text"] == context["style_analysis_text"]
    assert sum(part["type"] == "image_url" for part in parts) == 1
    measured = await runtime.measure("style_analysis", context, schema, 24000)
    assert measured["image_token_reserve"] == IMAGE_TOKEN_RESERVE
    assert measured["characters"] == runtime.prompt_chars("style_analysis", context, schema)
    assert measured["estimated_input_tokens"] == measured["characters"] + IMAGE_TOKEN_RESERVE
