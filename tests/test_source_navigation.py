"""Model-facing navigation stays compact and maps back to canonical sources."""

import json
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pypdf import PdfWriter

from folio.research import source_index
from folio.research.navigation import annotate_nodes, render_navigation
from folio.research.source_index import SourceIndex, SourceIndexError, prepare_uploaded_source
from folio.research.workspace import WorkspaceError, WorkspaceManager


@pytest.fixture
def sources(tmp_path):
    root = tmp_path / "sources"
    root.mkdir()
    (root / "rules.md").write_text("# 交易要求\n金额必须大于零\n## 审批\n双人复核\n", encoding="utf-8")
    return root


def test_markdown_is_persisted_from_structure_and_does_not_include_body_or_warnings(sources):
    index = SourceIndex(sources)
    index.prepare()
    original = index.structure("rules.md")
    original["warnings"].append("图片尚未提取")
    text = render_navigation(original, "rules.md")
    assert text == "# rules.md\n\n- N000 Full text\n- N001 交易要求\n  - N002 审批\n"
    assert "金额" not in text and "图片" not in text and "source_version" not in text
    assert index.navigation("rules.md") == text
    assert index.navigation_path("rules.md").read_text() == text
    assert index.files["rules.md"]["navigation_ref"].endswith(".md")
    assert all("short_id" not in node for node in index.structure("rules.md")["nodes"])


def test_old_index_and_corrupt_navigation_repaired_without_reparsing(sources, monkeypatch):
    index = SourceIndex(sources)
    index.prepare()
    expected = index.navigation("rules.md")
    target = index.navigation_path("rules.md")
    target.unlink()
    manifest = index.manifest
    for key in ("navigation_ref", "navigation_hash", "navigation_version"):
        manifest["files"]["rules.md"].pop(key)
    (index.index_root / "manifest.json").write_text(json.dumps(manifest))

    def reject(*args, **kwargs):
        pytest.fail("derived Markdown repair reparsed source data")

    monkeypatch.setattr(source_index, "read_source", reject)
    current = SourceIndex(sources)
    current.prepare()
    assert current.navigation("rules.md") == expected
    target.write_text("伪造的目录")
    assert current.navigation("rules.md") == expected
    assert target.read_text() == expected


def test_navigation_does_not_load_canonical_body(sources, tmp_path, monkeypatch):
    workspace = WorkspaceManager(sources, tmp_path / "task")
    original = Path.open

    def reject_content(path, *args, **kwargs):
        if "content" in path.parts or "objects" in path.parts:
            pytest.fail("navigation opened body cache")
        return original(path, *args, **kwargs)

    monkeypatch.setattr(Path, "open", reject_content)
    result = workspace.inspect_navigation("rules.md")
    assert result["structure_complete"] and result["structure_total"] == 2
    assert result["text"] == workspace.source_index.navigation("rules.md")


def test_short_node_ids_read_same_evidence_and_child_navigation(sources, tmp_path):
    workspace = WorkspaceManager(sources, tmp_path / "task")
    nav = workspace.inspect_navigation("rules.md")
    canonical = nav["structure"][1]["node_id"]
    assert nav["structure"][1]["short_id"] == "N002"
    assert workspace.resolve_navigation_node("rules.md", "N002") == canonical
    assert workspace.read("rules.md", node_id="N002") == workspace.read("rules.md", node_id=canonical)
    subtree = workspace.inspect_navigation("rules.md", node_id="N001")
    assert [node["short_id"] for node in subtree["structure"]] == ["N002"]
    assert "N001 交易要求" in subtree["text"] and "N002 审批" in subtree["text"]
    assert workspace.inspect("rules.md", node_id="N001")["structure_total"] == 1


def test_more_than_fifty_nodes_returned_together_and_paging_covers_each_once(sources, tmp_path):
    (sources / "many.md").write_text("\n".join(f"# 第 {i} 章 " + "标题" * 20 for i in range(150)))
    workspace = WorkspaceManager(sources, tmp_path / "task")
    full = workspace.inspect_navigation("many.md")
    assert full["structure_complete"] and len(full["structure"]) == 150
    assert full["text"] == workspace.source_index.navigation("many.md")
    explicit = workspace.inspect_navigation("many.md", limit=8)
    assert len(explicit["structure"]) == 8 and explicit["next_offset"] == 8
    assert "offset=8" in explicit["text"] and not explicit["structure_complete"]
    offset, seen = 0, []
    while True:
        page = workspace.inspect_navigation("many.md", offset=offset, max_chars=2000)
        assert len(page["text"]) <= 2000
        assert page["structure"]
        seen.extend(node["short_id"] for node in page["structure"])
        if page["next_offset"] is None:
            break
        assert page["next_offset"] > offset
        assert f"offset={page['next_offset']}" in page["text"]
        offset = page["next_offset"]
    assert seen == [node["short_id"] for node in full["structure"]]


def test_excel_structure_keeps_range_and_header_candidates_not_business_rows(sources, tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "功能需求"
    sheet.append([])
    sheet.append(["编号", "模块", "功能名称"])
    sheet.append(["R1", "交易", "订单撤销"])
    book.save(sources / "requirements.xlsx")
    workspace = WorkspaceManager(sources, tmp_path / "task")
    text = workspace.inspect_navigation("requirements.xlsx")["text"]
    assert "N001 Worksheet: 功能需求" in text
    assert "Range: A1:C3" in text and "First nonempty row: 2" in text
    assert "Candidate headers: A：编号；B：模块；C：功能名称" in text
    assert "订单撤销" not in text and "header_basis" not in text


def test_oversized_single_sheet_fails_explicitly_and_larger_budget_progresses(sources, tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.append(["列" + "超长标题" * 30] * 20)
    book.save(sources / "wide.xlsx")
    workspace = WorkspaceManager(sources, tmp_path / "task")
    with pytest.raises(WorkspaceError, match='structure node'):
        workspace.inspect_navigation("wide.xlsx", max_chars=2000)
    result = workspace.inspect_navigation("wide.xlsx", max_chars=4000)
    assert result["structure_complete"] and len(result["structure"]) == 1


def test_pdf_without_bookmarks_uses_compact_pages_with_resolvable_aliases(sources, tmp_path):
    writer = PdfWriter()
    for _ in range(300):
        writer.add_blank_page(width=600, height=800)
    writer.write(sources / "pages.pdf")
    workspace = WorkspaceManager(sources, tmp_path / "task")
    result = workspace.inspect_navigation("pages.pdf", max_chars=2000)
    assert "N001–N300: pages 1–300" in result["text"]
    assert result["structure_complete"] and len(result["structure"]) == 300
    assert len(result["text"]) < 150
    assert workspace.resolve_navigation_node("pages.pdf", "N299") == result["structure"][298]["node_id"]
    assert "OCR" not in result["text"] and "书签" not in result["text"]


def test_pdf_bookmark_titles_and_actual_page_ranges_are_retained(sources, tmp_path):
    writer = PdfWriter()
    for _ in range(4):
        writer.add_blank_page(width=600, height=800)
    first = writer.add_outline_item("业务范围", 0)
    writer.add_outline_item("订单", 1, parent=first)
    writer.add_outline_item("接口", 2)
    writer.write(sources / "outline.pdf")
    workspace = WorkspaceManager(sources, tmp_path / "task")
    result = workspace.inspect_navigation("outline.pdf")
    assert "业务范围 (page 1 through page 3)" in result["text"]
    assert "订单 (page 2 through page 3)" in result["text"]
    assert "接口 (page 3 through page 4)" in result["text"]


def test_word_ppt_and_image_use_real_headings_slides_and_frames(sources, tmp_path):
    word = Document()
    word.add_heading("项目背景", 1)
    word.add_heading("业务目标", 2)
    word.add_paragraph("这是正文")
    word.save(sources / "scope.docx")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "实施计划"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(sources / "plan.pptx")
    Image.new("RGB", (80, 60)).save(sources / "diagram.png")
    workspace = WorkspaceManager(sources, tmp_path / "task")
    assert "- N001 项目背景\n  - N002 业务目标" in workspace.inspect_navigation("scope.docx")["text"]
    slides = workspace.inspect_navigation("plan.pptx")["text"]
    assert "N001 Slide 1: 实施计划" in slides and "N002 Slide 2" in slides
    image = workspace.inspect_navigation("diagram.png")["text"]
    assert "N001 Image（80 × 60）" in image and "OCR" not in image


def test_changed_source_and_navigation_symlinks_rejected(sources, tmp_path):
    index = SourceIndex(sources)
    index.prepare()
    target = index.navigation_path("rules.md")
    target.unlink()
    external = tmp_path / "outside.md"
    external.write_text("不能修改")
    target.symlink_to(external)
    with pytest.raises(SourceIndexError, match='symbolic link'):
        index.navigation("rules.md")
    assert external.read_text() == "不能修改"
    target.unlink()
    (sources / "rules.md").write_text("已修改")
    with pytest.raises(SourceIndexError, match='changed'):
        index.navigation("rules.md")


def test_aliases_are_derived_without_mutating_existing_index(sources):
    index = SourceIndex(sources)
    index.prepare()
    structure = index.structure("rules.md")
    before = json.dumps(structure, ensure_ascii=False)
    assert [node["short_id"] for node in annotate_nodes(structure)] == ["N000", "N001", "N002"]
    assert json.dumps(structure, ensure_ascii=False) == before


def test_upload_preparation_creates_navigation_from_shared_parse_once(tmp_path, monkeypatch):
    upload = tmp_path / "uploads" / "upload-id"
    upload.mkdir(parents=True)
    path = upload / "source.md"
    path.write_text("# 原始章节\n这段正文不应复制到结构文件", encoding="utf-8")
    cache = tmp_path / "source-cache"
    calls = []
    original = source_index.read_source

    def counting(path):
        calls.append(path)
        return original(path)

    monkeypatch.setattr(source_index, "read_source", counting)
    document = prepare_uploaded_source(path, cache, display_name="业务需求.md")
    assert len(calls) == 1
    structure = json.loads((upload / "structure.json").read_text())
    assert structure["source_version"] == document.source_version
    assert structure["path"] == "source.md" and structure["display_name"] == "业务需求.md"
    markdown = (upload / "structure.md").read_text()
    assert markdown == "# 业务需求.md\n\n- N000 Full text\n- N001 原始章节\n"
    prepare_uploaded_source(path, cache, display_name="业务需求.md")
    assert len(calls) == 1
    assert (upload / "structure.md").read_text() == markdown


def test_upload_navigation_writes_cannot_follow_symlink_or_replace_original(tmp_path):
    upload = tmp_path / "upload"
    upload.mkdir()
    source = upload / "source.md"
    source.write_text("# 原始资料")
    external = tmp_path / "outside.md"
    external.write_text("保留外部数据")
    (upload / "structure.md").symlink_to(external)
    with pytest.raises(SourceIndexError, match='symbolic link'):
        prepare_uploaded_source(source, tmp_path / "cache")
    assert external.read_text() == "保留外部数据"
    (upload / "structure.md").unlink()
    (upload / "structure.md").write_text("# 原件")
    with pytest.raises(SourceIndexError, match='path conflicts'):
        prepare_uploaded_source(upload / "structure.md", tmp_path / "cache")
    assert (upload / "structure.md").read_text() == "# 原件"


def test_office_ocr_nodes_keep_aliases_when_earlier_page_is_inserted_and_restarted(sources, tmp_path):
    word = Document()
    word.add_heading("原始章节", 1)
    word.add_paragraph("原始正文")
    word.save(sources / "office.docx")
    task = tmp_path / "task"
    workspace = WorkspaceManager(sources, task)
    original = workspace.inspect_navigation("office.docx")
    native_id = original["structure"][0]["node_id"]
    assert workspace.navigation_node_id("office.docx", native_id) == "N001"
    workspace.register_extraction(
        "office.docx",
        {"operation": "office_ocr", "page_count": 2, "blocks": [{"page": 2, "text": "第二页识别结果"}]},
        execution_id="page_two", image_digest="sha256:two",
    )
    first = workspace.inspect_navigation("office.docx")
    page_two = next(node for node in first["structure"] if node["metadata"].get("page") == 2)
    assert page_two["short_id"] == "N002"
    workspace.register_extraction(
        "office.docx",
        {"operation": "office_ocr", "page_count": 2, "blocks": [{"page": 1, "text": "第一页识别结果"}]},
        execution_id="page_one", image_digest="sha256:one",
    )
    second = workspace.inspect_navigation("office.docx")
    assert [node["short_id"] for node in second["structure"]] == ["N001", "N003", "N002"]
    assert "N003 Converted page 1" in second["text"]
    assert "N002 Converted page 2" in second["text"]
    assert workspace.resolve_navigation_node("office.docx", "N002") == page_two["node_id"]
    assert [item["text"] for item in workspace.read("office.docx", node_id="N002")["evidence"]] == ["第二页识别结果"]
    assert [item["text"] for item in workspace.read("office.docx", node_id="N003")["evidence"]] == ["第一页识别结果"]
    restored = WorkspaceManager(sources, task)
    after_restart = restored.inspect_navigation("office.docx")
    assert after_restart["text"] == second["text"]
    assert restored.navigation_node_id("office.docx", page_two["node_id"]) == "N002"
    assert restored.read("office.docx", node_id="N002") == workspace.read("office.docx", node_id="N002")
    record = json.loads((task / ".workspace" / "navigation-references.json").read_text())
    assert record["snapshot_id"] == workspace.source_snapshot_id
    assert record["files"]["office.docx"]["nodes"][page_two["node_id"]] == "N002"
    assert "short_id" not in json.dumps(workspace.source_index.structure("office.docx"))


def test_first_overlay_navigation_seeds_original_aliases_and_does_not_reuse_removed_ids(sources, tmp_path):
    word = Document()
    word.add_heading("原始章节", 1)
    word.save(sources / "office.docx")
    workspace = WorkspaceManager(sources, tmp_path / "task")

    def replace_page(page):
        workspace.register_extraction(
            "office.docx",
            {"operation": "office_ocr", "page_count": 3, "blocks": [{"page": page, "text": f"识别第 {page} 页"}]},
            execution_id="same_execution", image_digest="sha256:same",
        )

    replace_page(2)
    # First navigation happens after an overlay already exists.
    first = workspace.inspect_navigation("office.docx")
    assert [node["short_id"] for node in first["structure"]] == ["N001", "N002"]
    replace_page(1)
    second = workspace.inspect_navigation("office.docx")
    assert [node["short_id"] for node in second["structure"]] == ["N001", "N003"]
    with pytest.raises(WorkspaceError, match='node does not exist'):
        workspace.read("office.docx", node_id="N002")
    replace_page(3)
    third = workspace.inspect_navigation("office.docx")
    assert [node["short_id"] for node in third["structure"]] == ["N001", "N004"]
    replace_page(2)
    assert workspace.inspect_navigation("office.docx")["structure"][1]["short_id"] == "N002"


def test_navigation_renderer_accepts_explicit_stable_aliases(sources):
    index = SourceIndex(sources)
    index.prepare()
    structure = index.structure("rules.md")
    nodes = structure["nodes"]
    aliases = {nodes[0]["node_id"]: "N000", nodes[1]["node_id"]: "N007", nodes[2]["node_id"]: "N003"}
    text = render_navigation(structure, "rules.md", aliases)
    assert "N007 交易要求" in text and "N003 审批" in text
    assert "N001 交易要求" in index.navigation("rules.md")


def test_corrupt_alias_mapping_does_not_silently_renumber_sources(sources, tmp_path):
    workspace = WorkspaceManager(sources, tmp_path / "task")
    workspace.inspect_navigation("rules.md")
    path = workspace.task_root / ".workspace" / "navigation-references.json"
    saved = json.loads(path.read_text())
    aliases = saved["files"]["rules.md"]["nodes"]
    for node, alias in list(aliases.items()):
        if alias == "N001":
            aliases[node] = "N002"
        elif alias == "N002":
            aliases[node] = "N001"
    path.write_text(json.dumps(saved))
    restored = WorkspaceManager(sources, workspace.task_root)
    with pytest.raises(WorkspaceError, match='Invalid short structure-ID mapping'):
        restored.read("rules.md", node_id="N001")
