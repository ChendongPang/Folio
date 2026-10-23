"""Business research integration with controlled models and real isolation."""

import asyncio
import hashlib
import json
import shutil
import subprocess
from pathlib import Path

import pytest
from docx import Document
from native_calls import research_call
from openpyxl import Workbook

from folio.models import Evidence, JobInput, Limits
from folio.orchestration.engine import Engine
from folio.orchestration.work import WorkPlan, WorkTask
from folio.research.artifacts import derived_evidence_error
from folio.research.contracts import ResearchConfig
from folio.research.ledger import ResearchLedger
from folio.research.sandbox import SandboxError, SandboxManager
from folio.research.workflow import research_findings
from folio.runtime.base import RuntimeResult
from folio.runtime.demo import DemoRuntime

SCRIPT = """import json, os
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
root = Path(os.environ['FOLIO_SOURCES'])
requirements = '\\n'.join(p.text for p in Document(root / 'requirements.docx').paragraphs)
assert '去重' in requirements
book = load_workbook(root / 'fields.xlsx', read_only=True, data_only=False)
rows = list(book['Fields'].values)
names = sorted({row[0] for row in rows[1:] if row[1] == '必填'})
book.close()
output = {'required_unique_fields': len(names), 'field_names': names}
Path(os.environ['FOLIO_ARTIFACTS'], 'count.json').write_text(json.dumps(output, ensure_ascii=False), encoding='utf-8')
"""


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class FixtureSandbox(SandboxManager):
    """Persist the real receipt format without executing generated code on host."""

    executed = []

    async def execute(self, code, *, execution_id, profile="base", timeout_seconds=None):
        self.executed.append(execution_id)
        if code == "RAISE_UNAVAILABLE":
            raise SandboxError("测试环境暂不可用，请改读原始资料", code="sandbox_unavailable")
        previous = self.get_execution(execution_id)
        if previous:
            return self.verify_execution(execution_id)
        script = self.scripts / (execution_id + ".py")
        script.write_text(code, encoding="utf-8")
        output = self.output_root / execution_id / "count.json"
        output.parent.mkdir(parents=True)
        output.write_text(json.dumps({"required_unique_fields": 2, "field_names": ["amount", "id"]}), encoding="utf-8")
        outputs = [{"path": output.relative_to(self.task_root).as_posix(), "sha256": sha(output), "size": output.stat().st_size}]
        record = {
            "execution_id": execution_id, "status": "completed", "exit_code": 0,
            "image_digest": "sha256:" + "1" * 64, "profile": profile, "dependencies": [],
            "script_path": script.relative_to(self.task_root).as_posix(), "script_sha256": sha(script),
            "source_manifest": self._source_manifest(), "artifacts": outputs, "output_files": outputs,
            "stdout": "", "stderr": "", "duration_ms": 1,
        }
        (self.executions / (execution_id + ".json")).write_text(json.dumps(record), encoding="utf-8")
        return record


class ScriptedResearchRuntime(DemoRuntime):
    def __init__(self, *, mode="computed", review_passed=True, cover_all=False):
        self.mode, self.review_passed, self.cover_all = mode, review_passed, cover_all
        self.calls = []
        self.writer_gap_raised = False

    def action(self, context):
        step = context["step"]
        if self.mode == "self_correct":
            return [
                ("sandbox.exec", {"code": "RAISE_UNAVAILABLE"}),
                ("fs.read", {"path": "requirements.docx"}),
                ("finish", {}),
            ][min(step, 2)]
        if self.mode == "read_only":
            return [("fs.read", {"path": "requirements.docx"}), ("finish", {})][min(step, 1)]
        if self.mode == "knowledge_gap":
            if context["writing_gaps"]:
                return [("fs.search", {"scope": "knowledge", "query": "审批"}), ("finish", {})][min(step, 1)]
            return [("fs.read", {"path": "requirements.docx"}), ("finish", {})][min(step, 1)]
        if self.mode == "coverage":
            return [
                ("fs.read", {"path": "fields.xlsx", "sheet": "Fields"}),
                ("artifact.save", {"operation": "coverage.create", "path": "fields.xlsx"}),
                ("finish", {}),
            ][min(step, 2)]
        steps = [
            ("fs.list", {}),
            ("fs.read", {"path": "requirements.docx"}),
            ("fs.read", {"path": "fields.xlsx", "sheet": "Fields"}),
            ("sandbox.exec", {"code": SCRIPT, "profile": "data"}),
        ]
        if step < len(steps):
            return steps[step]
        if step == 4:
            run = next(item["result"] for item in reversed(context["observations"]) if item["tool"] == "sandbox.exec")
            return "artifact.save", {
                "execution_id": run["execution_id"], "output_path": run["artifacts"][0]["path"],
                "evidence_ids": [item["evidence_id"] for item in context["evidence_catalog"]],
                "methodology": "按 Word 要求，以 Excel 第二列的必填标志筛选，再按第一列字段名称去重",
                "description": "共有 2 个不同的必填字段",
            }
        return "finish", {}

    async def complete(self, name, context, schema, session):
        self.calls.append((name, context))
        if name == "work_plan":
            plan = WorkPlan(tasks=[WorkTask(
                task_id="fields", title="字段设计", objective="明确字段规则并计算去重后的必填字段数量",
                section_ids=[section["section_id"] for section in context["sections"]],
                questions=["哪些字段必填？"], completion_criteria=["结论可追溯原始资料"],
                rationale="同一业务数据范围",
            )])
            return RuntimeResult(plan.model_dump_json())
        if name == "research_step":
            tool, arguments = ("research.save", {}) if context.get("memory_checkpoint_required") else self.action(context)
            return research_call(tool, arguments, summary="核对字段资料",
                                 notes="根据原始材料研究，计数按字段名去重。", missing_information=[],
                                 research_notes=self.research_notes(context))
        if name == "verify_derivation":
            result = json.loads(context["result"])
            assert result["required_unique_fields"] == 2
            assert result["field_names"] == ["amount", "id"]
            assert "去重" in context["methodology"]
            return RuntimeResult(json.dumps({
                "supported": self.review_passed, "source_matches": True,
                "computation_matches": self.review_passed, "business_scope_matches": True,
                "reason": "受控测试核对原始行、字段去重脚本及实际结果" if self.review_passed else "结果未通过独立核验",
            }))
        if name in {"generate", "repair"}:
            evidence = context["evidence"]
            derived = [item for item in evidence if item["evidence_id"].startswith("derived:")]
            selected = evidence if self.cover_all else derived or evidence[:1]
            pending = []
            if self.mode == "knowledge_gap" and name == "generate" and not self.writer_gap_raised:
                self.writer_gap_raised = True
                pending = ["补查审批规则"]
            return RuntimeResult(json.dumps({
                "blocks": [{"block_id": f"body{index}", "text": item["text"], "origin": "evidence", "evidence_ids": [item["evidence_id"]]} for index, item in enumerate(selected)],
                "pending_items": pending,
            }))
        return await super().complete(name, context, schema, session)


@pytest.fixture
def research_inputs(inputs, tmp_path, monkeypatch):
    monkeypatch.setenv("FOLIO_ARTIFACT_KEY_PATH", str(tmp_path / "host-control" / "signing.key"))
    inputs["user_request"] = "# 字段设计\n说明字段规则。"
    root = Path(inputs["knowledge_path"])
    (root / "rules.md").write_text("审批由主管复核。", encoding="utf-8")
    document = Document()
    document.add_paragraph("必填字段按字段名称去重后统计。")
    document.save(root / "requirements.docx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fields"
    for row in [
        ("字段", "要求"),
        ("id", "必填"),
        ("amount", "必填"),
        ("id", "必填"),
        ("name", "选填"),
    ]:
        sheet.append(row)
    workbook.save(root / "fields.xlsx")
    FixtureSandbox.executed = []
    return inputs


def make_request(inputs, **research):
    return JobInput(**inputs, runtime="demo", research=ResearchConfig(**research), limits=Limits(max_repairs=0))


def ledger_for(store, job_id):
    return ResearchLedger(store.root / "jobs" / job_id / "research" / "ledger.sqlite")


def use_fixture_sandbox(monkeypatch):
    monkeypatch.setattr("folio.research.workflow.SandboxManager", FixtureSandbox)


async def test_computed_research_reaches_body_business_receipt_and_review(store, research_inputs, monkeypatch):
    use_fixture_sandbox(monkeypatch)
    runtime = ScriptedResearchRuntime()
    engine = Engine(store, runtime)
    job_id = store.create(make_request(research_inputs))
    job = await engine.run(job_id)
    assert job["status"] == "draft"
    assert job["snapshot"]["review_passed"]
    work = store.work_states(job_id)["fields"]
    assert work.research["version"] == 2 and work.research["usage"]["calls"] == 5
    notes_path = Path(work.research["research_notes_path"])
    assert notes_path.name == "research.md" and "首段摘录" in notes_path.read_text(encoding="utf-8")
    assert work.research["research_memory"]["notes"]
    assert work.review_key
    state = store.states(job_id)[0]
    derived = next(item for item in state.evidence if item.metadata.get("provenance") == "derived")
    assert derived_evidence_error(derived) is None
    assert state.content.blocks[0].evidence_ids == [derived.evidence_id]
    assert engine.shared_state()["records"]
    names = [name for name, _ in runtime.calls]
    assert {"research_step", "verify_derivation", "understand_work", "generate", "extract", "validate", "review_work"} <= set(names)
    assert names.index("verify_derivation") < names.index("generate") < names.index("review_work")
    generation = next(context for name, context in runtime.calls if name == "generate")
    assert any(derived.evidence_id in fact["evidence_ids"] for fact in generation["business_record"]["facts"])
    receipts = list((store.root / "jobs" / job_id / "contexts" / "agent_results").glob("*.json"))
    assert receipts
    receipt = json.loads(receipts[0].read_text())
    assert receipt["research"]["version"] == 2 and receipt["review_key"]


@pytest.mark.skipif(shutil.which("docker") is None, reason="Docker is not installed")
async def test_real_docker_word_excel_research_produces_verified_document(store, research_inputs):
    probe = subprocess.run(["docker", "image", "inspect", "folio-research-data:v2"], capture_output=True, timeout=10)
    if probe.returncode:
        pytest.skip("The research data image is not installed")
    runtime = ScriptedResearchRuntime()
    engine = Engine(store, runtime)
    job_id = store.create(make_request(research_inputs, max_seconds_per_execution=30))
    job = await engine.run(job_id)
    assert job["status"] == "draft" and job["snapshot"]["review_passed"]
    run = next(item for item in ledger_for(store, job_id).runs("fields") if item["tool"] == "sandbox.exec")
    assert run["result"]["status"] == "completed", run["result"]
    assert run["result"]["image_digest"].startswith("sha256:")
    work = store.work_states(job_id)["fields"]
    output = Path(work.research["task_root"]) / run["result"]["artifacts"][0]["path"]
    assert json.loads(output.read_text())["required_unique_fields"] == 2
    body = "\n".join(p.text for p in Document(job["artifact"]).paragraphs)
    assert "共有 2 个不同的必填字段" in body


async def test_tool_failure_is_observed_and_business_agent_changes_strategy(store, research_inputs, monkeypatch):
    use_fixture_sandbox(monkeypatch)
    runtime = ScriptedResearchRuntime(mode="self_correct")
    job_id = store.create(make_request(research_inputs))
    job = await Engine(store, runtime).run(job_id)
    assert job["status"] == "draft" and job["snapshot"]["review_passed"]
    tools = ledger_for(store, job_id).runs("fields")
    assert [item["status"] for item in tools] == ["failed", "completed"]
    assert "暂不可用" in tools[0]["result"]["error"]
    next_context = next(context for name, context in runtime.calls if name == "research_step" and context["step"] == 1)
    assert next_context["observations"][-1]["result"]["status"] == "failed"
    failed_call = next_context["observations"][-1]["tool_call"]
    assert failed_call["id"] and failed_call["function"]["name"] == "sandbox_exec"
    assert json.loads(failed_call["function"]["arguments"])["code"] == "RAISE_UNAVAILABLE"


@pytest.mark.parametrize("legacy_pending", [False, True])
async def test_resume_after_tool_commit_reuses_result_without_charging_twice(store, research_inputs, monkeypatch, legacy_pending):
    use_fixture_sandbox(monkeypatch)
    runtime = ScriptedResearchRuntime(mode="read_only")
    job_id = store.create(make_request(research_inputs))
    original_finish = ResearchLedger.finish
    interrupted = False

    def interrupt_after_commit(self, request_id, result, **kwargs):
        nonlocal interrupted
        original_finish(self, request_id, result, **kwargs)
        if not interrupted:
            interrupted = True
            raise asyncio.CancelledError()

    monkeypatch.setattr(ResearchLedger, "finish", interrupt_after_commit)
    with pytest.raises(asyncio.CancelledError):
        await Engine(store, runtime).run(job_id)
    before = ledger_for(store, job_id).runs("fields")
    assert len(before) == 1 and before[0]["status"] == "completed"
    ledger = ledger_for(store, job_id)
    with ledger.connect() as conn:
        saved = conn.execute("SELECT round_id, data FROM research_sessions WHERE task_id='fields'").fetchone()
    session = json.loads(saved["data"])
    original_call = session["pending_batch"]["steps"][0]["tool_call"]
    assert original_call["function"]["name"] == "fs_read"
    if legacy_pending:
        session["pending_step"] = session.pop("pending_batch")["steps"][0]
        session["pending_step"].pop("tool_call")
        session["pending_step"]["arguments_json"] = json.dumps(session["pending_step"].pop("arguments"))
        ledger.save_session("fields", saved["round_id"], session)
    job = await Engine(store, runtime).run(job_id)
    assert job["status"] == "draft"
    after = ledger_for(store, job_id).runs("fields")
    assert after == before
    assert ledger_for(store, job_id).usage("fields")["calls"] == 1
    assert store.work_states(job_id)["fields"].evidence
    restored = ledger.load_session("fields", saved["round_id"])
    delivered = next(context for name, context in runtime.calls if name == "research_step" and context["step"] == 1)
    read = next(item for item in delivered["observations"] if item["tool"] == "fs.read")
    assert read["tool_call"] == ({} if legacy_pending else original_call)
    assert restored["research_memory"]["batches"][read["request_id"]]["reviewed"]
    assert len([context for name, context in runtime.calls if name == "research_step" and context["step"] == 0]) == 1


async def test_tool_budget_finishes_research_with_explicit_unresolved_finding(store, research_inputs, monkeypatch):
    use_fixture_sandbox(monkeypatch)
    runtime = ScriptedResearchRuntime()
    job_id = store.create(make_request(research_inputs, max_tool_calls=2))
    job = await Engine(store, runtime).run(job_id)
    work = store.work_states(job_id)["fields"]
    assert work.research["usage"]["calls"] == 2
    assert 'tool-call budget' in work.research["incomplete"]
    assert job["status"] == "needs_confirmation"
    assert any(item["key"].startswith("research-incomplete:") for item in job["snapshot"]["blocking_issues"])
    assert not FixtureSandbox.executed


async def test_finish_without_notes_preserves_originals_without_forcing_more_model_calls(store, research_inputs, monkeypatch):
    use_fixture_sandbox(monkeypatch)

    class NoNotesRuntime(ScriptedResearchRuntime):
        async def complete(self, name, context, schema, session):
            if name != "research_step":
                return await super().complete(name, context, schema, session)
            self.calls.append((name, context))
            if context["step"] == 0:
                return research_call("fs.read", {"path": "requirements.docx"})
            return research_call("finish", summary="资料读完了")

    runtime = NoNotesRuntime(mode="read_only")
    job_id = store.create(make_request(research_inputs))
    job = await Engine(store, runtime).run(job_id)
    work = store.work_states(job_id)["fields"]
    steps = [context for name, context in runtime.calls if name == "research_step"]
    assert len(steps) == 2 and not any(context.get("memory_checkpoint_required") for context in steps)
    assert work.research["research_memory"]["notes"] == []
    assert work.research["incomplete"] == ""
    assert job["status"] == "draft" and job["snapshot"]["review_passed"]
    ledger = ledger_for(store, job_id)
    runs = ledger.runs("fields")
    assert sum(item["tool"] == "fs.read" for item in runs) == 1
    from folio.research.memory import ResearchMemory

    memory = ResearchMemory(work.research["research_memory"])
    assert memory.pending_requests()  # No implicit review or permission to discard originals.
    session = ledger.load_session("fields", work.research["round_id"])
    original = next(item for item in session["observations"] if item["tool"] == "fs.read")
    assert original["model_text"] and original["result"]["evidence"]
    assert not memory.can_drop(original["request_id"])
    assert any(item["tool"] == "finish" for item in session["observations"])
    assert "understand_work" in [name for name, _ in runtime.calls]  # Content assessment still runs.


async def test_identical_read_requires_review_reason_and_preserves_error_reply(store, research_inputs, monkeypatch):
    use_fixture_sandbox(monkeypatch)

    class ReviewRuntime(ScriptedResearchRuntime):
        async def complete(self, name, context, schema, session):
            if name != "research_step":
                return await super().complete(name, context, schema, session)
            self.calls.append((name, context))
            step = context["step"]
            if context.get("memory_checkpoint_required"):
                return research_call("research.save", research_notes=self.research_notes(context))
            if step <= 2:
                return research_call("fs.read", {"path": "requirements.docx"},
                                     review_reason="复核是否明确按字段名称去重" if step == 2 else "")
            return research_call("finish", research_notes=self.research_notes(context))

    runtime = ReviewRuntime(mode="read_only")
    job_id = store.create(make_request(research_inputs))
    job = await Engine(store, runtime).run(job_id)
    assert job["status"] == "draft"
    reads = [item for item in ledger_for(store, job_id).runs("fields") if item["tool"] == "fs.read"]
    assert [item["status"] for item in reads] == ["completed", "failed", "completed"]
    next_context = next(context for name, context in runtime.calls if name == "research_step" and context["step"] == 2)
    error = next(item for item in next_context["observations"] if item["result"].get("status") == "failed")
    assert error["tool_call"]["function"]["name"] == "fs_read"
    assert "review_reason" in error["result"]["error"]


async def test_native_pending_finish_resumes_without_legacy_note_checkpoint_gate(store, research_inputs, monkeypatch):
    from folio.research.workflow import ResearchTools

    use_fixture_sandbox(monkeypatch)

    class FinishRuntime(ScriptedResearchRuntime):
        async def complete(self, name, context, schema, session):
            if name != "research_step":
                return await super().complete(name, context, schema, session)
            self.calls.append((name, context))
            if context["step"] == 0:
                return research_call("fs.read", {"path": "requirements.docx"})
            return research_call("finish", summary="读取已结束")

    original, interrupted = ResearchTools.model, False

    async def cancel_after_pending_saved(self, action, context, schema):
        nonlocal interrupted
        result = await original(self, action, context, schema)
        if action == "research_step" and result.steps[0].tool == "finish" and not interrupted:
            interrupted = True
            self.session["memory_finish_pending"] = True  # An older server could persist this gate.
            self.save_memory()
            raise asyncio.CancelledError()
        return result

    monkeypatch.setattr(ResearchTools, "model", cancel_after_pending_saved)
    runtime = FinishRuntime(mode="read_only")
    job_id = store.create(make_request(research_inputs))
    with pytest.raises(asyncio.CancelledError):
        await Engine(store, runtime).run(job_id)
    result = await Engine(store, runtime).run(job_id)
    assert result["status"] == "draft"
    work = store.work_states(job_id)["fields"]
    assert not work.research["research_memory"]["notes"]
    steps = [context for name, context in runtime.calls if name == "research_step"]
    assert len(steps) == 2  # The accepted pending finish is not requested again.
    assert not any(context.get("memory_checkpoint_required") for context in steps)
    session = ledger_for(store, job_id).load_session("fields", work.research["round_id"])
    assert "memory_finish_pending" not in session
    assert any(item["tool"] == "fs.read" and item["model_text"] for item in session["observations"])
    rejected = [item for item in ledger_for(store, job_id).runs("fields") if item["tool"] == "finish"]
    assert not rejected


async def test_rejected_computation_never_enters_writing_or_shared_facts(store, research_inputs, monkeypatch):
    use_fixture_sandbox(monkeypatch)
    runtime = ScriptedResearchRuntime(review_passed=False)
    engine = Engine(store, runtime)
    job_id = store.create(make_request(research_inputs))
    job = await engine.run(job_id)
    work = store.work_states(job_id)["fields"]
    assert job["status"] == "needs_confirmation"
    assert not job["snapshot"]["review_passed"]
    assert any(item["key"].startswith("research-artifact:") for item in job["snapshot"]["blocking_issues"])
    assert not any(item.metadata.get("provenance") == "derived" for item in work.evidence)
    writing = next(context for name, context in runtime.calls if name == "generate")
    assert not any(item["evidence_id"].startswith("derived:") for item in writing["evidence"])
    assert not any(eid.startswith("derived:") for fact in engine.shared_state()["records"].values() for eid in fact["evidence_ids"])
    result = next(item["result"] for item in ledger_for(store, job_id).runs("fields") if item["tool"] == "artifact.save")
    assert result["status"] == "needs_confirmation"


@pytest.mark.parametrize("cover_all", [False, True])
async def test_coverage_matches_current_written_blocks_not_research_claims(store, research_inputs, monkeypatch, cover_all):
    use_fixture_sandbox(monkeypatch)
    runtime = ScriptedResearchRuntime(mode="coverage", cover_all=cover_all)
    engine = Engine(store, runtime)
    job_id = store.create(make_request(research_inputs))
    job = await engine.run(job_id)
    work = store.work_states(job_id)["fields"]
    assert len(work.research["coverage_ids"]) == 1
    issues = research_findings(engine, engine.work_plan.tasks[0], work, store.states(job_id))
    assert bool(issues) is not cover_all
    if cover_all:
        assert job["snapshot"]["review_passed"]
    else:
        assert job["status"] == "needs_confirmation"
        assert any(item["key"].startswith("research-coverage:") for item in job["snapshot"]["blocking_issues"])


async def test_writing_gap_returns_to_research_and_existing_knowledge_provider(store, research_inputs, monkeypatch):
    use_fixture_sandbox(monkeypatch)
    external = Evidence(evidence_id="external-approval", source_id="knowledge", source_version="v1", locator="approval", text="审批由主管复核。")

    class Provider:
        calls = []

        async def search(self, queries, *, limit, max_chars):
            self.calls.append(queries)
            return [external]

    provider = Provider()
    runtime = ScriptedResearchRuntime(mode="knowledge_gap")
    engine = Engine(store, runtime, knowledge=provider)
    job_id = store.create(make_request(research_inputs))
    job = await engine.run(job_id)
    assert job["status"] == "draft" and job["snapshot"]["review_passed"]
    assert provider.calls == [["审批"]]
    work = store.work_states(job_id)["fields"]
    assert work.searches == [["审批"]]
    assert external.evidence_id in {item.evidence_id for item in work.evidence}
    gap_contexts = [context for name, context in runtime.calls if name == "research_step" and context["writing_gaps"]]
    assert gap_contexts and gap_contexts[0]["writing_gaps"] == ["补查审批规则"]
    runs = ledger_for(store, job_id).runs("fields")
    assert len({item["round_id"] for item in runs}) == 2


class ExecuteThenReadRuntime(ScriptedResearchRuntime):
    def action(self, context):
        return [
            ("sandbox.exec", {"code": SCRIPT, "profile": "data"}),
            ("fs.read", {"path": "requirements.docx"}),
            ("finish", {}),
        ][min(context["step"], 2)]


async def test_running_execution_at_full_reserved_budget_can_recover(store, research_inputs, monkeypatch):
    use_fixture_sandbox(monkeypatch)
    runtime = ExecuteThenReadRuntime()
    job_id = store.create(make_request(research_inputs, max_execution_seconds=1, max_seconds_per_execution=1))
    original_finish = ResearchLedger.finish
    interrupted = False

    def interrupt_before_commit(self, request_id, result, **kwargs):
        nonlocal interrupted
        if not interrupted:
            interrupted = True
            raise asyncio.CancelledError()
        original_finish(self, request_id, result, **kwargs)

    monkeypatch.setattr(ResearchLedger, "finish", interrupt_before_commit)
    with pytest.raises(asyncio.CancelledError):
        await Engine(store, runtime).run(job_id)
    before = ledger_for(store, job_id).runs("fields")
    assert before[0]["status"] == "running" and before[0]["reserved_seconds"] == 1
    job = await Engine(store, runtime).run(job_id)
    assert job["status"] == "draft"
    runs = ledger_for(store, job_id).runs("fields")
    assert [run["status"] for run in runs] == ["interrupted", "completed"]
    assert len(FixtureSandbox.executed) == 1
    assert ledger_for(store, job_id).usage("fields") == {"calls": 2, "seconds": 1}


async def test_nonzero_sandbox_result_is_failed_in_tool_ledger(store, research_inputs, monkeypatch):
    class TimedOutSandbox(FixtureSandbox):
        async def execute(self, code, *, execution_id, **kwargs):
            return {"execution_id": execution_id, "status": "timed_out", "exit_code": 137}

    monkeypatch.setattr("folio.research.workflow.SandboxManager", TimedOutSandbox)
    job_id = store.create(make_request(research_inputs))
    job = await Engine(store, ExecuteThenReadRuntime()).run(job_id)
    assert job["status"] == "draft"
    runs = ledger_for(store, job_id).runs("fields")
    assert runs[0]["status"] == "failed" and runs[0]["result"]["status"] == "timed_out"


@pytest.mark.parametrize("approved", [False, True])
async def test_coverage_exclusion_requires_independent_review(store, research_inputs, monkeypatch, approved):
    from folio.research.artifact_coverage import CoverageManager
    from folio.research.workspace import WorkspaceManager

    use_fixture_sandbox(monkeypatch)

    class ExclusionRuntime(ScriptedResearchRuntime):
        def action(self, context):
            if context["step"] < 2:
                return [("fs.read", {"path": "fields.xlsx", "sheet": "Fields"}),
                        ("artifact.save", {"operation": "coverage.create", "path": "fields.xlsx"})][context["step"]]
            if context["step"] == 2:
                record = context["observations"][-1]["result"]
                return "artifact.save", {
                    "operation": "coverage.resolve", "coverage_id": record["coverage_id"],
                    "resolutions": [{"item_id": record["scope"]["items"][0]["item_id"],
                                     "status": "not_applicable", "reason": "标题行属于列名，不是业务字段",
                                     "references": []}],
                }
            return "finish", {}

        async def complete(self, name, context, schema, session):
            if name == "review_coverage_exclusion":
                assert "item" in context and context["sections"]
                self.calls.append((name, context))
                return RuntimeResult(json.dumps({"approved": approved, "reason": "独立核对真实章节职责和原始标题行"}))
            if name in {"generate", "repair"}:
                context = dict(context, evidence=context["evidence"][1:])
            return await super().complete(name, context, schema, session)

    runtime = ExclusionRuntime(cover_all=True)
    engine = Engine(store, runtime)
    job_id = store.create(make_request(research_inputs))
    job = await engine.run(job_id)
    work = store.work_states(job_id)["fields"]
    root = Path(work.research["task_root"])
    manager = CoverageManager(root, WorkspaceManager(root.parents[2] / "sources", root))
    record = manager._load(work.research["coverage_ids"][0])
    first = record["dispositions"][record["scope"]["items"][0]["item_id"]]
    assert bool(first["review_key"]) is approved
    assert job["snapshot"]["review_passed"] is approved
    assert any(name == "review_coverage_exclusion" for name, _ in runtime.calls)


def test_recovered_ocr_evidence_stays_available_and_unverified(tmp_path):
    from types import SimpleNamespace

    from PIL import Image

    from folio.orchestration.work import WorkState
    from folio.research.workflow import ResearchTools
    from folio.research.workspace import WorkspaceManager

    sources = tmp_path / "sources"
    sources.mkdir()
    Image.new("RGB", (10, 10), "white").save(sources / "scan.png")
    workspace = WorkspaceManager(sources, tmp_path / "task")
    workspace.register_extraction(
        "scan.png", {"blocks": [{"locator": "page 1", "text": "限额100"}]},
        execution_id="ocr_run", image_digest="sha256:ocr",
    )
    evidence = workspace.read("scan.png")["evidence"]
    tools = ResearchTools(
        SimpleNamespace(request=SimpleNamespace(research=ResearchConfig())),
        SimpleNamespace(task_id="owner"), WorkState(task_id="owner"),
        WorkspaceManager(sources, workspace.task_root), FixtureSandbox(sources, workspace.task_root),
        ResearchLedger(tmp_path / "ledger.sqlite"), {"evidence": evidence},
    )
    restored = tools.evidence[evidence[0]["evidence_id"]]
    assert restored.text == "限额100"
    assert restored.metadata["requires_extraction_review"]
    assert derived_evidence_error(restored)


async def test_business_native_batch_reaches_next_request_with_all_correlated_results(store, research_inputs, monkeypatch):
    from folio.research.transport import research_messages

    use_fixture_sandbox(monkeypatch)

    class BatchRuntime(ScriptedResearchRuntime):
        async def complete(self, name, context, schema, session):
            if name != "research_step":
                return await super().complete(name, context, schema, session)
            self.calls.append((name, context))
            if context["step"] == 0:
                first = research_call("fs.read", {"path": "requirements.docx"})
                second = research_call("fs.read", {"path": "fields.xlsx", "sheet": "Fields"})
                return RuntimeResult("分别读取业务要求和字段清单。", tool_calls=first.tool_calls + second.tool_calls)
            return research_call("finish", research_notes=self.research_notes(context))

    runtime = BatchRuntime(mode="read_only")
    job_id = store.create(make_request(research_inputs))
    job = await Engine(store, runtime).run(job_id)
    assert job["status"] == "draft"
    calls = [context for name, context in runtime.calls if name == "research_step"]
    assert len(calls) == 2
    observations = calls[1]["observations"]
    assert len(observations) == 2
    assert all(item["result"]["evidence"] for item in observations)
    messages = research_messages(calls[1], {}, task="research_step")[2:]
    assert len(messages[0]["tool_calls"]) == 2
    assert messages[0]["content"] == "分别读取业务要求和字段清单。"
    assert [item["tool_call_id"] for item in messages[1:]] == [
        item["id"] for item in messages[0]["tool_calls"]
    ]
    assert ledger_for(store, job_id).usage("fields")["calls"] == 2


async def test_large_coverage_report_is_preserved_without_overflowing_finding(
    store, research_inputs, monkeypatch,
):
    from folio.research.artifact_coverage import CoverageManager

    use_fixture_sandbox(monkeypatch)
    original_report = CoverageManager.report

    def long_report(self, coverage_id, states=()):
        report = original_report(self, coverage_id, states)
        if report["missing"]:
            report["missing"] = [dict(item, reason="Uncovered original requirement. " * 200)
                                 for item in report["missing"]]
        return report

    monkeypatch.setattr(CoverageManager, "report", long_report)
    runtime = ScriptedResearchRuntime(mode="coverage", cover_all=False)
    engine = Engine(store, runtime)
    job_id = store.create(make_request(research_inputs))
    job = await engine.run(job_id)
    assert job["status"] == "needs_confirmation"
    work = store.work_states(job_id)["fields"]
    task = engine.work_plan.tasks[0]
    states = store.states(job_id)
    issues = research_findings(engine, task, work, states)
    coverage = [issue for issue in issues if issue.kind == "coverage"]
    assert coverage and all(len(issue.description) < 500 for issue in coverage)
    folder = Path(work.research["task_root"]) / ".research" / "coverage_reports"
    cid = work.research["coverage_ids"][0]
    report_path = folder / (cid + ".json")
    saved = json.loads(report_path.read_text())
    assert all(len(item["reason"]) > 3000 for item in saved["missing"])
    assert len(saved["missing"]) == saved["total"] - saved["covered"]
    assert saved["section_revisions"] == {state.section_id: state.revision for state in states}
    before = report_path.stat().st_mtime_ns
    research_findings(engine, task, work, states, update_coverage=False)
    assert report_path.stat().st_mtime_ns == before


def test_finding_description_limit_is_3500():
    from pydantic import ValidationError

    from folio.models import Finding

    data = {"kind": "coverage", "key": "coverage", "section_ids": ["s1"]}
    assert len(Finding(**data, description="x" * 3500).description) == 3500
    with pytest.raises(ValidationError):
        Finding(**data, description="x" * 3501)
