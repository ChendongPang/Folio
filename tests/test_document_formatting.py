import asyncio
import json

import pytest
from docx import Document
from pydantic import ValidationError

from folio.document.formatting import (
    FormattingResult,
    LayoutPatch,
    apply_patches,
    collect_targets,
    protected_signature,
    replay_formatting,
    target_view,
)
from folio.document.reader import visible_text, xpath
from folio.document.word import render
from folio.models import Block, DocumentPlan, Extraction, Section, SectionContent, SectionState
from folio.orchestration.formatting import format_document
from folio.orchestration.work import WorkPlan, WorkTask


@pytest.fixture
def rendered(tmp_path, request):
    plan = DocumentPlan(
        sections=[
            Section(
                section_id=f"section_{i}",
                node_id=f"outline_{i}",
                title=f"任务{i}",
                level=1,
                display_number=str(i),
                instruction="业务说明",
                expected_table_headers=["参数", "说明"],
            )
            for i in range(1, getattr(request, "param", 2) + 1)
        ]
    )
    states = [
        SectionState(
            section_id=section.section_id,
            content=SectionContent(
                blocks=[
                    Block(block_id="body", text=f"第{i}组已有材料的正文。", evidence_ids=["e1"]),
                    Block(
                        block_id="proposal",
                        text=f"第{i}组模型补写的业务建议。",
                        origin="model_proposal",
                    ),
                    Block(
                        block_id="list", text="1. 保持这个编号列表原样。", origin="model_proposal"
                    ),
                    Block(
                        block_id="table",
                        kind="table",
                        rows=[["参数", "说明"], ["编号", f"第{i}组编号"]],
                        origin="model_proposal",
                    ),
                ]
            ),
            extraction=Extraction(summary="业务摘要"),
            evidence=[],
        )
        for i, section in enumerate(plan.sections, 1)
    ]
    work = WorkPlan(
        tasks=[
            WorkTask(
                task_id=f"work_{i}",
                title=f"业务{i}",
                objective="完成业务说明",
                questions=["业务规则"],
                completion_criteria=["覆盖业务内容"],
                rationale="业务关联",
                section_ids=[section.section_id],
            )
            for i, section in enumerate(plan.sections, 1)
        ]
    )
    output = tmp_path / "result.docx"
    render(
        output,
        plan,
        states,
        title="封面标题",
        draft=True,
        cover={"document_title": "封面标题", "document_date": "2026-10-09"},
        defer_pagination=True,
    )
    return output, plan, states, work


def test_cosmetic_patch_preserves_text_red_color_fonts_headings_cover_and_lists(rendered):
    path, plan, states, _ = rendered
    doc = Document(path)
    targets = collect_targets(doc, plan, states)
    assert targets and all("封面" not in visible_text(t.element) for t in targets)
    assert not any("编号列表" in visible_text(t.element) for t in targets)
    signature = protected_signature(doc, targets)
    text = visible_text(doc.element)
    colors = xpath(doc.element, ".//w:rPr/w:color/@w:val")
    families = xpath(doc.element, ".//w:rPr/w:rFonts")
    family_xml = [node.xml for node in families]
    target = next(
        t for t in targets if t.kind == "paragraph" and "业务建议" in visible_text(t.element)
    )
    baseline_size = target.baseline["font_size_pt"]
    assert apply_patches(
        doc,
        targets,
        [
            LayoutPatch(
                target_id=target.target_id,
                alignment="right",
                font_delta_pt=1,
                space_after_pt=6,
            )
        ],
    )
    assert target_view(doc, target)["current_layout"]["font_size_pt"] == baseline_size + 1
    assert not apply_patches(
        doc, targets, [LayoutPatch(target_id=target.target_id, font_delta_pt=1)]
    )
    assert target_view(doc, target)["current_layout"]["font_size_pt"] == baseline_size + 1
    assert visible_text(doc.element) == text
    assert xpath(doc.element, ".//w:rPr/w:color/@w:val") == colors
    assert [node.xml for node in xpath(doc.element, ".//w:rPr/w:rFonts")] == family_xml
    assert protected_signature(doc, targets) == signature
    assert apply_patches(doc, targets, [LayoutPatch(target_id=target.target_id, font_delta_pt=0)])
    assert target_view(doc, target)["current_layout"]["font_size_pt"] == baseline_size


def test_invalid_target_property_or_content_change_is_rejected_before_edits(rendered):
    path, plan, states, _ = rendered
    doc = Document(path)
    targets = collect_targets(doc, plan, states)
    paragraph = next(t for t in targets if t.kind == "paragraph")
    before = doc.element.xml
    with pytest.raises(ValidationError):
        LayoutPatch.model_validate({"target_id": paragraph.target_id, "text": "替换业务内容"})
    with pytest.raises(ValueError, match="unknown"):
        apply_patches(doc, targets, [LayoutPatch(target_id="cover", alignment="center")])
    with pytest.raises(ValueError, match="not allowed"):
        apply_patches(
            doc, targets, [LayoutPatch(target_id=paragraph.target_id, repeat_table_header=True)]
        )
    with pytest.raises(ValueError, match="duplicate"):
        apply_patches(
            doc, targets, [LayoutPatch(target_id=paragraph.target_id, alignment="right")] * 2
        )
    assert doc.element.xml == before


def test_table_row_layout_changes_keep_every_cell_and_protected_property(rendered):
    path, plan, states, _ = rendered
    doc = Document(path)
    targets = collect_targets(doc, plan, states)
    table = next(t for t in targets if t.kind == "table")
    signature = protected_signature(doc, targets)
    cells = [cell.text for row in doc.tables[0].rows for cell in row.cells]
    apply_patches(
        doc,
        targets,
        [
            LayoutPatch(
                target_id=table.target_id,
                repeat_table_header=True,
                prevent_row_split=True,
            )
        ],
    )
    assert xpath(table.element, "./w:tr[1]/w:trPr/w:tblHeader")
    assert all(xpath(row, "./w:trPr/w:cantSplit") for row in xpath(table.element, "./w:tr"))
    assert [cell.text for row in doc.tables[0].rows for cell in row.cells] == cells
    assert protected_signature(doc, targets) == signature


async def test_groups_are_separate_and_second_round_reads_applied_layout(rendered, tmp_path):
    path, plan, states, work = rendered
    seen = []

    async def call(task, context, model):
        assert task == "format_document"
        group = context["work_task"]["task_id"]
        ids = {s["section_id"] for s in context["sections"]}
        assert len(ids) == 1
        assert all(t["section_id"] in ids for t in context["targets"])
        target = next(t for t in context["targets"] if t["kind"] == "paragraph")
        seen.append((group, context["iteration"], target["current_layout"]["alignment"]))
        return FormattingResult(
            done=False, patches=[LayoutPatch(target_id=target["target_id"], alignment="right")]
        )

    report = await format_document(path, plan, states, work, tmp_path, call, lambda _: None)
    assert seen == [
        ("work_1", 1, "both"),
        ("work_1", 2, "right"),
        ("work_2", 1, "both"),
        ("work_2", 2, "right"),
    ]
    assert report["changed"]
    assert all(group["calls"] == 2 for group in report["groups"])
    assert all(len(group["iterations"]) == 1 for group in report["groups"])


async def test_group_has_at_most_three_calls_and_formatting_is_durable(rendered, tmp_path):
    path, plan, states, work = rendered
    baseline = path.read_bytes()
    calls = []

    async def call(task, context, model):
        calls.append(context)
        target = next(t for t in context["targets"] if t["kind"] == "paragraph")
        return FormattingResult(
            done=False,
            patches=[
                LayoutPatch(
                    target_id=target["target_id"],
                    alignment=["right", "center", "left"][context["iteration"] - 1],
                )
            ],
        )

    report = await format_document(path, plan, states, work, tmp_path, call, lambda _: None)
    assert len(calls) == 6
    assert all(group["calls"] == 3 for group in report["groups"])
    final = Document(path).element.xml
    path.write_bytes(baseline)

    async def forbidden(*args):
        raise AssertionError("A saved formatting plan must replay without another model call")

    replayed = await format_document(path, plan, states, work, tmp_path, forbidden, lambda _: None)
    assert replayed["status"] == "replayed"
    assert Document(path).element.xml == final
    path.write_bytes(baseline)
    assert (
        replay_formatting(path, plan, states, tmp_path / "formatting.json")["status"] == "replayed"
    )
    assert Document(path).element.xml == final


async def test_replaying_on_formatted_document_does_not_accumulate_font_delta(rendered, tmp_path):
    path, plan, states, work = rendered

    async def call(task, context, model):
        target = next(t for t in context["targets"] if t["kind"] == "paragraph")
        return FormattingResult(
            done=True, patches=[LayoutPatch(target_id=target["target_id"], font_delta_pt=1)]
        )

    await format_document(path, plan, states, work, tmp_path, call, lambda _: None)
    final_xml = Document(path).element.xml
    for _ in range(2):
        replayed = replay_formatting(path, plan, states, tmp_path / "formatting.json")
        assert replayed["status"] == "replayed" and not replayed["changed"]
        assert Document(path).element.xml == final_xml

    async def forbidden(*args):
        raise AssertionError("Already formatted input should replay without another model call")

    result = await format_document(path, plan, states, work, tmp_path, forbidden, lambda _: None)
    assert result["status"] == "replayed" and not result["changed"]
    assert Document(path).element.xml == final_xml


@pytest.mark.parametrize("stage_limit", [False, True])
async def test_timeout_keeps_earlier_safe_changes_and_does_not_block(
    rendered, tmp_path, stage_limit
):
    path, plan, states, work = rendered
    calls = 0

    async def call(task, context, model):
        nonlocal calls
        calls += 1
        if calls > 1:
            await asyncio.sleep(1)
        target = next(t for t in context["targets"] if t["kind"] == "paragraph")
        return FormattingResult(
            done=False, patches=[LayoutPatch(target_id=target["target_id"], alignment="right")]
        )

    report = await format_document(
        path,
        plan,
        states,
        work,
        tmp_path,
        call,
        lambda _: None,
        group_seconds=1 if stage_limit else 0.04,
        stage_seconds=0.04 if stage_limit else 1,
    )
    assert report["status"] in {"partial", "timeout"}
    assert report["changed"]
    doc = Document(path)
    target = next(t for t in collect_targets(doc, plan, states) if t.kind == "paragraph")
    assert target_view(doc, target)["current_layout"]["alignment"] == "right"
    assert json.loads((tmp_path / "formatting.json").read_text())["changed"]


async def test_invalid_patch_is_optional_and_other_business_groups_continue(rendered, tmp_path):
    path, plan, states, work = rendered
    baseline = path.read_bytes()

    async def call(task, context, model):
        return FormattingResult(
            done=True, patches=[LayoutPatch(target_id="not-in-this-group", alignment="left")]
        )

    report = await format_document(path, plan, states, work, tmp_path, call, lambda _: None)
    assert report["status"] == "partial" and not report["changed"]
    assert len(report["groups"]) == 2 and all(g["status"] == "failed" for g in report["groups"])
    assert all("unknown" in g["error"] for g in report["groups"])
    assert path.read_bytes() == baseline


async def test_oversized_groups_use_bounded_pages_and_shared_three_call_budget(
    rendered, tmp_path, monkeypatch
):
    from folio.orchestration import formatting

    path, plan, states, work = rendered
    monkeypatch.setattr(formatting, "PAGE_TARGETS", 1)
    calls = []

    async def call(task, context, model):
        assert len(context["targets"]) == 1
        calls.append(context)
        target = context["targets"][0]
        property_name = "prevent_row_split" if target["kind"] == "table" else "alignment"
        value = True if target["kind"] == "table" else "right"
        return FormattingResult(
            done=True,
            patches=[
                LayoutPatch.model_validate(
                    {
                        "target_id": target["target_id"],
                        property_name: value,
                    }
                )
            ],
        )

    report = await format_document(path, plan, states, work, tmp_path, call, lambda _: None)
    assert len(calls) <= 6
    assert all(group["calls"] <= 3 for group in report["groups"])
    assert all(group["remaining_targets"] > 0 for group in report["groups"])


async def test_empty_single_page_finishes_and_demo_does_not_call_model(rendered, tmp_path):
    path, plan, states, work = rendered
    calls = []

    async def call(task, context, model):
        calls.append(context)
        return FormattingResult(done=True, patches=[])

    result = await format_document(path, plan, states, work, tmp_path, call, lambda _: None)
    assert len(calls) == 2 and not result["changed"]
    calls.clear()
    result = await format_document(
        path, plan, states, work, tmp_path, call, lambda _: None, enabled=False
    )
    assert not calls and result["reason"] == "offline_demo"


@pytest.mark.parametrize("rendered", [5], indirect=True)
async def test_four_groups_overlap_but_own_iterations_and_protected_content_are_serial(
    rendered, tmp_path
):
    path, plan, states, work = rendered
    document = Document(path)
    signature = protected_signature(document, collect_targets(document, plan, states))
    started, release = asyncio.Event(), asyncio.Event()
    active = peak = 0
    seen = []

    async def call(task, context, model):
        nonlocal active, peak
        active += 1
        peak = max(peak, active)
        group, iteration = context["work_task"]["task_id"], context["iteration"]
        target = next(t for t in context["targets"] if t["kind"] == "paragraph")
        seen.append((group, iteration, target["current_layout"]["alignment"]))
        if active == 4:
            started.set()
        try:
            await release.wait()
            await asyncio.sleep(0.01)
            return FormattingResult(
                done=False, patches=[LayoutPatch(target_id=target["target_id"], alignment="right")]
            )
        finally:
            active -= 1

    runner = asyncio.create_task(
        format_document(path, plan, states, work, tmp_path, call, lambda _: None)
    )
    await asyncio.wait_for(started.wait(), 2)
    assert {item[0] for item in seen} == {f"work_{i}" for i in range(1, 5)}
    release.set()
    report = await runner
    assert peak == 4 and active == 0
    for i in range(1, 6):
        assert [item[1:] for item in seen if item[0] == f"work_{i}"] == [(1, "both"), (2, "right")]
    assert [group["task_id"] for group in report["groups"]] == [f"work_{i}" for i in range(1, 6)]
    assert all(group["calls"] == 2 for group in report["groups"])
    document = Document(path)
    assert protected_signature(document, collect_targets(document, plan, states)) == signature


async def test_waiting_for_group_slot_does_not_consume_group_timeout(
    rendered, tmp_path, monkeypatch
):
    from folio.orchestration import formatting

    monkeypatch.setattr(formatting, "MAX_PARALLEL_GROUPS", 1)
    path, plan, states, work = rendered

    async def call(*args):
        await asyncio.sleep(0.12)
        return FormattingResult(done=True, patches=[])

    report = await format_document(
        path,
        plan,
        states,
        work,
        tmp_path,
        call,
        lambda _: None,
        group_seconds=0.2,
        stage_seconds=2,
    )
    assert report["status"] == "complete"
    assert all(group["status"] == "unchanged" for group in report["groups"])


async def test_empty_patch_page_continues_to_later_pages_with_three_call_limit(
    rendered, tmp_path, monkeypatch
):
    from folio.orchestration import formatting

    monkeypatch.setattr(formatting, "PAGE_TARGETS", 1)
    path, plan, states, work = rendered
    calls = []

    async def call(task, context, model):
        calls.append((context["work_task"]["task_id"], context["targets"][0]["target_id"]))
        return FormattingResult(done=True, patches=[])

    report = await format_document(path, plan, states, work, tmp_path, call, lambda _: None)
    assert len(calls) == len(set(calls)) == 6
    assert not report["changed"]
    assert all(
        group["calls"] == 3 and group["status"] == "iteration_limit" for group in report["groups"]
    )


@pytest.mark.parametrize("external_cancel", [False, True])
async def test_formatting_timeout_and_external_cancel_drain_all_workers(
    rendered, tmp_path, external_cancel
):
    path, plan, states, work = rendered
    baseline = path.read_bytes()
    started = asyncio.Event()
    active = 0
    drained = []

    async def call(task, context, model):
        nonlocal active
        active += 1
        if active == 2:
            started.set()
        try:
            await asyncio.Future()
        finally:
            await asyncio.sleep(0.01)
            active -= 1
            drained.append(context["work_task"]["task_id"])

    runner = asyncio.create_task(
        format_document(
            path,
            plan,
            states,
            work,
            tmp_path,
            call,
            lambda _: None,
            stage_seconds=2 if external_cancel else 0.1,
        )
    )
    await asyncio.wait_for(started.wait(), 1)
    if external_cancel:
        runner.cancel()
        with pytest.raises(asyncio.CancelledError):
            await runner
    else:
        assert (await runner)["status"] == "timeout"
    assert active == 0 and len(drained) == 2
    manifest = tmp_path / "formatting.json"
    saved = manifest.read_bytes()
    report = json.loads(saved)
    assert report["status"] == ("cancelled" if external_cancel else "timeout")
    assert all(group["status"] == "cancelled" for group in report["groups"])
    await asyncio.sleep(0.02)
    assert path.read_bytes() == baseline and manifest.read_bytes() == saved


async def test_protected_failure_cancels_other_groups_and_discards_all_edits(
    rendered, tmp_path, monkeypatch
):
    from folio.orchestration import formatting

    path, plan, states, work = rendered
    baseline = path.read_bytes()
    both_started = asyncio.Event()
    active = 0
    cancelled = []

    async def call(task, context, model):
        nonlocal active
        active += 1
        if active == 2:
            both_started.set()
        group = context["work_task"]["task_id"]
        try:
            if group == "work_2":
                await asyncio.Future()
            await both_started.wait()
            target = next(t for t in context["targets"] if t["kind"] == "paragraph")
            return FormattingResult(
                done=True, patches=[LayoutPatch(target_id=target["target_id"], alignment="right")]
            )
        except asyncio.CancelledError:
            cancelled.append(group)
            raise
        finally:
            active -= 1

    def corrupt(document, targets, patches):
        xpath(targets[0].element, ".//w:t")[0].text = "Unexpected content edit"
        raise ValueError("Formatting changed protected document content or typography")

    monkeypatch.setattr(formatting, "apply_patches", corrupt)
    report = await format_document(path, plan, states, work, tmp_path, call, lambda _: None)
    assert active == 0 and cancelled == ["work_2"]
    assert report["status"] == "failed" and not report["changed"]
    assert path.read_bytes() == baseline
    assert (
        replay_formatting(path, plan, states, tmp_path / "formatting.json")["status"] == "discarded"
    )


async def test_engine_final_export_survives_invalid_formatting_without_model_retry(
    rendered, inputs, store
):
    from folio.document.references import ReferenceDocuments
    from folio.models import JobInput, Limits
    from folio.orchestration.engine import Engine
    from folio.runtime.base import RuntimeResult

    _, plan, states, work = rendered

    class InvalidFormatter:
        calls = 0

        async def complete(self, task, context, schema, session_key):
            assert task == "format_document"
            self.calls += 1
            return RuntimeResult("invalid json")

    runtime = InvalidFormatter()
    job = store.create(JobInput(**inputs, runtime="llm", limits=Limits(call_attempts=5)))
    engine = Engine(store, runtime)
    engine.job_id = job
    engine.request = JobInput.model_validate(store.get(job)["input"])
    engine.limits = engine.request.limits
    engine.plan, engine.work_plan = plan, work
    engine.references = ReferenceDocuments([])
    store.save_plan(job, plan)
    store.save_work_plan(job, work)
    for state in states:
        store.save_section(job, state)
    await engine.finish(states, draft=True)
    final = store.get(job)
    assert final["status"] == "draft"
    assert runtime.calls == 2  # One failed attempt per business group; no automatic retries.
    report = json.loads((store.root / "jobs" / job / "formatting.json").read_text())
    assert report["status"] == "partial" and not report["changed"]
    assert all(group["status"] == "failed" and group["error_type"] for group in report["groups"])
    assert "第1组已有材料的正文。" in visible_text(Document(final["artifact"]).element)
    with store.connect() as connection:
        rows = connection.execute(
            "SELECT attempts FROM tasks WHERE job_id=? AND task_key LIKE 'format_document:%'",
            (job,),
        ).fetchall()
    assert len(rows) == 2 and all(row[0] == 1 for row in rows)
