import json
from pathlib import Path

import pytest
from docx import Document
from pydantic import ValidationError

from folio.document.samples import sample_profile
from folio.models import JobInput
from folio.orchestration.engine import Engine
from folio.runtime.demo import DemoRuntime


def make_sample(path):
    doc = Document()
    doc.add_paragraph("秘密旧项目名称")
    doc.add_heading("外汇交易", 1)
    doc.add_paragraph("样例规则：全部业务免审批。")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "秘密字段"
    table.cell(0, 1).text = "秘密数据"
    doc.save(path)
    return path


def test_sample_is_not_a_knowledge_source(inputs, tmp_path):
    path = make_sample(tmp_path / "sample.docx")
    data = inputs | {"knowledge_path": None, "sample_paths": [str(path)]}
    with pytest.raises(ValidationError, match="Provide knowledge_path"):
        JobInput(**data)
    profile = sample_profile([path])
    text = json.dumps(profile, ensure_ascii=False)
    for secret in ["秘密旧项目名称", "全部业务免审批", "秘密字段", "秘密数据", "外汇交易"]:
        assert secret not in text
    assert profile["documents"][0]["table_column_counts"] == [2]


async def test_sample_snapshot_does_not_leak_into_any_model_task(store, inputs, tmp_path):
    path = make_sample(tmp_path / "sample.docx")
    original = path.read_bytes()
    job_id = store.create(JobInput(**inputs, sample_paths=[str(path)], runtime="demo"))
    frozen = store.get(job_id)["input"]
    assert Path(frozen["sample_paths"][0]).read_bytes() == original
    path.unlink()  # The run must use the frozen example.

    class InspectRuntime(DemoRuntime):
        async def complete(self, task, context, schema, session_key):
            text = json.dumps(context, ensure_ascii=False)
            for secret in ["秘密旧项目名称", "全部业务免审批", "秘密字段", "秘密数据"]:
                assert secret not in text
            return await super().complete(task, context, schema, session_key)

    result = await Engine(store, InspectRuntime()).run(job_id)
    assert result["status"] == "draft"
    assert result["snapshot"]["source_audit"] == []
    assert all(e.source_id != "sample.docx" for s in store.states(job_id) for e in s.evidence)
