"""Completeness is scoped and visual evidence needs an independent receipt."""
import json
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject

from folio.models import Evidence
from folio.research.artifact_coverage import CoverageManager
from folio.research.artifacts import ArtifactError, derived_evidence_error
from folio.research.extraction_review import ExtractionReview, ExtractionReviewManager
from folio.research.formats import read_source
from folio.research.workspace import WorkspaceManager


def mixed_pdf(path, *, blank_last=False):
    from test_source_structure import pdf

    pdf(path, pages=2)
    writer = PdfWriter(clone_from=PdfReader(path))
    page = writer.pages[0]
    image = DecodedStreamObject()
    image.update({NameObject('/Type'): NameObject('/XObject'), NameObject('/Subtype'): NameObject('/Image'),
                  NameObject('/Width'): NumberObject(80), NameObject('/Height'): NumberObject(20),
                  NameObject('/ColorSpace'): NameObject('/DeviceRGB'), NameObject('/BitsPerComponent'): NumberObject(8)})
    image.set_data(b'\xff' * 80 * 20 * 3)
    page['/Resources'][NameObject('/XObject')] = DictionaryObject({NameObject('/Im1'): writer._add_object(image)})
    stream = DecodedStreamObject()
    stream.set_data(page.get_contents().get_data() + b' q 300 0 0 100 30 400 cm /Im1 Do Q')
    page[NameObject('/Contents')] = writer._add_object(stream)
    if blank_last:
        del writer.pages[1]['/Contents']
    writer.write(path)


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    monkeypatch.setenv('FOLIO_ARTIFACT_KEY_PATH', str(tmp_path / 'host-key'))
    sources = tmp_path / 'sources'
    sources.mkdir()
    mixed_pdf(sources / 'rules.pdf')
    return WorkspaceManager(sources, tmp_path / 'task')


def review(workspace, source_id, pages):
    manager = ExtractionReviewManager(workspace.task_root, workspace)
    prepared = manager.prepare(source_id, pages)
    images = []
    for page in pages:
        relative = f'artifacts/review-page-{page}.png'
        Image.new('RGB', (20, 20), 'white').save(workspace.task_root / relative)
        images.append(relative)
    return manager.record(prepared, approved_pages=pages, review_key='independent-vision-call',
                          reason='Original page and complete extraction agree', image_paths=images)


def test_native_title_does_not_hide_unread_pdf_image(workspace):
    document = read_source(workspace.source_root / 'rules.pdf')
    assert document.blocks[0].text.strip() == 'Page 1 body'
    assert document.metadata['unread_visual_pages'] == [1]
    assert document.metadata['unread_pages'] == [1]
    assert not document.metadata['complete']
    with pytest.raises(ArtifactError, match='unread'):
        CoverageManager(workspace.task_root, workspace).create('rules.pdf')


def test_complete_page_can_be_covered_when_another_page_has_unread_visuals(workspace):
    manager = CoverageManager(workspace.task_root, workspace)
    page = next(node for node in workspace.inspect('rules.pdf')['structure'] if node['metadata'].get('page') == 2)
    scope = manager.create('rules.pdf', node_id=page['node_id'])['scope']
    assert [item['text'] for item in scope['items']] == ['Page 2 body']
    assert manager.create('rules.pdf', 'page 2')['scope']['items'] == scope['items']
    with pytest.raises(ArtifactError):
        manager.create('rules.pdf', 'page 1')


def test_excel_image_marks_only_affected_sheet_unread(tmp_path, monkeypatch):
    monkeypatch.setenv('FOLIO_ARTIFACT_KEY_PATH', str(tmp_path / 'key'))
    image = tmp_path / 'rules.png'
    Image.new('RGB', (160, 40), 'white').save(image)
    sources = tmp_path / 'sources'
    sources.mkdir()
    book = Workbook()
    book.active.title = 'Rules'
    book.active.append(['Approval rules'])
    book.active.add_image(SheetImage(image), 'A2')
    book.create_sheet('Readable').append(['Limit', 1000])
    book.save(sources / 'rules.xlsx')
    workspace = WorkspaceManager(sources, tmp_path / 'task')
    document = workspace.source_index.load('rules.xlsx')
    assert document.metadata['unread_sheets'] == ['Rules']
    assert not document.metadata['complete']
    nodes = workspace.inspect('rules.xlsx')['structure']
    manager = CoverageManager(workspace.task_root, workspace)
    readable = next(node for node in nodes if node['title'] == 'Readable')
    assert manager.create('rules.xlsx', node_id=readable['node_id'])['scope']['items']
    with pytest.raises(ArtifactError):
        manager.create('rules.xlsx', node_id=next(node for node in nodes if node['title'] == 'Rules')['node_id'])


def test_raw_ocr_cannot_close_coverage_but_signed_visual_review_can(workspace):
    from test_research_coverage_fragments import state

    workspace.register_extraction('rules.pdf', {'blocks': [{'page': 1, 'text': 'Approval limit = 1000'}]},
                                  execution_id='ocr', image_digest='sha256:converter')
    raw = Evidence.model_validate(workspace.read('rules.pdf', page=1)['evidence'][-1])
    assert derived_evidence_error(raw)
    manager = CoverageManager(workspace.task_root, workspace)
    with pytest.raises(ArtifactError):
        manager.create('rules.pdf')
    result = review(workspace, 'rules.pdf', [1])
    verified = Evidence.model_validate(result['evidence'][0])
    assert verified.evidence_id == raw.evidence_id
    assert derived_evidence_error(verified) is None
    scope = manager.create('rules.pdf')['scope']
    assert len(scope['items']) == 3
    item = next(item for item in scope['items'] if '1000' in item['text'])
    assert manager.covering_references(scope, item, [state([verified])])
    restored = WorkspaceManager(workspace.source_root, workspace.task_root)
    assert derived_evidence_error(restored.resolve_evidence(verified)) is None
    assert CoverageManager(workspace.task_root, restored).create('rules.pdf')['scope'] == scope
    # Other tasks still have no authority to count this OCR as complete.
    other = WorkspaceManager(workspace.source_root, workspace.task_root.with_name('other'))
    with pytest.raises(ArtifactError):
        CoverageManager(other.task_root, other).create('rules.pdf')


@pytest.mark.parametrize('tamper', ['text', 'metadata', 'receipt', 'image', 'source'])
def test_review_receipt_rejects_modified_evidence_or_provenance(workspace, tamper):
    workspace.register_extraction('rules.pdf', {'blocks': [{'page': 1, 'text': 'Approval limit = 1000'}]},
                                  execution_id='ocr', image_digest='sha256:converter')
    verified = Evidence.model_validate(review(workspace, 'rules.pdf', [1])['evidence'][0])
    if tamper == 'text':
        verified.text = 'Approval limit = 9999'
    elif tamper == 'metadata':
        verified.metadata['requires_extraction_review'] = False
    elif tamper == 'receipt':
        path = Path(verified.metadata['extraction_review']['manifest_path'])
        value = json.loads(path.read_bytes())
        value['registration']['reason'] = 'self approved'
        path.write_text(json.dumps(value))
    elif tamper == 'image':
        (workspace.task_root / 'artifacts/review-page-1.png').write_bytes(b'changed')
    else:
        (workspace.source_root / 'rules.pdf').write_bytes(b'changed')
    assert derived_evidence_error(verified)


def test_independently_verified_blank_page_can_close_empty_scope(tmp_path, monkeypatch):
    monkeypatch.setenv('FOLIO_ARTIFACT_KEY_PATH', str(tmp_path / 'key'))
    sources = tmp_path / 'sources'
    sources.mkdir()
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(sources / 'blank.pdf')
    workspace = WorkspaceManager(sources, tmp_path / 'task')
    with pytest.raises(ArtifactError):
        CoverageManager(workspace.task_root, workspace).create('blank.pdf')
    review(workspace, 'blank.pdf', [1])
    manager = CoverageManager(workspace.task_root, workspace)
    record = manager.create('blank.pdf')
    assert manager.report(record['coverage_id'])['complete']
    assert record['scope']['items'] == []


def test_visual_verdict_cannot_approve_a_failed_page():
    with pytest.raises(ValueError):
        ExtractionReview.model_validate({'approved_pages': [1], 'reason': 'ok', 'pages': [
            {'page': 1, 'source_matches': True, 'complete': False, 'blank_or_decorative': False, 'reason': 'missing rules'},
        ]})


def test_office_coverage_requires_all_render_pages_and_keeps_sheet_mapping_honest(tmp_path, monkeypatch):
    from docx import Document

    monkeypatch.setenv('FOLIO_ARTIFACT_KEY_PATH', str(tmp_path / 'key'))
    sources = tmp_path / 'sources'
    sources.mkdir()
    image = tmp_path / 'rules.png'
    Image.new('RGB', (100, 60), 'white').save(image)
    word = Document()
    word.add_heading('Rules', 1)
    word.add_paragraph('Visible title')
    word.add_picture(str(image))
    word.save(sources / 'rules.docx')
    workspace = WorkspaceManager(sources, tmp_path / 'task')
    workspace.register_extraction('rules.docx', {
        'operation': 'office_ocr', 'page_count': 2,
        'blocks': [{'page': 1, 'text': 'Approval limit = 1000'}, {'page': 2, 'text': 'Audit required'}],
    }, execution_id='office', image_digest='sha256:office')
    manager = CoverageManager(workspace.task_root, workspace)
    review(workspace, 'rules.docx', [1])
    with pytest.raises(ArtifactError):
        manager.create('rules.docx')
    review(workspace, 'rules.docx', [2])
    texts = [item['text'] for item in manager.create('rules.docx')['scope']['items']]
    assert 'Approval limit = 1000' in texts and 'Audit required' in texts
    heading = next(node for node in workspace.inspect('rules.docx')['structure'] if node['kind'] == 'section')
    # Whole rendered coverage does not invent a mapping of image rules to this chapter.
    with pytest.raises(ArtifactError):
        manager.create('rules.docx', node_id=heading['node_id'])


def test_repeated_independent_reviews_do_not_invalidate_authentic_coverage(workspace):
    from test_research_coverage_fragments import state

    workspace.register_extraction('rules.pdf', {'blocks': [{'page': 1, 'text': 'Approval limit = 1000'}]},
                                  execution_id='ocr', image_digest='sha256:converter')
    first = Evidence.model_validate(review(workspace, 'rules.pdf', [1])['evidence'][0])
    manager = CoverageManager(workspace.task_root, workspace)
    first_scope = manager.create('rules.pdf')['scope']
    reviewer = ExtractionReviewManager(workspace.task_root, workspace)
    second = reviewer.record(reviewer.prepare('rules.pdf', [1]), approved_pages=[1],
                             review_key='another-review', reason='Same original and extraction',
                             image_paths=['artifacts/review-page-1.png'])
    latest = Evidence.model_validate(second['evidence'][0])
    item = next(item for item in first_scope['items'] if '1000' in item['text'])
    assert manager.covering_references(first_scope, item, [state([latest])])
    latest_scope = manager.create('rules.pdf')['scope']
    assert manager.covering_references(latest_scope, item, [state([first])])


def test_pdf_author_marked_decorative_image_does_not_block_text_scope(tmp_path):
    path = tmp_path / 'decorative.pdf'
    mixed_pdf(path)
    writer = PdfWriter(clone_from=PdfReader(path))
    page = writer.pages[0]
    stream = DecodedStreamObject()
    stream.set_data(page.get_contents().get_data().replace(b'/Im1 Do', b'/Artifact BMC /Im1 Do EMC'))
    page[NameObject('/Contents')] = writer._add_object(stream)
    writer.write(path)
    document = read_source(path)
    assert document.metadata['complete']
    assert document.metadata['unread_visual_pages'] == []
