from copy import deepcopy

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

from folio.document.formatting import (
    LayoutPatch,
    allowed_properties,
    apply_patches,
    collect_targets,
)
from folio.document.reader import heading_level, visible_text, xpath
from folio.document.template_contract import read_template
from folio.document.template_layout import validate_template_output
from folio.document.typography import prop, resolved_format
from folio.document.word import render
from folio.models import Block, DocumentPlan, Extraction, Section, SectionContent, SectionState


@pytest.fixture
def template_case(tmp_path):
    path = tmp_path / "公司规范模板.docx"
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(13)
    prop(doc.styles["Normal"].element.get_or_add_rPr(), "rFonts", eastAsia="仿宋")
    doc.styles["Heading 1"].font.name = "Times New Roman"
    doc.styles["Heading 1"].font.size = Pt(21)
    doc.sections[0].left_margin = Cm(3.3)
    doc.sections[0].header.paragraphs[0].text = "公司受控文件"
    doc.sections[0].footer.paragraphs[0].text = "内部使用"
    doc.add_paragraph("{{document_title}}", "Title")
    doc.add_paragraph("公司固定封面文字")
    doc.add_heading("1. 项目说明", 1)
    paragraph = doc.add_paragraph("填写说明：请介绍项目背景。")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.7
    paragraph.paragraph_format.keep_together = False
    paragraph.runs[0].font.size = Pt(14)
    doc.add_paragraph("旧项目名称甲项目，旧金额 999999 元，不应进入正文。")
    doc.add_heading("2. 接口要求", 1)
    doc.add_paragraph("填写说明：请填写接口表。")
    table = doc.add_table(rows=2, cols=2)
    table.style = "Light Shading Accent 1"
    table.autofit = False
    table.columns[0].width = Cm(4)
    table.columns[1].width = Cm(10)
    for cell, value in zip(table.rows[0].cells, ["参数", "说明"], strict=True):
        cell.text = value
        cell.paragraphs[0].runs[0].font.size = Pt(12)
    for cell, value in zip(table.rows[1].cells, ["旧接口", "示例数据必须删除"], strict=True):
        cell.text = value
        cell.paragraphs[0].runs[0].font.size = Pt(10)
    prop(table.rows[0]._tr.get_or_add_trPr(), "tblHeader")
    doc.save(path)
    contract = read_template(path)
    plan = DocumentPlan(
        template=contract,
        sections=[
            Section(
                section_id=f"sec_{i}",
                node_id=f"node_{i}",
                title=section.title,
                level=section.level,
                display_number=str(i),
                instruction=section.instruction,
                template_section_id=section.section_id,
                required_tables=section.tables,
                expected_table_headers=section.tables[0].headers if section.tables else [],
            )
            for i, section in enumerate(contract.sections, 1)
        ],
    )
    states = [
        SectionState(
            section_id="sec_1",
            content=SectionContent(
                blocks=[Block(block_id="body", text="新项目仅提供费用报销服务。")]
            ),
            extraction=Extraction(summary="新项目"),
            evidence=[],
        ),
        SectionState(
            section_id="sec_2",
            content=SectionContent(
                blocks=[
                    Block(
                        block_id="table",
                        kind="table",
                        rows=[["参数", "说明"], ["request_id", "请求唯一标识"]],
                    )
                ]
            ),
            extraction=Extraction(summary="接口"),
            evidence=[],
        ),
    ]
    return path, plan, states


def test_render_fills_template_preserving_cover_page_styles_and_tables(template_case, tmp_path):
    source, plan, states = template_case
    target = tmp_path / "result.docx"
    source_bytes = source.read_bytes()
    original_states = deepcopy([state.model_dump() for state in states])
    render(target, plan, states, title="新项目需求", draft=False, defer_pagination=True)
    doc = Document(target)
    text = visible_text(doc.element)
    assert "新项目需求" in text and "公司固定封面文字" in text
    assert "999999" not in text and "旧项目名称" not in text and "填写说明" not in text
    assert "示例数据必须删除" not in text and "旧接口" not in text
    assert doc.sections[0].left_margin == Document(source).sections[0].left_margin
    assert doc.sections[0].header.paragraphs[0].text == "公司受控文件"
    assert doc.sections[0].footer.paragraphs[0].text == "内部使用"
    assert [paragraph.text for paragraph in doc.paragraphs if heading_level(paragraph._p, doc)] == [
        "1. 项目说明",
        "2. 接口要求",
    ]
    body = next(paragraph for paragraph in doc.paragraphs if "新项目仅" in paragraph.text)
    properties, font = resolved_format(doc, body._p)
    assert xpath(properties, "./w:jc/@w:val") == ["both"]
    assert xpath(properties, "./w:spacing/@w:line") == ["408"]
    assert xpath(font, "./w:sz/@w:val") == ["28"]
    assert xpath(properties, "./w:keepLines/@w:val") == ["0"]
    assert [[cell.text for cell in row.cells] for row in doc.tables[0].rows] == [
        ["参数", "说明"],
        ["request_id", "请求唯一标识"],
    ]
    assert doc.tables[0].style.name == "Light Shading Accent 1"
    assert doc.tables[0].rows[1].cells[0].paragraphs[0].runs[0].font.size == Pt(10)
    assert [state.model_dump() for state in states] == original_states
    assert source.read_bytes() == source_bytes
    validate_template_output(target, plan)


def test_template_preview_retains_required_tables_but_never_old_examples(template_case, tmp_path):
    _, plan, states = template_case
    states[1].content.blocks = [Block(block_id="pending", text="尚未生成正文。")]
    target = tmp_path / "preview.docx"
    render(target, plan, states, title="预览", draft=True, preview=True, defer_pagination=True)
    doc = Document(target)
    assert doc.tables[0].rows[1].cells[0].text == "待填写"
    assert "旧接口" not in visible_text(doc.element)
    with pytest.raises(ValueError, match="required template table"):
        render(
            tmp_path / "final.docx", plan, states, title="最终", draft=False, defer_pagination=True
        )


def test_template_hash_prevents_reusing_a_changed_file(template_case, tmp_path):
    source, plan, states = template_case
    source.write_bytes(source.read_bytes() + b"changed")
    with pytest.raises(ValueError, match="template has changed"):
        render(
            tmp_path / "result.docx", plan, states, title="测试", draft=False, defer_pagination=True
        )


def test_template_blocks_formatter_from_changing_font_spacing_or_explicit_pagination(
    template_case, tmp_path
):
    _, plan, states = template_case
    target = tmp_path / "result.docx"
    render(target, plan, states, title="测试", draft=False, defer_pagination=True)
    doc = Document(target)
    targets = collect_targets(doc, plan, states)
    paragraph = next(item for item in targets if item.kind == "paragraph")
    assert "keep_together" not in allowed_properties(paragraph)
    assert "keep_with_next" in allowed_properties(paragraph)
    for field, value in [
        ("font_delta_pt", 1),
        ("alignment", "left"),
        ("space_after_pt", 12),
        ("keep_together", True),
    ]:
        with pytest.raises(ValueError, match="not allowed"):
            apply_patches(
                doc, targets, [LayoutPatch(target_id=paragraph.target_id, **{field: value})]
            )
    apply_patches(doc, targets, [LayoutPatch(target_id=paragraph.target_id, keep_with_next=True)])
    doc.save(target)
    validate_template_output(target, plan)


@pytest.mark.parametrize(
    "mutation",
    [
        "font",
        "page",
        "header",
        "heading",
        "table",
        "body_font",
        "body_spacing",
        "cell_font",
        "order",
    ],
)
def test_final_template_guard_detects_post_render_corruption(template_case, tmp_path, mutation):
    _, plan, states = template_case
    target = tmp_path / "result.docx"
    render(target, plan, states, title="测试", draft=False, defer_pagination=True)
    doc = Document(target)
    if mutation == "font":
        doc.styles["Normal"].font.size = Pt(11)
    elif mutation == "page":
        doc.sections[0].left_margin = Cm(2)
    elif mutation == "header":
        doc.sections[0].header.paragraphs[0].text = "被覆盖的页眉"
    elif mutation == "heading":
        next(paragraph for paragraph in doc.paragraphs if heading_level(paragraph._p, doc)).runs[
            0
        ].font.size = Pt(10)
    elif mutation == "table":
        prop(doc.tables[0]._tbl.tblPr, "tblW", w=10, type="pct")
    elif mutation == "body_font":
        next(paragraph for paragraph in doc.paragraphs if "新项目仅" in paragraph.text).runs[
            0
        ].font.size = Pt(11)
    elif mutation == "body_spacing":
        next(
            paragraph for paragraph in doc.paragraphs if "新项目仅" in paragraph.text
        ).paragraph_format.line_spacing = 1.0
    elif mutation == "cell_font":
        doc.tables[0].rows[1].cells[0].paragraphs[0].runs[0].font.size = Pt(11)
    else:
        headings = [paragraph for paragraph in doc.paragraphs if heading_level(paragraph._p, doc)]
        headings[0]._p.addprevious(headings[1]._p)
    doc.save(target)
    with pytest.raises(ValueError, match="[Tt]emplate"):
        validate_template_output(target, plan)


def test_template_preserves_multiple_physical_header_rows(template_case, tmp_path):
    source, plan, states = template_case
    doc = Document(source)
    table = doc.tables[0]
    row = deepcopy(table.rows[0]._tr)
    table.rows[0]._tr.addnext(row)
    for cell in table.rows[0].cells:
        cell.text = "接口"
    doc.save(source)
    plan.template = read_template(source)
    definition = plan.template.sections[1].tables[0]
    assert definition.header_rows == 2
    plan.sections[1].required_tables = [definition]
    plan.sections[1].expected_table_headers = definition.headers
    states[1].content.blocks[0].rows[0] = definition.headers
    target = tmp_path / "result.docx"
    render(target, plan, states, title="测试", draft=False, defer_pagination=True)
    result = Document(target).tables[0]
    assert [[cell.text for cell in row.cells] for row in result.rows] == [
        ["接口", "接口"],
        ["参数", "说明"],
        ["request_id", "请求唯一标识"],
    ]
    assert xpath(result._tbl, "./w:tr[1]/w:trPr/w:tblHeader")
    assert xpath(result._tbl, "./w:tr[2]/w:trPr/w:tblHeader")
    assert not xpath(result._tbl, "./w:tr[3]/w:trPr/w:tblHeader")
    assert qn("w:tbl") == result._tbl.tag


def test_header_only_merged_template_gets_unmerged_data_cells(template_case, tmp_path):
    source, plan, states = template_case
    doc = Document(source)
    table = doc.tables[0]
    table._tbl.remove(table.rows[1]._tr)
    table.rows[0].cells[0].merge(table.rows[0].cells[1]).text = "接口字段"
    doc.save(source)
    plan.template = read_template(source)
    definition = plan.template.sections[1].tables[0]
    plan.sections[1].required_tables = [definition]
    plan.sections[1].expected_table_headers = definition.headers
    states[1].content.blocks[0].rows[0] = definition.headers
    target = tmp_path / "result.docx"
    render(target, plan, states, title="测试", draft=False, defer_pagination=True)
    result = Document(target).tables[0]
    assert len(xpath(result.rows[0]._tr, "./w:tc")) == 1
    assert len(xpath(result.rows[1]._tr, "./w:tc")) == 2
    assert [cell.text for cell in result.rows[1].cells] == ["request_id", "请求唯一标识"]


def test_cover_placeholder_preserves_artwork_in_same_paragraph(template_case, tmp_path):
    source, plan, states = template_case
    doc = Document(source)
    image = tmp_path / "logo.png"
    Image.new("RGB", (4, 4), "black").save(image)
    paragraph = doc.paragraphs[0]
    paragraph.clear()
    paragraph.add_run("{{document_")
    paragraph.add_run().add_picture(str(image), width=Cm(1))
    paragraph.add_run("title}}")
    doc.save(source)
    plan.template = read_template(source)
    target = tmp_path / "result.docx"
    render(target, plan, states, title="新项目名称", draft=False, defer_pagination=True)
    result = Document(target)
    assert len(result.inline_shapes) == 1
    assert result.paragraphs[0].text == "新项目名称"


def test_section_break_and_distinct_page_settings_survive_body_replacement(template_case, tmp_path):
    source, plan, states = template_case
    doc = Document(source)
    next_heading = next(
        paragraph._p for paragraph in doc.paragraphs if paragraph.text == "2. 接口要求"
    )
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    boundary = doc.paragraphs[-1]._p
    next_heading.addprevious(boundary)
    doc.sections[1].left_margin = Cm(2)
    doc.sections[1].right_margin = Cm(2.4)
    doc.save(source)
    plan.template = read_template(source)
    target = tmp_path / "result.docx"
    render(target, plan, states, title="测试", draft=False, defer_pagination=True)
    original, result = Document(source), Document(target)
    assert len(result.sections) == 2
    assert [section._sectPr.xml for section in result.sections] == [
        section._sectPr.xml for section in original.sections
    ]


def test_two_required_tables_with_same_headers_each_keep_own_layout(template_case, tmp_path):
    source, plan, states = template_case
    doc = Document(source)
    duplicate = deepcopy(doc.tables[0]._tbl)
    doc.tables[0]._tbl.addnext(duplicate)
    doc.tables[1].style = "Table Grid"
    doc.save(source)
    plan.template = read_template(source)
    plan.sections[1].required_tables = plan.template.sections[1].tables
    states[1].content.blocks.append(
        Block(block_id="second", kind="table", rows=[["参数", "说明"], ["amount", "金额，分"]])
    )
    target = tmp_path / "result.docx"
    render(target, plan, states, title="测试", draft=False, defer_pagination=True)
    result = Document(target)
    assert len(result.tables) == 2
    assert [table.rows[1].cells[0].text for table in result.tables] == ["request_id", "amount"]
    assert [table.style.name for table in result.tables] == ["Light Shading Accent 1", "Table Grid"]


def test_additional_unapproved_headings_and_automatic_appendix_are_rejected(
    template_case, tmp_path
):
    _, plan, states = template_case
    target = tmp_path / "result.docx"
    with pytest.raises(ValueError, match="references appendix"):
        render(
            target,
            plan,
            states,
            title="测试",
            draft=False,
            include_evidence_appendix=True,
            defer_pagination=True,
        )
    render(target, plan, states, title="测试", draft=False, defer_pagination=True)
    doc = Document(target)
    doc.add_heading("额外的章节", 1)
    doc.save(target)
    with pytest.raises(ValueError, match="template does not permit"):
        validate_template_output(target, plan)


def test_pagination_is_followed_by_template_validation(template_case, tmp_path, monkeypatch):
    _, plan, states = template_case

    def corrupt(path, **kwargs):
        doc = Document(path)
        doc.styles["Normal"].font.size = Pt(9)
        doc.save(path)
        return {"status": "paginated"}

    monkeypatch.setattr("folio.document.word.paginate", corrupt)
    with pytest.raises(ValueError, match="template style"):
        render(tmp_path / "result.docx", plan, states, title="测试", draft=False)


def test_blank_required_table_rows_are_not_completed_data(template_case, tmp_path):
    _, plan, states = template_case
    target = tmp_path / "result.docx"
    render(target, plan, states, title="测试", draft=False, defer_pagination=True)
    doc = Document(target)
    for cell in doc.tables[0].rows[1].cells:
        # Retain cell formatting so the failure specifically identifies missing data.
        for run in cell.paragraphs[0].runs:
            run.text = "  "
    doc.save(target)
    with pytest.raises(ValueError, match="insufficient valid data rows"):
        validate_template_output(target, plan)
    states[1].content.blocks[0].rows[1] = ["  ", "\n"]
    with pytest.raises(ValueError, match="insufficient valid data rows"):
        render(target, plan, states, title="测试", draft=False, defer_pagination=True)
