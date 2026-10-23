"""The synthetic acceptance corpus must be uploadable and have trustworthy oracles."""

import importlib
import json
from collections import defaultdict
from datetime import datetime
from decimal import ROUND_CEILING, Decimal
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from native_calls import native_call, research_call
from pypdf import PdfReader

from folio.models import JobInput, Limits
from folio.orchestration.engine import Engine
from folio.research.contracts import ResearchConfig
from folio.research.formats import read_source
from folio.research.structure import build_structure
from folio.runtime.base import RuntimeResult
from folio.runtime.demo import DemoRuntime

ROOT = Path(__file__).resolve().parents[1] / "examples" / "acceptance_cases"


@pytest.fixture
def fixture_tools(monkeypatch):
    monkeypatch.syspath_prepend(str(ROOT))
    return importlib.import_module("build_cases"), importlib.import_module("check_output")


def recipe(prefix):
    path = next((ROOT / "recipes").glob(prefix + "*.json"))
    return json.loads(path.read_text(encoding="utf-8"))


@pytest.mark.parametrize("number", range(1, 11))
def test_synthetic_sources_match_manifest_and_shared_reader(number, fixture_tools):
    build, _ = fixture_tools
    case = recipe(f"{number:02d}_")
    build.validate_recipe(case)
    folder = ROOT / "cases" / case["slug"]
    manifest = json.loads((folder / "manifest.json").read_text(encoding="utf-8"))
    with ZipFile(folder / "upload_sources.zip") as archive:
        assert sorted(archive.namelist()) == sorted(s["path"] for s in case["sources"])
    assert len(manifest["files"]) <= 200
    for item in manifest["files"]:
        path = folder / "sources" / item["path"]
        assert path.stat().st_size == item["bytes"] > 0
        assert build.hashlib.sha256(path.read_bytes()).hexdigest() == item["sha256"]
        parsed = read_source(path)
        structure = build_structure(parsed)
        assert structure["nodes"]
        assert parsed.blocks or parsed.metadata.get("needs_ocr")
    # Expectations are readable, but never included in upload_sources.zip.
    assert Document(folder / "expected" / "交付预期.docx").paragraphs


def test_inventory_gold_numbers_are_derived_from_raw_rows():
    case = recipe("05_")
    sources = {source["path"]: source for source in case["sources"]}
    sales = sources["02_九月销售与退款.xlsx"]["sheets"]
    inventory = sources["03_库存与采购.xlsx"]["sheets"]
    start = datetime.fromisoformat("2026-09-01T00:00:00+08:00")
    end = datetime.fromisoformat("2026-10-01T00:00:00+08:00")
    included = {
        row[0]: row
        for row in sales[0]["rows"]
        if row[4] == "已完成" and start <= datetime.fromisoformat(row[1]) < end
    }
    units, revenue = defaultdict(int), defaultdict(int)
    for row in included.values():
        units[row[3]] += row[5]
        revenue[row[3]] += row[6]
    for row in sales[1]["rows"]:
        if (
            row[1] in included
            and row[4] == "已完成"
            and start <= datetime.fromisoformat(row[7]) < end
        ):
            units[row[2]] -= row[5]
            revenue[row[2]] -= row[6]
    params = {row[0]: row for row in inventory[2]["rows"]}
    cost = sum(count * params[sku][3] for sku, count in units.items())
    net = sum(revenue.values())
    assert net == 1234000 and net - cost == 488000
    positions = defaultdict(int)
    for _, sku, physical, reserved in inventory[0]["rows"]:
        positions[sku] += physical - reserved
    for row in inventory[1]["rows"]:
        if row[3] == "已确认未到货":
            positions[row[1]] += row[4]
    replenishments = {}
    for sku, row in params.items():
        daily = Decimal(units[sku]) / 30
        trigger = daily * (row[4] + row[5])
        amount = Decimal(0)
        if positions[sku] < trigger:
            amount = ((daily * row[6] - positions[sku]) / row[7]).to_integral_value(
                rounding=ROUND_CEILING
            ) * row[7]
        replenishments[sku] = int(amount)
    assert replenishments == {"A": 140, "B": 0, "C": 150, "D": 0}
    budget = sum(count * params[sku][3] for sku, count in replenishments.items())
    gold = {item["id"]: item["result"] for item in case["expected"]["calculations"]}
    assert cost == gold["C04"]
    assert net - cost == gold["C05"]
    assert budget == gold["C07"] == 327000


def test_scan_sources_have_no_hidden_answer_text_layer():
    case = recipe("07_")
    folder = ROOT / "cases" / case["slug"] / "sources"
    scanned = 0
    native = 0
    for source in case["sources"]:
        if source["kind"] not in {"pdf", "scan_pdf"}:
            continue
        pdf = PdfReader(folder / source["path"])
        assert len(pdf.pages) == len(source["pages"])
        for description, page in zip(source["pages"], pdf.pages, strict=True):
            if source["kind"] == "scan_pdf" or description.get("scan"):
                scanned += 1
                assert not page.extract_text().strip()
                assert list(page.images)
            else:
                native += 1
                assert page.extract_text().strip()
    assert scanned >= 2 and native >= 1


def test_traceability_fixture_keeps_formulas_and_hidden_dictionary():
    path = ROOT / "cases" / "03_traceability_srs" / "sources" / "02_追溯数据与样例.xlsx"
    parsed = read_source(path)
    assert parsed.metadata["formula_count"] == 10
    sheets = {item["name"]: item for item in parsed.metadata["sheets"]}
    assert len(sheets) == 13
    assert sheets["状态字典"]["hidden"]
    assert sheets["字段字典"]["rows"] - 1 == 98


def test_precheck_excludes_cover_headings_and_counts_merged_cells_once(tmp_path, fixture_tools):
    _, checker = fixture_tools
    doc = Document()
    doc.add_paragraph("封面不计数ABC")
    doc.add_heading("正文第一章", 1)
    doc.add_paragraph("正文甲乙123，空格 !")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合并格"
    doc.sections[0].header.paragraphs[0].text = "页眉不计数"
    path = tmp_path / "result.docx"
    doc.save(path)
    result = checker.inspect_document(path)
    assert result["effective_body_chars"] == 12
    assert result["headings"] == ["正文第一章"]


async def test_operations_fixture_runs_three_local_two_offline_web_then_completion(
    store,
    inputs,
    monkeypatch,
    fixture_tools,
):
    import httpx

    def forbidden_http(*args, **kwargs):
        raise AssertionError("The synthetic acceptance test must never perform HTTP")

    monkeypatch.setattr(httpx.AsyncClient, "send", forbidden_http)
    web_module = importlib.import_module("mock_web")
    folder = ROOT / "cases" / "06_operations_manual"
    web = web_module.FixtureWebKnowledge(folder / "mock_web.json")
    events = []
    seen = defaultdict(int)

    class ControlledResearch(DemoRuntime):
        async def complete(self, task, context, schema, session_key):
            if task == "research_step":
                number = context["local_round"]
                seen[number] += 1
                events.append(("local", number))
                action = "fs.read" if seen[number] == 1 else "finish"
                paths = [
                    "01_运维交付范围.docx",
                    "03_告警与指标.xlsx",
                    "06_演练与容量计划_草稿.pptx",
                ]
                arguments = {"path": paths[number - 1]} if action == "fs.read" else {}
                if context.get("memory_checkpoint_required"):
                    action, arguments = "research.save", {}
                return research_call(action, arguments, missing_information=["项目RTO/RPO尚未批准"],
                                     research_notes=self.research_notes(context))
            if task == "understand_work":
                return RuntimeResult(
                    json.dumps(
                        {
                            "summary": "公开方法不能确认内部恢复指标",
                            "facts": [],
                            "open_questions": ["项目RTO/RPO尚未批准"],
                                # A pending internal decision alone must not trigger
                                # more research. This scenario requests concrete
                                # local reads and public-method verification.
                                "research_actions": [
                                    *([{
                                        "channel": "local_sources",
                                        "question": "告警指标与容量草稿尚未核清",
                                        "action": "继续检查告警与指标表、演练与容量计划草稿",
                                    }] if context.get("research_phase") != "web"
                                        and context.get("local_round", 0) < 3 else []),
                                    {
                                        "channel": "public_web",
                                        "question": "公开恢复目标与阈值制定方法有哪些适用条件",
                                        "action": "查找恢复目标与阈值制定方法的公开原文",
                                    },
                                ],
                        }
                    )
                )
            if task == "web_query":
                assert set(seen) == {1, 2, 3}
                return native_call("search_web", {"query": f"恢复目标与阈值制定方法 {len(context['previous_queries'])}"})
            if task == "complete_gaps":
                events.append(("completion", len(web.calls)))
            return await super().complete(task, context, schema, session_key)

    inputs["user_request"] = "# 恢复目标与阈值\n区分已有事实、未批准目标与模型建议。"
    payload = inputs | {"knowledge_path": str(folder / "sources")}
    job = store.create(
        JobInput(
            **payload,
            runtime="demo",
            research=ResearchConfig(max_steps_per_round=2),
            limits=Limits(max_repairs=0),
        )
    )
    result = await Engine(store, ControlledResearch(), web_knowledge=web).run(job)
    assert result["status"] in {"draft", "needs_confirmation"}
    assert list(seen) == [1, 2, 3] and len(web.calls) == web.attempts == 2
    assert ("completion", 2) in events
    work = next(iter(store.work_states(job).values()))
    assert work.research["local_rounds"] == 3 and work.research["web_rounds"] == 2
    assert work.proposals and not work.record.facts
    public = [e for e in work.evidence if e.metadata.get("role") == "public_web"]
    assert public and all(
        e.metadata["synthetic_fixture"] and not e.metadata["project_fact"] for e in public
    )
    with pytest.raises(AssertionError, match="third web round"):
        await web.search("unexpected third attempt", limit=1, max_chars=1000)
    assert web.attempts == 3  # Even an exhausted request remains observable.
