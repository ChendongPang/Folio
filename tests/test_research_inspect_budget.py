"""Structure navigation keeps complete JSON and forward progress under budgets."""

import json
from copy import deepcopy

import pytest
from openpyxl import Workbook

from folio.research.formats import SourceBlock, SourceDocument
from folio.research.structure import build_structure
from folio.research.workspace import WorkspaceError, WorkspaceManager


@pytest.fixture
def metadata_heavy_workspace(tmp_path):
    sources = tmp_path / "sources"
    sources.mkdir()
    book = Workbook()
    for index in range(3):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = f"Sheet{index + 1}"
        sheet.append(["h" * 120] * 20)
        for row in range(2, 502):
            sheet.merge_cells(f"A{row}:B{row}")
    book.save(sources / "many.xlsx")
    return WorkspaceManager(sources, tmp_path / "task")


def bounded_json(result, max_chars):
    encoded = json.dumps(result, ensure_ascii=False)
    assert len(encoded) <= max_chars
    assert json.loads(encoded) == result
    assert "block_indices" not in encoded
    return result


def test_real_xlsx_merge_and_header_metadata_paginate_by_json_budget(metadata_heavy_workspace):
    workspace = metadata_heavy_workspace
    expected = [node for node in workspace.source_index.structure("many.xlsx")["nodes"] if node["kind"] == "sheet"]
    complete = bounded_json(workspace.inspect("many.xlsx"), 600000)
    assert complete["structure"] == expected and complete["structure_complete"]
    assert complete["next_offset"] is None and not complete["preview_truncated"]
    # An explicitly smaller page budget still preserves every header and merge.
    offset, found, pages = 0, [], []
    while True:
        result = bounded_json(workspace.inspect("many.xlsx", offset=offset, max_chars=12000), 12000)
        pages.append(result)
        assert result["structure"]
        assert result["complete"]  # Source readability is independent of pagination.
        found.extend(result["structure"])
        following = result["next_offset"]
        if following is None:
            assert result["structure_complete"]
            break
        assert following > offset and not result["structure_complete"]
        offset = following
    assert len(pages) == 3
    assert found == expected
    assert not any(page["preview_truncated"] for page in pages)


def test_single_node_inspection_compresses_only_preview_and_full_detail_is_recoverable(metadata_heavy_workspace):
    workspace = metadata_heavy_workspace
    saved = deepcopy(workspace.source_index.structure("many.xlsx"))
    sheet = next(node for node in saved["nodes"] if node["kind"] == "sheet")
    result = bounded_json(workspace.inspect("many.xlsx", node_id=sheet["node_id"], max_chars=2000), 2000)
    assert result["preview_truncated"] and result["node"]["preview_truncated"]
    assert "max_chars" in result["preview_note"] and "node_id" in result["preview_note"]
    for key in ("node_id", "parent_id", "kind", "level", "range", "locator"):
        assert result["node"][key] == sheet[key]
    metadata = result["node"]["metadata"]
    assert metadata["rows"] == 501 and metadata["columns"] == 20
    assert metadata["merged_cells_total"] == 500
    assert metadata["header_basis"] == "first_nonempty_row"
    assert result["structure"] == [] and result["next_offset"] is None
    assert result["structure_complete"]
    detailed = bounded_json(workspace.inspect("many.xlsx", node_id=sheet["node_id"], max_chars=100000), 100000)
    assert detailed["node"] == sheet and not detailed["preview_truncated"]
    assert len(detailed["node"]["metadata"]["merged_cells"]) == 500
    assert len(detailed["node"]["metadata"]["headers"]) == 20
    assert workspace.source_index.structure("many.xlsx") == saved


def test_minimum_budget_never_returns_empty_page_at_same_offset(metadata_heavy_workspace):
    workspace = metadata_heavy_workspace
    offset, identifiers = 0, []
    while True:
        result = bounded_json(workspace.inspect("many.xlsx", offset=offset, max_chars=2000), 2000)
        assert result["structure"] and result["preview_truncated"]
        identifiers.extend(node["node_id"] for node in result["structure"])
        if result["next_offset"] is None:
            break
        assert result["next_offset"] > offset
        offset = result["next_offset"]
    assert len(set(identifiers)) == len(identifiers) == 3
    exhausted = bounded_json(workspace.inspect("many.xlsx", offset=50, max_chars=2000), 2000)
    assert exhausted["structure"] == [] and exhausted["next_offset"] is None
    assert exhausted["structure_complete"]


def test_unread_status_survives_large_warnings_and_noncontiguous_selectors(metadata_heavy_workspace, monkeypatch):
    workspace = metadata_heavy_workspace
    document = SourceDocument("a" * 64, {"format": "pdf", "page_count": 2000, "complete": False,
                                          "needs_ocr": True, "unread_pages": list(range(3, 2001))},
                              blocks=[SourceBlock("page 1/block 1", "native", metadata={"page": 1}),
                                      SourceBlock("page 2/block 1", "other page", metadata={"page": 2}),
                                      SourceBlock("page 1/ocr/block 1", "ocr", metadata={"page": 1})],
                              warnings=["未读取扫描页" * 500] * 20)
    structure = build_structure(document)
    assert any("block_indices" in node for node in structure["nodes"])
    monkeypatch.setattr(workspace, "_structure", lambda _: structure)
    result = bounded_json(workspace.inspect("many.xlsx", max_chars=2000), 2000)
    assert not result["complete"] and result["metadata"]["needs_ocr"]
    assert result["metadata"]["unread_pages_truncated"]
    assert result["warnings_truncated"]
    assert result["structure"] and result["next_offset"] > 0
    assert not result["structure_complete"]
    assert result["preview_truncated"]


@pytest.mark.parametrize("budget", [0, -1, 1.5, True, "12000"])
def test_inspect_rejects_invalid_character_budget(metadata_heavy_workspace, budget):
    with pytest.raises(WorkspaceError, match="max_chars"):
        metadata_heavy_workspace.inspect("many.xlsx", max_chars=budget)
