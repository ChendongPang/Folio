"""Readability comes from enlarged views, never dropped pixels or tiny final text."""

import hashlib
from copy import deepcopy
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from test_word_figures import figure_inputs, output, png  # noqa: F401

from folio.document.figure_layout import detail_views


def test_real_run_dimensions_get_readable_detail_pages_with_identical_original_png(
    tmp_path, request
):
    inputs = request.getfixturevalue("figure_inputs")
    _, _, original_figures = inputs
    figures = deepcopy(original_figures)
    item = figures["s1"][1]
    image = png(tmp_path / "wide.png", 1142, 1041)
    item.update(kind="flowchart", width=1142, height=1041, image_path=str(image), text_font_px=18)
    target = output(tmp_path, inputs, figures)
    doc = Document(target)
    crop_shapes = [s for s in doc.inline_shapes if s._inline.xpath(".//a:srcRect")]
    assert len(crop_shapes) >= 2
    original_hash = hashlib.sha256(image.read_bytes()).digest()
    with ZipFile(target) as archive:
        assert (
            sum(
                hashlib.sha256(archive.read(n)).digest() == original_hash
                for n in archive.namelist()
                if n.startswith("word/media/")
            )
            == 1
        )
    for shape in crop_shapes:
        crop = shape._inline.xpath(".//a:srcRect")[0]
        visible_width = 1142 * (1 - (int(crop.get("l")) + int(crop.get("r"))) / 100000)
        visible_height = 1041 * (1 - (int(crop.get("t")) + int(crop.get("b"))) / 100000)
        actual_font_pt = 13.5 * (shape.width / 914400 * 96 / visible_width)
        assert actual_font_pt >= 10
        assert shape.width / shape.height == pytest.approx(
            visible_width / visible_height, rel=0.001
        )
        assert (
            shape.width
            <= doc.sections[0].page_width
            - doc.sections[0].left_margin
            - doc.sections[0].right_margin
        )
        p = shape._inline.getparent().getparent().getparent()
        assert p.find("./" + qn("w:pPr") + "/" + qn("w:pageBreakBefore")) is not None
    assert any("总览" in p.text for p in doc.paragraphs)
    assert sum(p.text == "s1业务正文开始。" for p in doc.paragraphs) == 1
    assert sum(p.text == "s1业务正文结束。" for p in doc.paragraphs) == 1


def test_viewports_cover_every_pixel_and_overlap_for_large_flowchart():
    item = {"kind": "flowchart", "width": 2400, "height": 2400, "text_font_px": 18}
    views = detail_views(item, Inches(5.9), Inches(9))
    assert views
    xs = sorted({(v["box"][0], v["box"][2]) for v in views})
    ys = sorted({(v["box"][1], v["box"][3]) for v in views})
    for spans in (xs, ys):
        assert spans[0][0] == 0 and spans[-1][1] == 2400
        assert all(a[1] > b[0] for a, b in zip(spans, spans[1:], strict=False))
    assert len(views) == len(xs) * len(ys)


def test_readable_flowchart_and_ui_do_not_gain_unnecessary_pages():
    base = {"kind": "flowchart", "width": 600, "height": 400, "text_font_px": 18}
    assert detail_views(base, Inches(5.9), Inches(9)) == []
    assert detail_views(base | {"kind": "ui", "width": 2000}, Inches(5.9), Inches(9)) == []
