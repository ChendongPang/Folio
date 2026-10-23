import shutil
import subprocess
from copy import deepcopy

import pytest
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

from folio.document.lists import ProseListLayout
from folio.document.presentation import color_proposal
from folio.document.reader import xpath
from folio.models import Block


def render_items(items):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(12)
    ind = doc.styles["Normal"].element.get_or_add_pPr().get_or_add_ind()
    ind.set(qn("w:firstLineChars"), "200")
    layout = ProseListLayout()
    for identifier, text in items:
        p = doc.add_paragraph(text)
        layout.apply(p, text, identifier)
    return doc


def test_legacy_split_lists_return_to_parent_and_align_wrapped_lines():
    doc = render_items(
        [
            ("b10", "10. 原有业务。"),
            ("b11", "11. 具体功能包括："),
            ("b11_p1", "1. 行情展示。"),
            ("b11_p2", "2. 监控看板。"),
            ("b11_p3", "3. 异常预警。"),
            ("b12", "12. 具体规则如下："),
            ("b12_p1", "1. 角色定义。"),
        ]
    )
    assert [p.text.lstrip("\t") for p in doc.paragraphs] == [
        "10.\t原有业务。",
        "11.\t具体功能包括：",
        "（1）\t行情展示。",
        "（2）\t监控看板。",
        "（3）\t异常预警。",
        "12.\t具体规则如下：",
        "（1）\t角色定义。",
    ]
    parents = [doc.paragraphs[i] for i in (0, 1, 5)]
    children = [doc.paragraphs[i] for i in (2, 3, 4, 6)]
    assert {p.paragraph_format.left_indent for p in parents} == {Pt(24.6)}
    assert {p.paragraph_format.left_indent for p in children} == {Pt(60.6)}
    for p in doc.paragraphs:
        child = p.text.startswith("\t（")
        assert p.paragraph_format.first_line_indent == (Pt(-36.6) if child else Pt(-24.6))
        assert xpath(p._p, "./w:pPr/w:numPr/w:numId/@w:val") == ["0"]
        assert xpath(p._p, "./w:pPr/w:ind/@w:firstLineChars") == ["0"]
        assert xpath(p._p, "./w:pPr/w:ind/@w:hangingChars") == (["305"] if child else ["205"])
        tabs = list(p.paragraph_format.tab_stops)
        assert [tab.alignment for tab in tabs] == [WD_TAB_ALIGNMENT.RIGHT, WD_TAB_ALIGNMENT.LEFT]
        assert tabs[1].position == p.paragraph_format.left_indent
        assert tabs[1].position - tabs[0].position == Pt(4.8)


def test_explicit_children_and_legacy_separate_blocks():
    doc = render_items(
        [
            ("a", "11. 功能包括："),
            ("b", "1. 行情。"),
            ("c", "2. 监控。"),
            ("d", "12. 权限包括："),
            ("e", "（1）角色。"),
            ("f", "（2）权限。"),
            ("g", "13. 最后一个主项。"),
        ]
    )
    assert [p.text.split("\t")[1] for p in doc.paragraphs] == [
        "11.",
        "（1）",
        "（2）",
        "12.",
        "（1）",
        "（2）",
        "13.",
    ]


@pytest.mark.parametrize(
    "text", ["利率1.25%，版本V1.0。", "1.2.3 章节引用", "2026.09.29", "普通段落。"]
)
def test_non_list_prose_unchanged(text):
    doc = Document()
    p = doc.add_paragraph(text)
    before = deepcopy(p._p.xml)
    ProseListLayout().apply(p, text, "b")
    assert p._p.xml == before


def test_red_list_number_and_text_remain_red_without_changing_metadata():
    block = Block(block_id="b", text="1. 建议业务规则。", origin="model_proposal")
    original = block.model_dump()
    doc = render_items([(block.block_id, block.text)])
    p = doc.paragraphs[0]
    color_proposal(p._p, block)
    assert all(str(r.font.color.rgb) == "FF0000" for r in p.runs)
    assert p.style.font.size == Pt(12)
    assert block.model_dump() == original


def test_list_reset_does_not_carry_parent_into_next_section():
    layout = ProseListLayout()
    doc = Document()
    layout.apply(doc.add_paragraph(), "11. 内容包括：", "b")
    layout.reset()
    p = doc.add_paragraph()
    layout.apply(p, "1. 新章节项目。", "c")
    assert p.text == "\t1.\t新章节项目。"


def test_explicit_child_marker_returns_to_parent_even_when_numbers_overlap():
    doc = render_items([("a", "1. 内容包括："), ("b", "（1）唯一子项。"), ("c", "2. 下一主项。")])
    assert [p.text.split("\t")[1] for p in doc.paragraphs] == ["1.", "（1）", "2."]


def test_later_large_numbers_align_all_items_without_widening_the_number_to_text_gap():
    doc = render_items([("a", "1. 首项"), ("b", "10. 中项"), ("c", "100. 末项")])
    assert len({p.paragraph_format.left_indent for p in doc.paragraphs}) == 1
    for p in doc.paragraphs:
        right, left = p.paragraph_format.tab_stops
        assert left.position - right.position == Pt(4.8)
        assert left.position == p.paragraph_format.left_indent


def test_inherited_tab_stops_are_cleared_locally_without_changing_template_styles():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.size = Pt(13)
    normal.paragraph_format.tab_stops.add_tab_stop(Pt(8))
    normal.paragraph_format.tab_stops.add_tab_stop(Pt(120), WD_TAB_ALIGNMENT.RIGHT)
    normal.paragraph_format.first_line_indent = Pt(26)
    p = doc.add_paragraph("1. 业务正文。")
    p.paragraph_format.tab_stops.add_tab_stop(Pt(240))
    before = normal.element.xml
    ProseListLayout().apply(p, p.text, "b")
    assert set(xpath(p._p, './w:pPr/w:tabs/w:tab[@w:val="clear"]/@w:pos')) == {
        "160",
        "2400",
        "4800",
    }
    assert normal.element.xml == before
    active = [
        tab for tab in p.paragraph_format.tab_stops if tab.alignment != WD_TAB_ALIGNMENT.CLEAR
    ]
    assert len(active) == 2
    assert active[1].position - active[0].position == Pt(5.2)


def test_inherited_and_direct_fonts_are_preserved_and_used_for_list_geometry():
    doc = Document()
    doc.styles["Normal"].font.size = Pt(15)
    p = doc.add_paragraph("1. 模板字号和字体。")
    run = p.runs[0]
    run.font.name = "DejaVu Serif"
    run.font.size = Pt(14)
    run.italic = True
    before = run._r.rPr.xml
    ProseListLayout().apply(p, p.text, "b")
    assert p.runs[0]._r.rPr.xml == before
    right, left = p.paragraph_format.tab_stops
    assert left.position - right.position == Pt(5.6)


def test_rendered_number_gap_and_wrapped_body_alignment(tmp_path):
    from folio.document.pagination import office_runtime, paginate

    if office_runtime() is None or not shutil.which("pdftotext"):
        pytest.skip("Local Office and PDF text geometry tools are required")
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "DejaVu Serif"
    style.font.size = Pt(12)
    style.paragraph_format.tab_stops.add_tab_stop(Pt(6))
    style.paragraph_format.tab_stops.add_tab_stop(Pt(100))
    style.paragraph_format.first_line_indent = Pt(24)
    layout = ProseListLayout()
    for number, word in [(1, "Alpha"), (2, "Bravo"), (10, "Charlie"), (100, "Delta")]:
        text = f"{number}. {word} " + "The transaction requires approval before settlement. " * 5
        paragraph = doc.add_paragraph(text)
        layout.apply(paragraph, text, f"b{number}")
    target, pdf = tmp_path / "lists.docx", tmp_path / "lists.pdf"
    doc.save(target)
    paginate(target, pdf=pdf, required=True)
    bbox = subprocess.run(
        ["pdftotext", "-bbox", str(pdf), "-"], check=True, capture_output=True
    ).stdout
    root = etree.fromstring(bbox, parser=etree.XMLParser(no_network=True))
    words = root.xpath('//*[local-name()="word"]')
    starts = []
    for number, word in [(1, "Alpha"), (2, "Bravo"), (10, "Charlie"), (100, "Delta")]:
        label = next(value for value in words if value.text == f"{number}.")
        first = next(value for value in words if value.text == word)
        assert float(first.get("xMin")) - float(label.get("xMax")) == pytest.approx(4.8, abs=0.15)
        starts.append(float(first.get("xMin")))
    assert max(starts) - min(starts) < 0.1
    rows = {}
    for word in words:
        rows.setdefault(round(float(word.get("yMin")), 1), []).append(word)
    wrapped = [
        row
        for row in rows.values()
        if not any(word.text in {"1.", "2.", "10.", "100."} for word in row)
    ]
    assert len(wrapped) >= 4
    for row in wrapped:
        assert min(float(word.get("xMin")) for word in row) == pytest.approx(starts[0], abs=0.15)
