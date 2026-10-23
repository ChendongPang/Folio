import asyncio
import json

import httpx
import pytest

from folio.model_settings import ModelSettingsInput, ModelSettingsStore
from folio.models import JobInput, LocalReview, Query, Review, SectionContent
from folio.orchestration.context import (
    ContextSelection,
    GapMaterialRead,
    deduplicate_dependency_facts,
    project,
    scope_work_review,
)
from folio.orchestration.engine import Engine
from folio.orchestration.work import BusinessRecord, GapCompletion
from folio.runtime.base import RuntimeResult
from folio.runtime.demo import DemoRuntime
from folio.runtime.llm import LLMRuntime


def engine_for(store, inputs, runtime):
    engine = Engine(store, runtime)
    engine.request = JobInput(**inputs, runtime="demo")
    engine.limits = engine.request.limits
    engine.job_id = store.create(engine.request)
    return engine


def evidence(index, text="完整原始业务规则"):
    return {
        "evidence_id": f"e{index}",
        "source_id": "document",
        "source_version": "v1",
        "locator": str(index),
        "text": text,
        "metadata": {"ingestion_debug": "huge metadata"},
    }


def test_duplicate_global_fact_is_removed_only_for_identical_source_and_revision():
    fact = {"claim_id": "c1", "statement": "先审批后交易", "evidence_ids": ["e1"]}
    duplicate = fact | {"section_id": "a", "section_revision": 2, "record_id": "a:c1"}
    changed = duplicate | {"statement": "无需审批"}
    older = duplicate | {"section_revision": 1}
    other = duplicate | {"section_id": "b", "record_id": "b:c1"}
    context = {
        "dependencies": [
            {"sections": [{"section_id": "a", "revision": 2, "supported_facts": [fact]}]}
        ],
        "global_facts": [duplicate, changed, older, other],
    }
    deduplicate_dependency_facts(context)
    assert context["global_facts"] == [changed, older, other]
    assert context["dependencies"][0]["sections"][0]["supported_facts"] == [fact]


def test_repeated_context_references_are_idempotent_but_unknown_refs_fail(store, inputs):
    engine = engine_for(store, inputs, DemoRuntime())
    context = {
        "catalog": [{"context_ref": "f", "kind": "fact"}, {"context_ref": "r", "kind": "risk"}]
    }
    selection = ContextSelection(
        selected_refs=["f", "f"], risk_refs=["r", "r"], risk_summary="风险"
    )
    assert selection.selected_refs == ["f"]
    assert selection.risk_refs == ["r"]
    engine.check_result(selection, context)
    with pytest.raises(ValueError):
        engine.check_result(ContextSelection(selected_refs=["missing", "missing"]), context)
    with pytest.raises(ValueError):
        engine.check_result(ContextSelection(selected_refs=["r", "r"]), context)


async def test_dependency_reports_are_selected_not_blindly_propagated(store, inputs):
    class Selective(DemoRuntime):
        def __init__(self):
            self.catalog = []
            self.reads = []

        async def complete(self, task, context, schema, session):
            if task == "select_context":
                self.catalog += context["catalog"]
                chosen = [
                    x["context_ref"]
                    for x in context["catalog"]
                    if x["kind"] == "risk" and x["value"]["key"] == "applicable"
                ]
                return RuntimeResult(
                    json.dumps(
                        {
                            "selected_refs": [
                                x["context_ref"] for x in context["catalog"] if x["kind"] != "risk"
                            ],
                            "risk_refs": chosen,
                            "risk_summary": "审批规则尚未确认" if chosen else "",
                        }
                    )
                )
            if task == "understand_work":
                self.reads.append(context)
            return await super().complete(task, context, schema, session)

    runtime = Selective()
    # This exercises selection under an intentionally small request budget;
    # the application default now retains this entire report without selection.
    engine = engine_for(store, inputs | {"limits": {"max_prompt_chars": 48000}}, runtime)
    source = {
        "work_task": {"objective": "审批业务"},
        "sections": [],
        "evidence": [evidence(1)],
        "dependencies": [
            {
                "task_id": "upstream",
                "title": "基础规则",
                "open_questions": [],
                "model_proposals": [],
                "sections": [
                    {
                        "section_id": "a",
                        "revision": 1,
                        "supported_facts": [
                            {"claim_id": "c", "statement": "事实", "evidence_ids": ["e1"]}
                        ],
                        "findings": [],
                    }
                ],
                "findings": [
                    {
                        "kind": "unsupported",
                        "key": "applicable" if i == 0 else f"layout{i}",
                        "description": "完整报告" * 300,
                        "section_ids": ["a"],
                    }
                    for i in range(60)
                ],
            }
        ],
    }
    original = json.dumps(source, sort_keys=True)
    await engine.call("understand_work", source, BusinessRecord)
    assert json.dumps(source, sort_keys=True) == original
    assert len(runtime.catalog) == 61  # Every original item was available for semantic selection.
    view = runtime.reads[0]["dependencies"][0]
    assert "findings" not in view and "findings" not in view["sections"][0]
    assert view["sections"][0]["supported_facts"][0]["claim_id"] == "c"
    assert len(view["risk_briefs"]) == 1
    assert view["risk_briefs"][0]["status"] == "unresolved_risk"
    assert len(json.dumps(view)) < 2000
    files = list((store.root / "jobs" / engine.job_id / "contexts" / "sources").glob("*.json"))
    assert json.loads(files[0].read_text())["context"] == source
    calls = len(runtime.catalog)
    await engine.call("understand_work", source, BusinessRecord)
    assert (
        len(runtime.catalog) == calls
    )  # Reuse semantic read decisions, including risk provenance.


async def test_reader_batches_preserve_all_facts_and_resume_independently(store, inputs):
    class Reader(DemoRuntime):
        calls = 0
        interrupt = True

        async def measure(self, task, context, schema, char_limit):
            size = sum(len(e["text"]) for e in context.get("evidence", []))
            return {"characters": size, "fits": task != "understand_work" or size <= 6000}

        async def complete(self, task, context, schema, session):
            if task == "understand_work":
                self.calls += 1
                if self.calls == 2 and self.interrupt:
                    raise asyncio.CancelledError()
            return await super().complete(task, context, schema, session)

    runtime = Reader()
    engine = engine_for(store, inputs, runtime)
    context = {
        "work_task": {"objective": "业务任务边界不变"},
        "evidence": [evidence(i, str(i) + "规则" * 300) for i in range(60)],
    }
    with pytest.raises(asyncio.CancelledError):
        await engine.call("understand_work", context, BusinessRecord)
    runtime.interrupt = False
    result = await engine.call("understand_work", context, BusinessRecord)
    assert len(result.facts) == 60
    assert len({f.fact_id for f in result.facts}) == 60
    assert {e for f in result.facts for e in f.evidence_ids} == {f"e{i}" for i in range(60)}
    # Eight completed leaves, one interrupted attempt; successful first leaf is reused.
    assert runtime.calls == 9
    # Aggregate is persistable and not constrained by a single call's 50-fact schema.
    assert BusinessRecord.model_validate_json(result.model_dump_json()) == result


@pytest.mark.parametrize("capacity", [None, 3])
async def test_gap_completion_only_partitions_when_measured_capacity_requires_it(store, inputs, capacity):
    class Gaps(DemoRuntime):
        batches = []

        async def measure(self, task, context, schema, char_limit):
            return {
                "characters": len(json.dumps(context)),
                "fits": task != "complete_gaps" or capacity is None or len(context["gaps"]) <= capacity,
            }

        async def complete(self, task, context, schema, session):
            if task == "complete_gaps":
                self.batches.append(context["gaps"])
            return await super().complete(task, context, schema, session)

    runtime = Gaps()
    engine = engine_for(store, inputs, runtime)
    context = {
        "gaps": [{"gap_id": f"g{i}", "question": f"业务规则{i}"} for i in range(9)],
        "sections": [{"section_id": "a", "title": "业务规则"}],
    }
    result = await engine.call("complete_gaps", context, GapCompletion)
    assert {p.gap_id for p in result.proposals} == {f"g{i}" for i in range(9)}
    if capacity is None:
        assert runtime.batches == [context["gaps"]]
    else:
        assert len(runtime.batches) > 1
        assert all(len(page) <= capacity for page in runtime.batches)
    assert len(result.proposals) == 9


async def test_work_review_partitions_sources_instead_of_repeating_them_per_pair(store, inputs):
    class Reviewer(DemoRuntime):
        def __init__(self):
            self.focuses, self.reads = [], []

        async def measure(self, task, context, schema, char_limit):
            return {"characters": 100, "fits": len(context.get("definitions", [])) <= 2}

        async def complete(self, task, context, schema, session):
            if task == "read_review_materials":
                self.reads.extend(context["catalog"])
            if task == "review_focus":
                assert "evidence" not in context and "business_record" not in context
                self.reads.extend(context.get("original_materials", []))
                assert context["reading_complete"]
                self.focuses.append(context["review_focus"]["section_ids"])
            return await super().complete(task, context, schema, session)

    runtime = Reviewer()
    engine = engine_for(store, inputs, runtime)
    source = {
        "sections": [
            {"section_id": str(i), "content": {"blocks": [{"text": "完整正文"}]}} for i in range(4)
        ],
        "definitions": [{"section_id": str(i), "instruction": "完整要求"} for i in range(4)],
        "model_proposals": [
            {"gap_id": f"g{i}", "section_ids": [str(i)], "content": "原方案"} for i in range(4)
        ],
        "evidence": [evidence(i) for i in range(4)],
        "business_record": {"conflicts": ["材料冲突"]},
    }
    snapshot = json.dumps(source, sort_keys=True)
    await engine.call("review_work", source, Review)
    assert sorted(runtime.focuses) == [[str(i)] for i in range(4)]
    assert {m["value"]["evidence_id"] for m in runtime.reads if m["kind"] == "evidence"} == {
        f"e{i}" for i in range(4)
    }
    assert json.dumps(source, sort_keys=True) == snapshot


def test_work_review_keeps_proposals_referenced_by_legacy_shared_blocks():
    context = {
        "sections": [{"section_id": "a", "content": {"blocks": [{"proposal_ref": "p1"}]}}],
        "model_proposals": [{"gap_id": "p1", "section_ids": ["b"], "content": "原文"}],
    }
    assert scope_work_review(context)["model_proposals"] == context["model_proposals"]


async def test_drill_down_requests_complete_flat_decision_and_keeps_repair_guards(store, inputs):
    from folio.models import DrillDown

    class Decision(DemoRuntime):
        async def complete(self, task, context, schema, session):
            assert task == "drill_down" and "anyOf" not in schema
            assert set(schema["required"]) == {
                "confirmed",
                "reason",
                "affected_section_ids",
                "repair_instruction",
                "needs_confirmation",
            }
            assert schema["properties"]["affected_section_ids"]["items"]["enum"] == ["a"]
            return RuntimeResult(
                json.dumps(
                    {
                        "confirmed": True,
                        "reason": "原文确认存在遗漏",
                        "affected_section_ids": ["a"],
                        "repair_instruction": "补充材料明确的原始规则",
                        "needs_confirmation": False,
                    }
                )
            )

    engine = engine_for(store, inputs, Decision())
    context = {"finding": {"section_ids": ["a"]}, "sections": []}
    result = await engine.call("drill_down", context, DrillDown)
    assert result.confirmed and result.reason
    with pytest.raises(ValueError):
        engine.check_result(result.model_copy(update={"affected_section_ids": []}), context)
    with pytest.raises(ValueError):
        engine.check_result(result.model_copy(update={"repair_instruction": ""}), context)


async def test_large_single_gap_reads_all_sources_in_pages_and_resumes(store, inputs):
    class Reader(DemoRuntime):
        def __init__(self):
            self.reads = []
            self.interrupt = True
            self.final = None

        async def measure(self, task, context, schema, char_limit):
            fits = True
            if task == "complete_gaps":
                fits = len(context.get("evidence", [])) <= 3
            if task == "read_gap_materials":
                fits = len(context["catalog"]) <= 5
            return {"characters": 100, "fits": fits}

        async def complete(self, task, context, schema, session):
            if task == "read_gap_materials":
                self.reads.append(context["catalog"])
                if len(self.reads) == 2 and self.interrupt:
                    raise asyncio.CancelledError()
                return RuntimeResult(
                    json.dumps(
                        {
                            "constraints": [
                                {
                                    "source_refs": [item["context_ref"]],
                                    "source_kind": item["kind"],
                                    "text": item["value"]["text"][:20],
                                }
                                for item in context["catalog"]
                            ]
                        }
                    )
                )
            if task == "complete_gaps":
                self.final = context
            return await super().complete(task, context, schema, session)

    runtime = Reader()
    engine = engine_for(store, inputs, runtime)
    context = {
        "gaps": [{"gap_id": "g1", "question": "完整业务流程缺口"}],
        "sections": [{"section_id": "a", "title": "业务流程"}],
        "evidence": [evidence(i, f"原始规则{i}:" + "完整内容" * 500) for i in range(18)],
    }
    with pytest.raises(asyncio.CancelledError):
        await engine.call("complete_gaps", context, GapCompletion)
    runtime.interrupt = False
    result = await engine.call("complete_gaps", context, GapCompletion)
    assert [p.gap_id for p in result.proposals] == ["g1"]
    assert len(runtime.reads) == 5  # Four completed leaves plus one interrupted call.
    notes = runtime.final["material_constraints"]
    assert len(notes) == 18
    assert len({ref for n in notes for ref in n["source_refs"]}) == 18
    assert all(n["source_kind"] == "evidence" for n in notes)
    assert "evidence" not in runtime.final
    read_sources = {
        item["value"]["evidence_id"]: item["value"]["text"]
        for page in runtime.reads
        for item in page
    }
    assert read_sources == {e["evidence_id"]: e["text"] for e in context["evidence"]}
    assert not store.states(engine.job_id)  # Notes never become document/global facts.
    before = len(runtime.reads)
    assert await engine.call("complete_gaps", context, GapCompletion) == result
    assert len(runtime.reads) == before


def test_gap_material_notes_cannot_promote_risks_or_invent_sources(store, inputs):
    engine = engine_for(store, inputs, DemoRuntime())
    context = {"catalog": [{"context_ref": "r", "kind": "risk"}]}
    for ref, kind in (("unknown", "risk"), ("r", "fact")):
        value = GapMaterialRead(
            constraints=[{"source_refs": [ref], "source_kind": kind, "text": "约束"}]
        )
        with pytest.raises(ValueError):
            engine.check_result(value, context)
    engine.check_result(
        GapMaterialRead(
            constraints=[{"source_refs": ["r"], "source_kind": "risk", "text": "未确认风险"}]
        ),
        context,
    )


async def test_independent_material_pages_run_concurrently_and_merge_in_source_order(store, inputs):
    class Reader(DemoRuntime):
        active = 0
        peak = 0

        async def measure(self, task, context, schema, char_limit):
            return {"characters": 100, "fits": len(context.get("catalog", [])) <= 2}

        async def complete(self, task, context, schema, session):
            assert task == "read_gap_materials"
            self.active += 1
            self.peak = max(self.peak, self.active)
            try:
                # Complete later source pages first; merge order must remain stable.
                index = int(context["catalog"][0]["value"]["evidence_id"][1:])
                await asyncio.sleep((8 - index) * 0.002)
                return RuntimeResult(
                    json.dumps(
                        {
                            "constraints": [
                                {
                                    "source_refs": [item["context_ref"]],
                                    "source_kind": item["kind"],
                                    "text": item["value"]["text"],
                                }
                                for item in context["catalog"]
                            ]
                        }
                    )
                )
            finally:
                self.active -= 1

    runtime = Reader()
    engine = engine_for(store, inputs, runtime)
    result = await engine.read_gap_materials(
        {
            "gaps": [{"gap_id": "g", "question": "业务约束"}],
            "sections": [{"section_id": "a", "title": "业务流程"}],
            "evidence": [evidence(i, f"完整规则{i}") for i in range(8)],
        }
    )
    assert runtime.peak == 4 and runtime.active == 0
    assert [item["text"] for item in result["material_constraints"]] == [
        f"完整规则{i}" for i in range(8)
    ]


@pytest.mark.parametrize("gap_count,capacity", [(6, 3), (2, 1)])
async def test_gap_batches_select_their_own_materials_and_preserve_sources(
    store, inputs, gap_count, capacity
):
    class ScopedGaps(DemoRuntime):
        def __init__(self):
            self.batches = []
            self.selections = []

        async def measure(self, task, context, schema, char_limit):
            return {
                "characters": 100,
                "fits": task != "complete_gaps" or len(context["evidence"]) <= capacity,
            }

        async def complete(self, task, context, schema, session):
            if task == "select_context":
                topics = {g["question"] for g in context["purpose"]["gaps"]}
                self.selections.append(topics)
                # Select facts, not evidence directly: their full originals must follow.
                selected = [
                    item["context_ref"]
                    for item in context["catalog"]
                    if item["kind"] == "fact" and item["value"]["statement"] in topics
                ]
                risks = [x["context_ref"] for x in context["catalog"] if x["kind"] == "risk"]
                return RuntimeResult(
                    json.dumps(
                        {
                            "selected_refs": selected,
                            "risk_refs": risks,
                            "risk_summary": "保留原始依赖风险" if risks else "",
                        }
                    )
                )
            if task == "complete_gaps":
                self.batches.append(context)
            return await super().complete(task, context, schema, session)

    runtime = ScopedGaps()
    engine = engine_for(store, inputs, runtime)
    context = {
        "gaps": [{"gap_id": f"g{i}", "question": f"topic-{i}"} for i in range(gap_count)],
        "sections": [{"section_id": "a", "title": "业务需求"}],
        "evidence": [evidence(i, f"topic-{i} 的完整原始规则") for i in range(12)],
        "business_record": {
            "facts": [
                {"fact_id": f"f{i}", "statement": f"topic-{i}", "evidence_ids": [f"e{i}"]}
                for i in range(12)
            ],
            "conflicts": ["不可自行裁决的冲突"],
        },
        "global_facts": [{"record_id": f"r{i}", "statement": f"topic-{i}"} for i in range(12)],
        "dependencies": [
            {
                "task_id": "upstream",
                "title": "上游任务",
                "sections": [],
                "findings": [
                    {"kind": "pending", "key": "business_rule", "description": "依赖风险原文"}
                ],
            }
        ],
    }
    original = json.dumps(context, sort_keys=True)
    result = await engine.call("complete_gaps", context, GapCompletion)
    assert len(result.proposals) == gap_count
    assert {p.gap_id for p in result.proposals} == {g["gap_id"] for g in context["gaps"]}
    for batch in runtime.batches:
        topics = {g["question"] for g in batch["gaps"]}
        assert len(topics) <= capacity
        assert {f["statement"] for f in batch["business_record"]["facts"]} == topics
        assert {f["statement"] for f in batch["global_facts"]} == topics
        assert {e["text"].split()[0] for e in batch["evidence"]} == topics
        assert batch["business_record"]["conflicts"] == ["不可自行裁决的冲突"]
        risk = batch["dependencies"][0]["risk_briefs"][0]
        assert risk["summary"] == "保留原始依赖风险"
        assert risk["source_refs"]
    assert json.dumps(context, sort_keys=True) == original
    archived = list((store.root / "jobs" / engine.job_id / "contexts" / "sources").glob("*.json"))
    assert any(json.loads(p.read_text())["context"] == context for p in archived)
    before = (len(runtime.selections), len(runtime.batches))
    assert await engine.call("complete_gaps", context, GapCompletion) == result
    assert (len(runtime.selections), len(runtime.batches)) == before


@pytest.mark.parametrize("capacity,expected_calls", [(None, 1), (24, 4)])
async def test_validate_capacity_batches_keep_body_and_every_claim(store, inputs, capacity, expected_calls):
    class Validator(DemoRuntime):
        bodies = []

        async def measure(self, task, context, schema, char_limit):
            return {
                "characters": len(json.dumps(context)),
                "fits": task != "validate" or capacity is None
                or len(context["extraction"]["claims"]) <= capacity,
            }

        async def complete(self, task, context, schema, session):
            if task == "validate":
                self.bodies.append(context["content"])
                return RuntimeResult(
                    json.dumps(
                        {
                            "verdicts": [
                                {"claim_id": c["claim_id"], "supported": True, "reason": "依据匹配"}
                                for c in context["extraction"]["claims"]
                            ]
                        }
                    )
                )
            return await super().complete(task, context, schema, session)

    runtime = Validator()
    engine = engine_for(store, inputs, runtime)
    body = {"blocks": [{"block_id": "b1", "text": "原始正文", "evidence_ids": ["e1"]}]}
    context = {
        "section": {"section_id": "a"},
        "content": body,
        "evidence": [evidence(1), evidence(2)],
        "extraction": {"claims": [{"claim_id": f"c{i}"} for i in range(50)]},
    }
    result = await engine.call("validate", context, LocalReview)
    assert len(result.verdicts) == 50
    assert len({v.claim_id for v in result.verdicts}) == 50
    assert len(runtime.bodies) == expected_calls
    assert all(b == body for b in runtime.bodies)


def test_each_operation_has_its_own_read_contract():
    source = {
        "evidence": [evidence(1), evidence(2)],
        "content": {"blocks": [{"evidence_ids": ["e1"]}]},
        "required_project_evidence_ids": ["e2"],
        "peer_sections": [{"full": "body"}],
        "previous_record": {"summary": "redundant"},
    }
    assert "text" not in project("next_action", source)["evidence"][0]
    assert "previous_record" not in project("understand_work", source)
    assert "peer_sections" not in project("generate", source)
    assert len(project("validate", source)["evidence"]) == 2
    assert "ingestion_debug" not in project("validate", source)["evidence"][0]["metadata"]
    assert "ingestion_debug" in source["evidence"][0]["metadata"]


def test_risk_cannot_be_promoted_to_fact_or_fabricated(store, inputs):
    engine = engine_for(store, inputs, DemoRuntime())
    context = {"catalog": [{"context_ref": "r", "kind": "risk"}]}
    with pytest.raises(ValueError, match="declared kind"):
        engine.check_result(ContextSelection(selected_refs=["r"]), context)
    with pytest.raises(ValueError, match="references"):
        engine.check_result(ContextSelection(risk_refs=["missing"], risk_summary="风险"), context)
    with pytest.raises(ValueError, match="summary"):
        engine.check_result(ContextSelection(risk_refs=["r"]), context)


async def test_actual_token_window_replaces_character_guess(tmp_path):
    settings = ModelSettingsStore(tmp_path)
    settings.save(
        ModelSettingsInput(
            base_url="http://model.test/v1", model="local", no_auth=True, enable_thinking=False,
            max_tokens=1024,
        )
    )
    runtime = LLMRuntime(settings.path)
    await runtime.client.aclose()
    requests = []

    def handler(request):
        requests.append(request)
        if request.url.path == "/v1/models":
            return httpx.Response(200, json={"data": [{"id": "local", "max_model_len": 4096}]})
        body = json.loads(request.content)
        assert body["chat_template_kwargs"] == {"enable_thinking": False}
        assert body["add_generation_prompt"] is True
        return httpx.Response(200, json={"count": 2600 if "too many" in str(body) else 2500})

    runtime.client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    try:
        good = await runtime.measure(
            "query", {"data": "x" * 60000}, Query.model_json_schema(), 48000
        )
        assert good["characters"] > 48000 and good["fits"]
        assert good["input_tokens"] == 2500 and good["output_tokens"] == 1024
        bad = await runtime.measure("query", {"data": "too many"}, Query.model_json_schema(), 48000)
        assert not bad["fits"] and bad["characters"] < 48000
        assert [r.url.path for r in requests] == ["/v1/models", "/tokenize", "/tokenize"]
    finally:
        await runtime.aclose()


@pytest.mark.parametrize("task", ["generate", "repair", "validate", "review_work", "review"])
async def test_remaining_window_keeps_minimum_output_and_matches_actual_request(tmp_path, task):
    settings = ModelSettingsStore(tmp_path)
    settings.save(
        ModelSettingsInput(
            base_url="http://model.test/v1", model="local", no_auth=True, max_tokens=8192
        )
    )
    runtime = LLMRuntime(settings.path)
    await runtime.client.aclose()
    sent = []

    def handler(request):
        if request.url.path == "/v1/models":
            return httpx.Response(200, json={"data": [{"id": "local", "max_model_len": 32768}]})
        if request.url.path == "/tokenize":
            return httpx.Response(
                200,
                json={
                    "count": 31000
                    if "too large" in request.content.decode()
                    else 29000
                    if "review allowance" in request.content.decode()
                    else 26532
                },
            )
        sent.append(json.loads(request.content))
        chunk = {"choices": [{"index": 0, "delta": {"content": "{}"}, "finish_reason": "stop"}]}
        return httpx.Response(
            200,
            headers={"content-type": "text/event-stream"},
            text="data: " + json.dumps(chunk) + "\n\ndata: [DONE]\n\n",
        )

    runtime.client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    try:
        measured = await runtime.measure(task, {"section": "current"}, {}, 48000)
        assert measured["fits"] and measured["output_tokens"] == 5724
        assert measured["requested_output_tokens"] == 8192
        await runtime.complete(task, {"section": "current"}, {}, "test")
        assert sent[0]["max_tokens"] == measured["output_tokens"]
        too_large = await runtime.measure(task, {"section": "too large"}, {}, 48000)
        assert not too_large["fits"] and too_large["output_tokens"] == 8192
        assert runtime.config["max_tokens"] == 8192
        allowance = await runtime.measure(task, {"section": "review allowance"}, {}, 48000)
        assert allowance["fits"] == (task in {"validate", "review_work", "review"})
        if allowance["fits"]:
            assert allowance["output_tokens"] == 3256
        query = await runtime.measure("query", {}, {}, 48000)
        assert query["output_tokens"] == 8192
    finally:
        await runtime.aclose()


async def test_tokenizer_unavailable_uses_default_token_budget(tmp_path):
    settings = ModelSettingsStore(tmp_path)
    settings.save(ModelSettingsInput(base_url="http://model.test/v1", model="local", no_auth=True))
    runtime = LLMRuntime(settings.path)
    await runtime.client.aclose()
    runtime.client = httpx.AsyncClient(transport=httpx.MockTransport(lambda r: httpx.Response(404)))
    try:
        result = await runtime.measure(
            "query", {"data": "x" * 60000}, Query.model_json_schema(), 48000
        )
        assert result["measurement"] == "default_tokenizer_estimate"
        assert result["input_tokens"] is None and result["fits"]
        assert result["max_input_tokens"] == 600000
        assert result["tokenizer_encoding"] == "o200k_base"
    finally:
        await runtime.aclose()


@pytest.mark.parametrize("failure", ["connection", "503", "429"])
async def test_tokenizer_outage_does_not_disable_token_capacity_checks(
    tmp_path, failure, monkeypatch
):
    from folio.runtime.errors import ModelCallFailed

    settings = ModelSettingsStore(tmp_path)
    settings.save(ModelSettingsInput(base_url="http://model.test/v1", model="local", no_auth=True))
    runtime = LLMRuntime(settings.path)
    await runtime.client.aclose()
    failing = True
    calls = 0

    async def no_wait(_):
        return

    monkeypatch.setattr("folio.runtime.capacity.asyncio.sleep", no_wait)

    def handler(request):
        nonlocal calls
        if request.url.path == "/v1/models":
            return httpx.Response(200, json={"data": [{"id": "local", "max_model_len": 4096}]})
        calls += 1
        if failing:
            if failure == "connection":
                raise httpx.ReadError("temporary failure", request=request)
            return httpx.Response(int(failure))
        return httpx.Response(200, json={"count": 3000})

    runtime.client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    try:
        with pytest.raises(ModelCallFailed) as error:
            await runtime.measure("query", {"data": "short text"}, Query.model_json_schema(), 48000)
        assert error.value.code == "capacity_unavailable"
        assert calls == 4 and runtime.capacity.counter
        failing = False
        result = await runtime.measure(
            "query", {"data": "short text"}, Query.model_json_schema(), 48000
        )
        assert result["measurement"] == "server_tokenizer"
        assert result["input_tokens"] == 3000 and not result["fits"]
        assert calls == 5
    finally:
        await runtime.aclose()


async def test_length_failure_changes_validation_action_and_is_not_replayed(store, inputs):
    from folio.runtime.errors import ModelCallFailed

    class LongReview(DemoRuntime):
        scopes = []

        async def complete(self, task, context, schema, session):
            if task == "validate":
                scope = context.get("validation_scope")
                self.scopes.append(scope)
                if scope is None:
                    raise ModelCallFailed("output exhausted", "length_limit")
                if scope == "coverage":
                    assert len(context["evidence"]) == 1
                    assert len(context["content"]["blocks"]) == 2
                return RuntimeResult(
                    json.dumps(
                        {
                            "verdicts": [
                                {"claim_id": c["claim_id"], "supported": True, "reason": "来源支持"}
                                for c in context["extraction"]["claims"]
                            ]
                        }
                    )
                )
            return await super().complete(task, context, schema, session)

    runtime = LongReview()
    engine = engine_for(store, inputs, runtime)
    context = {
        "section": {"section_id": "a"},
        "content": {
            "blocks": [
                {"block_id": f"b{i}", "text": "业务正文", "evidence_ids": [f"e{i}"]} for i in (1, 2)
            ]
        },
        "evidence": [evidence(1), evidence(2)],
        "required_project_evidence_ids": ["e1", "e2"],
        "extraction": {
            "claims": [
                {"claim_id": f"c{i}", "block_ids": [f"b{i}"], "evidence_ids": [f"e{i}"]}
                for i in (1, 2)
            ]
        },
    }
    result = await engine.call("validate", context, LocalReview)
    assert len(result.verdicts) == 2
    assert runtime.scopes == [None, "claims", "body", "coverage", "coverage"]
    await engine.call("validate", context, LocalReview)
    assert (
        len(runtime.scopes) == 5
    )  # Parent failure and successful children are reusable checkpoints.


async def test_large_pair_review_preserves_every_cross_node_combination(store, inputs):
    from folio.models import Review

    class Reviewer(DemoRuntime):
        pairs = set()

        async def measure(self, task, context, schema, char_limit):
            count = sum(len(s["claims"]) for s in context.get("sections", []))
            return {"characters": count, "fits": count <= 2}

        async def complete(self, task, context, schema, session):
            if task == "review":
                a, b = context["sections"]
                assert context["review_slice"] is True
                self.pairs.add((a["claims"][0]["claim_id"], b["claims"][0]["claim_id"]))
                return RuntimeResult('{"findings": []}')
            return await super().complete(task, context, schema, session)

    runtime = Reviewer()
    engine = engine_for(store, inputs, runtime)
    context = {
        "sections": [
            {
                "section_id": sid,
                "claims": [{"claim_id": sid + str(i), "statement": "完整事实"} for i in range(3)],
            }
            for sid in ("a", "b")
        ]
    }
    result = await engine.call("review", context, Review)
    assert not result.findings
    assert runtime.pairs == {(f"a{i}", f"b{j}") for i in range(3) for j in range(3)}


async def test_writer_reads_selected_originals_and_retains_mandatory_requirements(store, inputs):
    class SelectOne(DemoRuntime):
        async def complete(self, task, context, schema, session):
            if task == "select_context":
                return RuntimeResult(
                    json.dumps(
                        {
                            "selected_refs": [
                                x["context_ref"]
                                for x in context["catalog"]
                                if x["value"].get("evidence_id") == "e2"
                            ]
                        }
                    )
                )
            return await super().complete(task, context, schema, session)

    engine = engine_for(store, inputs, SelectOne())
    source = {
        "section": {"section_id": "a"},
        "evidence": [evidence(1), evidence(2), evidence(3)],
        "required_project_evidence_ids": ["e1"],
    }
    view = await engine.prepare_context("generate", source)
    assert [e["evidence_id"] for e in view["evidence"]] == ["e1", "e2"]
    assert [e["text"] for e in view["evidence"]] == [e["text"] for e in source["evidence"][:2]]
    assert len(source["evidence"]) == 3


def test_backend_project_sources_can_exceed_one_read_budget(tmp_path):
    from docx import Document

    from folio.document.references import ReferenceDocuments
    from folio.models import Section

    path = tmp_path / "requirements.docx"
    doc = Document()
    doc.add_heading("交易处理", 1)
    doc.add_paragraph("第一组完整规则" * 1000)
    doc.add_paragraph("第二组完整规则" * 1000)
    doc.save(path)
    references = ReferenceDocuments([path])
    section = Section(
        section_id="s",
        node_id="s",
        title="交易处理",
        heading_path=["交易处理"],
        level=1,
        display_number="1",
        instruction="生成正文",
    )
    with pytest.raises(ValueError, match="evidence budget"):
        references.evidence(section, 12000)
    result = references.evidence(section, None)
    assert len(result) == 2
    assert sum(len(e.text) for e in result) == 14000


async def test_unknown_fact_reference_retry_names_bad_field_and_archives_rejected_result(
    store, inputs
):
    class Writer(DemoRuntime):
        calls = 0

        async def complete(self, task, context, schema, session):
            assert task == "understand_work"
            self.calls += 1
            assert schema["$defs"]["BusinessFact"]["properties"]["evidence_ids"]["items"][
                "enum"
            ] == ["e1"]
            if self.calls == 2:
                feedback = context["format_feedback"]
                assert "/facts/0/evidence_ids" in feedback
                assert "wrong_section_ref" in feedback and "e1" in feedback
                assert "do not guess replacement identifiers" in feedback
            return RuntimeResult(
                json.dumps(
                    {
                        "summary": "业务摘要",
                        "facts": [
                            {
                                "fact_id": "f1",
                                "statement": "交易需要审批",
                                "evidence_ids": ["wrong_section_ref" if self.calls == 1 else "e1"],
                            }
                        ],
                    }
                )
            )

    runtime = Writer()
    engine = engine_for(store, inputs, runtime)
    context = {
        "work_task": {"task_id": "t"},
        "sections": [],
        "evidence": [evidence(1)],
        "dependencies": [],
    }
    result = await engine._call("understand_work", context, BusinessRecord)
    assert runtime.calls == 2 and result.facts[0].evidence_ids == ["e1"]
    rejected = list((store.root / "jobs" / engine.job_id / "contexts/rejected").glob("*.json"))
    assert len(rejected) == 1
    record = json.loads(rejected[0].read_text())
    assert record["candidate"]["facts"][0]["evidence_ids"] == ["wrong_section_ref"]
    cached = await engine._call("understand_work", context, BusinessRecord)
    assert cached == result and runtime.calls == 2


async def test_single_table_review_splits_all_rows_after_output_exhaustion(store, inputs):
    from folio.runtime.errors import ModelCallFailed

    class Reviewer(DemoRuntime):
        read_rows = []

        async def complete(self, task, context, schema, session):
            assert task == "validate"
            assert schema["$defs"]["Finding"]["properties"]["description"]["maxLength"] == 240
            rows = context["content"]["blocks"][0]["rows"]
            if len(rows) > 1:
                raise ModelCallFailed("output exhausted", "length_limit")
            position = context["body_slice"]
            assert position["total_rows"] == 5
            assert context["evidence"] == [evidence(1)]
            self.read_rows.append((position["row_start"], rows[0]))
            return RuntimeResult(
                json.dumps(
                    {
                        "verdicts": [],
                        "findings": [
                            {
                                "kind": "unsupported",
                                "description": "本行字段需要核对",
                                "section_ids": ["s"],
                                "key": "row-" + str(position["row_start"]),
                            }
                        ],
                    }
                )
            )

    runtime = Reviewer()
    engine = engine_for(store, inputs, runtime)
    rows = [[f"角色{i}", f"职责{i}"] for i in range(5)]
    context = {
        "section": {"section_id": "s"},
        "validation_scope": "body",
        "content": {
            "blocks": [{"block_id": "table", "kind": "table", "rows": rows, "evidence_ids": ["e1"]}]
        },
        "extraction": {"claims": []},
        "evidence": [evidence(1)],
    }
    result = await engine.call_bounded("validate", context, LocalReview)
    assert sorted(runtime.read_rows) == list(enumerate(rows))
    assert len(result.findings) == 5
    assert context["content"]["blocks"][0]["rows"] == rows


@pytest.mark.parametrize("always_exhausted", [False, True])
async def test_atomic_review_output_retry_is_changed_and_bounded(store, inputs, always_exhausted):
    from folio.runtime.errors import ModelCallFailed

    class Reviewer(DemoRuntime):
        calls = []

        async def complete(self, task, context, schema, session):
            self.calls.append(context)
            if always_exhausted or not context.get("compact_review_retry"):
                raise ModelCallFailed("output exhausted", "length_limit")
            return RuntimeResult('{"verdicts":[],"findings":[]}')

    runtime = Reviewer()
    engine = engine_for(store, inputs, runtime)
    context = {
        "section": {"section_id": "s"},
        "validation_scope": "body",
        "content": {
            "blocks": [
                {
                    "block_id": "p",
                    "kind": "paragraph",
                    "text": "完整的一段正文",
                    "evidence_ids": ["e1"],
                }
            ]
        },
        "extraction": {"claims": []},
        "evidence": [evidence(1)],
    }
    if always_exhausted:
        with pytest.raises(ModelCallFailed):
            await engine.call_bounded("validate", context, LocalReview)
    else:
        result = await engine.call_bounded("validate", context, LocalReview)
        assert result.findings == []
    assert len(runtime.calls) == 2
    assert runtime.calls[1]["compact_review_retry"]
    assert runtime.calls[0]["content"] == runtime.calls[1]["content"]
    assert runtime.calls[0]["evidence"] == runtime.calls[1]["evidence"]


async def test_writer_selection_uses_node_scope_and_keeps_required_originals(store, inputs):
    class NodeReader(DemoRuntime):
        purposes = []

        async def complete(self, task, context, schema, session):
            if task == "select_context":
                self.purposes.append(context["purpose"])
                chosen = []
                for item in context["catalog"]:
                    value = item["value"]
                    if isinstance(value, dict) and value.get("statement") == "本节规则":
                        chosen.append(item["context_ref"])
                risks = [x["context_ref"] for x in context["catalog"] if x["kind"] == "risk"]
                return RuntimeResult(
                    json.dumps(
                        {
                            "selected_refs": chosen,
                            "risk_refs": risks,
                            "risk_summary": "适用约束待核实" if risks else "",
                        }
                    )
                )
            return await super().complete(task, context, schema, session)

    runtime = NodeReader()
    engine = engine_for(store, inputs, runtime)
    source = {
        "work_task": {"objective": "整个交易业务", "title": "所有产品"},
        "section": {"section_id": "option", "title": "外汇期权"},
        "business_record": {
            "facts": [
                {"statement": "本节规则", "evidence_ids": ["e1"]},
                {"statement": "其他产品", "evidence_ids": ["e2"]},
            ],
            "conflicts": ["约束冲突仍需处理"],
        },
        "global_facts": [{"statement": "其他产品", "evidence_ids": ["e2"]}],
        "evidence": [evidence(1), evidence(2), evidence(3)],
        "required_project_evidence_ids": ["e3"],
        "dependencies": [
            {
                "task_id": "upstream",
                "title": "上游规则",
                "sections": [
                    {
                        "section_id": "base",
                        "revision": 1,
                        "supported_facts": [{"statement": "其他产品", "evidence_ids": ["e2"]}],
                    }
                ],
            }
        ],
    }
    before = json.dumps(source, sort_keys=True)
    view = await engine.prepare_context("generate", source)
    assert all(p["section"]["section_id"] == "option" for p in runtime.purposes)
    assert all("work_task" not in p for p in runtime.purposes)
    assert {e["evidence_id"] for e in view["evidence"]} == {"e1", "e3"}
    assert len(view["business_record"]["facts"]) == 1
    assert view["business_record"]["conflicts"] == ["约束冲突仍需处理"]
    assert view["global_facts"] == []
    assert json.dumps(source, sort_keys=True) == before


@pytest.mark.parametrize("operation", ["generate", "repair"])
@pytest.mark.parametrize("choose_relevant", [True, False])
async def test_fitting_writer_still_selects_node_sources_and_preserves_audit_and_risks(
    store, inputs, operation, choose_relevant
):
    class ScopedWriter(DemoRuntime):
        def __init__(self):
            self.purposes = []
            self.writing = []

        async def measure(self, task, context, schema, char_limit):
            return {"characters": 100, "fits": True}

        async def complete(self, task, context, schema, session):
            if task == "select_context":
                self.purposes.append(context["purpose"])
                selected = [
                    item["context_ref"]
                    for item in context["catalog"]
                    if choose_relevant
                    and item["kind"] == "fact"
                    and item["value"].get("statement") == "静态数据维护规则"
                ]
                risks = [
                    item["context_ref"] for item in context["catalog"] if item["kind"] == "risk"
                ]
                return RuntimeResult(
                    json.dumps(
                        {
                            "selected_refs": selected,
                            "risk_refs": risks,
                            "risk_summary": "适用约束待核实" if risks else "",
                        }
                    )
                )
            if task == operation:
                self.writing.append(context)
                return RuntimeResult('{"blocks":[]}')
            raise AssertionError("Unexpected task: " + task)

    runtime = ScopedWriter()
    engine = engine_for(store, inputs, runtime)
    source = {
        "section": {
            "section_id": "static",
            "title": "数据维护功能",
            "heading_path": ["通用静态数据", "数据维护功能"],
            "instruction": "描述静态数据录入、修改和查询。",
        },
        "evidence": [
            evidence(1, "静态数据维护的完整规则"),
            evidence(2, "另一节点的交易录入及现金流处理规则"),
            evidence(3, "当前节点强制覆盖的项目约束"),
        ],
        "business_record": {
            "facts": [
                {"statement": "静态数据维护规则", "evidence_ids": ["e1"]},
                {"statement": "交易维护规则", "evidence_ids": ["e2"]},
            ],
            "conflicts": ["现有来源存在冲突，不能自行裁决"],
            "open_questions": ["本节点尚缺权限规则"],
        },
        "global_facts": [],
        "required_project_evidence_ids": ["e3"],
    }
    if operation == "repair":
        source.update(current={"blocks": []}, repair_instruction="纠正业务对象错位")
    original = json.dumps(source, sort_keys=True)
    await engine.call(operation, source, SectionContent)
    assert runtime.purposes and len(runtime.writing) == 1
    assert all(purpose["section"] == source["section"] for purpose in runtime.purposes)
    view = runtime.writing[0]
    expected_sources = {"e1", "e3"} if choose_relevant else {"e3"}
    assert {item["evidence_id"] for item in view["evidence"]} == expected_sources
    assert len(view["business_record"]["facts"]) == int(choose_relevant)
    for field in ("conflicts", "open_questions"):
        assert view["business_record"][field] == source["business_record"][field]
    assert json.dumps(source, sort_keys=True) == original
    archives = (store.root / "jobs" / engine.job_id / "contexts" / "sources").glob("*.json")
    assert any(json.loads(path.read_text())["context"] == source for path in archives)
    before = (len(runtime.purposes), len(runtime.writing))
    await engine.call(operation, source, SectionContent)
    assert (len(runtime.purposes), len(runtime.writing)) == before


async def test_research_pages_run_concurrently_and_merge_all_sources(store, inputs):
    class Reader(DemoRuntime):
        active = 0
        peak = 0

        async def measure(self, task, context, schema, char_limit):
            return {"characters": 100, "fits": len(context.get("evidence", [])) <= 2}

        async def complete(self, task, context, schema, session):
            if task == "understand_work":
                self.active += 1
                self.peak = max(self.peak, self.active)
                try:
                    await asyncio.sleep(0.01)
                    return await super().complete(task, context, schema, session)
                finally:
                    self.active -= 1
            return await super().complete(task, context, schema, session)

    runtime = Reader()
    engine = engine_for(store, inputs, runtime)
    result = await engine.call_bounded(
        "understand_work",
        {
            "work_task": {"objective": "并发读取所有规则"},
            "evidence": [evidence(i) for i in range(8)],
        },
        BusinessRecord,
    )
    assert runtime.peak == 4
    assert runtime.active == 0
    assert {eid for fact in result.facts for eid in fact.evidence_ids} == {
        f"e{i}" for i in range(8)
    }


async def test_gap_batches_share_snapshot_and_merge_in_input_order(store, inputs):
    from folio.orchestration.work import GapProposal

    engine = engine_for(store, inputs, DemoRuntime())
    active = peak = 0
    seen = []
    context = {
        "gaps": [{"gap_id": str(i), "question": f"业务缺口{i}"} for i in range(6)],
        "sections": [{"section_id": "s"}],
        "existing_proposals": [{"gap_id": "prior", "content": "已有研究方案"}],
    }
    original = json.dumps(context, sort_keys=True)

    async def complete(task, child, model):
        nonlocal active, peak
        active += 1
        peak = max(peak, active)
        seen.append(child)
        try:
            await asyncio.sleep(0.02 if child["gaps"][0]["gap_id"] == "0" else 0)
            return GapCompletion(
                proposals=[
                    GapProposal(
                        gap_id=g["gap_id"],
                        section_ids=["s"],
                        content="可供本节使用的完整业务设计方案内容" * 2,
                    )
                    for g in child["gaps"]
                ]
            )
        finally:
            active -= 1

    engine.call = complete
    result = await engine.split_gaps(context, GapCompletion)
    assert peak == 2 and active == 0
    assert [p.gap_id for p in result.proposals] == [str(i) for i in range(6)]
    assert all(c["existing_proposals"] == context["existing_proposals"] for c in seen)
    assert json.dumps(context, sort_keys=True) == original


async def test_capacity_recovers_after_three_transient_failures(monkeypatch):
    from types import SimpleNamespace

    from folio.runtime.capacity import Capacity

    calls, waits = 0, []

    async def request(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls <= 3:
            raise httpx.ConnectError("temporarily unavailable")
        return httpx.Response(200, json={"count": 100})

    async def sleep(seconds):
        waits.append(seconds)

    monkeypatch.setattr("folio.runtime.capacity.asyncio.sleep", sleep)
    capacity = Capacity(SimpleNamespace(config={}, client=SimpleNamespace(request=request)))
    result = await capacity.request("POST", "http://model.test/tokenize")
    assert result.json()["count"] == 100
    assert calls == 4 and waits == [2, 4, 8]


async def test_capacity_wait_budget_failure_stays_explicit():
    from types import SimpleNamespace

    from folio.runtime.capacity import Capacity
    from folio.runtime.errors import ModelCallFailed

    async def request(*args, **kwargs):
        raise TimeoutError("deadline")

    capacity = Capacity(SimpleNamespace(config={}, client=SimpleNamespace(request=request)))
    with pytest.raises(ModelCallFailed) as error:
        await capacity.request("POST", "http://model.test/tokenize")
    assert error.value.code == "capacity_unavailable"


@pytest.mark.parametrize("all_required", [False, True])
async def test_overbroad_decision_selection_gets_one_bounded_refinement(
    store, inputs, monkeypatch, all_required
):
    from folio.models import SectionContent

    monkeypatch.setattr(
        "folio.model_settings.feature_enabled", lambda _, component: component == "decision"
    )

    class Reader(DemoRuntime):
        config = {}

        def __init__(self):
            self.selections = []
            self.writes = []

        async def measure(self, task, context, schema, char_limit):
            fits = task != "generate" or len(context.get("evidence", [])) <= 2
            return {"characters": 100 if fits else 50000, "fits": fits}

        async def complete(self, task, context, schema, session):
            if task == "select_context":
                retry = context.get("selection_capacity_retry", False)
                self.selections.append(retry)
                selected = [
                    x["context_ref"]
                    for x in context["catalog"]
                    if not retry or all_required or x["kind"] == "fact"
                ]
                return RuntimeResult(json.dumps({"selected_refs": selected}))
            if task == "generate":
                self.writes.append(context)
                return RuntimeResult(
                    json.dumps(
                        {
                            "blocks": [
                                {
                                    "block_id": "b1",
                                    "text": "本节业务规则",
                                    "origin": "evidence",
                                    "evidence_ids": ["e1"],
                                }
                            ]
                        }
                    )
                )
            return await super().complete(task, context, schema, session)

    runtime = Reader()
    engine = engine_for(store, inputs, runtime)
    source = {
        "section": {"section_id": "s", "title": "本节"},
        "evidence": [evidence(1), evidence(2), evidence(3)],
        "required_project_evidence_ids": ["e3"],
        "business_record": {"facts": [{"statement": "必需业务约束", "evidence_ids": ["e1"]}]},
    }
    original = json.dumps(source, sort_keys=True)
    if all_required:
        with pytest.raises(ValueError, match="smallest safe request"):
            await engine.call("generate", source, SectionContent)
        assert not runtime.writes
    else:
        result = await engine.call("generate", source, SectionContent)
        assert result.blocks[0].evidence_ids == ["e1"]
        assert {e["evidence_id"] for e in runtime.writes[0]["evidence"]} == {"e1", "e3"}
    assert runtime.selections == [False, True]
    assert json.dumps(source, sort_keys=True) == original
