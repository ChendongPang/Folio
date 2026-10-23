"""Execution barriers and proposal provenance, not just scheduler implementation details."""

import asyncio
import json

import pytest
from docx import Document
from test_business_work import GroupRuntime, task

from folio.models import DocumentPlan, JobInput, Limits
from folio.orchestration.engine import Engine
from folio.orchestration.state import hierarchical_state
from folio.orchestration.work import WorkPlan
from folio.runtime.base import RuntimeResult


async def test_ready_siblings_overlap_child_starts_before_unrelated_task_finishes(store, inputs):
    inputs["user_request"] = "# 交易规则\n# 交易权限\n# 交易审批"

    class Concurrent(GroupRuntime):
        def __init__(self):
            super().__init__()
            self.slow_started = asyncio.Event()
            self.child_started = asyncio.Event()
            self.parent_reviewed = False
            self.slow_finished = False

        async def complete(self, name, context, schema, session):
            if name == "work_plan":
                ids = [s["section_id"] for s in context["sections"]]
                return RuntimeResult(
                    WorkPlan(
                        tasks=[
                            task("child", [ids[2]], ["parent"]),
                            task("parent", [ids[0]]),
                            task("slow", [ids[1]]),
                        ]
                    ).model_dump_json()
                )
            owner = context.get("work_task", {}).get("task_id")
            if name == "generate":
                if owner == "parent":
                    await asyncio.wait_for(self.slow_started.wait(), 3)
                elif owner == "slow":
                    self.slow_started.set()
                    await asyncio.wait_for(self.child_started.wait(), 3)
                    self.slow_finished = True
                elif owner == "child":
                    assert self.parent_reviewed and not self.slow_finished
                    assert context["dependencies"][0]["sections"][0]["supported_facts"]
                    self.child_started.set()
            if name == "review_work" and owner == "parent":
                self.parent_reviewed = True
            return await super().complete(name, context, schema, session)

    runtime = Concurrent()
    job_id = store.create(JobInput(**inputs, runtime="demo", limits=Limits(max_parallel_tasks=2)))
    job = await Engine(store, runtime).run(job_id)
    assert job["status"] == "draft"
    assert len(store.states(job_id)) == 3
    assert {r["section_id"] for r in job["snapshot"]["records"].values()} == {
        s.section_id for s in store.states(job_id)
    }  # All concurrent commits survive in the derived global registry.
    work = store.work_states(job_id)
    assert all(w.review_key for w in work.values())
    assert all(not facts for facts in work["parent"].global_facts.values())
    assert all(not facts for facts in work["slow"].global_facts.values())
    child_facts = [f for fs in work["child"].global_facts.values() for f in fs]
    assert child_facts and all(
        f["section_id"] in store.work_plan(job_id).tasks[1].section_ids for f in child_facts
    )
    # Conflicting values coexist; there is no last-writer-wins dictionary by business key.
    states = store.states(job_id)
    for i, state in enumerate(states):
        claim = state.extraction.claims[0]
        claim.subject, claim.predicate, claim.scope = "交易", "审批角色", "一期"
        claim.value = str(i)
        claim.cardinality = "single"
    view = hierarchical_state(DocumentPlan.model_validate(job["plan"]), states)
    assert len(view["disputed_keys"]) == 1
    assert len(view["fact_index"][view["disputed_keys"][0]]) == 3
    for state in states:
        state.extraction.claims[0].cardinality = "multiple"
    enumerations = hierarchical_state(DocumentPlan.model_validate(job["plan"]), states)
    assert not enumerations["disputed_keys"]
    assert len(enumerations["records"]) == len(view["records"])


async def test_parallel_failure_cancels_and_drains_siblings_before_return(store, inputs):
    class Failure(GroupRuntime):
        def __init__(self):
            super().__init__()
            self.waiting = asyncio.Event()
            self.drained = False

        async def complete(self, name, context, schema, session):
            if name == "work_plan":
                ids = [s["section_id"] for s in context["sections"]]
                return RuntimeResult(
                    WorkPlan(
                        tasks=[task("fail", [ids[0]]), task("wait", [ids[1]])]
                    ).model_dump_json()
                )
            if name == "generate":
                if context["work_task"]["task_id"] == "fail":
                    await asyncio.wait_for(self.waiting.wait(), 3)
                    raise RuntimeError("planned failure")
                self.waiting.set()
                try:
                    await asyncio.Event().wait()
                finally:
                    self.drained = True
            return await super().complete(name, context, schema, session)

    runtime = Failure()
    job_id = store.create(JobInput(**inputs, runtime="demo", limits=Limits(call_attempts=1)))
    with pytest.raises(RuntimeError, match="planned failure"):
        await Engine(store, runtime).run(job_id)
    assert runtime.drained
    assert not store.states(job_id)
    assert store.get(job_id)["status"] == "failed"
    assert not any(t.get_name() in {"wait", "fail"} for t in asyncio.all_tasks())


async def test_concurrency_limit_one_keeps_execution_serial(store, inputs):
    class Sequential(GroupRuntime):
        def __init__(self):
            super().__init__()
            self.first_reviewed = False

        async def complete(self, name, context, schema, session):
            if name == "work_plan":
                ids = [s["section_id"] for s in context["sections"]]
                return RuntimeResult(
                    WorkPlan(tasks=[task("a", [ids[0]]), task("b", [ids[1]])]).model_dump_json()
                )
            owner = context.get("work_task", {}).get("task_id")
            if name == "next_action" and owner == "b":
                assert self.first_reviewed
            if name == "review_work" and owner == "a":
                self.first_reviewed = True
            return await super().complete(name, context, schema, session)

    job_id = store.create(JobInput(**inputs, runtime="demo", limits=Limits(max_parallel_tasks=1)))
    assert (await Engine(store, Sequential()).run(job_id))["status"] == "draft"


async def test_completion_retries_incomplete_coverage_and_cannot_promote_suggestions(store, inputs):
    class Proposals(GroupRuntime):
        attempts = 0

        async def complete(self, name, context, schema, session):
            result = await super().complete(name, context, schema, session)
            if result.tool_calls:
                if name == "next_action" and context["evidence"]:
                    data = json.loads(result.tool_calls[0]["function"]["arguments"])
                    data["missing_information"] = ["如何复核", "如何处理异常"]
                    result.tool_calls[0]["function"]["arguments"] = json.dumps(data)
                return result
            data = json.loads(result.content)
            if name == "complete_gaps":
                self.attempts += 1
                if self.attempts == 1:
                    data["proposals"] = data["proposals"][:1]
                else:
                    assert "every requested gap" in context["format_feedback"]
            if name == "extract":
                proposal = next(
                    b for b in context["content"]["blocks"] if b["origin"] == "model_proposal"
                )
                data["claims"].append(
                    {
                        "claim_id": "fake_fact",
                        "block_ids": [proposal["block_id"]],
                        "statement": proposal["text"],
                        "subject": "交易",
                        "predicate": "规则",
                        "value": "建议",
                        "evidence_ids": [context["evidence"][0]["evidence_id"]],
                        "status": "supported",
                    }
                )
            return RuntimeResult(json.dumps(data))

    runtime = Proposals()
    job_id = store.create(JobInput(**inputs, runtime="llm"))
    job = await Engine(store, runtime).run(job_id)
    assert runtime.attempts == 2 and job["status"] == "draft"
    assert not job["snapshot"]["blocking_issues"]
    assert all(c.claim_id != "fake_fact" for s in store.states(job_id) for c in s.extraction.claims)
    doc = Document(job["artifact"])
    assert "【模型建议】" not in "".join(p.text for p in doc.paragraphs)
    assert any(str(r.font.color.rgb) == "FF0000" for p in doc.paragraphs for r in p.runs)
    with store.connect() as conn:
        assert conn.execute("SELECT published_version FROM documents").fetchone()[0] is None


async def test_completion_checkpoint_resume_reuses_proposals_and_writing_gaps(store, inputs):
    class WritingGap(GroupRuntime):
        async def complete(self, name, context, schema, session):
            result = await super().complete(name, context, schema, session)
            if name == "generate":
                data = json.loads(result.content)
                data["pending_items"] = ["重复提交如何处理"]
                return RuntimeResult(json.dumps(data))
            return result

    class Interrupted(Engine):
        async def complete_gaps(self, *args, **kwargs):
            await super().complete_gaps(*args, **kwargs)
            raise asyncio.CancelledError()

    runtime = WritingGap()
    job_id = store.create(JobInput(**inputs, runtime="llm"))
    with pytest.raises(asyncio.CancelledError):
        await Interrupted(store, runtime).run(job_id)
    checkpoint = store.work_states(job_id)["business"]
    assert len(checkpoint.proposals) == 1
    original_proposal = checkpoint.proposals[0].model_copy(deep=True)
    job = await Engine(store, runtime).run(job_id)
    work = store.work_states(job_id)["business"]
    assert job["status"] == "draft" and not job["snapshot"]["blocking_issues"]
    assert work.missing == ["重复提交如何处理"]
    assert len(work.proposals) == 1 and len(work.proposals[0].section_ids) == 2
    assert work.proposals[0].gap_id == original_proposal.gap_id
    assert work.proposals[0].content == original_proposal.content
    assert len(work.searches) == 3
    assert len([c for name, c in runtime.calls if name == "complete_gaps"]) == 1
    assert all(not state.content.pending_items for state in store.states(job_id))
    integrations = [c for name, c in runtime.calls if name == "integrate_supplement"]
    # Recovery writes the already-persisted proposal verbatim into each node.
    # The repeated gap must not trigger a second editor call for identical text.
    assert integrations == []
    for state in store.states(job_id):
        assert any(
            block.text == original_proposal.content
            and block.origin == "model_proposal"
            and not block.evidence_ids
            for block in state.content.blocks
        )
    records = list(
        (store.root / "jobs" / job_id / "contexts" / "already_integrated_proposals").glob("*.json")
    )
    matches = [item for p in records for item in json.loads(p.read_text())["matches"]]
    assert {item["section_id"] for item in matches} == {s.section_id for s in store.states(job_id)}
    assert all(item["gap_id"] == original_proposal.gap_id for item in matches)
    for state in store.states(job_id):
        adopted = [
            block for block in state.content.blocks if block.text == original_proposal.content
        ]
        assert len(adopted) == 1
        assert adopted[0].origin == "model_proposal" and adopted[0].evidence_ids == []


async def test_suggested_table_keeps_required_headers(store, inputs):
    from test_business_work import document

    class Tables(GroupRuntime):
        async def complete(self, name, context, schema, session):
            if name in {"generate", "repair"}:
                return RuntimeResult(
                    json.dumps(
                        {
                            "blocks": [
                                {
                                    "block_id": "table",
                                    "kind": "table",
                                    "origin": "model_proposal",
                                    "rows": [
                                        ["字段", "处理规则"],
                                        ["交易编号", "建议重复提交时返回原处理结果"],
                                    ],
                                }
                            ]
                        }
                    )
                )
            return await super().complete(name, context, schema, session)

    plan = document(["a"])
    plan.sections[0].expected_table_headers = ["字段", "处理规则"]
    job_id = store.create(JobInput(**inputs, runtime="llm"), plan=plan)
    runtime = Tables()
    job = await Engine(store, runtime).run(job_id)
    assert job["status"] == "draft" and not job["snapshot"]["blocking_issues"]
    blocks = store.states(job_id)[0].content.blocks
    tables = [block for block in blocks if block.kind == "table"]
    assert len(tables) == 1 and tables[0].origin == "model_proposal"
    assert tables[0].rows == [["字段", "处理规则"], ["交易编号", "建议重复提交时返回原处理结果"]]
    assert tables[0].evidence_ids == []
    integrations = [c for name, c in runtime.calls if name == "integrate_supplement"]
    assert integrations
    supplemented = [block for block in blocks if block.kind != "table"]
    assert supplemented
    assert {block.text for block in supplemented} <= {
        c["integration_item"]["content"] for c in integrations
    }
    assert all(
        block.origin == "model_proposal" and block.evidence_ids == [] for block in supplemented
    )
    assert Document(job["artifact"]).tables[0].cell(0, 0).text == "字段"
