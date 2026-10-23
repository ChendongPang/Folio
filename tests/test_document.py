from docx import Document

from folio.document.planner import build_plan
from folio.document.word import render
from folio.models import Block, Extraction, PlanItem, PlanProposal, SectionContent, SectionState


def test_directory_identity_and_hierarchy_do_not_require_source_references():
    proposal = PlanProposal(
        sections=[
            PlanItem(node_id="a", title="交易", level=1, instruction="描述需求"),
            PlanItem(node_id="b", title="撤销", level=2, instruction="说明审批"),
        ]
    )
    first = build_plan("doc1", proposal)
    assert first == build_plan("doc1", proposal)
    assert first.sections[1].parent_id == first.sections[0].section_id
    assert first.sections[1].display_number == "1.1"


def test_table_content_is_rendered_with_builtin_document_style(tmp_path):
    plan = build_plan(
        "doc",
        PlanProposal(
            sections=[PlanItem(node_id="trade", title="交易", level=1, instruction="说明字段")]
        ),
    )
    state = SectionState(
        extraction=Extraction(summary="表格"),
        evidence=[],
        section_id=plan.sections[0].section_id,
        content=SectionContent(
            blocks=[
                Block(block_id="table", kind="table", rows=[["字段", "说明"], ["金额", "交易金额"]])
            ]
        ),
    )
    target = tmp_path / "output.docx"
    render(target, plan, [state], title="测试需求说明书", draft=True)
    document = Document(target)
    assert any(table.cell(1, 1).text == "交易金额" for table in document.tables)
    assert any("交易" in p.text for p in document.paragraphs)
    assert "{{folio.content}}" not in "\n".join(p.text for p in document.paragraphs)
