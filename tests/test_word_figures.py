"""Real PNGs retain anchors, source colors and bounded sizes in built-in Word output."""

import struct
import zlib
from copy import deepcopy
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from folio.document.reader import heading_level, xpath
from folio.document.typography import BODY
from folio.document.word import render
from folio.models import Block, DocumentPlan, Extraction, Section, SectionContent, SectionState


def png(path, width, height, color=(40, 90, 150)):
    def chunk(kind, data):
        return (
            struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
        )

    scanline = b"\x00" + bytes(color) * width
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(scanline * height))
        + chunk(b"IEND", b"")
    )
    return path


@pytest.fixture
def figure_inputs(tmp_path):
    plan = DocumentPlan(
        sections=[
            Section(
                section_id=key,
                node_id=f"o_{key}",
                title=f"业务{key}",
                level=1,
                display_number=str(index),
                instruction="业务说明",
            )
            for index, key in enumerate(("s1", "s2"), 1)
        ]
    )
    states = [
        SectionState(
            section_id=key,
            evidence=[],
            extraction=Extraction(summary="说明"),
            content=SectionContent(
                blocks=[
                    Block(block_id="start", text=f"{key}业务正文开始。"),
                    Block(
                        block_id="table",
                        kind="table",
                        rows=[["字段", "说明"], ["金额", "输入金额"]],
                    ),
                    Block(block_id="end", text=f"{key}业务正文结束。"),
                ]
            ),
        )
        for key in ("s1", "s2")
    ]
    image = png(tmp_path / "界面预览.png", 1600, 800)
    tall = png(tmp_path / "流程图.png", 400, 2400, (150, 50, 40))
    # Manifest ordering deliberately differs from insertion/document ordering.
    figures = {
        "s2": [
            dict(
                figure_id="last",
                caption="最终界面",
                image_path=str(image),
                width=1600,
                height=800,
                origin="evidence",
            )
        ],
        "s1": [
            dict(
                figure_id="tail",
                caption="完成状态",
                image_path=str(image),
                width=1600,
                height=800,
                origin="evidence",
            ),
            dict(
                figure_id="flow",
                caption="审批流程",
                image_path=str(tall),
                width=400,
                height=2400,
                origin="model_proposal",
                after_block_id="table",
            ),
            dict(
                figure_id="ui",
                caption="交易录入",
                image_path=str(image),
                width=1600,
                height=800,
                origin="evidence",
                after_block_id="start",
            ),
        ],
    }
    return plan, states, figures


def output(tmp_path, inputs, figures):
    plan, states, _ = inputs
    target = tmp_path / "output.docx"
    render(target, plan, states, title="需求说明书", draft=True, figures=figures)
    return target


def test_figures_embed_at_real_blocks_with_fonts_color_alt_and_page_bounds(
    tmp_path, figure_inputs
):
    target = output(tmp_path, figure_inputs, figure_inputs[2])
    doc = Document(target)
    captions = [p for p in doc.paragraphs if p.text.startswith("图 ")]
    assert [p.text for p in captions] == [
        "图 1  交易录入",
        "图 2  审批流程（拟议设计）",
        "图 3  完成状态",
        "图 4  最终界面",
    ]
    for p in captions:
        assert p.style.name == BODY
        assert p.style.font.name == "Calibri" and p.style.font.size == Pt(11)
        assert xpath(p.style.element, "./w:rPr/w:rFonts/@w:eastAsia") == ["宋体"]
        assert heading_level(p._p, doc) is None
        assert p.paragraph_format.keep_with_next is False
        assert not p._p.xpath(".//w:bookmarkStart")
    assert str(captions[1].runs[0].font.color.rgb) == "FF0000"
    assert captions[0].runs[0].font.color.rgb is None
    assert len(doc.inline_shapes) == 4
    with ZipFile(target) as archive:
        assert len([name for name in archive.namelist() if name.startswith("word/media/")]) == 2
    section = doc.sections[-1]
    for shape, expected_ratio in zip(doc.inline_shapes, [2, 1 / 6, 2, 2], strict=True):
        assert shape.width <= section.page_width - section.left_margin - section.right_margin
        assert shape.height < section.page_height - section.top_margin - section.bottom_margin
        assert shape.width / shape.height == pytest.approx(expected_ratio, rel=1e-5)
        assert shape._inline.docPr.get("descr").startswith("图 ")
        assert str(tmp_path) not in shape._inline.xml
    elements = list(doc.element.body)
    first_caption = elements.index(captions[0]._p)
    assert "s1业务正文开始。" in "".join(elements[first_caption - 2].itertext())
    assert elements[first_caption - 1].xpath(".//w:drawing")
    assert elements[first_caption + 1].tag == qn("w:tbl")
    flow_caption = elements.index(captions[1]._p)
    assert elements[flow_caption - 2].tag == qn("w:tbl")
    assert "s1业务正文结束。" in "".join(elements[flow_caption + 1].itertext())
    # Existing body text and table values are untouched.
    for key in ("s1", "s2"):
        assert sum(p.text == f"{key}业务正文开始。" for p in doc.paragraphs) == 1
        assert sum(p.text == f"{key}业务正文结束。" for p in doc.paragraphs) == 1
    assert all(t.cell(1, 1).text == "输入金额" for t in doc.tables)


@pytest.mark.parametrize(
    "invalid", ["missing", "dimensions", "block", "section", "duplicate", "origin", "caption"]
)
def test_invalid_figure_manifest_fails_explicitly_before_saving(
    tmp_path, figure_inputs, invalid
):
    figures = deepcopy(figure_inputs[2])
    first = figures["s1"][0]
    if invalid == "missing":
        first["image_path"] = str(tmp_path / "missing.png")
    elif invalid == "dimensions":
        first["width"] += 1
    elif invalid == "block":
        first["after_block_id"] = "no_such_block"
    elif invalid == "section":
        figures["unknown"] = figures.pop("s1")
    elif invalid == "duplicate":
        first["figure_id"] = figures["s1"][1]["figure_id"]
    elif invalid == "origin":
        first["origin"] = "untraced"
    else:
        first["caption"] = " "
    with pytest.raises(ValueError):
        output(tmp_path, figure_inputs, figures)
    assert not (tmp_path / "output.docx").exists()


def test_empty_manifest_preserves_the_existing_output(tmp_path, figure_inputs):
    target = output(tmp_path, figure_inputs, None)
    with ZipFile(target) as archive:
        before = archive.read("word/document.xml")
    plan, states, _ = figure_inputs
    render(target, plan, states, title="需求说明书", draft=True, figures={})
    with ZipFile(target) as archive:
        assert archive.read("word/document.xml") == before
    assert not Document(target).inline_shapes


def test_figures_respect_builtin_page_size(tmp_path, figure_inputs):
    plan, states, figures = figure_inputs
    target = tmp_path / "sections.docx"
    render(target, plan, states, title="测试", draft=True, figures=figures)
    document = Document(target)
    section = document.sections[0]
    usable = section.page_width - section.left_margin - section.right_margin
    assert all(shape.width <= usable for shape in document.inline_shapes)


def test_render_embeds_manifest_and_returns_pagination_status(tmp_path, figure_inputs, monkeypatch):
    import folio.document.word as word

    monkeypatch.setattr(word, "paginate", lambda *args, **kwargs: {"status": "test"})
    plan, states, figures = figure_inputs
    target = tmp_path / "delegated.docx"
    result = word.render(target, plan, states, title="测试", draft=False, figures=figures)
    assert result == {"status": "test"}
    assert len(Document(target).inline_shapes) == 4


def test_unresolved_figure_keeps_text_and_visible_position_without_inserting_bad_image(
    tmp_path, figure_inputs
):
    figures = figure_inputs[2]
    failed = next(item for item in figures["s1"] if item["figure_id"] == "flow")
    failed.update(status="unresolved", source_block_ids=["table"], reason="内部节点错误 N_Book")
    failed.pop("image_path")
    target = output(tmp_path, figure_inputs, figures)
    doc = Document(target)
    assert len(doc.inline_shapes) == 3
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "配图待确认：未通过校验，暂未插入" in text
    assert "s1业务正文开始" in text and "s1业务正文结束" in text
    assert "N_Book" not in text
    assert "图 2  审批流程（拟议设计）" in text


def test_figure_anchor_follows_all_paragraphs_of_one_business_block(
    tmp_path, figure_inputs
):
    figure_inputs[1][0].content.blocks[0].text = "业务包括：1. 录入金额；2. 确认提交。"
    target = output(tmp_path, figure_inputs, figure_inputs[2])
    paragraphs = Document(target).paragraphs
    caption_index = next(i for i, p in enumerate(paragraphs) if p.text == "图 1  交易录入")
    assert paragraphs[caption_index - 1]._p.xpath(".//w:drawing")
    assert paragraphs[caption_index - 2].text.replace("\t", " ").strip() == "2. 确认提交。"


def test_figure_after_parent_lead_in_is_not_lost(tmp_path, figure_inputs):
    plan, states, figures = figure_inputs
    plan.sections[1].parent_id = "s1"
    plan.sections[1].level = 2
    plan.sections[1].display_number = "1.1"
    states[0].content.blocks[0] = Block(block_id="start", kind="heading", text="业务概述")
    target = tmp_path / "parent.docx"
    render(target, plan, states, title="测试", draft=False, figures=figures)
    doc = Document(target)
    assert len(doc.inline_shapes) == 4
    index = next(i for i, p in enumerate(doc.paragraphs) if p.text == "图 1  交易录入")
    assert doc.paragraphs[index - 2].text == "业务概述"
    assert heading_level(doc.paragraphs[index - 2]._p, doc) is None


def test_unavailable_section_body_does_not_silently_drop_figures(tmp_path, figure_inputs):
    plan, states, figures = figure_inputs
    with pytest.raises(ValueError, match="no body state"):
        render(
            tmp_path / "missing-body.docx",
            plan,
            states[:1],
            title="测试",
            draft=False,
            figures=figures,
        )


def test_figure_ids_are_local_to_each_document_section(tmp_path, figure_inputs):
    figures = deepcopy(figure_inputs[2])
    figures["s2"][0]["figure_id"] = figures["s1"][0]["figure_id"]
    target = output(tmp_path, figure_inputs, figures)
    doc = Document(target)
    assert len(doc.inline_shapes) == 4
    assert [p.text for p in doc.paragraphs if p.text.startswith("图 ")] == [
        "图 1  交易录入",
        "图 2  审批流程（拟议设计）",
        "图 3  完成状态",
        "图 4  最终界面",
    ]
