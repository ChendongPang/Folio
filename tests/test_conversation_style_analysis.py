import asyncio
import json
import re
import threading
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt

from folio.api import style_analysis
from folio.api.style_analysis import ConversationStyles, profile_context
from folio.runtime.base import RuntimeResult
from folio.runtime.errors import ModelCallBudgetExceeded

CID = "a" * 32
OTHER = "b" * 32


class FakeRuntime:
    config = {"supports_image_input": False}

    def __init__(self, config):
        self.calls = []
        self.closed = False

    async def complete(self, task, context, schema, session_key):
        self.calls.append(task)
        text = context.get("style_analysis_text", "")
        ids = re.findall(r"### (D\d+-S\d+) /", text)
        return RuntimeResult(json.dumps({
            "rules": [{"dimension": dimension, "scope": "body", "description": description, "sample_ids": ids[:2], "confidence": "observed"}
                      for dimension, description in {
                          "writing": "使用正式客观的语言，避免宣传性表达。",
                          "structure": "正文分成多个层次清楚的小节。",
                          "organization": "按主题展开背景与说明。",
                          "detail": "使用完整语句展开内容。",
                          "headings": "标题以简明主题词命名。",
                          "expression": "同一类概念采用一致的称谓。",
                      }.items()] if ids else [],
            "dimension_checks": [{"dimension": "figures_tables", "status": "not_present", "reason": "已读文字未出现图表说明。", "sample_ids": ids[:2]}] if ids else [],
            "request_sample_ids": [], "uncertainties": [],
        }, ensure_ascii=False))

    async def aclose(self):
        self.closed = True


@pytest.fixture
def runtimes(monkeypatch):
    instances = []

    def factory(config):
        runtime = FakeRuntime(config)
        instances.append(runtime)
        return runtime

    monkeypatch.setattr(style_analysis, "LLMRuntime", factory)
    return instances


@pytest.fixture
def attachment(tmp_path):
    path = tmp_path / "示例.md"
    path.write_text("# 秘密公司原始项目\n" + "\n".join(f"示例业务正文{i}，说明材料写法与形式。" for i in range(12)), encoding="utf-8")
    return {"id": "1" * 32, "name": path.name, "path": str(path), "role": "style_reference"}


async def test_standalone_analysis_freezes_sources_and_emits_bounded_safe_summary(tmp_path, attachment, runtimes):
    service = ConversationStyles(tmp_path / "state")
    events = []
    report = await service.analyze("unused.json", CID, [attachment], on_progress=events.append)
    assert report["status"] == "analyzed"
    assert report["style_url"] == f"/conversations/{CID}/styles/{report['revision']}"
    assert report["document_names"] == ["示例.md"]
    assert len(report["source_versions"][0]["sha256"]) == 64
    assert "写作风格" in report["summary"]
    assert "专业表达习惯" in report["summary"]
    assert "秘密公司" not in report["summary"]
    assert "示例业务正文" not in report["summary"]
    assert len(report["summary"]) <= 6000
    saved = service.load(report["revision"], CID)
    assert saved["sample_paths"][0].read_bytes() == Path(attachment["path"]).read_bytes()
    assert (saved["folder"] / "style" / "style.md").is_file()
    assert (saved["folder"] / "style" / "audit" / "request-01.json").is_file()
    assert events
    assert runtimes[0].calls == ["style_analysis"]
    assert runtimes[0].closed
    assert attachment["role"] == "style_reference"


async def test_repeat_uses_content_cache_and_closes_unused_runtime(tmp_path, attachment, runtimes):
    service = ConversationStyles(tmp_path / "state")
    first = await service.analyze(None, CID, [attachment])
    second = await service.analyze(None, CID, [attachment])
    assert first["revision"] != second["revision"]
    assert second["cache_hit"] is True
    assert runtimes[0].calls == ["style_analysis"]
    assert runtimes[1].calls == []
    assert all(runtime.closed for runtime in runtimes)
    assert list((tmp_path / "state" / "source-cache" / "styles").glob("*.json"))


async def test_context_requires_live_role_id_hash_and_owner(tmp_path, attachment, runtimes):
    service = ConversationStyles(tmp_path / "state")
    report = await service.analyze(None, CID, [attachment])
    text = service.context([report], [attachment], conversation_id=CID)
    assert "正式客观" in text
    assert "line 1" not in text
    assert "秘密公司" not in text
    assert service.context([report], [attachment | {"role": "source"}], conversation_id=CID) == ""
    assert service.context([report], [attachment | {"id": "2" * 32}], conversation_id=CID) == ""
    assert service.context([report], [attachment], conversation_id=OTHER) == ""
    assert len(service.context([report], [attachment], max_chars=300, conversation_id=CID)) <= 300
    Path(attachment["path"]).write_text("内容已经更换", encoding="utf-8")
    assert service.context([report], [attachment], conversation_id=CID) == ""
    assert service.load(report["revision"], CID)["profile"]["source_hashes"] == [report["source_versions"][0]["sha256"]]


async def test_load_rejects_wrong_owner_traversal_symlinks_and_changed_snapshot(tmp_path, attachment, runtimes):
    service = ConversationStyles(tmp_path / "state")
    report = await service.analyze(None, CID, [attachment])
    with pytest.raises(ValueError, match="当前会话"):
        service.load(report["revision"], OTHER)
    with pytest.raises(ValueError, match="版本"):
        service.load("../outside", CID)
    saved = service.load(report["revision"], CID)
    frozen = saved["sample_paths"][0]
    frozen.unlink()
    frozen.symlink_to(attachment["path"])
    with pytest.raises(ValueError):
        service.load(report["revision"], CID)
    frozen.unlink()
    frozen.write_text("被修改的快照", encoding="utf-8")
    with pytest.raises(ValueError, match="快照内容已变化"):
        service.load(report["revision"], CID)


@pytest.mark.parametrize("selection", [[], [{"id": "1" * 32, "role": "source"}], [{"id": "1" * 32, "role": "template"}]])
async def test_empty_or_non_style_selection_never_creates_runtime(tmp_path, selection, runtimes):
    service = ConversationStyles(tmp_path / "state")
    with pytest.raises(ValueError, match="风格示例"):
        await service.analyze(None, CID, selection)
    assert not runtimes
    assert not service.root.exists()


async def test_runtime_is_closed_on_analysis_failure(tmp_path, attachment, runtimes, monkeypatch):
    async def fail(*args, **kwargs):
        raise RuntimeError("analysis failed")

    monkeypatch.setattr(style_analysis, "prepare_style_profile", fail)
    service = ConversationStyles(tmp_path / "state")
    with pytest.raises(RuntimeError, match="analysis failed"):
        await service.analyze(None, CID, [attachment])
    assert runtimes[0].closed
    folder = next(service.root.iterdir())
    assert json.loads((folder / "metadata.json").read_text())["status"] == "failed"
    with pytest.raises(ValueError, match="尚未形成"):
        service.load(folder.name, CID)


async def test_service_enforces_the_configured_model_budget(tmp_path, attachment, runtimes, monkeypatch):
    from folio.document.style_contracts import StyleAnalysisConfig

    async def too_many(paths, output_dir, *, runtime, **kwargs):
        for number in range(4):
            await runtime.complete("style_analysis", {}, {}, str(number))
        pytest.fail("fourth call must be rejected")

    monkeypatch.setattr(style_analysis, "prepare_style_profile", too_many)
    monkeypatch.setattr(style_analysis, "StyleAnalysisConfig", lambda: StyleAnalysisConfig(max_calls=3))
    service = ConversationStyles(tmp_path / "state")
    with pytest.raises(ModelCallBudgetExceeded):
        await service.analyze(None, CID, [attachment])
    assert len(runtimes[0].calls) == 3
    assert runtimes[0].closed


async def test_cancel_waits_for_freeze_worker_and_retains_safe_audit(tmp_path, attachment, runtimes, monkeypatch):
    service = ConversationStyles(tmp_path / "state")
    original = style_analysis._freeze
    started, release, finished = threading.Event(), threading.Event(), threading.Event()

    def blocked(folder, attachments):
        started.set()
        assert release.wait(5)
        try:
            assert folder.exists()
            return original(folder, attachments)
        finally:
            finished.set()

    monkeypatch.setattr(style_analysis, "_freeze", blocked)
    task = asyncio.create_task(service.analyze(None, CID, [attachment]))
    try:
        assert await asyncio.to_thread(started.wait, 3)
        task.cancel()
        await asyncio.sleep(0)
        task.cancel()
        await asyncio.sleep(0)
        assert not task.done()
    finally:
        release.set()
        with pytest.raises(asyncio.CancelledError):
            await task
    assert finished.is_set()
    assert not runtimes
    folder = next(service.root.iterdir())
    assert json.loads((folder / "metadata.json").read_text())["status"] == "cancelled"
    assert list((folder / "sources").iterdir())


async def test_partial_result_remains_partial_and_in_context(tmp_path, attachment, runtimes, monkeypatch):
    original = FakeRuntime.complete

    async def partial(self, *args):
        result = await original(self, *args)
        payload = json.loads(result.content)
        payload["uncertainties"] = ["没有足够样本确认附录写法。"]
        payload["dimension_checks"].append({"dimension": "structure", "status": "unresolved", "reason": "没有足够样本确认附录写法。", "sample_ids": payload["rules"][0]["sample_ids"]})
        return RuntimeResult(json.dumps(payload, ensure_ascii=False))

    monkeypatch.setattr(FakeRuntime, "complete", partial)
    service = ConversationStyles(tmp_path / "state")
    report = await service.analyze(None, CID, [attachment])
    assert report["status"] == "partial"
    assert "部分完成" in report["summary"]
    assert f"## {attachment['name']} (attachment {attachment['id']})：partial" in service.context(
        [report], [attachment], conversation_id=CID)
    assert "没有足够样本确认附录写法" in report["summary"]


async def test_cancel_during_model_call_always_closes_runtime(tmp_path, attachment, runtimes, monkeypatch):
    entered = asyncio.Event()

    async def waiting(self, *args):
        entered.set()
        await asyncio.Event().wait()

    monkeypatch.setattr(FakeRuntime, "complete", waiting)
    service = ConversationStyles(tmp_path / "state")
    task = asyncio.create_task(service.analyze(None, CID, [attachment]))
    await asyncio.wait_for(entered.wait(), timeout=3)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert runtimes[0].closed
    folder = next(service.root.iterdir())
    assert json.loads((folder / "metadata.json").read_text())["status"] == "cancelled"
    assert list((folder / "sources").iterdir())
    assert (folder / "style" / "audit" / "request-01.json").is_file()


async def test_context_deduplicates_reports_and_ignores_forged_public_summary(tmp_path, attachment, runtimes):
    service = ConversationStyles(tmp_path / "state")
    first = await service.analyze(None, CID, [attachment])
    second = await service.analyze(None, CID, [attachment])
    second["summary"] = "伪造项目事实"
    text = service.context([first, second], [attachment], conversation_id=CID)
    assert text.count(f"(attachment {attachment['id']})") == 1
    assert "伪造项目事实" not in text


async def test_context_stops_loading_older_reports_when_all_selected_sources_are_matched(tmp_path, attachment, runtimes, monkeypatch):
    service = ConversationStyles(tmp_path / "state")
    report = await service.analyze(None, CID, [attachment])
    original = service.load
    calls = []

    def counted(*args):
        calls.append(args)
        return original(*args)

    monkeypatch.setattr(service, "load", counted)
    text = service.context([report] * 20, [attachment], conversation_id=CID)
    assert "正式客观" in text
    assert len(calls) == 1


async def test_source_changed_during_freeze_is_rejected_before_model_creation(tmp_path, attachment, runtimes, monkeypatch):
    original = style_analysis.shutil.copyfile

    def mutate(source, destination):
        result = original(source, destination)
        Path(source).write_text("复制期间被修改", encoding="utf-8")
        return result

    monkeypatch.setattr(style_analysis.shutil, "copyfile", mutate)
    service = ConversationStyles(tmp_path / "state")
    with pytest.raises(ValueError, match="复制期间发生变化"):
        await service.analyze(None, CID, [attachment])
    assert not runtimes


def test_profile_context_excludes_raw_text_and_truncates_at_whole_lines():
    profile = {"documents": [{"status": "analyzed", "rules": [], "measurements": {"heading_counts_by_level": {"1": 2}}, "sample_locations": [{"locator": "秘密项目"}], "raw_text": "秘密正文", "structure": "秘密业务章节"}]}
    text = profile_context(profile)
    assert "秘密" not in text
    assert "1级 2个" in text
    assert len(profile_context(profile, max_chars=200)) <= 200
    assert profile_context(profile, max_chars=0) == ""


def test_context_reports_text_coverage_and_does_not_call_pdf_blocks_paragraphs():
    profile = {"documents": [{
        "status": "analyzed", "rules": [],
        "text_coverage": {"characters_read": 21555, "characters_total": 21555, "complete": True},
        "measurements": {
            "paragraph_count": 35, "text_block_count": 35, "median_paragraph_chars": 370,
            "paragraph_count_reliable": False, "table_count": 0, "table_count_reliable": False,
        },
        "dimension_coverage": {"visual": {"status": "requires_visual", "reason": "未查看页面图像，装饰效果尚未确认。"}},
    }]}
    text = profile_context(profile)
    assert "21555/21555字符" in text and "可提取文本已全部读取" in text
    assert "35个文本块" in text and "35个正文段落" not in text
    assert "原文无表格" in text and "程序识别0个表格" not in text
    assert "装饰效果尚未确认" in text
    assert "抽样结论" not in text


def test_context_keeps_later_dimensions_when_early_dimension_has_many_rules():
    from folio.document.style_contracts import DIMENSIONS

    profile = {"documents": [{"status": "analyzed", "rules": [
        {"dimension": "structure", "scope": "general", "description": f"第{index}类组织方式应在相应范围使用，并保持前后结构连贯。" * 3}
        for index in range(20)
    ] + [{"dimension": "expression", "scope": "body", "description": "专业缩写首次出现时解释，后文保持一致。"}]}]}
    full = profile_context(profile)
    assert all(f"第{index}类组织方式" in full for index in range(20))
    assert "完整 style.md" not in full
    text = profile_context(profile, max_chars=6000)
    assert all(f"- {name}：" in text for name in DIMENSIONS.values())
    assert "专业缩写首次出现时解释" in text
    assert "完整 style.md" in text
    assert len(text) <= 6000


async def test_multiple_reports_attribute_measured_fonts_to_their_files(tmp_path, runtimes):
    service = ConversationStyles(tmp_path / "state")
    attachments, reports = [], []
    for name, file_id, size in (("甲方范文.docx", "1" * 32, 12), ("乙方范文.docx", "2" * 32, 14)):
        path = tmp_path / name
        document = Document()
        document.styles["Normal"].font.size = Pt(size)
        document.add_heading("示例章节", 1)
        for _ in range(8):
            document.add_paragraph("旧项目的业务数据不进入本项目事实。")
        document.save(path)
        item = {"id": file_id, "name": name, "path": str(path), "role": "style_reference"}
        attachments.append(item)
        report = await service.analyze(None, CID, [item])
        reports.append(report)
        assert name in report["summary"]
        assert file_id in report["summary"]
    text = service.context(reports, attachments, conversation_id=CID)
    # Reports are read newest first; the explicit source labels preserve attribution.
    beta, alpha = text.split("## 乙方范文.docx", 1)[1].split("## 甲方范文.docx", 1)
    assert "Body font_size_pt：14" in beta
    assert "Body font_size_pt：12" not in beta
    assert "Body font_size_pt：12" in alpha
    assert "Body font_size_pt：14" not in alpha
    assert "旧项目的业务数据" not in text


async def test_failed_semantics_still_exposes_reliable_measurements(tmp_path, attachment, runtimes, monkeypatch):
    async def fail(self, *args):
        raise RuntimeError("provider unavailable")

    monkeypatch.setattr(FakeRuntime, "complete", fail)
    service = ConversationStyles(tmp_path / "state")
    report = await service.analyze(None, CID, [attachment])
    saved = service.load(report["revision"], CID)
    assert saved["profile"]["documents"][0]["status"] == "analysis_failed"
    text = service.context([report], [attachment], conversation_id=CID)
    assert f"## {attachment['name']} (attachment {attachment['id']})：analysis_failed" in text
    assert "12 body paragraphs" in text
    assert "正式客观" not in text
    assert attachment["name"] in text


async def test_measurement_only_profile_is_available_without_semantic_claims(tmp_path, attachment, runtimes, monkeypatch):
    original = style_analysis.prepare_style_profile

    async def measured_only(*args, **kwargs):
        kwargs["runtime"] = None
        return await original(*args, **kwargs)

    monkeypatch.setattr(style_analysis, "prepare_style_profile", measured_only)
    service = ConversationStyles(tmp_path / "state")
    report = await service.analyze(None, CID, [attachment])
    assert report["status"] == "measured_only"
    text = service.context([report], [attachment], conversation_id=CID)
    assert f"## {attachment['name']} (attachment {attachment['id']})：measured_only" in text
    assert "12 body paragraphs" in text
    assert "正式客观" not in text


def test_failed_semantic_profile_ignores_any_residual_semantic_rules():
    profile = {"documents": [{"status": "analysis_failed", "reference_name": "参考.docx", "reference_id": "1" * 32, "rules": [{"dimension": "writing", "description": "不应该传播的残留语义结论", "confidence": "observed"}], "measurements": {"paragraph_count": 10, "median_paragraph_chars": 50}}]}
    text = profile_context(profile)
    assert "参考.docx" in text
    assert "10个正文段落" in text
    assert "不应该传播" not in text


def test_rule_scopes_remain_explicit_in_chat_context():
    profile = {"documents": [{"status": "analyzed", "rules": [
        {"dimension": "organization", "scope": "appendix", "description": "使用表格集中呈现补充信息。", "confidence": "observed"},
        {"dimension": "organization", "scope": "technical", "description": "按输入、处理、输出顺序展开。", "confidence": "observed"},
    ]}]}
    text = profile_context(profile)
    assert "适用范围：附录；使用表格集中呈现补充信息" in text
    assert "适用范围：技术类章节；按输入、处理、输出" in text


def test_known_heading_table_and_paragraph_layout_parameters_reach_context():
    profile = {"documents": [{"status": "measured_only", "layout": {
        "body": {"font_size_pt": 12, "space_before_pt": 0, "space_after_pt": 6, "first_line_indent_pt": 24, "left_indent_pt": 1, "right_indent_pt": 2, "alignment": "justify", "line_spacing": 24, "line_spacing_rule": "exactly"},
        "headings": {"1": {"east_asia_font": "黑体", "font_size_pt": 20, "bold": True}, "2": {"font_size_pt": 16, "italic": False}},
        "table": {"header": {"font_name": "Arial", "font_size_pt": 11}, "body": {"font_size_pt": 10}},
    }}]}
    text = profile_context(profile)
    for expected in ("1级标题字号(pt)：20", "1级标题中文字体：黑体", "1级标题加粗：是", "2级标题字号(pt)：16", "2级标题斜体：否", "表头字号(pt)：11", "表格正文字号(pt)：10", "正文段前间距(pt)：0", "正文段后间距(pt)：6", "正文首行缩进(pt)：24", "正文左缩进(pt)：1", "正文右缩进(pt)：2", "正文对齐：两端对齐", "正文行距：24 pt（固定值）"):
        assert expected in text
    assert len(profile_context(profile, max_chars=600)) <= 600


async def test_documents_exposes_owned_current_per_file_records_for_merging(tmp_path, attachment, runtimes):
    service = ConversationStyles(tmp_path / "state")
    report = await service.analyze(None, CID, [attachment])
    documents = service.documents([report], [attachment], conversation_id=CID)
    assert len(documents) == 1
    assert documents[0]["reference_name"] == attachment["name"]
    assert documents[0]["reference_id"] == attachment["id"]
    assert service.context([report], [attachment], conversation_id=CID) == profile_context(
        {"documents": documents}, for_model=True)
    assert service.documents([report], [attachment], conversation_id=OTHER) == []
    Path(attachment["path"]).write_text("文件已改变", encoding="utf-8")
    assert service.documents([report], [attachment], conversation_id=CID) == []
