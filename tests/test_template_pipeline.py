"""Uploaded templates constrain writing without becoming factual research inputs."""

import json
import shutil
from pathlib import Path
from unittest.mock import AsyncMock

import pytest
from docx import Document

from folio.document.planner import build_plan
from folio.document.template_contract import (
    TemplateTable,
    bind_template,
    read_template,
    validate_template_plan,
)
from folio.models import (
    Block,
    DocumentPlan,
    Extraction,
    JobInput,
    PlanItem,
    PlanProposal,
    Section,
    SectionContent,
)
from folio.orchestration.context import project
from folio.orchestration.engine import Engine
from folio.orchestration.local_checks import hard_findings
from folio.orchestration.review_scope import section_review_scope
from folio.runtime.base import RuntimeResult
from folio.runtime.demo import DemoRuntime


def template_plan(tmp_path):
    path = tmp_path / "user-template.docx"
    document = Document()
    document.add_heading("交易要求", 1)
    document.add_paragraph("填写说明：根据资料填写交易与审批规则。")
    document.add_paragraph("示例项目审批额度为 987654 元，不能用于实际项目。")
    document.save(path)
    plan = build_plan(
        "template-test",
        PlanProposal(sections=[PlanItem(node_id="requirements", title="交易要求", level=1,
                                        instruction="研究交易和审批规则")]),
    )
    return bind_template(plan, read_template(path))


def requirement(table_id="table_a", *, min_rows=1, physical_headers=1):
    return TemplateTable(
        table_id=table_id, body_index=2, headers=["名称", "说明"],
        header_rows=physical_headers, min_data_rows=min_rows,
    )


def section(**updates):
    return Section(
        section_id="section", node_id="node", title="交易要求", level=1,
        display_number="1", instruction="填写交易规则", **updates,
    )


def table(block_id, rows=1):
    return Block(
        block_id=block_id, kind="table", origin="model_proposal",
        rows=[["名称", "说明"], *[[f"交易{i}", "拟议规则"] for i in range(rows)]],
    )


def findings(node, blocks):
    return hard_findings(node, SectionContent(blocks=blocks), [], Extraction(summary="交易"))


def test_store_freezes_template_separately_and_excludes_duplicate_evidence(store, inputs, tmp_path):
    plan = template_plan(tmp_path)
    original = plan.template.path
    original_bytes = Path(original).read_bytes()
    duplicate = Path(inputs["knowledge_path"]) / "example-as-source.docx"
    shutil.copyfile(original, duplicate)
    job_id = store.create(
        JobInput(**inputs, reference_paths=[original], runtime="demo"), plan=plan
    )
    saved = store.get(job_id)
    frozen = DocumentPlan.model_validate(saved["plan"])
    assert frozen.template.path == str(store.root / "jobs" / job_id / "inputs" / "template.docx")
    assert Path(frozen.template.path).read_bytes() == original_bytes
    assert plan.template.path == original
    assert saved["input"]["reference_paths"] == []
    assert {p.name for p in Path(saved["input"]["knowledge_path"]).iterdir()} == {"rules.md"}
    Path(original).write_bytes(b"original replaced after confirmation")
    validate_template_plan(frozen)


def test_changed_template_is_rejected_before_a_job_is_created(store, inputs, tmp_path):
    plan = template_plan(tmp_path)
    Path(plan.template.path).write_bytes(b"changed template")
    with pytest.raises(ValueError, match="[Tt]emplate"):
        store.create(JobInput(**inputs, runtime="demo"), plan=plan)
    assert not list((store.root / "jobs").iterdir())
    assert store.candidates() == []


async def test_changed_frozen_template_stops_before_any_model_call(store, inputs, tmp_path):
    plan = template_plan(tmp_path)
    job_id = store.create(JobInput(**inputs, runtime="demo"), plan=plan)
    frozen = DocumentPlan.model_validate(store.get(job_id)["plan"])
    Path(frozen.template.path).write_bytes(b"changed frozen template")
    runtime = DemoRuntime()
    runtime.complete = AsyncMock(side_effect=AssertionError("No model calls before template checks"))
    with pytest.raises(ValueError, match="[Tt]emplate"):
        await Engine(store, runtime).run(job_id)
    runtime.complete.assert_not_awaited()
    assert store.get(job_id)["status"] == "failed"


async def test_template_job_runs_research_writing_and_final_word_without_sample_leakage(store, inputs, tmp_path):
    plan = template_plan(tmp_path)
    job_id = store.create(JobInput(**inputs, runtime="demo"), plan=plan)
    result = await Engine(store, DemoRuntime()).run(job_id)
    assert result["status"] == "draft"
    document = Document(result["artifact"])
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "交易要求" in text
    assert "987654" not in text and "填写说明" not in text
    assert "987654" not in json.dumps(result["snapshot"], ensure_ascii=False)
    assert all("987654" not in e.text for s in store.states(job_id) for e in s.evidence)


async def test_template_format_change_after_render_cannot_be_delivered(store, inputs, tmp_path, monkeypatch):
    from docx.shared import Pt

    async def corrupt_format(path, *args, **kwargs):
        document = Document(path)
        document.styles["Normal"].font.size = Pt(42)
        document.save(path)
        return {"status": "applied", "changed": True}

    monkeypatch.setattr("folio.orchestration.formatting.format_document", corrupt_format)
    plan = template_plan(tmp_path)
    job_id = store.create(JobInput(**inputs, runtime="demo"), plan=plan)
    with pytest.raises(ValueError, match="[Tt]emplate"):
        await Engine(store, DemoRuntime()).run(job_id)
    assert store.get(job_id)["status"] == "failed"
    assert store.get(job_id)["artifact"] is None


def test_publication_rechecks_template_snapshot(store, inputs, tmp_path):
    plan = template_plan(tmp_path)
    job_id = store.create(JobInput(**inputs, runtime="llm"), plan=plan)
    frozen = DocumentPlan.model_validate(store.get(job_id)["plan"])
    Path(frozen.template.path).write_bytes(b"changed before publication")
    artifact = store.root / "jobs" / job_id / "document.docx"
    artifact.write_bytes(b"an existing artifact")
    store.update(job_id, snapshot={"review_passed": True, "blocking_issues": []})
    with pytest.raises(ValueError, match="[Tt]emplate"):
        store.publish(job_id, artifact)


def test_distinct_template_tables_cannot_reuse_one_generated_table():
    node = section(required_tables=[requirement("first"), requirement("second")])
    failures = findings(node, [table("only")])
    assert [f.key for f in failures] == ["required-table:second"]
    assert findings(node, [table("first_block"), table("second_block")]) == []


def test_template_table_matching_preserves_stricter_row_requirements():
    node = section(required_tables=[requirement("short"), requirement("long", min_rows=3)])
    assert findings(node, [table("long", 3), table("short", 1)]) == []
    failures = findings(node, [table("short_a", 1), table("short_b", 1)])
    assert [f.key for f in failures] == ["required-table:long"]


def test_required_table_needs_data_and_uses_one_logical_header_row():
    node = section(required_tables=[requirement(physical_headers=2)])
    assert [f.key for f in findings(node, [table("empty", 0)])] == ["required-table:table_a"]
    assert findings(node, [table("filled", 1)]) == []
    blank = table("blank", 1)
    blank.rows[1] = ["", "  "]
    assert [f.key for f in findings(node, [blank])] == ["required-table:table_a"]
    reordered = table("wrong_order", 1)
    reordered.rows[0].reverse()
    assert [f.key for f in findings(node, [reordered])] == ["required-table:table_a"]


def test_fixed_template_directory_rejects_new_body_headings():
    heading = Block(block_id="new_heading", kind="heading", heading_level=2,
                    text="擅自新增子节", origin="model_proposal")
    failures = findings(section(allow_subheadings=False), [heading])
    assert [f.key for f in failures] == ["template-heading:new_heading"]
    assert failures[0].block_refs[0].block_id == "new_heading"
    assert findings(section(allow_subheadings=True), [heading]) == []


def test_template_contract_survives_writer_and_review_scope_projection():
    parent = section()
    child = section(template_section_id="fixed", required_tables=[requirement()],
                    allow_subheadings=False).model_copy(update={"section_id": "child", "parent_id": "section", "level": 2})
    source = {"section": parent.model_dump(), "child_sections": [child.model_dump()]}
    projected = project("generate", source)
    scoped = section_review_scope("section", [parent, child])
    for view in (projected["child_sections"][0], scoped["child_sections"][0]):
        assert view["required_tables"] == [requirement().model_dump()]
        assert view["allow_subheadings"] is False
        assert view["template_section_id"] == "fixed"


async def test_template_table_repair_preserves_data_and_does_not_restart_research(store, inputs):
    job_id = store.create(JobInput(**inputs, runtime="demo"))
    engine = Engine(store, DemoRuntime())
    await engine.run(job_id)
    node = engine.plan.sections[0].model_copy(update={
        "required_tables": [requirement("first"), requirement("second")],
    })
    previous = next(s for s in store.states(job_id) if s.section_id == node.section_id)
    previous.local_findings = hard_findings(node, previous.content, previous.evidence,
                                           previous.extraction)
    finding = next(f for f in previous.local_findings if f.key == "required-table:first")
    candidate = SectionContent(blocks=[table("a"), table("b")])
    engine.research_work = AsyncMock(side_effect=AssertionError("Structural repair needs no research"))
    engine.call = AsyncMock(return_value=candidate)
    engine.validate_section = AsyncMock(return_value=previous)
    await engine.generate(node, previous, finding, "不要继承这条无关指令")
    engine.research_work.assert_not_awaited()
    context = engine.call.call_args.args[1]
    assert "first" in context["repair_instruction"] and "second" in context["repair_instruction"]
    assert "Do not invent facts to fill rows" in context["repair_instruction"]
    assert "无关指令" not in context["repair_instruction"]
    assert context["section"]["required_tables"] == [t.model_dump() for t in node.required_tables]


async def test_template_table_failure_cannot_be_dismissed_by_model(store, inputs):
    class DismissingRuntime(DemoRuntime):
        async def complete(self, task, context, schema, session_key):
            if task == "drill_down":
                return RuntimeResult(json.dumps({
                    "confirmed": False, "reason": "模型认为可以忽略模板表格",
                    "needs_confirmation": False, "affected_section_ids": [], "repair_instruction": "",
                }))
            return await super().complete(task, context, schema, session_key)

    job_id = store.create(JobInput(**inputs, runtime="demo"))
    engine = Engine(store, DemoRuntime())
    await engine.run(job_id)
    engine.plan.sections[0].required_tables = [requirement()]
    store.save_plan(job_id, engine.plan)
    store.update(job_id, status="failed", stage="review")
    store.resume_job(job_id)
    restored = Engine(store, DismissingRuntime())
    # The assertions concern deterministic template validation, not business review.
    restored.business_findings = AsyncMock(return_value=[])
    restored.global_review = AsyncMock(return_value=[])
    result = await restored.run(job_id)
    assert result["status"] == "needs_confirmation"
    assert result["snapshot"]["review_passed"] is False
    assert any(f["key"] == "required-table:table_a" for f in result["snapshot"]["blocking_issues"])
    assert restored._dismissed() == {}
