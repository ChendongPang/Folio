from copy import deepcopy
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Mm, Pt

from folio.document.formatting import (
    LayoutPatch,
    allowed_properties,
    apply_patches,
    collect_targets,
)
from folio.document.reader import xpath
from folio.document.style_layout import (
    StyleLayout,
    apply_style_layout,
    document_style_layout,
    resolve_style_policy,
)
from folio.document.template_contract import read_template
from folio.document.typography import BODY, resolved_format
from folio.document.word import render
from folio.models import Block, DocumentPlan, Extraction, Section, SectionContent, SectionState
from folio.orchestration.layout_review import PageLayoutReview, review_rendered_document


@pytest.fixture
def case(tmp_path):
    plan = DocumentPlan(
        sections=[
            Section(
                section_id="s1",
                node_id="n1",
                title="当前项目",
                level=1,
                display_number="1",
                instruction="写当前项目",
                expected_table_headers=["项目", "说明"],
            )
        ]
    )
    states = [
        SectionState(
            section_id="s1",
            content=SectionContent(
                blocks=[
                    Block(block_id="b1", text="当前项目已有材料。"),
                    Block(
                        block_id="t1", kind="table", rows=[["项目", "说明"], ["服务", "处理请求"]]
                    ),
                ]
            ),
            extraction=Extraction(summary="概况"),
            evidence=[],
        )
    ]
    return tmp_path / "styled.docx", plan, states


@pytest.fixture
def profile():
    return {
        "role": "style_only",
        "style_markdown": "# 风格\n## 排版\n正文仿宋，12 磅，段后 9 磅。",
        "layout": {
            "body": {
                "font_name": "Arial",
                "east_asia_font": "仿宋",
                "font_size_pt": 12,
                "color": "234567",
                "alignment": "justify",
                "space_after_pt": 9,
                "line_spacing": 1.7,
                "first_line_indent_pt": 24,
            },
            "headings": {"1": {"east_asia_font": "黑体", "font_size_pt": 21, "bold": True}},
            "page": {"width_mm": 210, "height_mm": 297, "margin_left_mm": 33},
            "table": {
                "header": {"font_size_pt": 11, "bold": True},
                "body": {"font_size_pt": 10},
                "header_fill": "DDEEFF",
                "border_color": "123456",
                "border_size_pt": 1,
            },
        },
    }


def test_sample_layout_reaches_saved_docx_and_is_protected_after_reload(case, profile):
    path, plan, states = case
    before = deepcopy(profile)
    result = render(
        path,
        plan,
        states,
        title="当前需求",
        draft=False,
        defer_pagination=True,
        writing_profile=profile,
    )
    assert profile == before
    assert result["style_application"]["status"] == "applied"
    document = Document(path)
    paragraph = next(p for p in document.paragraphs if p.text == "当前项目已有材料。")
    ppr, rpr = resolved_format(document, paragraph._p)
    assert xpath(rpr, "./w:rFonts/@w:eastAsia") == ["仿宋"]
    assert xpath(rpr, "./w:rFonts/@w:ascii") == ["Arial"]
    assert xpath(rpr, "./w:sz/@w:val") == ["24"]
    assert xpath(rpr, "./w:color/@w:val") == ["234567"]
    assert xpath(ppr, "./w:spacing/@w:line") == ["408"]
    assert xpath(ppr, "./w:spacing/@w:after") == ["180"]
    assert xpath(ppr, "./w:ind/@w:firstLine") == ["480"]
    assert document.styles["Folio Heading 1"].font.size == Pt(21)
    assert document.sections[0].page_width.mm == pytest.approx(210, abs=0.02)
    assert document.sections[0].left_margin.mm == pytest.approx(33, abs=0.02)
    assert xpath(document.tables[0]._tbl, "./w:tr[1]/w:tc/w:tcPr/w:shd/@w:fill") == [
        "DDEEFF",
        "DDEEFF",
    ]
    assert xpath(document.tables[0]._tbl, "./w:tblPr/w:tblBorders/w:top/@w:sz") == ["8"]
    assert document_style_layout(document).body.font_size_pt == 12
    targets = collect_targets(document, plan, states)
    body = next(target for target in targets if target.kind == "paragraph")
    assert not {"alignment", "space_after_pt", "font_delta_pt"} & set(allowed_properties(body))
    assert "space_before_pt" in allowed_properties(body)
    with pytest.raises(ValueError, match="not allowed"):
        apply_patches(document, targets, [LayoutPatch(target_id=body.target_id, space_after_pt=6)])
    cell = next(target for target in targets if target.kind == "cell")
    assert "font_delta_pt" not in allowed_properties(cell)


def test_previous_brand_layout_contract_is_read_without_changing_original(case, profile):
    from docx.oxml.ns import qn

    path, plan, states = case
    render(path, plan, states, title="既有文档", draft=False,
           defer_pagination=True, writing_profile=profile)
    document = Document(path)
    variable = xpath(document.settings.element, "./w:docVars/w:docVar")[0]
    variable.set(qn("w:name"), "PreviousBrandStyleLayout")
    before = document.settings.element.xml
    assert document_style_layout(document).body.font_size_pt == 12
    assert document.settings.element.xml == before


def test_template_retains_format_over_conflicting_sample(case, profile, tmp_path):
    path, plan, states = case
    template = tmp_path / "template.docx"
    document = Document()
    document.styles["Normal"].font.size = Pt(14)
    document.sections[0].left_margin = Mm(41)
    document.add_heading("当前项目", 1)
    document.add_paragraph("填写项目说明。")
    document.save(template)
    plan.template = read_template(template)
    plan.sections[0].template_section_id = plan.template.sections[0].section_id
    result = render(
        path,
        plan,
        states,
        title="新文档",
        draft=False,
        defer_pagination=True,
        writing_profile=profile,
    )
    output = Document(path)
    assert output.sections[0].left_margin.mm == pytest.approx(41, abs=0.02)
    paragraph = next(p for p in output.paragraphs if p.text == "当前项目已有材料。")
    assert xpath(resolved_format(output, paragraph._p)[1], "./w:sz/@w:val") == ["28"]
    assert result["style_application"]["status"] == "template_priority"
    assert result["style_application"]["applied"] == []
    assert document_style_layout(output).body.font_size_pt is None


def test_empty_profile_keeps_default_behavior(case):
    path, plan, states = case
    result = render(path, plan, states, title="需求", draft=False, defer_pagination=True)
    assert result == {"status": "deferred"}
    document = Document(path)
    assert document.styles[BODY].font.size == Pt(11)
    targets = collect_targets(document, plan, states)
    assert "space_after_pt" in allowed_properties(next(t for t in targets if t.kind == "paragraph"))


def test_invalid_sample_geometry_is_rejected_before_export():
    document = Document()
    with pytest.raises(ValueError, match="writable"):
        apply_style_layout(
            document,
            {"layout": {"page": {"width_mm": 100, "margin_left_mm": 60, "margin_right_mm": 60}}},
        )
    with pytest.raises(ValueError):
        StyleLayout.model_validate({"body": {"line_spacing": float("nan")}})
    with pytest.raises(ValueError):
        StyleLayout.model_validate({"body": {"line_spacing": 20}})


def test_exact_line_spacing_retains_point_units(case):
    path, plan, states = case
    render(
        path,
        plan,
        states,
        title="需求",
        draft=False,
        defer_pagination=True,
        writing_profile={"layout": {"body": {"line_spacing": 20, "line_spacing_rule": "exactly"}}},
    )
    document = Document(path)
    paragraph = next(p for p in document.paragraphs if p.text == "当前项目已有材料。")
    ppr, _ = resolved_format(document, paragraph._p)
    assert xpath(ppr, "./w:spacing/@w:line") == ["400"]
    assert xpath(ppr, "./w:spacing/@w:lineRule") == ["exact"]


def test_explicit_user_format_overrides_sample_without_mutating_original(profile):
    before = deepcopy(profile)
    output = resolve_style_policy(
        profile,
        "正文采用宋体，小四号，1.5倍行距；一级标题采用黑体，字号 18 磅；左边距2.5厘米；纸张A4。",
    )
    assert profile == before
    assert output["layout"]["body"]["east_asia_font"] == "宋体"
    assert output["layout"]["body"]["font_size_pt"] == 12
    assert output["layout"]["body"]["line_spacing"] == 1.5
    assert output["layout"]["headings"]["1"]["font_size_pt"] == 18
    assert output["layout"]["page"]["margin_left_mm"] == 25
    assert "margin_right_mm" not in output["layout"]["page"]
    assert "本次明确的格式要求" in output["markdown"]
    assert output["markdown"] == output["style_markdown"]


@pytest.mark.parametrize(
    "instruction",
    [
        "正文不要宋体。",
        "取消正文用宋体的要求。",
        "正文采用宋体。取消正文用宋体的要求。",
        "正文采用宋体。取消上述格式要求。",
        "正文用宋体可以吗？",
        "例如正文采用宋体。",
        "示例写着“正文采用宋体”。",
        "> 正文用宋体\n实际仍按示例。",
        "```\n正文用宋体\n```",
    ],
)
def test_questions_negation_and_quoted_examples_do_not_become_format_instructions(
    profile, instruction
):
    output = resolve_style_policy(profile, instruction)
    assert output["layout"] == profile["layout"]
    assert not output["user_format_policy"]["overrides"]


async def test_visual_review_receives_style_and_disabled_review_reports_adoption(
    case, profile, tmp_path
):
    path, plan, states = case
    render(
        path,
        plan,
        states,
        title="需求",
        draft=False,
        defer_pagination=True,
        writing_profile=profile,
    )
    seen = []

    async def call(task, context, schema):
        seen.append(context)
        return PageLayoutReview()

    report = await review_rendered_document(
        path, plan, states, tmp_path, call, lambda _: None, writing_profile=profile
    )
    assert not seen
    assert report["style_application"]["status"] == "applied"
    # Reuse the existing offline renderer rather than requiring Office or a model.
    from test_layout_review import preview

    report = await review_rendered_document(
        path,
        plan,
        states,
        tmp_path,
        call,
        lambda _: None,
        writing_profile=profile,
        enabled=True,
        supports_images=True,
        renderer=preview,
    )
    assert report["status"] == "passed"
    assert seen[0]["style_reference"]["style_markdown"] == profile["style_markdown"]
    assert "space_after_pt" not in seen[0]["targets"][0]["allowed_properties"]
    assert Path(report["pdf"]).is_file()
