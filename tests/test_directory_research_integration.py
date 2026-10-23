"""Planning reads the late original requirements and hands only owned evidence to writing."""

import json
from copy import deepcopy
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from folio.api.app import create_app
from folio.model_settings import ModelSettingsInput, ModelSettingsStore
from folio.models import JobInput
from folio.orchestration.engine import Engine
from folio.orchestration.work import WorkState, WorkTask
from folio.previews import PreviewLibrary
from folio.research.planning_handoff import seed_directory_research
from folio.research.workspace import WorkspaceManager
from folio.runtime.base import RuntimeResult
from folio.runtime.demo import DemoRuntime
from folio.source_documents import SourceDirectoryDrafts
from folio.store import Store

TECHNICAL_REQUIREMENTS = {
    "var_backtest": ("VaR 回测", "每日执行 VaR 回测，分别比较实际损益和假设损益，并记录超限次数。"),
    "stress_test": ("压力测试", "压力测试必须同时支持历史情景和假设情景，结果按交易台汇总。"),
}


class TailReadingRuntime(DemoRuntime):
    """Choose late sections by their actual inspected titles, then read their original nodes."""

    partial = False
    all_calls = []

    def __init__(self, config=None):
        self.seen = set()
        self.nodes = {}
        self.inspection = None
        self.reads = {}
        self.evidence = {}

    def research_action(self, context):
        if self.partial:
            return "finish", {"summary": "尚未读取技术需求。", "blocking_gaps": ["技术需求原文尚未读取。"]}
        path = context["manifest"]["files"][0]["path"]
        for item in context["observations"]:
            if item["request_id"] in self.seen:
                continue
            self.seen.add(item["request_id"])
            result = item["result"]
            assert "error" not in result, result
            if item["tool"] == "fs.inspect":
                self.inspection = result
                self.nodes.update({node["title"]: node for node in result["structure"]})
            elif item["tool"] == "fs.read":
                self.reads[item["arguments"]["node_id"]] = result
                self.evidence.update({entry["evidence_id"]: entry for entry in result["evidence"]})
        if self.inspection is None or self.inspection["next_offset"] is not None:
            return "fs.inspect", {"path": path, "offset": self.inspection["next_offset"] if self.inspection else 0}
        for title, _ in TECHNICAL_REQUIREMENTS.values():
            node = self.nodes[title]
            previous = self.reads.get(node["node_id"])
            if previous is None or previous["next_offset"] is not None:
                return "fs.read", {"path": path, "node_id": node["node_id"],
                                   "offset": previous["next_offset"] if previous else 0}
        requirements = []
        for requirement_id, (_, original) in TECHNICAL_REQUIREMENTS.items():
            source = next(item for item in self.evidence.values() if item["text"] == original)
            requirements.append({"requirement_id": requirement_id, "text": source["text"],
                                 "evidence_ids": [source["evidence_id"]]})
        return "finish", {
            "summary": "已浏览完整结构并读取 VaR 回测和压力测试原文。",
            "requirements": requirements,
            "selected_ranges": [{"range_id": result["range_id"], "reason": "技术标需响应的实际功能要求"}
                                for result in self.reads.values()],
        }

    async def complete(self, task, context, schema, session_key):
        self.all_calls.append((task, deepcopy(context)))
        if task == "directory_research":
            tool, arguments = self.research_action(context)
            return RuntimeResult("", tool_calls=[{
                "id": f"call_{len(self.all_calls)}", "type": "function", "function": {
                    "name": tool.replace(".", "_"),
                    "arguments": json.dumps({"summary": "按文件结构读取技术需求原文", **arguments}, ensure_ascii=False),
                },
            }])
        elif task == "plan":
            researched = {item["requirement_id"]: item for item in context["directory_research"]["requirements"]}
            result = {"sections": [
                {"node_id": key, "title": title, "level": 1, "content_mode": "generate",
                 "instruction": researched[key]["text"] if key in researched else "待研究技术需求后补充写作范围"}
                for key, (title, _) in TECHNICAL_REQUIREMENTS.items()
            ]}
        elif task == "review_directory_research":
            instructions = {item["node_id"]: item["instruction"] for item in context["directory"]}
            originals = {item["evidence_id"]: item["text"] for item in context["evidence"]}
            checks = []
            for item in context["researched_requirements"]:
                assert instructions[item["requirement_id"]] == item["text"]
                assert originals[item["evidence_ids"][0]] == item["text"]
                checks.append({"requirement_id": item["requirement_id"], "disposition": "covered",
                               "node_ids": [item["requirement_id"]], "evidence_ids": item["evidence_ids"],
                               "reason": "章节写作说明明确承接对应原文要求。"})
            result = {"checks": checks}
        else:
            return await super().complete(task, context, schema, session_key)
        return RuntimeResult(json.dumps(result, ensure_ascii=False))


@pytest.fixture
def planning_setup(tmp_path, monkeypatch):
    source = tmp_path / "长招标文件.docx"
    document = Document()
    for index in range(1, 9):
        document.add_heading(f"投标须知 {index}", 1)
        document.add_paragraph("本段仅描述报名地点、文件装订和递交手续，不构成本次系统功能需求。" * 75)
    for title, text in TECHNICAL_REQUIREMENTS.values():
        document.add_heading(title, 1)
        document.add_paragraph(text)
    document.save(source)
    TailReadingRuntime.all_calls = []
    monkeypatch.setattr(TailReadingRuntime, "partial", False)
    monkeypatch.setattr("folio.directory_drafts.LLMRuntime", TailReadingRuntime)
    monkeypatch.setattr("folio.document.word.paginate", lambda *a, **k: {"status": "not_required"})
    monkeypatch.setattr("folio.document.pagination.paginate", lambda *a, **k: {"status": "not_required"})
    return tmp_path / "state", source


@pytest.fixture
async def researched_directory(planning_setup):
    root, source = planning_setup
    drafts = SourceDirectoryDrafts(root, PreviewLibrary(root))
    result = await drafts.create_from_sources(None, "tender", [
        {"id": "a" * 32, "name": source.name, "path": str(source), "role": "source"},
    ], "编制本项目技术标，完整响应资料中实际提出的系统功能。")
    loaded = drafts.load(result["revision"], "tender")
    return root, source, result, loaded


def start_writing(root, loaded):
    store = Store(root)
    job_id = store.create(JobInput(
        title="技术标", user_request="按已确认目录编制技术标", runtime="demo",
        knowledge_path=str(loaded["knowledge_path"]),
    ), plan=loaded["plan"], directory_research=loaded["folder"])
    engine = Engine(store, DemoRuntime())
    engine.job_id, engine.plan = job_id, loaded["plan"]
    return store, job_id, engine


def writing_task(section):
    return WorkTask(task_id=section.node_id, title=section.title, objective=section.instruction,
                    section_ids=[section.section_id], questions=["原文规定了哪些要求？"],
                    completion_criteria=["准确响应本章节所承担的技术要求"], rationale="各功能独立编写")


async def test_late_original_requirements_drive_directory_and_task_specific_handoff(researched_directory):
    root, source, result, loaded = researched_directory
    original = "\n".join(paragraph.text for paragraph in Document(source).paragraphs)
    assert original.index("VaR 回测") > 12000
    assert [p.text for p in Document(source).paragraphs if p.style.name.startswith("Heading")][8:] == [
        "VaR 回测", "压力测试",
    ]
    assert result["confirmable"] is True
    assert result["research"]["status"] == "ready"
    task_names = [task for task, _ in TailReadingRuntime.all_calls]
    plan_index = task_names.index("plan")
    assert all(task == "directory_research" for task in task_names[:plan_index])
    assert task_names[plan_index + 1] == "review_directory_research"
    planning = TailReadingRuntime.all_calls[plan_index][1]
    assert "excerpt" not in planning["source_catalog"][0]
    evidence_text = "\n".join(item["text"] for item in planning["directory_research"]["evidence"])
    for _, text in TECHNICAL_REQUIREMENTS.values():
        assert text in evidence_text
    assert "文件装订" not in evidence_text
    assert {s.node_id: s.instruction for s in loaded["plan"].sections} == {
        key: text for key, (_, text) in TECHNICAL_REQUIREMENTS.items()
    }
    tool_events = [json.loads(path.read_text()) for path in (loaded["folder"] / "research/events").glob("*-tool.json")]
    actual_reads = [event for event in tool_events if event["tool"] == "fs.read"]
    assert len(actual_reads) == 2
    assert all(event["arguments"].get("node_id") for event in actual_reads)

    store, job_id, engine = start_writing(root, loaded)
    frozen_root = root / "jobs" / job_id / "inputs" / "directory-research"
    frozen_report = json.loads((frozen_root / "report.json").read_text())
    assert frozen_report["source_snapshot_id"] == result["research"]["source_snapshot_id"]
    assert (frozen_root / "coverage.json").is_file()
    seeds = []
    for section in loaded["plan"].sections:
        task = writing_task(section)
        workspace = WorkspaceManager(Path(store.get(job_id)["input"]["knowledge_path"]),
                                     root / "jobs" / job_id / "research" / task.task_id)
        work = WorkState(task_id=task.task_id)
        inherited = seed_directory_research(engine, task, work, workspace)
        assert [item["requirement_id"] for item in inherited["requirements"]] == [task.task_id]
        assert [item.text for item in work.evidence] == [TECHNICAL_REQUIREMENTS[task.task_id][1]]
        assert work.record.facts == [] and work.global_facts == {}
        # Restored work must not override a freshly verified original with a
        # stale or altered quote bearing the same evidence identifier.
        work.evidence[0] = work.evidence[0].model_copy(update={"text": "旧缓存篡改了原文。"})
        seed_directory_research(engine, task, work, workspace)
        assert [item.text for item in work.evidence] == [TECHNICAL_REQUIREMENTS[task.task_id][1]]
        seeds.append(set(inherited["evidence_ids"]))
    assert seeds[0].isdisjoint(seeds[1]), "Separate writing tasks must not inherit each other's original quotes"


@pytest.mark.parametrize("change", ["source", "quote"])
async def test_handoff_rejects_changed_sources_or_forged_original_quotes(researched_directory, change):
    root, _, _, loaded = researched_directory
    store, job_id, engine = start_writing(root, loaded)
    knowledge = Path(store.get(job_id)["input"]["knowledge_path"])
    if change == "source":
        path = next(knowledge.glob("*.docx"))
        document = Document(path)
        document.add_paragraph("正文阶段擅自改变了已确认的资料。")
        document.save(path)
    else:
        report_path = root / "jobs" / job_id / "inputs/directory-research/report.json"
        report = json.loads(report_path.read_text())
        original = next(item for item in report["evidence"] if item["text"] == TECHNICAL_REQUIREMENTS["var_backtest"][1])
        original["text"] = "无需执行 VaR 回测。"
        report_path.write_text(json.dumps(report, ensure_ascii=False))
    section = next(section for section in loaded["plan"].sections if section.node_id == "var_backtest")
    task = writing_task(section)
    workspace = WorkspaceManager(knowledge, root / "jobs" / job_id / "research" / task.task_id)
    work = WorkState(task_id=task.task_id)
    with pytest.raises(ValueError, match='source version|does not match the original'):
        seed_directory_research(engine, task, work, workspace)
    assert work.evidence == []


async def test_job_creation_rejects_source_changes_after_directory_confirmation(researched_directory):
    root, _, _, loaded = researched_directory
    path = next(loaded["knowledge_path"].glob("*.docx"))
    document = Document(path)
    document.add_paragraph("确认目录之后原始资料已改变。")
    document.save(path)
    with pytest.raises(ValueError, match='snapshots differ'):
        start_writing(root, loaded)
    with Store(root).connect() as connection:
        assert connection.execute("SELECT count(*) FROM jobs").fetchone()[0] == 0
    assert list((root / "jobs").iterdir()) == []


async def test_job_creation_rejects_changed_writing_scope_even_with_identical_node_ids(researched_directory):
    root, _, _, loaded = researched_directory
    changed = loaded["plan"].model_copy(deep=True)
    changed.sections[0].instruction = "仅介绍风险计量背景，无需描述实际损益和假设损益的回测比较。"
    assert [section.node_id for section in changed.sections] == [section.node_id for section in loaded["plan"].sections]
    with pytest.raises(ValueError, match='outline differs from the researched and reviewed version'):
        start_writing(root, loaded | {"plan": changed})
    assert loaded["plan"].sections[0].instruction == TECHNICAL_REQUIREMENTS["var_backtest"][1]
    with Store(root).connect() as connection:
        assert connection.execute("SELECT count(*) FROM jobs").fetchone()[0] == 0
    assert list((root / "jobs").iterdir()) == []


@pytest.mark.parametrize("mode", ["ready", "partial", "legacy"])
def test_api_confirmation_enforces_saved_research_and_freezes_ready_handoff(planning_setup, monkeypatch, mode):
    root, source = planning_setup
    monkeypatch.setattr(TailReadingRuntime, "partial", mode == "partial")
    ModelSettingsStore(root).save(ModelSettingsInput(base_url="http://unused.test/v1", model="fixture", no_auth=True))
    executed = []

    async def stop_after_freezing(store, job_id):
        executed.append(job_id)
        store.update(job_id, status="failed", error="Test stops after verifying frozen research inputs")

    monkeypatch.setattr("folio.api.workspace.execute", stop_after_freezing)
    app = create_app(root)
    with TestClient(app) as client:
        cid = client.post("/conversations").json()["id"]
        uploaded = client.post(f"/conversations/{cid}/sources", params={"filename": source.name}, content=source.read_bytes())
        assert uploaded.status_code == 201
        generated = client.post(f"/conversations/{cid}/messages", json={"action": "generate", "text": "编制本项目技术标"})
        assert "directory.review" in generated.text and "message.failed" not in generated.text
        selected = client.get(f"/conversations/{cid}").json()["directory_review"]
        if mode == "partial":
            assert selected["directory"]["research"]["status"] == "partial"
            assert selected["directory"]["confirmable"] is False
            assert selected["directory"]["confirmation_issues"][0]["code"] == "research_incomplete"
            metadata = json.loads((root / "directory-drafts" / selected["directory"]["revision"] / "metadata.json").read_text())
            assert metadata["confirmation_issues"] == selected["directory"]["confirmation_issues"]
        elif mode == "legacy":
            metadata_path = root / "directory-drafts" / selected["directory"]["revision"] / "metadata.json"
            metadata = json.loads(metadata_path.read_text())
            metadata.pop("research")
            metadata.pop("research_review")
            metadata.pop("confirmable")
            metadata_path.write_text(json.dumps(metadata, ensure_ascii=False))
            restored = client.get(f"/conversations/{cid}").json()["directory_review"]
            assert restored["directory"]["confirmable"] is False
        response = client.post(f"/conversations/{cid}/messages", json={
            "action": "confirm_directory", "text": "确认目录", "selection_revision": selected["revision"],
        })
        if mode == "ready":
            assert "job.created" in response.text and len(executed) == 1
            frozen = root / "jobs" / executed[0] / "inputs/directory-research"
            assert json.loads((frozen / "report.json").read_text())["status"] == "ready"
            assert json.loads((frozen / "coverage.json").read_text())["status"] == "passed"
        else:
            assert response.status_code == 409
            if mode == "partial":
                assert selected["directory"]["confirmation_issues"][0]["message"] in response.json()["detail"]
            else:
                assert "检查尚未完成" in response.json()["detail"]
            assert executed == []
            with app.state.store.connect() as connection:
                assert connection.execute("SELECT count(*) FROM jobs").fetchone()[0] == 0
