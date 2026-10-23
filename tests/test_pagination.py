import json
import re
import shutil
import subprocess
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from folio.document.pagination import office_runtime, paginate
from folio.document.reader import xpath
from folio.document.toc import field_run


def bookmark(paragraph, name, text, identity):
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    start.set(qn("w:id"), str(identity))
    start.set(qn("w:name"), name)
    end.set(qn("w:id"), str(identity))
    paragraph._p.append(start)
    paragraph.add_run(text)
    paragraph._p.append(end)


def field(paragraph, instruction, cached="999"):
    paragraph._p.extend(
        [field_run("begin"), field_run("instruction", instruction), field_run("separate")]
    )
    paragraph.add_run(cached)
    paragraph._p.append(field_run("end"))


@pytest.mark.parametrize("nested_toc", [False, True])
def test_real_pagination_uses_restarted_formatted_page_labels(tmp_path, nested_toc):
    if office_runtime() is None or shutil.which("pdftotext") is None:
        pytest.skip("LibreOffice/UNO and pdftotext are required")
    document = Document()
    numbering = OxmlElement("w:pgNumType")
    numbering.set(qn("w:start"), "1")
    numbering.set(qn("w:fmt"), "upperRoman")
    document.sections[0]._sectPr.append(numbering)
    bookmark(document.add_paragraph(), "front", "FRONT_MARKER", 1)
    entries = []
    for label, name in (("FRONT_REF", "front"), ("BODY_REF", "body"), ("LATER_REF", "later")):
        paragraph = document.add_paragraph(label + " = ")
        field(paragraph, f" PAGEREF {name} \\h ")
        entries.append(paragraph._p)
    if nested_toc:
        # Writer imports these nested references as a DocumentIndex, so they
        # cannot be found by enumerating GetReference fields alone.
        for index, run in enumerate(
            [
                field_run("begin"),
                field_run("instruction", ' TOC \\o "1-3" \\h '),
                field_run("separate"),
            ]
        ):
            entries[0].insert(index, run)
        entries[-1].append(field_run("end"))
        sdt, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
        sdt.append(content)
        entries[0].addprevious(sdt)
        for entry in entries:
            content.append(entry)
    section = document.add_section(WD_SECTION_START.NEW_PAGE)
    numbering = section._sectPr.find(qn("w:pgNumType"))
    numbering.set(qn("w:start"), "1")
    numbering.set(qn("w:fmt"), "decimal")
    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0]
    footer.add_run("BODY_FOOTER = ")
    field(footer, " PAGE ")
    bookmark(document.add_paragraph(), "body", "BODY_FIRST_MARKER", 2)
    document.add_page_break()
    bookmark(document.add_paragraph(), "later", "BODY_SECOND_MARKER", 3)
    source, pdf = tmp_path / "restarted.docx", tmp_path / "restarted.pdf"
    document.save(source)
    before_sections = [section._sectPr.xml for section in Document(source).sections]
    with ZipFile(source) as archive:
        before_parts = {name: archive.read(name) for name in archive.namelist()}

    result = paginate(source, pdf=pdf, required=True)

    assert result["display_pages"] == {"front": "I", "body": "1", "later": "2"}
    assert result["pages"]["body"] > 1  # Office layout index is a separate diagnostic.
    assert result["toc_entries"] == 3
    assert [section._sectPr.xml for section in Document(source).sections] == before_sections
    text = subprocess.run(
        ["pdftotext", "-layout", str(pdf), "-"], check=True, capture_output=True, text=True
    ).stdout
    assert re.search(r"FRONT_REF\s*=\s*I\b", text)
    assert re.search(r"BODY_REF\s*=\s*1\b", text)
    assert re.search(r"LATER_REF\s*=\s*2\b", text)
    for marker, displayed in (("BODY_FIRST_MARKER", "1"), ("BODY_SECOND_MARKER", "2")):
        page = next(page for page in text.split("\f") if marker in page)
        assert re.search(r"BODY_FOOTER\s*=\s*" + displayed + r"\b", page)
    cached = Document(source)
    cached_text = "".join(xpath(cached.element, ".//w:t/text()"))
    for reference in ("FRONT_REF = I", "BODY_REF = 1", "LATER_REF = 2"):
        assert reference in cached_text
    with ZipFile(source) as archive:
        for name, original in before_parts.items():
            if name not in {"word/document.xml", "word/settings.xml"}:
                assert archive.read(name) == original


def test_pagination_waits_for_changed_toc_layout_and_publishes_matching_pdf(tmp_path, monkeypatch):
    source, pdf = tmp_path / "toc.docx", tmp_path / "toc.pdf"
    document = Document()
    field(document.add_paragraph(), " PAGEREF body \\h ", cached="0")
    document.save(source)
    seen = []

    class Worker:
        pid = 999999
        returncode = 0

        def __init__(self, command, **kwargs):
            request = json.loads(open(command[-1], encoding="utf-8").read())
            current = "".join(xpath(Document(request["source"]).element, ".//w:t/text()"))
            seen.append(current)
            # Caching the first measured value changes TOC wrapping and therefore
            # the next layout. Only the third pass sees the same map as its input.
            displayed = "2" if len(seen) == 1 else "3"
            from pathlib import Path

            Path(request["result"]).write_text(
                json.dumps(
                    {
                        "pages": {"body": 7},
                        "display_pages": {"body": displayed},
                        "labels": {},
                        "total_pages": 7,
                    }
                )
            )
            Path(request["pdf"]).write_text("PDF rendered with TOC " + current)

        def communicate(self, timeout):
            return "", ""

        def wait(self):
            return 0

    monkeypatch.setattr(
        "folio.document.pagination.office_runtime",
        lambda: {"python": "unused", "soffice": "unused"},
    )
    monkeypatch.setattr(
        "folio.document.pagination.subprocess",
        SimpleNamespace(Popen=Worker, PIPE=subprocess.PIPE),
    )
    monkeypatch.setattr("folio.document.pagination.os.killpg", lambda *args: None)
    result = paginate(source, pdf=pdf, required=True)
    assert seen == ["0", "2", "3"]
    assert result["display_pages"] == {"body": "3"}
    assert "".join(xpath(Document(source).element, ".//w:t/text()")) == "3"
    assert pdf.read_text() == "PDF rendered with TOC 3"
