"""Reader-facing Word output keeps built-in typography and internal records separate."""

import json
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt

from folio.document.presentation import (
    normalize_content,
    prose_paragraphs,
    validate_document_content,
)
from folio.document.reader import xpath
from folio.document.typography import TABLE_BODY, TABLE_HEADER
from folio.document.word import render
from folio.models import (
    Block,
    DocumentPlan,
    Evidence,
    Extraction,
    JobInput,
    Section,
    SectionContent,
    SectionState,
)
from folio.orchestration.engine import Engine
from folio.runtime.base import RuntimeResult


def test_numbered_items_are_paragraphs_without_splitting_decimals_or_versions():
    text = "涉及以下角色：1. 交易员录入；2. 复核员核对；3. 清算员处理。"
    assert prose_paragraphs(text) == [
        "涉及以下角色：",
        "1. 交易员录入；",
        "2. 复核员核对；",
        "3. 清算员处理。",
    ]
    assert prose_paragraphs("利率1.25%，版本V1.0，日期2026.09.28，章节1.2.3。") == [
        "利率1.25%，版本V1.0，日期2026.09.28，章节1.2.3。"
    ]
    source = SectionContent(
        blocks=[
            Block(block_id="b", text=text, origin="model_proposal"),
            Block(block_id="b_p1", text="保留本段。"),
        ]
    )
    output = normalize_content(source)
    assert len(output.blocks) == 5
    assert len({b.block_id for b in output.blocks}) == 5
    assert all(b.origin == "model_proposal" for b in output.blocks[:4])
    assert len(source.blocks) == 2


@pytest.mark.parametrize(
    "text", ["参考ev_abcdef012345", "依据 sfref12", "model_proposal", "【模型建议】填写办法"]
)
@pytest.mark.parametrize("table", [False, True])
def test_internal_notation_is_rejected_in_body_and_cells(text, table):
    block = (
        Block(block_id="b", kind="table", rows=[["说明"], [text]])
        if table
        else Block(block_id="b", text=text)
    )
    with pytest.raises(ValueError):
        validate_document_content(SectionContent(blocks=[block]))


def test_evidence_fields_remain_valid_but_audit_columns_are_not_prose():
    validate_document_content(
        SectionContent(blocks=[Block(block_id="b", text="交易员录入。", evidence_ids=["ev_123"])])
    )
    with pytest.raises(ValueError, match="evidence-source"):
        validate_document_content(
            SectionContent(
                blocks=[
                    Block(
                        block_id="b",
                        kind="table",
                        rows=[["名称", "证据来源"], ["规则", "内部编号"]],
                    )
                ]
            )
        )


def test_red_supplements_preserve_builtin_fonts_and_table_sizes(tmp_path):
    target = tmp_path / "out.docx"
    plan = DocumentPlan(
        sections=[
            Section(
                section_id="s",
                node_id="o",
                title="业务要求",
                level=1,
                display_number="1",
                instruction="描述业务",
            )
        ]
    )
    state = SectionState(
        section_id="s",
        evidence=[
            Evidence(
                evidence_id="ev_123",
                source_id="后台来源",
                source_version="1",
                locator="1",
                text="交易员录入。",
            )
        ],
        extraction=Extraction(summary="业务"),
        content=SectionContent(
            blocks=[
                Block(block_id="fact", text="交易员录入。", evidence_ids=["ev_123"]),
                Block(
                    block_id="proposal",
                    text="1. 校验失败返回提示；2. 允许修正后重试。",
                    origin="model_proposal",
                ),
                Block(
                    block_id="table",
                    kind="table",
                    origin="model_proposal",
                    rows=[["名称", "说明", "备注"], ["校验", "逐项检查", ""]],
                ),
            ]
        ),
    )
    render(target, plan, [state], title="测试", draft=True)
    doc = Document(target)
    assert "【模型建议】" not in "".join(p.text for p in doc.paragraphs)
    assert "证据来源" not in "".join(p.text for p in doc.paragraphs)
    for label in ["1. 校验失败返回提示；", "2. 允许修正后重试。"]:
        paragraph = next(p for p in doc.paragraphs if p.text.replace("\t", " ").strip() == label)
        assert paragraph.style.font.size == Pt(11)
        assert xpath(paragraph.style.element, "./w:rPr/w:rFonts/@w:eastAsia") == ["宋体"]
        assert all(str(r.font.color.rgb) == "FF0000" for r in paragraph.runs)
    fact = next(p for p in doc.paragraphs if p.text == "交易员录入。")
    assert all(r.font.color.rgb is None for r in fact.runs)
    generated = doc.tables[-1]
    for i, row in enumerate(generated.rows):
        for cell in row.cells:
            style = cell.paragraphs[0].style
            assert style.name == (TABLE_HEADER if i == 0 else TABLE_BODY)
            assert style.font.size == Pt(11)
            assert xpath(style.element, "./w:rPr/w:rFonts/@w:eastAsia") == ["宋体"]
            assert all(str(r.font.color.rgb) == "FF0000" for p in cell.paragraphs for r in p.runs)


async def test_internal_gap_reports_are_not_appended_to_written_business_content(store, inputs):
    from test_business_work import GroupRuntime

    class Reports(GroupRuntime):
        async def complete(self, name, context, schema, session):
            result = await super().complete(name, context, schema, session)
            if result.tool_calls:
                if name == "next_action" and context["evidence"]:
                    data = json.loads(result.tool_calls[0]["function"]["arguments"])
                    data["missing_information"] = ["本文件审批人未知"]
                    result.tool_calls[0]["function"]["arguments"] = json.dumps(data)
                return result
            data = json.loads(result.content)
            if name == "complete_gaps":
                for proposal in data["proposals"]:
                    proposal["disposition"] = "record_only"
                    proposal["content"] = (
                        "编写人和审批人信息缺失，应仅在后台记录，不能编造或附加收集信息的工作指南。"
                    )
            return RuntimeResult(json.dumps(data))

    job_id = store.create(JobInput(**inputs, runtime="llm"))
    job = await Engine(store, Reports()).run(job_id)
    work = next(iter(store.work_states(job_id).values()))
    assert len(work.searches) == 3 and work.proposals[0].disposition == "record_only"
    assert all("工作指南" not in b.text for s in store.states(job_id) for b in s.content.blocks)
    assert "审批人信息缺失" not in "".join(p.text for p in Document(job["artifact"]).paragraphs)
    assert Path(job["artifact"]).exists()


def test_long_lists_keep_checkpoints_valid_and_domain_field_names_are_not_internal_ids():
    original = SectionContent(
        blocks=[Block(block_id="list", text="".join(f"{i}. 项目说明；" for i in range(1, 130)))]
    )
    normalized = normalize_content(original)
    SectionContent.model_validate_json(normalized.model_dump_json())
    assert len(prose_paragraphs(normalized.blocks[0].text)) == 129
    legitimate = SectionContent(
        blocks=[
            Block(
                block_id="field",
                text="接口字段 task_id 表示业务调度任务编号，ev_status 表示事件状态。",
            )
        ]
    )
    validate_document_content(legitimate)
    with pytest.raises(ValueError):
        validate_document_content(
            SectionContent(blocks=[Block(block_id="ref", text="参考 ev_demo")]), ["ev_demo"]
        )


def test_confirmed_children_keep_numbering_generated_leadins_do_not_compete(tmp_path):
    target = tmp_path / "out.docx"
    plan = DocumentPlan(
        sections=[
            Section(
                section_id="parent",
                node_id="p",
                title="功能",
                level=1,
                display_number="1",
                instruction="总述",
            ),
            Section(
                section_id="child",
                parent_id="parent",
                node_id="c",
                title="已确认子节",
                level=2,
                display_number="1.1",
                instruction="详情",
            ),
        ]
    )
    states = [
        SectionState(
            section_id="parent",
            extraction=Extraction(summary="总述"),
            evidence=[],
            content=SectionContent(
                blocks=[
                    Block(block_id="h", kind="heading", text="1.1 范围说明", heading_level=2),
                    Block(block_id="p", text="这里是概括内容。"),
                ]
            ),
        ),
        SectionState(
            section_id="child",
            extraction=Extraction(summary="详情"),
            evidence=[],
            content=SectionContent(blocks=[Block(block_id="p", text="这里是具体内容。")]),
        ),
    ]
    render(target, plan, states, title="测试", draft=True)
    doc = Document(target)
    lead = next(p for p in doc.paragraphs if p.text == "范围说明")
    assert lead.style.name == "Folio Body"
    assert not xpath(lead._p, "./w:pPr/w:numPr")
    assert len(xpath(doc.element, ".//w:bookmarkStart")) == 2


def test_provenance_gaps_distinguish_structure_from_assertions():
    from folio.document.presentation import missing_provenance

    assert missing_provenance(Block(block_id="b", text="维护窗口为星期日凌晨。"))
    assert not missing_provenance(
        Block(block_id="b", text="维护窗口为星期日凌晨。", origin="model_proposal")
    )
    assert not missing_provenance(Block(block_id="b", text="交易员录入。", evidence_ids=["ev_1"]))
    assert not missing_provenance(Block(block_id="h", kind="heading", text="业务流程"))
    assert not missing_provenance(
        Block(block_id="t", kind="table", rows=[["编写人", "日期"], ["", ""]])
    )
    assert missing_provenance(Block(block_id="t", kind="table", rows=[["可用性"], ["99.9%"]]))


async def test_missing_citations_trigger_three_searches_then_proposal_rewrite(store, inputs):
    from test_business_work import GroupRuntime

    class MissingCitation(GroupRuntime):
        def __init__(self):
            super().__init__()
            self.generated = {}

        async def complete(self, name, context, schema, session):
            response = await super().complete(name, context, schema, session)
            if response.tool_calls:
                return response
            data = json.loads(response.content)
            if name == "generate":
                data["blocks"].append(
                    {
                        "block_id": "untraced",
                        "text": "系统支持维护窗口配置。",
                        "origin": "evidence",
                        "evidence_ids": [],
                    }
                )
                self.generated[context["section"]["section_id"]] = SectionContent.model_validate(
                    data
                )
            return RuntimeResult(json.dumps(data))

    runtime = MissingCitation()
    job_id = store.create(JobInput(**inputs, runtime="llm"))
    result = await Engine(store, runtime).run(job_id)
    assert result["status"] in {"draft", "needs_confirmation", "published"}
    assert all(len(w.searches) == 3 for w in store.work_states(job_id).values())
    integrations = [c for name, c in runtime.calls if name == "integrate_supplement"]
    assert integrations
    assert {c["section"]["section_id"] for c in integrations} == set(runtime.generated)
    assert all(len(c["model_proposals"]) == 1 for c in integrations)
    assert all(
        c["integration_item"]["gap_id"] == c["model_proposals"][0]["gap_id"] for c in integrations
    )
    from folio.document.presentation import missing_provenance

    assert not any(missing_provenance(b) for s in store.states(job_id) for b in s.content.blocks)
    for state in store.states(job_id):
        blocks = {block.block_id: block for block in state.content.blocks}
        assert blocks["untraced"].text == "系统支持维护窗口配置。"
        assert blocks["untraced"].origin == "model_proposal"
        assert blocks["untraced"].evidence_ids == []
        for original in runtime.generated[state.section_id].blocks:
            if original.block_id != "untraced":
                assert blocks[original.block_id] == original


def test_provenance_report_does_not_relabel_or_promote_old_content():
    from folio.document.presentation import provenance_summary

    state = SectionState(
        section_id="s",
        extraction=Extraction(summary="test"),
        evidence=[],
        content=SectionContent(
            blocks=[
                Block(block_id="missing", text="未知依据"),
                Block(block_id="invalid", text="错误引用", evidence_ids=["ev_absent"]),
                Block(block_id="proposal", text="拟议业务设计", origin="model_proposal"),
            ]
        ),
    )
    before = state.model_dump_json()
    report = provenance_summary([state])
    assert report["total_blocks"] == 3 and report["proposal_blocks"] == 1
    assert report["untraced_blocks"] == [{"section_id": "s", "block_id": "missing"}]
    assert report["invalid_references"] == [
        {"section_id": "s", "block_id": "invalid", "evidence_ids": ["ev_absent"]}
    ]
    assert state.model_dump_json() == before
