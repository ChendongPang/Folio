from docx import Document

from folio.document.planner import build_plan
from folio.document.word import render
from folio.models import PlanItem, PlanProposal


def test_builtin_cover_uses_explicit_title_and_preview_status(tmp_path):
    plan = build_plan(
        "document",
        PlanProposal(
            sections=[PlanItem(node_id="body", title="业务需求", level=1, instruction="说明功能")]
        ),
    )
    target = tmp_path / "preview.docx"
    render(
        target,
        plan,
        [],
        title="需求说明书",
        draft=True,
        preview=True,
        cover={"document_title": "交易系统需求说明书"},
    )
    document = Document(target)
    assert document.paragraphs[0].text == "交易系统需求说明书"
    assert document.paragraphs[1].text == "目录预览"
    assert any("业务需求" in p.text for p in document.paragraphs)


def test_english_preview_uses_separate_compact_contents_style(tmp_path):
    plan = build_plan(
        "document",
        PlanProposal(
            document_title="Website Redesign Proposal",
            sections=[PlanItem(node_id="body", title="Business Requirements", level=1, instruction="Describe functions")],
        ),
    )
    target = tmp_path / "preview.docx"
    render(
        target,
        plan,
        [],
        title="Fallback title",
        draft=True,
        preview=True,
        defer_pagination=True,
        language="en",
    )
    document = Document(target)
    title = document.paragraphs[0]
    contents = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Contents")
    assert title.text == "Website Redesign Proposal"
    assert document.paragraphs[1].text == "OUTLINE PREVIEW"
    assert contents.style.name == "Folio TOC Heading"
    assert contents.style.font.size < title.style.font.size
