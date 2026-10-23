"""Model-visible research must deliver source text, not metadata-heavy wrappers."""

import json
import re

import pytest
from docx import Document

from folio.research.planning import (
    PlanningResearchLimits,
    PlanningResearchStep,
    _DirectoryResearch,
    research_for_directory,
)
from folio.research.transport import research_messages, research_prompt_chars
from folio.research.workspace import WorkspaceManager
from folio.runtime.base import RuntimeResult


class ReadableRuntime:
    def __init__(self, policy):
        self.policy = policy
        self.calls = []

    def prompt_chars(self, task, context, schema):
        return research_prompt_chars(context, schema)

    def request_parts(self, task, context, schema):
        return research_messages(context, schema), schema, {}, None, {}

    async def complete(self, task, context, schema, session_key):
        self.calls.append(context)
        value = self.policy(context)
        return RuntimeResult("", tool_calls=[{
            "id": f"call_{len(self.calls)}", "type": "function", "function": {
                "name": value["tool"].replace(".", "_"),
                "arguments": json.dumps({"summary": "按资料结构核对原文", **value["arguments"]}, ensure_ascii=False),
            },
        }])


def step(tool, **arguments):
    return {"tool": tool, "arguments": arguments}


@pytest.mark.asyncio
async def test_large_canonical_result_is_delivered_as_text_and_short_refs_resolve(tmp_path):
    sources = tmp_path / "sources"
    sources.mkdir()
    name = "a" * 32 + "__" + "市场风险管理系统信创改造招标技术要求" * 4 + ".docx"
    document = Document()
    document.add_heading("技术要求", 1)
    for i in range(80):
        document.add_paragraph(f"第{i + 1}项：" + "必须按原始需求进行验收并提供对应证明。" * 6)
    document.save(sources / name)

    def policy(context):
        if not context["observations"]:
            assert "F1" in context["manifest_text"]
            return step("fs.inspect", file="F1")
        observation = context["observations"][-1]
        if observation["tool"] == "fs.inspect":
            assert "N001 技术要求" in observation["model_text"]
            return step("fs.read", file="F1", node="N001", limit=100)
        text = observation["model_text"]
        assert "第80项" in text
        assert "source_version" not in text and "metadata" not in text
        ids = re.findall(r"\[(E\d+)\]", text)
        assert "读取范围：" not in text
        return step("finish", summary="已读取技术要求。",
                    requirements=[{"requirement_id": "req1", "text": "逐项安排技术响应", "evidence_ids": ids[:1]}])

    runtime = ReadableRuntime(policy)
    report = await research_for_directory(runtime, sources, tmp_path / "research", "编制技术标")
    assert report["status"] == "ready", report["validation_issues"]
    assert report["calls"] == 3
    read = runtime.calls[2]["observations"][-1]
    assert len(json.dumps(read["result"], ensure_ascii=False)) > 48000
    assert len(read["model_text"]) < 16000
    assert report["files"][0]["characters"] == report["files"][0]["delivered_characters"]
    assert report["files"][0]["structure_delivered"]
    assert report["requirements"][0]["evidence_ids"][0].startswith("ev_")
    recorded = json.loads((tmp_path / "research/events/research_003-request.json").read_text())
    assert recorded["messages"] == runtime.request_parts("directory_research", runtime.calls[2],
                                                        PlanningResearchStep.model_json_schema())[0]
    visible = "\n".join(str(message.get("content") or "") for message in recorded["messages"])
    assert "finish_contract" not in visible and "arguments_json" not in visible
    assert "第80项" in visible


@pytest.mark.asyncio
async def test_last_extracted_result_does_not_count_as_delivered(tmp_path):
    sources = tmp_path / "sources"
    sources.mkdir()
    (sources / "rules.md").write_text("# 业务要求\n必须支持库存校验。", encoding="utf-8")
    runtime = ReadableRuntime(lambda ctx: step("fs.inspect", file="F1") if not ctx["observations"]
                              else step("fs.read", file="F1", node="N001"))
    report = await research_for_directory(runtime, sources, tmp_path / "research", "编写目录",
                                          limits=PlanningResearchLimits(max_calls=2))
    assert report["files"][0]["characters"] > 0
    assert report["files"][0]["delivered_characters"] == 0
    assert not report["files"][0]["read_ranges"][0]["delivered_complete"]
    assert report["status"] == "partial"


@pytest.mark.asyncio
async def test_history_read_counts_only_complete_quotes_and_full_range_after_all_pages(tmp_path):
    sources = tmp_path / "sources"
    sources.mkdir()
    (sources / "rules.md").write_text("# 要求\n" + "库存校验。" * 500, encoding="utf-8")
    output = tmp_path / "research"
    output.mkdir()
    workspace = WorkspaceManager(sources, output / "workspace")
    state = _DirectoryResearch(workspace, output, "目录", None, PlanningResearchLimits(), None)
    initial = workspace.list()
    state.register("fs.list", {}, initial)
    args = {"path": "rules.md", "limit": 100}
    result = await state.dispatch("fs.read", args)
    observation = {"request_id": "research_001", "tool": "fs.read", "arguments": args,
                   "result": result, "model_text": state.presentation.result_text("fs.read", args, result)}
    state.observations.append(observation)
    state.history["research_001"] = result
    offset = 0
    count = 0
    while True:
        history_args = {"scope": "history", "request_id": "research_001", "offset": offset, "limit": 1000}
        page = await state.dispatch("fs.read", history_args)
        assert not page["text"].startswith('{"evidence"')
        count += 1
        state.delivered({"manifest": initial, "observations": [{
            "request_id": f"research_{count + 1:03}", "tool": "fs.read", "arguments": history_args,
            "result": page, "model_text": page["text"],
        }]}, f"model_{count}")
        if page["next_offset"] is None:
            break
        assert not next(iter(state.ranges.values()))["delivered_complete"]
        assert all(state.evidence[key]["text"] in page["text"] for key in state.visible_evidence({
            "model_text": page["text"], "result": page,
        }))
        offset = page["next_offset"]
    assert state.delivered_evidence == set(state.evidence)
    assert next(iter(state.ranges.values()))["delivered_complete"]


@pytest.mark.asyncio
async def test_full_history_upgrades_trimmed_same_request_in_one_context(tmp_path):
    from copy import deepcopy

    sources = tmp_path / "sources"
    sources.mkdir()
    (sources / "rules.md").write_text("# 要求\n" + "库存校验。" * 500, encoding="utf-8")
    output = tmp_path / "research"
    output.mkdir()
    workspace = WorkspaceManager(sources, output / "workspace")
    state = _DirectoryResearch(workspace, output, "目录", None, PlanningResearchLimits(), None)
    initial = workspace.list()
    state.register("fs.list", {}, initial)
    args = {"path": "rules.md", "limit": 100}
    result = await state.dispatch("fs.read", args)
    original = {"request_id": "research_001", "tool": "fs.read", "arguments": args,
                "result": result, "model_text": state.presentation.result_text("fs.read", args, result)}
    state.observations.append(original)
    state.history["research_001"] = result
    trimmed = deepcopy(original)
    trimmed["result"]["evidence"] = trimmed["result"]["evidence"][:1]
    trimmed["result"]["read_segments"] = trimmed["result"]["read_segments"][:1]
    trimmed["model_text"] = state.presentation.result_text("fs.read", args, trimmed["result"])
    history_args = {"scope": "history", "request_id": "research_001"}
    history = await state.dispatch("fs.read", history_args)
    assert history["complete"]
    state.delivered({"manifest": initial, "observations": [trimmed, {
        "request_id": "research_002", "tool": "fs.read", "arguments": history_args,
        "result": history, "model_text": history["text"],
    }]}, "model_3")
    assert state.delivered_evidence == set(state.evidence)
    assert next(iter(state.ranges.values()))["delivered_complete"]


def test_engine_preserves_paged_research_requirements():
    from types import SimpleNamespace

    from folio.orchestration.requests import RequestWorkflow

    engine = SimpleNamespace(request=SimpleNamespace(user_request="完整用户要求" * 10000, writing_guide=None))
    page = {"text": "首段要求", "offset": 0, "next_offset": 4, "total": 60000}
    context = RequestWorkflow.request_context(engine, "directory_research", {"requirements": page})
    assert context["requirements"] == page
    assert "user_request" not in context
    business = RequestWorkflow.request_context(engine, "research_step", {"user_request": page})
    assert business["user_request"] == page


@pytest.mark.asyncio
async def test_cursor_reads_continue_exact_source_range(tmp_path):
    sources = tmp_path / "sources"
    sources.mkdir()
    (sources / "rules.md").write_text("# 要求\n" + "规则。" * 1800, encoding="utf-8")
    calls = 0
    ids = []

    def policy(context):
        nonlocal calls
        calls += 1
        if calls == 1:
            return step("fs.inspect", file="F1")
        if calls == 2:
            return step("fs.read", file="F1", node="N001", limit=1)
        text = context["observations"][-1]["model_text"]
        ids.extend(re.findall(r"\[(E\d+)\]", text))
        cursor = re.search(r"cursor=(C\d+)", text)
        if cursor:
            return step("fs.read", cursor=cursor.group(1), limit=1)
        return step("finish", summary="已读取全部要求",
                    requirements=[{"requirement_id": "r1", "text": "覆盖规则", "evidence_ids": list(dict.fromkeys(ids))}],
                    selected_ranges=[{"range_id": "R1", "reason": "本次范围"}])

    report = await research_for_directory(ReadableRuntime(policy), sources, tmp_path / "research", "写目录")
    assert report["status"] == "ready", report["validation_issues"]
    assert len(report["files"][0]["read_ranges"]) == 1
    assert report["files"][0]["characters"] == report["files"][0]["delivered_characters"]


def _text_pdf(path, pages):
    """A real multipage PDF with extractable source text, without rendering deps."""
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    for number in range(1, pages + 1):
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        lines = [f"Page {number}, clause {clause}: Cancelled orders release stock; approval logs identify the operator."
                 for clause in range(1, 27)]
        stream = DecodedStreamObject()
        stream.set_data(("BT /F1 9 Tf 12 TL 30 750 Td "
                         + " ".join(f"({line}) Tj T*" for line in lines) + " ET").encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as output:
        writer.write(output)


@pytest.mark.asyncio
async def test_43_page_research_keeps_notes_and_reuses_old_evidence_across_files(tmp_path):
    sources = tmp_path / "sources"
    sources.mkdir()
    _text_pdf(sources / "tender.pdf", 43)
    (sources / "supplement.md").write_text("# Additional acceptance requirement\nCancellation audit logs must include approval timestamps.")
    files = ["tender.pdf", "supplement.md"]
    inspected, reads, ranges, seen, evidence = {}, {}, {}, set(), {}
    native_reads = []
    seen_old_note_without_raw = False

    def policy(context):
        nonlocal seen_old_note_without_raw
        notes = []
        for observation in context["observations"]:
            key, result = observation["request_id"], observation["result"]
            if key in seen or "error" in result:
                continue
            seen.add(key)
            args = observation["arguments"]
            if observation["tool"] == "fs.inspect":
                inspected[args["path"]] = result["next_offset"]
            elif observation["tool"] == "fs.read" and args.get("scope", "sources") == "sources":
                path = args["path"]
                reads[path] = (None if result["next_offset"] is None else
                               (result["next_offset"], result["next_character_offset"]))
                ranges[path] = result["range_id"]
                evidence.setdefault(path, []).extend(raw["evidence_id"] for raw in result["evidence"])
                notes.append({"topic": path + " · " + result["evidence"][0]["locator"],
                              "text": "取消订单释放库存，验收需要按来源逐项核对审批操作日志。",
                              "kind": "source_statement",
                              "evidence_ids": [item["evidence_id"] for item in result["evidence"]]})
        if "supplement.md" in inspected:
            assert "tender.pdf" in context["research_notes"]
            assert "取消订单释放库存" in context["research_notes"]
            assert "tender.pdf" in context["history_text"]
            seen_old_note_without_raw |= not any(
                observation["tool"] == "fs.read" and observation["arguments"].get("path") == "tender.pdf"
                and observation["arguments"].get("offset", 0) == 0
                for observation in context["observations"])
        if context.get("memory_checkpoint_required"):
            return step("research.save", research_notes=notes)
        for path in files:
            if path not in inspected or inspected[path] is not None:
                return step("fs.inspect", path=path, offset=inspected.get(path) or 0, research_notes=notes)
            if path not in reads or reads[path] is not None:
                selector = (path, *(reads.get(path) or (0, 0)))
                assert selector not in native_reads, "a previously delivered page must not be fetched again"
                native_reads.append(selector)
                return step("fs.read", path=path, offset=selector[1], character_offset=selector[2],
                            limit=8, max_chars=14000,
                            research_notes=notes)
        return step("finish", summary="两份资料已研究，按取消规则与审批审计组织目录。", research_notes=notes,
                    requirements=[{"requirement_id": "cancel-and-audit", "text": "覆盖撤单释放库存与审批时间戳。",
                                   "evidence_ids": [evidence[path][0] for path in files]}],
                    selected_ranges=[{"range_id": ranges[path], "reason": "决定验收与业务响应范围"} for path in files])

    runtime = ReadableRuntime(policy)
    output = tmp_path / "research"
    report = await research_for_directory(runtime, sources, output, "研究43页技术要求和补充材料，生成响应目录。",
                                          limits=PlanningResearchLimits(max_context_chars=48000))
    assert report["status"] == "ready", report["blocking_gaps"]
    assert seen_old_note_without_raw
    assert len(native_reads) > 5
    assert {raw["metadata"]["page"] for raw in report["evidence"] if raw["source_id"] == "tender.pdf"} == set(range(1, 44))
    assert all(runtime.prompt_chars("directory_research", context, PlanningResearchStep.model_json_schema()) <= 48000
               for context in runtime.calls)
    assert report["research_memory"]["notes"]
    assert all(not batch["evidence_ids"] or batch["reviewed"] for batch in report["research_memory"]["batches"].values())
    saved = (output / "research.md").read_text()
    assert "tender.pdf" in saved and "supplement.md" in saved
    assert "not original evidence" in saved


@pytest.mark.asyncio
async def test_partial_batch_advances_past_summarized_prefix_without_fabricating_delivery(tmp_path):
    sources = tmp_path / "sources"
    sources.mkdir()
    # Independent source paragraphs make this a whole-evidence prefix test;
    # source readers no longer manufacture fixed-size chunks from a long block.
    (sources / "rules.md").write_text("# Requirements\n" + "\n\n".join(
        f"Clause {index}: " + "Cancellation releases stock. " * 50 for index in range(5)))
    output = tmp_path / "research"
    output.mkdir()
    workspace = WorkspaceManager(sources, output / "workspace")
    state = _DirectoryResearch(workspace, output, "目录", None,
                               PlanningResearchLimits(max_context_chars=4000), None)
    initial = workspace.list()
    state.register("fs.list", {}, initial)
    args = {"path": "rules.md", "limit": 100}
    result = await state.dispatch("fs.read", args)
    call = {"id": "native_original", "type": "function", "function": {
        "name": "fs_read", "arguments": json.dumps(args)}}
    original = {"request_id": "research_001", "tool": "fs.read", "arguments": args, "tool_call": call,
                "result": result, "model_text": state.presentation.result_text("fs.read", args, result)}
    state.observe(original)
    state.history["research_001"] = result

    class PayloadCapacity:
        @staticmethod
        def prompt_chars(task, context, schema):
            # Simulate capacity for exactly one complete quote, independent of
            # the quote's character length, so no partial quote is required.
            return 500 + len(context["research_notes"]) + 2000 * sum(
                len(item["result"].get("evidence", [])) for item in context["observations"])

    shown = []
    for index in range(10):
        context = state.context(initial, [], PayloadCapacity(), {})
        if not context["observations"]:
            break
        observation = context["observations"][-1]
        assert observation["tool_call"] == call
        ids = state.visible_evidence(observation)
        assert ids and not set(ids) & set(shown), "already summarized prefix must not be replayed forever"
        shown.extend(ids)
        state.delivered(context, f"model_{index}")
        assert state.delivered_evidence == set(shown)
        state.memory.apply([{"topic": f"Batch {index}", "text": "已核对本批取消规则。",
                             "kind": "source_statement", "evidence_ids": ids,
                             "request_ids": ["research_001"]}], evidence=state.evidence, call_id=f"save_{index}")
    assert len(shown) == len(state.evidence)
    assert state.memory.can_drop("research_001")
    assert next(iter(state.ranges.values()))["delivered_complete"]
    assert len(state.observations) == 1  # no fake tool executions to advance a context page
    for field in ("research_notes", "research_history"):
        parts, offset = [], 0
        while True:
            page = await state.dispatch("fs.read", {"scope": "context", "field": field, "offset": offset, "limit": 80})
            parts.append(page["text"])
            if page["next_offset"] is None:
                break
            offset = page["next_offset"]
        assert "".join(parts) == state.memory_context(field)
    assert not state.context_pages  # reading evolving notes is not a new finish requirement


@pytest.mark.asyncio
async def test_model_token_capacity_uses_a_note_checkpoint_before_continuing(tmp_path):
    sources = tmp_path / "sources"
    sources.mkdir()
    (sources / "rules.md").write_text("# Requirements\nCancelled orders release stock.")
    calls = 0

    def policy(context):
        nonlocal calls
        calls += 1
        if calls == 1:
            return step("fs.inspect", file="F1")
        if calls == 2:
            return step("fs.read", file="F1")
        if context.get("memory_checkpoint_required"):
            item = context["observations"][0]
            return step("research.save", research_notes=[{
                "topic": "取消规则", "text": "取消订单释放库存。", "kind": "source_statement",
                "evidence_ids": [raw["evidence_id"] for raw in item["result"]["evidence"]],
                "request_ids": [item["request_id"]],
            }])
        assert "取消订单释放库存" in context["research_notes"]
        return step("finish", requirements=[{"requirement_id": "cancel", "text": "取消订单释放库存。", "evidence_ids": ["E1"]}],
                    selected_ranges=[{"range_id": "R1", "reason": "业务范围"}])

    class TokenCapacityRuntime(ReadableRuntime):
        measurements = []

        async def measure(self, task, context, schema, char_limit):
            has_raw = any(item["tool"] == "fs.read" for item in context["observations"])
            fits = not has_raw or bool(context.get("memory_checkpoint_required"))
            self.measurements.append(fits)
            return {"fits": fits, "measurement": "test_token_counter", "characters": self.prompt_chars(task, context, schema)}

    runtime = TokenCapacityRuntime(policy)
    report = await research_for_directory(runtime, sources, tmp_path / "research", "写目录")
    assert report["status"] == "ready", report["blocking_gaps"]
    assert report["calls"] == 4
    assert runtime.measurements == [True, True, False, True, False, True]
    assert runtime.calls[2]["memory_checkpoint_required"]
    request = json.loads((tmp_path / "research/events/research_003-request.json").read_text())
    assert [tool["function"]["name"] for tool in request["tools"]] == ["research_save"]


@pytest.mark.asyncio
async def test_capacity_timeout_does_not_spend_a_generation_call_or_drop_research(tmp_path):
    import asyncio

    sources = tmp_path / "sources"
    sources.mkdir()
    (sources / "rules.md").write_text("# Requirements\nCancelled orders release stock.")

    class SlowCapacityRuntime(ReadableRuntime):
        async def measure(self, task, context, schema, char_limit):
            await asyncio.sleep(10)

    runtime = SlowCapacityRuntime(lambda context: pytest.fail("must not send an unmeasured request"))
    report = await research_for_directory(runtime, sources, tmp_path / "research", "写目录",
                                          limits=PlanningResearchLimits(call_timeout_seconds=0.01))
    assert report["status"] == "partial"
    assert report["calls"] == 0
    assert any("capacity checking timed out" in gap for gap in report["blocking_gaps"])
    assert (tmp_path / "research/research.md").exists()


@pytest.mark.asyncio
async def test_capacity_available_keeps_native_navigation_reads_and_full_notes(tmp_path):
    sources = tmp_path / "sources"
    sources.mkdir()
    (sources / "rules.md").write_text("# 审批规则\n必须由主管审批。")
    output = tmp_path / "research"
    output.mkdir()
    workspace = WorkspaceManager(sources, output / "workspace")
    state = _DirectoryResearch(workspace, output, "目录", None, PlanningResearchLimits(), None)
    initial = workspace.list()
    state.register("fs.list", {}, initial)
    for number, (tool, args) in enumerate([
        ("fs.list", {}), ("fs.inspect", {"path": "rules.md"}), ("fs.read", {"path": "rules.md"}),
    ], 1):
        result = await state.dispatch(tool, args)
        state.observe({"request_id": f"research_{number:03}", "tool": tool, "arguments": args,
                       "tool_call": {"id": f"call_{number}", "type": "function", "function": {
                           "name": tool.replace(".", "_"), "arguments": json.dumps(args)}},
                       "result": result, "model_text": state.presentation.result_text(tool, args, result)})
        context = state.context(initial, [], ReadableRuntime(None), {})
        state.delivered(context, f"model_{number}")
    evidence_id = next(iter(state.evidence))
    state.memory.apply([{"topic": f"规则 {index}", "text": (f"审批要求 {index}。" * 150),
                         "kind": "source_statement", "evidence_ids": [evidence_id],
                         "request_ids": ["research_003"]} for index in range(8)], evidence=state.evidence)
    # Even no-progress attempts must not invent capacity pressure.
    state.memory.record_attempt("retry-one", False)
    state.memory.record_attempt("retry-two", False)
    context = state.context(initial, [], ReadableRuntime(None), {})
    assert len(context["research_notes"]) > 6000
    assert context["research_notes"] == state.memory.markdown(evidence_label=state.presentation.evidence_id)
    assert context["history_text"] == state.memory.history_text(evidence_label=state.presentation.evidence_id)
    assert [item["tool"] for item in context["observations"]] == ["fs.list", "fs.inspect", "fs.read"]
    assert not context.get("memory_checkpoint_required")
    messages = research_messages(context, {})
    assert [item["tool_call_id"] for item in messages if item["role"] == "tool"] == ["call_1", "call_2", "call_3"]


@pytest.mark.asyncio
async def test_explicit_runtime_token_budget_overrides_legacy_character_limit(tmp_path):
    sources = tmp_path / "sources"
    sources.mkdir()
    (sources / "rules.md").write_text("# Requirements\n" + "审批规则 " * 800)
    output = tmp_path / "research"
    output.mkdir()
    workspace = WorkspaceManager(sources, output / "workspace")
    state = _DirectoryResearch(workspace, output, "目录", None,
                               PlanningResearchLimits(max_context_chars=4000), None)
    initial = workspace.list()
    state.register("fs.list", {}, initial)
    args = {"path": "rules.md"}
    result = await state.dispatch("fs.read", args)
    state.observe({"request_id": "r1", "tool": "fs.read", "arguments": args, "result": result,
                   "model_text": state.presentation.result_text("fs.read", args, result)})

    class ExplicitTokenBudget(ReadableRuntime):
        def fits_prompt_budget(self, task, context, schema, char_limit):
            assert char_limit == 4000
            return True

    runtime = ExplicitTokenBudget(None)
    context = state.context(initial, [], runtime, {})
    assert runtime.prompt_chars("directory_research", context, {}) > 4000
    assert state.fits_context(context, runtime, {})
    assert context["observations"] == state.observations
    assert not context.get("memory_checkpoint_required")
