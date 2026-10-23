"""CLI/direct jobs use the same research-before-plan gate and durable call ledger."""

import json
from pathlib import Path

import pytest

from folio.document.planner import build_plan
from folio.models import JobInput, Limits, PlanItem, PlanProposal
from folio.orchestration.engine import Engine
from folio.runtime.base import RuntimeResult
from folio.runtime.call_budget import ResearchModelDeferred
from folio.runtime.demo import DemoRuntime
from folio.runtime.errors import ModelCallBudgetExceeded


class ReachedWriting(Exception):
    pass


class PlanningOnlyEngine(Engine):
    async def run_business_tasks(self):
        raise ReachedWriting()


class RecordingRuntime(DemoRuntime):
    def __init__(self):
        self.calls = []
        self.inputs = []

    async def complete(self, task, context, schema, session_key):
        self.calls.append(task)
        self.inputs.append((task, context))
        return await super().complete(task, context, schema, session_key)


async def test_direct_job_researches_before_plan_and_counts_every_model_call(store, inputs):
    runtime = RecordingRuntime()
    job_id = store.create(JobInput(**inputs, runtime="demo"))
    with pytest.raises(ReachedWriting):
        await PlanningOnlyEngine(store, runtime).run(job_id)
    assert runtime.calls[:3] == ["directory_research"] * 3
    assert runtime.calls.index("plan") > runtime.calls.index("directory_research")
    assert runtime.calls.index("review_directory_research") > runtime.calls.index("plan")
    job = store.get(job_id)
    assert job["plan"] and job["calls"] == len(runtime.calls)
    assert store.call_budget_usage(job_id, "directory_research") == runtime.calls.count("directory_research")
    assert store.call_budget_usage(job_id, "review_directory_research") == 1
    assert not list((store.root / "jobs" / job_id).rglob("*.docx"))
    frozen = store.root / "jobs" / job_id / "inputs" / "directory-research"
    assert json.loads((frozen / "report.json").read_text())["status"] == "ready"
    assert json.loads((frozen / "coverage.json").read_text())["checks"]
    plan_context = next(context for task, context in runtime.inputs if task == "plan")
    assert plan_context["directory_research"]["evidence"]
    assert "excerpt" not in plan_context["source_catalog"][0]
    with store.connect() as connection:
        assert connection.execute("SELECT count(*) FROM tasks WHERE job_id=?", (job_id,)).fetchone()[0] == len(runtime.calls)


async def test_partial_research_never_saves_a_plan_or_enters_writing(store, inputs):
    class Partial(RecordingRuntime):
        async def complete(self, task, context, schema, session_key):
            if task == "directory_research":
                self.calls.append(task)
                return RuntimeResult("", tool_calls=[{
                    "id": f"call_{len(self.calls)}", "type": "function", "function": {
                        "name": "finish", "arguments": json.dumps({
                            "summary": "未读到技术范围", "blocking_gaps": ["尚未确定投标包件"],
                        }, ensure_ascii=False),
                    },
                }])
            return await super().complete(task, context, schema, session_key)

    runtime = Partial()
    job_id = store.create(JobInput(**inputs, runtime="demo"))
    with pytest.raises(ValueError, match="Source research for the outline is incomplete"):
        await PlanningOnlyEngine(store, runtime).run(job_id)
    assert runtime.calls == ["directory_research"] * 3
    assert store.get(job_id)["plan"] is None
    reports = list((store.root / "jobs" / job_id / "directory-planning").glob("*/research/report.json"))
    assert len(reports) == 1
    assert json.loads(reports[0].read_text())["status"] == "partial"


async def test_missing_coverage_is_repaired_with_engine_calls_before_writing(store, inputs):
    class MissingOnce(RecordingRuntime):
        async def complete(self, task, context, schema, session_key):
            response = await super().complete(task, context, schema, session_key)
            if task == "plan" and context.get("format_feedback"):
                parsed = json.loads(response.content)
                parsed["sections"][0]["instruction"] += "；明确说明来源审批例外。"
                return RuntimeResult(json.dumps(parsed, ensure_ascii=False))
            if task == "review_directory_research" and self.calls.count(task) == 1:
                parsed = json.loads(response.content)
                for check in parsed["checks"]:
                    check.update(disposition="missing", node_ids=[], reason="未安排来源审批例外")
                return RuntimeResult(json.dumps(parsed, ensure_ascii=False))
            return response

    runtime = MissingOnce()
    job_id = store.create(JobInput(**inputs, runtime="demo"))
    with pytest.raises(ReachedWriting):
        await PlanningOnlyEngine(store, runtime).run(job_id)
    assert runtime.calls.count("plan") == 2
    assert runtime.calls.count("review_directory_research") == 2
    plans = [context for task, context in runtime.inputs if task == "plan"]
    assert "未安排来源审批例外" in plans[-1]["format_feedback"]
    assert store.get(job_id)["calls"] == len(runtime.calls)


async def test_exhausted_coverage_repair_never_commits_plan(store, inputs):
    class MissingAlways(RecordingRuntime):
        async def complete(self, task, context, schema, session_key):
            response = await super().complete(task, context, schema, session_key)
            if task == "review_directory_research":
                value = json.loads(response.content)
                for check in value["checks"]:
                    check.update(disposition="missing", node_ids=[], reason="目录缺少对应范围")
                return RuntimeResult(json.dumps(value, ensure_ascii=False))
            return response

    runtime = MissingAlways()
    job_id = store.create(JobInput(**inputs, runtime="demo"))
    with pytest.raises(ValueError, match="outline does not yet cover source requirements"):
        await PlanningOnlyEngine(store, runtime).run(job_id)
    assert store.get(job_id)["plan"] is None
    assert not (store.root / "jobs" / job_id / "inputs" / "directory-research" / "report.json").exists()


async def test_overall_budget_blocks_more_research_and_provider_retries_are_counted(store, inputs):
    runtime = RecordingRuntime()
    job_id = store.create(JobInput(**inputs, runtime="demo", limits=Limits(max_calls=1)))
    with pytest.raises(ModelCallBudgetExceeded):
        await PlanningOnlyEngine(store, runtime).run(job_id)
    assert runtime.calls == ["directory_research"]
    assert store.get(job_id)["calls"] == 1
    assert store.call_budget_usage(job_id, "directory_research") == 1


async def test_no_local_sources_keep_existing_remote_or_request_only_boundary(store, tmp_path):
    runtime = RecordingRuntime()
    sources = tmp_path / "empty-sources"
    sources.mkdir()
    job_id = store.create(JobInput(user_request="编写市场分析", title="分析", runtime="demo",
                                  knowledge_path=str(sources)))
    with pytest.raises(ReachedWriting):
        await PlanningOnlyEngine(store, runtime).run(job_id)
    assert runtime.calls == ["plan"]
    assert not (store.root / "jobs" / job_id / "directory-planning").exists()


async def test_confirmed_plan_does_not_repeat_directory_research(store, inputs):
    plan = build_plan("confirmed", PlanProposal(sections=[
        PlanItem(node_id="n1", title="交易要求", level=1, instruction="说明审批规则"),
    ]))
    runtime = RecordingRuntime()
    job_id = store.create(JobInput(**inputs, runtime="demo"), plan=plan)
    with pytest.raises(ReachedWriting):
        await PlanningOnlyEngine(store, runtime).run(job_id)
    assert runtime.calls == []
    assert not (store.root / "jobs" / job_id / "directory-planning").exists()


async def test_reference_only_job_freezes_same_combined_source_snapshot(store, tmp_path):
    from docx import Document

    reference = tmp_path / "业务要求.docx"
    document = Document()
    document.add_heading("交易审批", 1)
    document.add_paragraph("大额交易需要复核。")
    document.save(reference)
    runtime = RecordingRuntime()
    job_id = store.create(JobInput(user_request="编写交易需求", title="交易",
                                  reference_paths=[str(reference)], runtime="demo"))
    with pytest.raises(ReachedWriting):
        await PlanningOnlyEngine(store, runtime).run(job_id)
    root = store.root / "jobs" / job_id
    report = json.loads((root / "inputs" / "directory-research" / "report.json").read_text())
    assert report["evidence"]
    assert all(item["source_id"].startswith("references/0/") for item in report["evidence"])
    assert Path(store.get(job_id)["input"]["reference_paths"][0]).exists()


async def test_provider_format_retry_is_in_both_durable_budgets(store, inputs):
    class InvalidFirst(RecordingRuntime):
        async def complete(self, task, context, schema, session_key):
            if task == "directory_research" and not self.calls:
                self.calls.append(task)
                return RuntimeResult("not JSON")
            return await super().complete(task, context, schema, session_key)

    runtime = InvalidFirst()
    job_id = store.create(JobInput(**inputs, runtime="demo"))
    with pytest.raises(ReachedWriting):
        await PlanningOnlyEngine(store, runtime).run(job_id)
    assert runtime.calls.count("directory_research") == 4
    assert store.call_budget_usage(job_id, "directory_research") == 4
    assert store.get(job_id)["calls"] == len(runtime.calls)


async def test_recovery_does_not_reset_directory_research_bucket(store, inputs):
    runtime = RecordingRuntime()
    job_id = store.create(JobInput(**inputs, runtime="demo"))
    # A previous interrupted planning run already consumed the phase allowance.
    with store.connect() as connection:
        connection.execute("UPDATE jobs SET status='failed', calls=36 WHERE id=?", (job_id,))
        connection.execute("INSERT INTO call_budgets(job_id,bucket,calls) VALUES (?,?,?)",
                           (job_id, "directory_research", 36))
    # Failed jobs require an explicit resume; it must preserve phase usage.
    store.resume_job(job_id)
    with pytest.raises(ResearchModelDeferred):
        await PlanningOnlyEngine(store, runtime).run(job_id)
    assert runtime.calls == []
    assert store.call_budget_usage(job_id, "directory_research") == 36
    assert store.get(job_id)["calls"] == 36
    assert store.get(job_id)["plan"] is None


async def test_exhausted_model_format_retries_do_not_restart_as_research_steps(store, inputs):
    class InvalidAlways(RecordingRuntime):
        async def complete(self, task, context, schema, session_key):
            self.calls.append(task)
            return RuntimeResult("not JSON")

    runtime = InvalidAlways()
    job_id = store.create(JobInput(**inputs, runtime="demo", limits=Limits(call_attempts=2)))
    with pytest.raises(ValueError, match="native tool_call"):
        await PlanningOnlyEngine(store, runtime).run(job_id)
    assert runtime.calls == ["directory_research", "directory_research"]
    assert store.get(job_id)["calls"] == 2
    assert store.get(job_id)["status"] == "failed"
    assert "native tool_call" in store.get(job_id)["error"]
    assert store.call_budget_usage(job_id, "directory_research") == 2
    assert store.get(job_id)["plan"] is None
    reports = list((store.root / "jobs" / job_id / "directory-planning").glob("*/research/report.json"))
    assert len(reports) == 1
    assert json.loads(reports[0].read_text())["status"] == "partial"
