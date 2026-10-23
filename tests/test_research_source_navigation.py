"""Shared preprocessing, progressive navigation and task-local OCR integration."""

import json

import pytest
from docx import Document
from PIL import Image

from folio.research.artifact_coverage import CoverageManager
from folio.research.artifacts import ArtifactError
from folio.research.workspace import WorkspaceError, WorkspaceManager


@pytest.fixture
def chapter_workspace(tmp_path):
    sources = tmp_path / "sources"
    sources.mkdir()
    document = Document()
    document.add_paragraph("文档引言")
    document.add_heading("交易流程", 1)
    document.add_paragraph("发起交易")
    document.add_heading("审批", 2)
    document.add_paragraph("经理复核")
    document.add_heading("审批", 2)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "角色"
    table.cell(0, 1).text = "主管"
    document.add_heading("审计流程", 1)
    document.add_paragraph("定期审计")
    document.save(sources / "业务.docx")
    return WorkspaceManager(sources, tmp_path / "task")


def test_inspect_and_cross_task_resume_never_reparse_or_open_body(chapter_workspace, monkeypatch):
    workspace = chapter_workspace
    monkeypatch.setattr(
        "folio.research.source_index.read_source",
        lambda *_: pytest.fail("cross-task source parsing is unnecessary"),
    )
    other = WorkspaceManager(workspace.source_root, workspace.task_root.with_name("other"))
    assert other.source_index.index_root == workspace.source_index.index_root
    assert other.source_snapshot_id == workspace.source_snapshot_id
    monkeypatch.setattr(other.source_index, "load", lambda *_: pytest.fail("inspect loads body"))
    index = other.inspect("业务.docx")
    assert [item["title"] for item in index["structure"]] == ["交易流程", "审批", "审批", "审计流程"]
    assert not index["evidence"] and "经理复核" not in json.dumps(index, ensure_ascii=False)
    manifest = json.loads((other.source_index.index_root / "manifest.json").read_text())
    entry = manifest["files"]["业务.docx"]
    assert entry["format"] == "docx" and entry["status"] == "ready"
    assert entry["structure_ref"] != entry["content_ref"]
    assert {path.name for path in (other.task_root / ".workspace").glob("*.json")} == {
        "manifest.json", "navigation-references.json",
    }
    references = json.loads((other.task_root / ".workspace" / "navigation-references.json").read_text())
    assert references["snapshot_id"] == other.source_snapshot_id
    assert "经理复核" not in json.dumps(references, ensure_ascii=False)


def test_node_reads_entire_parent_chapter_with_exact_pagination(chapter_workspace):
    workspace = chapter_workspace
    index = workspace.inspect("业务.docx")
    parent = index["structure"][0]
    children = workspace.inspect("业务.docx", node_id=parent["node_id"])["structure"]
    assert len(children) == 2 and children[0]["node_id"] != children[1]["node_id"]
    records, offset = [], 0
    while True:
        result = workspace.read("业务.docx", node_id=parent["node_id"], offset=offset, limit=2)
        records.extend(result["evidence"])
        if result["next_offset"] is None:
            break
        offset = result["next_offset"]
    assert [item["text"] for item in records] == ["交易流程", "发起交易", "审批", "经理复核", "审批", "角色 | 主管"]
    for record in records:
        assert workspace.resolve_evidence(record).text == record["text"]
    assert workspace.read("业务.docx", locator=parent["locator"])["total"] == 1
    child = workspace.read("业务.docx", node_id=children[0]["node_id"])
    assert [item["text"] for item in child["evidence"]] == ["审批", "经理复核"]
    assert workspace.read("业务.docx", node_id=index["node"]["node_id"])["total"] == 9
    with pytest.raises(WorkspaceError, match='cannot be combined|cannot be specified together'):
        workspace.read("业务.docx", node_id=parent["node_id"], locator=parent["locator"])
    with pytest.raises(WorkspaceError, match='does not exist'):
        workspace.read("业务.docx", node_id="invented")


def test_coverage_node_scope_enumerates_chapter_body_from_shared_original(chapter_workspace, monkeypatch):
    workspace = chapter_workspace
    monkeypatch.setenv("FOLIO_ARTIFACT_KEY_PATH", str(workspace.task_root / "host-key"))
    parent = workspace.inspect("业务.docx")["structure"][0]
    monkeypatch.setattr(
        "folio.research.formats.read_source", lambda *_: pytest.fail("must use shared original")
    )
    manager = CoverageManager(workspace.task_root, workspace)
    record = manager.create("业务.docx", node_id=parent["node_id"])
    texts = [item["text"] for item in record["scope"]["items"]]
    assert "经理复核" in texts and "角色 | 主管" in texts
    assert "文档引言" not in texts and "定期审计" not in texts
    assert record["scope"]["node_id"] == parent["node_id"]
    assert manager._load(record["coverage_id"])["scope"] == record["scope"]
    with pytest.raises(ArtifactError, match='cannot be combined|cannot be specified together'):
        manager.create("业务.docx", parent["locator"], node_id=parent["node_id"])


def test_ocr_overlay_does_not_change_other_task_or_original_completeness(tmp_path, monkeypatch):
    sources = tmp_path / "sources"
    sources.mkdir()
    Image.new("RGB", (20, 20), "white").save(sources / "scan.png")
    one = WorkspaceManager(sources, tmp_path / "one")
    index_root = one.source_index.index_root
    shared_before = {path.relative_to(index_root): path.read_bytes() for path in index_root.rglob("*.json")}
    one.register_extraction(
        "scan.png", {"blocks": [{"page": 1, "text": "审核结论"}]},
        execution_id="ocr_run", image_digest="sha256:image",
    )
    assert one.inspect("scan.png")["complete"]
    assert all((index_root / path).read_bytes() == content for path, content in shared_before.items())
    two = WorkspaceManager(sources, tmp_path / "two")
    assert not two.inspect("scan.png")["complete"]
    assert two.read("scan.png")["evidence"] == []
    restored = WorkspaceManager(sources, one.task_root)
    monkeypatch.setattr(restored, "_load", lambda *_: pytest.fail("OCR inspect loads body"))
    assert restored.inspect("scan.png")["complete"]
    monkeypatch.setenv("FOLIO_ARTIFACT_KEY_PATH", str(tmp_path / "host-key"))
    with pytest.raises(ArtifactError, match='unread or unverified content'):
        CoverageManager(one.task_root, one).create("scan.png")


def test_pdf_node_read_keeps_exact_page_after_discontiguous_ocr(tmp_path):
    from test_source_structure import pdf

    sources = tmp_path / "sources"
    sources.mkdir()
    pdf(sources / "guide.pdf", pages=2)
    workspace = WorkspaceManager(sources, tmp_path / "task")
    workspace.register_extraction(
        "guide.pdf", {"blocks": [{"page": 1, "text": "OCR first page"}]},
        execution_id="ocr_one", image_digest="sha256:image",
    )
    structure = workspace.inspect("guide.pdf")
    first_page = next(node for node in structure["structure"] if node["kind"] == "page" and node["metadata"]["page"] == 1)
    by_node = workspace.read("guide.pdf", node_id=first_page["node_id"])
    assert by_node["evidence"] == workspace.read("guide.pdf", page=1)["evidence"]
    assert [item["text"] for item in by_node["evidence"]] == ["Page 1 body", "OCR first page"]
    assert "block_indices" not in first_page
    restored = WorkspaceManager(sources, workspace.task_root)
    assert restored.read("guide.pdf", node_id=first_page["node_id"])["evidence"] == by_node["evidence"]


async def test_research_dispatch_enforces_structure_result_budget(tmp_path):
    from types import SimpleNamespace

    from folio.research.contracts import ResearchConfig, ResearchStep
    from folio.research.workflow import ResearchTools

    sources = tmp_path / "sources"
    sources.mkdir()
    (sources / "outline.md").write_text("\n".join(f"# Section {number} " + "Long title " * 20 for number in range(20)))
    workspace = WorkspaceManager(sources, tmp_path / "task")
    tools = SimpleNamespace(workspace=workspace, config=ResearchConfig(
        max_result_chars=2000, max_observation_chars=2000))
    result = await ResearchTools.dispatch(tools, ResearchStep(
        tool="fs.inspect", arguments_json=json.dumps({"path": "outline.md", "max_chars": 100000}),
    ), "inspect-request", 1)
    # Only the Markdown navigation is sent to the model. The canonical index
    # remains in the audit result and does not consume the request budget.
    assert len(result["text"]) <= 2000
    assert result["next_offset"] is not None and result["structure"]
