import json
import re

from docx import Document

from folio.document.style_batches import TEXT_DIMENSIONS
from folio.document.style_profile import prepare_style_profile, profile_for_stage
from folio.runtime.base import RuntimeResult


def test_large_style_context_preserves_each_writing_dimension_without_changing_full_result():
    dimensions = ("writing", "organization", "detail", "figures_tables", "expression")
    profile = {"documents": [{"document_id": "D1", "status": "analyzed", "rules": [
        {"dimension": dimension, "scope": "body", "description": f"观察{index}：" + "先简述主题，再具体解释适用范围及相关条件。" * 5,
         "sample_ids": [f"D1-S{index + 1}"], "confidence": "observed"}
        for dimension in dimensions for index in range(80)
    ]}]}
    before = json.dumps(profile, ensure_ascii=False)
    text = profile_for_stage(profile, "writing", max_chars=4000)
    assert len(text) <= 4000
    assert all(name in text for name in ("Writing style", "Content organization", "Level of detail", "Figure and table style", "Professional expression"))
    assert "complete rules, scope and evidence are stored in style.md" in text
    assert "未充分确认" not in text
    assert json.dumps(profile, ensure_ascii=False) == before


def test_compact_style_context_keeps_all_observations():
    profile = {"documents": [{"document_id": "D1", "status": "analyzed", "rules": [
        {"dimension": "writing", "scope": "body", "description": "通过简短句子明确表达观点。", "sample_ids": ["D1-S1"], "confidence": "observed"},
    ]}]}
    text = profile_for_stage(profile, "writing")
    assert "通过简短句子明确表达观点" in text
    assert "摘要已限长" not in text


def test_model_style_instructions_are_english_without_changing_report_or_rules():
    from folio.document.style_profile import _document_markdown

    document = {
        "document_id": "D1", "status": "analyzed",
        "measurements": {"paragraph_count": 3, "table_count_reliable": False},
        "rules": [{
            "dimension": "structure", "scope": "body",
            "description": "先说明用户原文，保留中文观察。",
            "sample_ids": ["D1-S1"], "confidence": "observed",
        }],
    }
    text = profile_for_stage({"documents": [document]}, "planning")
    assert "These statistics are not quotas for the new document" in text
    assert "do not introduce requirements on this basis" in text
    assert "先说明用户原文，保留中文观察。" in text
    assert not any(value in text for value in ("统计不是新文档配额", "未充分确认", "风格参考"))
    report = _document_markdown(document)
    assert "风格参考" in report
    assert "统计不是新文档配额" in report
    assert "先说明用户原文，保留中文观察。" in report


async def test_native_style_paraphrases_cannot_replace_measurements_or_block_complete_text(tmp_path):
    source = tmp_path / "example.docx"
    doc = Document()
    doc.add_paragraph("这是第一段参考资料，用于了解写作语气与叙述方式。")
    doc.add_paragraph("这是另一段参考资料，介绍流程和注意事项。")
    doc.save(source)

    class Runtime:
        async def complete(self, task, context, schema, session_key):
            ids = re.findall(r"### (D\d+-S\d+) /", context["style_analysis_text"])
            rules = [{
                "dimension": dimension, "scope": "body", "description": "叙述条理明确，通过不同侧面组织主题。",
                "sample_ids": ids[:2], "confidence": "observed",
            } for dimension in TEXT_DIMENSIONS]
            rules.append({
                "dimension": "typography", "scope": "body", "description": "正文使用模型凭空指定的字体。",
                "sample_ids": ids[:2], "confidence": "observed",
            })
            return RuntimeResult(json.dumps({"rules": rules, "dimension_checks": [], "request_sample_ids": [], "uncertainties": []}, ensure_ascii=False))

    profile = await prepare_style_profile([source], tmp_path / "style", runtime=Runtime())
    assert profile["status"] == "analyzed"
    assert profile["documents"][0]["layout"]["body"]
    assert "模型凭空指定" not in profile["markdown"]
    assert all(rule["dimension"] != "typography" for rule in profile["documents"][0]["rules"])


def test_default_style_contexts_keep_long_observations_and_layout():
    from folio.api.style_analysis import profile_context
    from folio.document.style_batches import observation_memory
    from folio.document.style_layout import layout_review_style

    rules = [
        {"dimension": "writing", "scope": "body", "description": f"观察{index}：" + "通过具体例子解释复杂概念。" * 9,
         "sample_ids": [f"D1-S{index + 1}"], "confidence": "observed"}
        for index in range(160)
    ]
    profile = {"documents": [{"document_id": "D1", "status": "analyzed", "rules": rules}]}
    for text in (profile_for_stage(profile, "writing"), profile_context(profile),
                 "\n".join(observation_memory(rules, {}))):
        assert len(text) > 16000
        assert "观察159：" in text
        assert "摘要已限长" not in text
    profile["style_markdown"] = "已核验的版式参数。" * 3000 + "末尾版式约束"
    assert layout_review_style(profile)["style_markdown"] == profile["style_markdown"]
