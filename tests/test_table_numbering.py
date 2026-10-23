"""Generated tables preserve literal values without adding Word list counters."""

from docx import Document

from folio.document.reader import xpath
from folio.document.word import render
from folio.models import Block, DocumentPlan, Extraction, Section, SectionContent, SectionState

ROWS = [
    ["序号", "类型", "说明"],
    ["1", "参考利率", "用于曲线构建。"],
    ["2", "利率曲线", "1. 保留版本\n2. 校验引用"],
    ["003", "汇率曲线", "007 是业务代码。"],
]


def effective_numbering(paragraph, table=None):
    """Inspect Word's direct/style overrides, not just visible paragraph text."""
    elements = [paragraph._p]
    style = paragraph.style
    while style is not None:
        elements.append(style.element)
        style = style.base_style
    if table is not None:
        style = table.style
        while style is not None:
            elements.append(style.element)
            style = style.base_style
    for element in elements:
        values = xpath(element, "./w:pPr/w:numPr/w:numId/@w:val")
        if values:
            return int(values[0])
    return 0


def test_generated_table_keeps_literal_numbers_and_disables_word_numbering(tmp_path):
    plan = DocumentPlan(
        sections=[
            Section(
                section_id="data",
                node_id="data",
                title="业务数据",
                level=1,
                display_number="1",
                instruction="说明字段",
                expected_table_headers=ROWS[0],
            )
        ]
    )
    state = SectionState(
        extraction=Extraction(summary="表格"),
        evidence=[],
        section_id="data",
        content=SectionContent(blocks=[Block(block_id="table", kind="table", rows=ROWS)]),
    )
    target = tmp_path / "output.docx"
    render(target, plan, [state], title="需求说明书", draft=False)
    table = Document(target).tables[0]
    assert [[cell.text for cell in row.cells] for row in table.rows] == ROWS
    assert all(
        effective_numbering(p, table) == 0
        for row in table.rows
        for cell in row.cells
        for p in cell.paragraphs
    )
    assert xpath(table._tbl, "./w:tr[1]/w:trPr/w:tblHeader")
    assert len(xpath(table._tbl, "./w:tr/w:trPr/w:cantSplit")) == len(ROWS)
