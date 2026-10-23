import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def test_pdf_preview_handles_bookmark_inside_table_cell(tmp_path):
    from folio.document.pagination import office_runtime, paginate

    if office_runtime() is None:
        pytest.skip("LibreOffice/UNO not installed")
    source = tmp_path / "table-bookmark.docx"
    doc = Document()
    paragraph = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "1")
    start.set(qn("w:name"), "cell_bookmark")
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "1")
    paragraph._p.append(start)
    paragraph.add_run("Table cell bookmark")
    paragraph._p.append(end)
    doc.save(source)
    result = paginate(source, pdf=source.with_suffix(".pdf"), required=True)
    assert result["status"] == "paginated"
    assert result["pages"]["cell_bookmark"] >= 1
    assert source.with_suffix(".pdf").read_bytes().startswith(b"%PDF")
