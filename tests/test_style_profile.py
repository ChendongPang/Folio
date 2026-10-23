import json
import re

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.shared import Mm, Pt
from PIL import Image

from folio.document.samples import sample_profile
from folio.document.style_batches import TEXT_DIMENSIONS, full_text_segments
from folio.document.style_contracts import StyleAnalysisConfig
from folio.document.style_profile import (
    prepare_style_profile,
    profile_for_stage,
    style_messages,
    validate_profile_sources,
)
from folio.document.style_sampling import read_style_source, sample_candidates
from folio.runtime.base import RuntimeResult


class StyleRuntime:
    def __init__(self, *, followup=False, fail=False, copy=False):
        self.calls = []
        self.followup, self.fail, self.copy = followup, fail, copy
        self.delivered = set()

    async def complete(self, task, context, schema, session_key):
        assert task == "style_analysis"
        text = context["style_analysis_text"]
        self.calls.append(text)
        if self.fail:
            raise RuntimeError("mock provider failed")
        if context.get("style_consolidation"):
            observations = re.findall(r"^- (R\d+) \| ([a-z_]+) \| ([a-z]+) \| ([a-z]+): (.*?); evidence (.*)$", text, re.M)
            rules = [{"dimension": dimension, "scope": scope, "description": description, "confidence": confidence, "sample_ids": re.findall(r"D\d+-[SP]\d+", evidence)} for _, dimension, scope, confidence, description, evidence in observations]
            return RuntimeResult(json.dumps({
                "rules": rules, "request_sample_ids": [], "uncertainties": [],
                "reviewed_rule_ids": [item[0] for item in observations],
                "dimension_checks": [{"dimension": "figures_tables", "status": "not_present", "reason": "全部文字中未出现可辨认的图表标记或表格结构。", "sample_ids": rules[0]["sample_ids"][:1]}],
            }, ensure_ascii=False))
        samples = re.findall(r"### (D\d+-S\d+) /[^\n]*\n\nType: (?:paragraph|row|table_row)", text)
        self.delivered.update(samples)
        all_ids = re.findall(r"^- (D\d+-S\d+)：", text, re.M)
        requested = [sample_id for sample_id in all_ids if sample_id not in self.delivered][:2] if self.followup else []
        return RuntimeResult(json.dumps({
            "rules": [{"dimension": dimension, "scope": "body", "description": "内部金额为123万元" if self.copy else description, "sample_ids": samples[:2], "confidence": "observed"} for dimension, description in {
                "structure": "采用层次化的章节排列，主体部分分层展开。",
                "writing": "使用正式客观的表达，避免口语化评价。",
                "organization": "逐项描述工作内容，并按主题组织说明。",
                "detail": "不同部分的篇幅有变化，说明详略具有差异。",
                "headings": "标题以简短名词短语命名，层级较明确。",
                "expression": "用词偏向业务说明，使用陈述式句子表达。",
            }.items()] if samples else [],
            "dimension_checks": [{"dimension": "figures_tables", "status": "not_present", "reason": "所读资料没有可辨识的图表标记或表格结构。", "sample_ids": samples[:1]}] if samples and not self.copy else [],
            "request_sample_ids": requested,
            "uncertainties": [],
        }, ensure_ascii=False))


def make_text(path, paragraphs=100):
    path.write_text("# 秘密旧项目名称\n" + "\n".join(f"## 模块{i}\n" + "样例正文体现不同的章节展开方式和工作描述。" * (1 + i % 9) for i in range(paragraphs)), encoding="utf-8")
    return path


async def test_style_writes_separate_structure_markdown_and_does_not_publish_examples(tmp_path):
    path = make_text(tmp_path / "reference.md")
    runtime = StyleRuntime()
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=runtime)
    assert profile["role"] == "style_only"
    assert profile["status"] == "analyzed"
    assert runtime.calls
    assert profile["coverage"]["text_complete"] is True
    assert (tmp_path / "style" / "D1" / "structure.md").exists()
    assert (tmp_path / "style" / "D1" / "style.md").exists()
    assert "秘密旧项目名称" in runtime.calls[0]
    assert "秘密旧项目名称" not in json.dumps(profile, ensure_ascii=False)
    assert "样例正文体现不同" not in profile["markdown"]
    assert "正式客观" in profile_for_stage(profile, "writing")
    assert "正式客观" not in profile_for_stage(profile, "planning")
    assert "已读样本位置" not in profile_for_stage(profile, "writing")
    validate_profile_sources(profile, [path])
    path.write_text("changed", encoding="utf-8")
    with pytest.raises(ValueError, match="does not match the reference-file snapshot"):
        validate_profile_sources(profile, [path])


async def test_large_reference_honors_explicit_call_and_real_input_limits(tmp_path):
    path = make_text(tmp_path / "large.md", paragraphs=3000)
    runtime = StyleRuntime(followup=True)
    config = StyleAnalysisConfig(max_calls=3, max_prompt_chars=15000, max_total_chars=30000, max_samples=24, sample_chars=800)
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=runtime, config=config)
    assert 0 < len(runtime.calls) <= 3
    assert profile["coverage"]["input_characters"] <= 30000
    assert profile["coverage"]["samples_delivered"] <= 24
    assert profile["status"] == "partial"
    assert all(len(call) < 15000 for call in runtime.calls)
    positions = [sample["locator"] for sample in profile["documents"][0]["sample_locations"]]
    assert "line 1" in positions
    assert profile["documents"][0]["text_coverage"]["characters_read"] < profile["documents"][0]["text_coverage"]["characters_total"]
    assert profile["coverage"]["text_complete"] is False
    assert profile["coverage"]["full_document_review"] is False


async def test_actual_runtime_budget_can_stop_before_any_call(tmp_path):
    path = make_text(tmp_path / "reference.md")
    runtime = StyleRuntime()
    runtime.prompt_chars = lambda task, context, schema: 600001
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=runtime)
    assert not runtime.calls
    assert profile["coverage"]["model_calls"] == 0
    assert profile["coverage"]["samples_delivered"] == 0
    assert profile["documents"][0]["status"] == "partial"


async def test_network_failure_is_not_reported_as_delivered_or_complete(tmp_path):
    path = make_text(tmp_path / "reference.md")
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=StyleRuntime(fail=True))
    assert profile["status"] == "partial"
    assert profile["documents"][0]["status"] == "analysis_failed"
    assert profile["coverage"]["samples_attempted"] > 0
    assert profile["coverage"]["samples_delivered"] == 0
    assert profile["documents"][0]["sample_locations"] == []


async def test_content_and_version_cache_reuses_analysis_but_not_failed_result(tmp_path):
    path = make_text(tmp_path / "reference.md")
    runtime = StyleRuntime()
    args = {"runtime": runtime, "cache_root": tmp_path / "cache"}
    first = await prepare_style_profile([path], tmp_path / "a", **args)
    first_calls = len(runtime.calls)
    second = await prepare_style_profile([path], tmp_path / "b", **args)
    assert len(runtime.calls) == first_calls
    assert second["cache_hit"] is True
    assert second["source_hashes"] == first["source_hashes"]
    path.write_text(path.read_text() + "\n新增段落", encoding="utf-8")
    await prepare_style_profile([path], tmp_path / "c", **args)
    assert len(runtime.calls) > first_calls
    failed = StyleRuntime(fail=True)
    await prepare_style_profile([path], tmp_path / "failed", runtime=failed, cache_root=tmp_path / "failure-cache")
    success = StyleRuntime()
    result = await prepare_style_profile([path], tmp_path / "retry", runtime=success, cache_root=tmp_path / "failure-cache")
    assert success.calls
    assert result["cache_hit"] is False


async def test_docx_measured_layout_respects_body_and_exact_spacing(tmp_path):
    path = tmp_path / "reference.docx"
    doc = Document()
    doc.sections[0].left_margin = Mm(25)
    doc.add_paragraph("秘密封面").runs[0].font.size = Pt(36)
    doc.add_heading("秘密项目背景", 1)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(12)
    doc.styles["Normal"].paragraph_format.line_spacing = Pt(24)
    doc.styles["Normal"].paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    for _ in range(6):
        doc.add_paragraph("正文内容需要排除封面的大字体。")
    doc.save(path)
    profile = await prepare_style_profile([path], tmp_path / "style")
    assert profile["status"] == "measured_only"
    assert profile["layout"]["body"]["font_name"] == "Arial"
    assert profile["layout"]["body"]["font_size_pt"] == 12
    assert profile["layout"]["body"]["line_spacing"] == 24
    assert profile["layout"]["body"]["line_spacing_rule"] == "exactly"
    assert profile["layout"]["page"]["margin_left_mm"] == pytest.approx(25, abs=0.05)
    assert "秘密" not in json.dumps(profile, ensure_ascii=False)
    assert "语义风格未分析" in profile["markdown"]


async def test_docx_wrapped_body_is_measured_and_main_section_weight_wins(tmp_path):
    path = tmp_path / "wrapped.docx"
    doc = Document()
    doc.sections[0].page_width = Mm(210)
    doc.add_paragraph("封面").runs[0].font.size = Pt(36)
    doc.add_heading("正文示例", 1)
    for _ in range(10):
        paragraph = doc.add_paragraph("这是较长的普通正文，用于测试主体页面样式。" * 10)
        paragraph.runs[0].font.size = Pt(12)
        wrapper, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
        paragraph._p.addprevious(wrapper)
        content.append(paragraph._p)
        wrapper.append(content)
    for _ in range(2):
        section = doc.add_section(WD_SECTION_START.NEW_PAGE)
        section.page_width = Mm(297)
        doc.add_paragraph("附录短句。")
    doc.save(path)
    profile = await prepare_style_profile([path], tmp_path / "style")
    assert profile["layout"]["body"]["font_size_pt"] == 12
    assert profile["layout"]["page"]["width_mm"] == pytest.approx(210, abs=0.05)
    assert profile["documents"][0]["layout_coverage"]["paragraphs_measured"] >= 12


def test_layout_stage_uses_adopted_user_parameters_not_original_measurements():
    profile = {
        "documents": [{"document_id": "D1", "status": "measured_only", "rules": [], "layout": {"body": {"font_size_pt": 12}}}],
        "layout": {"body": {"font_size_pt": 14}},
        "user_format_policy": {"overrides": {"body.font_size_pt": 14}, "unresolved": []},
    }
    text = profile_for_stage(profile, "layout")
    assert '"font_size_pt": 14' in text
    assert '"font_size_pt": 12' not in text
    assert "Explicit user formatting requirements take precedence" in text


def test_initial_candidates_include_body_when_positional_strata_are_headings(tmp_path):
    path = tmp_path / "many-headings.md"
    path.write_text("# 标题\n" + "\n".join(f"## 小节{i}\n正文示例{i}" for i in range(50)))
    candidates = sample_candidates(read_style_source(path))
    assert any(sample["kind"] == "paragraph" for sample in candidates[:8])
    assert any("章节展开" in sample["reasons"] for sample in candidates[:8])


async def test_business_rules_and_undelivered_citations_are_rejected(tmp_path):
    path = make_text(tmp_path / "reference.md")
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=StyleRuntime(copy=True))
    assert profile["documents"][0]["rules"] == []
    assert "123万元" not in json.dumps(profile, ensure_ascii=False)
    assert profile["status"] == "partial"


async def test_multiple_references_keep_order_and_total_budget(tmp_path):
    paths = [make_text(tmp_path / f"reference-{index}.md") for index in range(5)]
    runtime = StyleRuntime()
    profile = await prepare_style_profile(paths, tmp_path / "style", runtime=runtime, config={"max_calls": 3})
    assert len(runtime.calls) <= 3
    assert len(profile["documents"]) == 5
    assert len(profile["source_hashes"]) == 5
    assert profile["documents"][-1]["status"] == "partial"
    assert "D2" not in profile_for_stage(profile, "writing")
    validate_profile_sources(profile, paths)


def test_long_single_block_has_middle_and_tail_samples(tmp_path):
    path = tmp_path / "large.txt"
    path.write_text("A" * 5000 + "B" * 5000 + "C" * 5000)
    candidates = sample_candidates(read_style_source(path), sample_chars=300)
    assert any(sample["text"] == "A" * 300 for sample in candidates)
    assert any(sample["text"] == "B" * 300 for sample in candidates)
    assert any(sample["text"] == "C" * 300 for sample in candidates)
    assert all(sample["truncated"] for sample in candidates)
    assert "AAA" not in json.dumps(sample_profile([path]), ensure_ascii=False)


async def test_no_samples_does_not_write_files_or_call_model(tmp_path):
    runtime = StyleRuntime()
    profile = await prepare_style_profile([], tmp_path / "style", runtime=runtime)
    assert profile["status"] == "empty"
    assert not (tmp_path / "style").exists()
    assert not runtime.calls


def test_style_transport_keeps_text_readable_and_schema_once():
    context = {"style_analysis_text": "# 风格分析\n\n代表性正文"}
    schema = {"type": "object", "properties": {"rules": {"type": "array"}}}
    strict = style_messages(context, schema, structured_output=True)
    assert strict[1]["content"].startswith("# 风格分析")
    assert '"properties"' not in strict[0]["content"]
    assert '"properties"' in style_messages(context, schema)[0]["content"]


@pytest.mark.parametrize("user_request", ["分析这份示例的风格", "Analyze this example's style.", "用英文分析这个示例"])
def test_style_transport_uses_original_user_language_not_source_language(user_request):
    context = {"user_request": user_request, "style_analysis_text": "# Style analysis\n\nEnglish source"}
    messages = style_messages(context, {"type": "object"}, structured_output=True)
    assert user_request in messages[1]["content"]
    assert "English source" in messages[1]["content"]
    assert "language of the user's input" in messages[0]["content"]
    assert "explicitly requests a different output language" in messages[0]["content"]
    assert not re.search(r"[\u4e00-\u9fff]", messages[0]["content"])


async def test_style_language_request_reaches_batches_and_separates_cache(tmp_path):
    path = make_text(tmp_path / "reference.md", paragraphs=2)

    class LanguageRuntime(StyleRuntime):
        requests = []

        async def complete(self, task, context, schema, session_key):
            self.requests.append(context["user_request"])
            assert context["user_request"] in style_messages(context, schema)[1]["content"]
            return await super().complete(task, context, schema, session_key)

    runtime = LanguageRuntime()
    options = {"runtime": runtime, "cache_root": tmp_path / "cache"}
    first = await prepare_style_profile([path], tmp_path / "zh", user_request="分析风格", **options)
    assert first["status"] == "analyzed"
    calls = len(runtime.calls)
    second = await prepare_style_profile([path], tmp_path / "en", user_request="Analyze the style in English.", **options)
    assert second["cache_hit"] is False
    assert len(runtime.calls) > calls
    calls = len(runtime.calls)
    cached = await prepare_style_profile([path], tmp_path / "en-repeat", user_request="Analyze the style in English.", **options)
    assert cached["cache_hit"] is True
    assert len(runtime.calls) == calls
    assert set(runtime.requests) == {"分析风格", "Analyze the style in English."}


async def test_image_only_reference_can_be_analyzed_with_bounded_page_inputs(tmp_path):
    path = tmp_path / "reference.png"
    Image.new("RGB", (100, 200), "white").save(path)

    class ImageRuntime(StyleRuntime):
        config = {"supports_image_input": True}

        async def measure(self, task, context, schema, char_limit):
            assert len(context["page_images"]) <= 2
            return {"fits": True}

        async def complete(self, task, context, schema, session_key):
            self.calls.append(context)
            assert len(context["page_images"]) == 1
            assert "D1-P1" in context["style_analysis_text"]
            return RuntimeResult(json.dumps({"rules": [{"dimension": "visual", "scope": "cover", "description": "页面保留大面积空白，内容呈现较为简洁。", "sample_ids": ["D1-P1"], "confidence": "observed"}], "request_sample_ids": [], "uncertainties": []}, ensure_ascii=False))

    runtime = ImageRuntime()
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=runtime)
    assert len(runtime.calls) == 1
    assert profile["coverage"]["images_delivered"] == 1
    assert profile["documents"][0]["rules"][0]["dimension"] == "visual"
    assert "未查看页面图片" not in profile["markdown"]
    assert "少量页面样本" in profile["markdown"]
    assert not profile["layout"]


async def test_visual_followup_never_sends_three_images_and_rejects_unknown_page(tmp_path):
    path = tmp_path / "reference.tiff"
    pages = [Image.new("RGB", (50, 70), color) for color in ("white", "gray", "black")]
    pages[0].save(path, save_all=True, append_images=pages[1:])

    class VisualRuntime:
        config = {"supports_image_input": True}

        def __init__(self):
            self.image_counts = []

        async def measure(self, task, context, schema, char_limit):
            return {"fits": True}

        async def complete(self, task, context, schema, session_key):
            self.image_counts.append(len(context["page_images"]))
            first = len(self.image_counts) == 1
            return RuntimeResult(json.dumps({"rules": [{"dimension": "visual", "scope": "body", "description": "页面使用简单清晰的背景。", "sample_ids": ["D1-P1" if first else "D1-P999"], "confidence": "observed"}], "request_sample_ids": ["D1-P3"] if first else [], "uncertainties": []}, ensure_ascii=False))

    runtime = VisualRuntime()
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=runtime)
    assert runtime.image_counts == [2, 1]
    assert profile["coverage"]["images_delivered"] == 3
    assert len(profile["documents"][0]["rules"]) == 1
    assert profile["documents"][0]["rules"][0]["sample_ids"] == ["D1-P1"]
    assert profile["status"] == "partial"
    audit = (tmp_path / "style" / "audit" / "request-01.md").read_text()
    assert "SHA256" in audit
    assert "data:image" not in audit


async def test_visual_capacity_refusal_does_not_claim_images_delivered(tmp_path):
    path = tmp_path / "reference.png"
    Image.new("RGB", (50, 70), "white").save(path)
    runtime = StyleRuntime()
    runtime.config = {"supports_image_input": True}

    async def refuse(*args):
        return {"fits": False}

    runtime.measure = refuse
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=runtime)
    assert not runtime.calls
    assert profile["coverage"]["images_delivered"] == 0
    assert profile["status"] == "partial"


def test_full_text_segments_cover_long_blocks_without_overlap_or_lost_tail(tmp_path):
    path = tmp_path / "long.txt"
    path.write_text("开头甲" * 333 + "中间乙" * 417 + "结尾丙" * 333, encoding="utf-8")
    source = read_style_source(path)
    samples, coverage = full_text_segments(source, document_id="D1", sample_chars=300, max_samples=100)
    assert len(samples) > 8
    assert "".join(sample["text"] for sample in samples) == source.blocks[0].text
    assert samples[0]["char_start"] == 0
    assert samples[-1]["char_end"] == len(source.blocks[0].text)
    assert all(left["char_end"] == right["char_start"] for left, right in zip(samples, samples[1:], strict=False))
    assert coverage["characters_total"] == sum(sample["characters"] for sample in samples)
    assert coverage["segments_total"] == len(samples)
    assert not coverage["complete"]


async def test_empty_followups_cannot_stop_full_text_reading_and_each_real_prompt_fits(tmp_path):
    path = tmp_path / "large.txt"
    path.write_text("从头连续阅读每一个文字位置。" * 4500, encoding="utf-8")
    runtime = StyleRuntime()
    runtime.prompt_chars = lambda task, context, schema: sum(len(message["content"]) for message in style_messages(context, schema)) + 500
    config = StyleAnalysisConfig(max_prompt_chars=14000)
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=runtime, config=config)
    coverage = profile["documents"][0]["text_coverage"]
    assert len(runtime.calls) > 3  # The old global three-call ceiling would lose the tail.
    assert profile["status"] == "analyzed"
    assert profile["coverage"]["text_complete"] is True
    assert coverage["complete"] is True
    assert coverage["characters_read"] == coverage["characters_total"] == len(path.read_text())
    assert coverage["segments_read"] == coverage["segments_total"]
    texts = [part for request in runtime.calls for part in re.findall(r"<style_sample>\n\n(.*?)\n\n</style_sample>", request, re.S)]
    assert "".join(texts) == path.read_text()
    assert all(json.loads(audit.read_text())["prompt_characters"] <= config.max_prompt_chars for audit in (tmp_path / "style" / "audit").glob("request-*.json"))
    assert "全部可提取文字已读" in profile["markdown"]
    assert profile["coverage"]["full_document_review"] is False


async def test_later_batch_incremental_rules_do_not_erase_earlier_observations(tmp_path):
    path = tmp_path / "incremental.txt"
    path.write_text("每批提供新的原始文字。" * 2200, encoding="utf-8")

    class IncrementalRuntime(StyleRuntime):
        async def complete(self, task, context, schema, session_key):
            result = await super().complete(task, context, schema, session_key)
            if context.get("style_consolidation"):
                return result
            payload = json.loads(result.content)
            payload["rules"][1]["description"] = f"第{len(self.calls)}批展示独立的表达习惯。"
            return RuntimeResult(json.dumps(payload, ensure_ascii=False))

    runtime = IncrementalRuntime()
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=runtime, config={"max_prompt_chars": 18000})
    rules = [rule["description"] for rule in profile["documents"][0]["rules"] if rule["dimension"] == "writing"]
    assert len(runtime.calls) > 1
    assert rules == [f"第{index}批展示独立的表达习惯。" for index in range(1, len(runtime.calls))]
    assert profile["documents"][0]["consolidation"]["status"] == "completed"
    assert "第1批展示独立的表达习惯。" in runtime.calls[-1]


async def test_missing_text_dimensions_get_one_bounded_completion_check(tmp_path):
    path = make_text(tmp_path / "small.md", paragraphs=2)

    class UnresolvedRuntime(StyleRuntime):
        async def complete(self, task, context, schema, session_key):
            result = await super().complete(task, context, schema, session_key)
            payload = json.loads(result.content)
            # A broad writing rule cannot conceal a specifically unresolved text feature.
            payload["dimension_checks"].append({"dimension": "writing", "status": "unresolved", "reason": "尚未核对句子长度和客观表达方式。", "sample_ids": payload["rules"][0]["sample_ids"]})
            return RuntimeResult(json.dumps(payload, ensure_ascii=False))

    runtime = UnresolvedRuntime()
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=runtime)
    assert len(runtime.calls) == 2
    assert "Final targeted review" in runtime.calls[-1]
    assert profile["coverage"]["text_complete"] is True
    assert profile["documents"][0]["dimension_coverage"]["writing"]["status"] == "unresolved"
    assert profile["status"] == "partial"


async def test_completion_check_can_fill_missing_dimensions_and_clear_stale_uncertainty(tmp_path):
    path = make_text(tmp_path / "small.md", paragraphs=2)

    class RepairRuntime(StyleRuntime):
        async def complete(self, task, context, schema, session_key):
            result = await super().complete(task, context, schema, session_key)
            payload = json.loads(result.content)
            if len(self.calls) == 1:
                payload["rules"] = [rule for rule in payload["rules"] if rule["dimension"] == "writing"]
                payload["dimension_checks"] = []
                payload["uncertainties"] = ["其他文字维度还没有核对。"]
            return RuntimeResult(json.dumps(payload, ensure_ascii=False))

    runtime = RepairRuntime()
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=runtime)
    assert len(runtime.calls) == 2
    assert profile["status"] == "analyzed"
    assert "其他文字维度还没有核对" not in profile["markdown"]
    assert all(profile["documents"][0]["dimension_coverage"][dimension]["status"] in {"analyzed", "not_present"} for dimension in TEXT_DIMENSIONS)
    assert profile["documents"][0]["dimension_coverage"]["visual"]["status"] == "requires_visual"


async def test_first_batch_absence_is_not_promoted_to_whole_document_absence(tmp_path):
    path = tmp_path / "multiple.txt"
    path.write_text("每批都有需要进一步核对的文字。" * 1500, encoding="utf-8")

    class PartialAbsenceRuntime(StyleRuntime):
        async def complete(self, task, context, schema, session_key):
            result = await super().complete(task, context, schema, session_key)
            payload = json.loads(result.content)
            if len(self.calls) > 1:
                payload["dimension_checks"] = []
            return RuntimeResult(json.dumps(payload, ensure_ascii=False))

    runtime = PartialAbsenceRuntime()
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=runtime, config={"max_prompt_chars": 18000})
    assert len(runtime.calls) > 2
    assert sum("Final targeted review" in call for call in runtime.calls) == 1
    assert profile["coverage"]["text_complete"] is True
    assert profile["documents"][0]["dimension_coverage"]["figures_tables"]["status"] == "unresolved"
    assert profile["status"] == "partial"


async def test_explicit_small_segment_budget_reports_exact_unread_characters(tmp_path):
    path = tmp_path / "limited.txt"
    path.write_text("一段很长的风格示例原文。" * 500, encoding="utf-8")
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=StyleRuntime(), config={"sample_chars": 100, "max_samples": 3})
    coverage = profile["documents"][0]["text_coverage"]
    assert coverage["characters_read"] == 300
    assert coverage["characters_total"] == len(path.read_text())
    assert coverage["segments_read"] == 3
    assert coverage["segments_total"] > 3
    assert profile["status"] == "partial"
    assert profile["coverage"]["text_complete"] is False


async def test_multi_batch_final_review_can_restrict_general_rules_and_keeps_original_observations(tmp_path):
    path = tmp_path / "contrast.txt"
    path.write_text("每个章节都有不同的表达方式。" * 1800, encoding="utf-8")

    class ConsolidatingRuntime(StyleRuntime):
        async def complete(self, task, context, schema, session_key):
            result = await super().complete(task, context, schema, session_key)
            payload = json.loads(result.content)
            if context.get("style_consolidation"):
                payload["uncertainties"] = ["图形内部装饰效果仍需要页面图像确认。"]
                for rule in payload["rules"]:
                    if rule["dimension"] == "writing":
                        rule["scope"] = "technical"
                        rule["description"] = "技术章节使用客观表达，其他章节不据此作统一要求。"
            return RuntimeResult(json.dumps(payload, ensure_ascii=False))

    runtime = ConsolidatingRuntime()
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=runtime, config={"max_prompt_chars": 18000})
    document = profile["documents"][0]
    assert profile["status"] == "analyzed"
    assert sum("# Final whole-document style consolidation" in call for call in runtime.calls) == 1
    assert document["consolidation"]["status"] == "completed"
    assert "图形内部装饰效果仍需要页面图像确认。" in document["uncertainties"]
    assert next(rule for rule in document["rules"] if rule["dimension"] == "writing")["scope"] == "technical"
    assert next(rule for rule in document["batch_observations"] if rule["dimension"] == "writing")["scope"] == "body"


async def test_incomplete_final_review_keeps_all_batch_rules_and_marks_partial(tmp_path):
    path = tmp_path / "review.txt"
    path.write_text("连续正文用于最终观察核对。" * 1800, encoding="utf-8")

    class IncompleteReviewRuntime(StyleRuntime):
        async def complete(self, task, context, schema, session_key):
            result = await super().complete(task, context, schema, session_key)
            payload = json.loads(result.content)
            if context.get("style_consolidation"):
                payload["reviewed_rule_ids"] = payload["reviewed_rule_ids"][:1]
                payload["rules"] = payload["rules"][:1]
            return RuntimeResult(json.dumps(payload, ensure_ascii=False))

    profile = await prepare_style_profile([path], tmp_path / "style", runtime=IncompleteReviewRuntime(), config={"max_prompt_chars": 18000})
    document = profile["documents"][0]
    assert profile["coverage"]["text_complete"] is True
    assert profile["status"] == "partial"
    assert document["consolidation"]["status"] == "partial"
    assert document["rules"] == document["batch_observations"]
    assert len(document["rules"]) > 1


async def test_final_review_does_not_truncate_complete_observations_to_fit_context(tmp_path):
    path = tmp_path / "capacity.txt"
    path.write_text("阅读完成之后还需要核对所有观察。" * 1700, encoding="utf-8")
    runtime = StyleRuntime()
    runtime.prompt_chars = lambda task, context, schema: 600001 if context.get("style_consolidation") else sum(len(message["content"]) for message in style_messages(context, schema))
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=runtime, config={"max_prompt_chars": 18000})
    assert profile["coverage"]["text_complete"] is True
    assert profile["status"] == "partial"
    assert not any("# Final whole-document style consolidation" in call for call in runtime.calls)
    assert "未截断冒充全篇总结" in profile["markdown"]
    assert profile["documents"][0]["rules"] == profile["documents"][0]["batch_observations"]


async def test_single_text_batch_does_not_add_unnecessary_final_review(tmp_path):
    path = make_text(tmp_path / "short.md", paragraphs=3)
    runtime = StyleRuntime()
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=runtime)
    assert len(runtime.calls) == 1
    assert profile["status"] == "analyzed"
    assert "consolidation" not in profile["documents"][0]


def test_consolidation_contains_full_structure_counts_and_actual_native_measurements():
    from folio.document.style_consolidation import consolidation_text
    from folio.research.formats import SourceBlock, SourceDocument

    blocks = [
        SourceBlock("page 1/block 1", "目录\n一、主体说明 .......... 2\n二、结论 .......... 3\n"),
        SourceBlock("page 2/block 1", "一、主体说明\n" + "说明与案例。" * 24 + "\n"),
        SourceBlock("page 3/block 1", "二、结论\n概括前文。\n参考文献\n引用来源。"),
    ]
    source = SourceDocument("fake", {"format": "pdf"}, blocks)
    samples, coverage = full_text_segments(source, document_id="D1", sample_chars=20, max_samples=100)
    coverage.update(characters_read=coverage["characters_total"], segments_read=coverage["segments_total"], complete=True)
    doc = {
        "rules": [], "sample_locations": samples, "text_coverage": coverage,
        "dimension_coverage": {}, "layout": {"body": {"font_name": "TimesNewRomanPSMT", "font_size_pt": 11.04}, "page": {"width_mm": 215.9, "height_mm": 279.4}},
        "layout_coverage": {"observations": [{"dimension": "typography", "text": "已测量正文主字体。"}], "limitations": ["源 PDF 没有原始 Word 段落语义。"]},
    }
    text = consolidation_text(doc, {sample["sample_id"]: sample for sample in samples})
    assert '"font_size_pt": 11.04' in text
    assert '"width_mm": 215.9' in text
    assert "TimesNewRomanPSMT" in text
    assert "已测量正文主字体" in text
    assert f"{sum(len(block.text) for block in blocks)} characters reconstructed" in text
    assert f"主体说明: {len(blocks[1].text)} characters" in text
    assert f"结论: {len(blocks[2].text)} characters" in text
    assert "Page 2" in text
    assert "Original table of contents (only whitespace and dot leaders collapsed; no entries omitted)" in text
    assert "References：Matches at these original locations" in text
    assert "Abstract：These labels were not matched in all extractable text" in text
    assert "Native measurements are retained separately; do not duplicate them as typography/page/visual rules" in text
    assert "without marking completed textual analysis unresolved" in text


def test_consolidation_structure_rejoins_split_labels_without_claiming_unread_absence():
    from folio.document.style_consolidation import _structure_evidence
    from folio.research.formats import SourceBlock, SourceDocument

    source = SourceDocument("fake", {"format": "pdf"}, [SourceBlock("page 1/block 1", "正文材料\n参考文献\n文献说明\n附录\n补充说明")])
    samples, coverage = full_text_segments(source, document_id="D1", sample_chars=3, max_samples=100)
    coverage.update(characters_read=coverage["characters_total"], segments_read=coverage["segments_total"], complete=True)
    text = _structure_evidence({"text_coverage": coverage}, {sample["sample_id"]: sample for sample in samples})
    assert "References：Matches at these original locations" in text
    assert "Appendix：Matches at these original locations" in text
    missing = _structure_evidence({"text_coverage": coverage}, {sample["sample_id"]: sample for sample in samples[:1]})
    assert "Reconstruction or reading coverage is incomplete; do not infer absence across the entire document" in missing


async def test_style_input_uses_token_capacity_instead_of_character_cap(tmp_path):
    path = tmp_path / "large-token-fit.txt"
    path.write_text("The report explains the topic using complete paragraphs. " * 12000)
    runtime = StyleRuntime()
    measured = []

    async def measure(task, context, schema, char_limit):
        measured.append(context["style_analysis_text"])
        return {"fits": True, "estimated_input_tokens": 160000, "max_input_tokens": 600000}

    runtime.measure = measure
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=runtime)
    assert profile["coverage"]["text_complete"]
    assert len(runtime.calls) == len(measured) == 1
    assert len(runtime.calls[0]) > 600000
    texts = re.findall(r"<style_sample>\n\n(.*?)\n\n</style_sample>", runtime.calls[0], re.S)
    assert "".join(texts) == path.read_text()


async def test_style_input_still_respects_measured_token_capacity(tmp_path):
    path = make_text(tmp_path / "reference.md")
    runtime = StyleRuntime()

    async def measure(task, context, schema, char_limit):
        return {"fits": False, "estimated_input_tokens": 600001, "max_input_tokens": 600000}

    runtime.measure = measure
    profile = await prepare_style_profile([path], tmp_path / "style", runtime=runtime)
    assert not runtime.calls
    assert profile["coverage"]["model_calls"] == 0
    assert not profile["coverage"]["text_complete"]
