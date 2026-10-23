"""Drive the Vue document workflow against controlled HTTP and SSE responses."""

from copy import deepcopy
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
from threading import Thread
from urllib.parse import urlsplit

import pytest
from docx import Document
from playwright.async_api import expect

pytest_plugins = ("test_workspace_task_navigation",)


@pytest.fixture
def artifact_server():
    """Chromium native downloads bypass page.route; serve only a local fixture."""
    content = {"body": b""}
    requests = []

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            if self.path != "/jobs/job-flow/artifact":
                self.send_error(404)
                return
            requests.append(self.path)
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            self.send_header("Content-Disposition", 'attachment; filename="Folio.docx"')
            self.send_header("Content-Length", str(len(content["body"])))
            self.end_headers()
            self.wfile.write(content["body"])

        def log_message(self, *_args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_port}/jobs/job-flow/artifact", content, requests
    finally:
        server.shutdown()
        server.server_close()
        thread.join()


@pytest.mark.parametrize("viewport_width", [1280, 390], ids=["desktop", "mobile"])
async def test_sources_to_reviewed_document_and_download(task_page, viewport_width, artifact_server):
    page, calls, states, errors = task_page
    artifact_url, artifact_content, artifact_requests = artifact_server
    downloaded = BytesIO()
    document = Document()
    document.add_heading("开源项目技术方案", level=1)
    document.add_paragraph("本文档由受控测试生成，用于验证下载交付。")
    document.save(downloaded)
    artifact_bytes = downloaded.getvalue()
    artifact_content["body"] = artifact_bytes

    async def preview(route):
        path = urlsplit(route.request.url).path
        if path.endswith("/pages/1"):
            await route.fulfill(
                content_type="image/svg+xml",
                body='<svg xmlns="http://www.w3.org/2000/svg" width="300" height="400"><rect width="300" height="400" fill="white"/><text x="20" y="40">Document preview</text></svg>',
            )
        elif path.endswith("/pages"):
            await route.fulfill(json={"pages": [f"{path}/1"]})
        else:
            await route.fulfill(json={"blocks": [{"kind": "heading", "level": 1, "text": "技术方案目录"}]})

    await page.route("**/previews/**", preview)
    await page.route(artifact_url, lambda route: route.continue_())

    async def submit(button, action):
        previous_count = await page.evaluate("streams.length")
        await button.click()
        await page.wait_for_function("count => streams.length > count", arg=previous_count)
        payload = await page.evaluate("JSON.parse(streams.at(-1).options.body)")
        assert payload["action"] == action
        assert "files" not in payload, "Messages must use the server's bound source list"
        return payload

    async def selection(data):
        states["new"]["directory_review"] = deepcopy(data)
        await page.evaluate(
            """data => {
              streams.at(-1).emit('directory.review', data);
              streams.at(-1).control.close();
            }""",
            data,
        )
        await expect(page.locator("#messages .workflow-actions").last.locator("button").first).to_be_enabled()
        if viewport_width < 1150:
            await expect(page.get_by_role("button", name="打开文档预览")).to_have_attribute("aria-expanded", "false")

    await page.locator("#new-chat").click()
    await expect(page.locator("#message-input")).to_be_enabled()
    await page.set_viewport_size({"width": viewport_width, "height": 900})
    await page.locator("#file-input").set_input_files([
        {"name": "meeting-notes.md", "mimeType": "text/markdown", "buffer": b"# Meeting notes\nProject constraints"},
        {"name": "research.txt", "mimeType": "text/plain", "buffer": b"Research and implementation findings"},
        {"name": "reference.docx", "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "buffer": artifact_bytes},
    ])
    await expect(page.locator("#file-count")).to_have_text("3")
    await page.locator("#close-materials").click()
    await page.locator("#message-input").fill("我需要一份面向开源社区的技术方案。")
    await submit(page.locator("#send-button"), "chat")
    assert len(states["new"]["files"]) == 3
    assert all(item["role"] == "source" for item in states["new"]["files"])
    await expect(page.locator("#staged-files select")).to_have_count(0)
    await page.evaluate("""() => {
      streams.at(-1).emit('message.completed', {text: '可以围绕项目目标、架构设计与实施计划组织技术方案。'});
      streams.at(-1).control.close();
    }""")
    await expect(page.locator("#message-input")).to_be_enabled()
    await page.locator("#message-input").fill("请根据这些资料生成完整目录，重点描述架构与实施步骤。")
    payload = await submit(page.locator("#send-button"), "chat")
    assert "架构与实施步骤" in payload["text"]
    assert any(method == "PUT" and path == "/conversations/new/draft" for method, path, _ in calls)
    assert len(states["new"]["files"]) == 3

    directory = {
        "revision": "r1",
        "phase": "directory_review",
        "title": "开源项目技术方案",
        "reason": "完整目录已准备好。",
        "directory": {
            "sections": [{"number": "1", "title": "项目概述", "level": 1}, {"number": "1.1", "title": "目标与范围", "level": 2}],
            "preview": {"id": "flow-directory", "name": "技术方案目录", "download_url": "/previews/flow-directory/download"},
        },
    }
    await selection(directory)
    await expect(page.locator(".cover-form")).to_have_count(0)
    await expect(page.locator("#template-catalog")).to_have_count(0)
    if viewport_width < 1150:
        await page.get_by_role("button", name="打开文档预览").click()
    await expect(page.locator("#directory-toc")).to_contain_text("目标与范围")
    await expect(page.locator("#preview-pages img")).to_be_visible()
    await expect(page.locator("#preview-status")).to_contain_text("目录预览 · 1 页")
    if viewport_width < 1150:
        await page.get_by_role("button", name="关闭文档预览").click()

    await page.locator("#advice-r1").fill("增加独立的实施计划章节")
    payload = await submit(page.get_by_role("button", name="按建议修改目录").last, "revise_directory")
    assert payload["selection_revision"] == "r1"
    assert payload["advice"] == "增加独立的实施计划章节"
    revised = deepcopy(directory)
    revised.update(revision="r2", reason="已增加实施计划章节。")
    revised["directory"]["sections"].append({"number": "2", "title": "实施计划", "level": 1})
    await selection(revised)
    await expect(page.locator("#advice-r1")).to_be_disabled()
    payload = await submit(page.get_by_role("button", name="确认目录，生成文档").last, "confirm_directory")
    assert payload["selection_revision"] == "r2"

    await page.evaluate("streams.at(-1).emit('job.created', {job_id: 'job-flow'})")
    await expect(page.locator("#stop-button")).to_have_attribute("aria-label", "停止生成")
    await expect(page.locator("#message-input")).to_be_disabled()
    await page.evaluate("streams.at(-1).emit('work.update', {event_id: 1, text: '已完成正文与一致性检查'})")
    await expect(page.locator("#messages .activity-item").last).to_contain_text("已完成正文与一致性检查")
    await page.evaluate("""downloadUrl => {
      streams.at(-1).emit('artifact.ready', {
        filename: 'Folio.docx', download_url: downloadUrl,
        status: 'published', pending_count: 0, proposal_count: 0
      });
      streams.at(-1).control.close();
    }""", artifact_url)
    await expect(page.locator("#message-input")).to_be_enabled()
    await expect(page.locator("#stop-button")).to_be_hidden()
    await expect(page.locator("#messages .download-info")).to_contain_text("Word 文档 · 检查通过")

    async with page.expect_download() as pending_download:
        await page.locator("#messages .download-link").click()
    download = await pending_download.value
    assert download.suggested_filename == "Folio.docx"
    assert artifact_requests, await download.failure()
    downloaded_path = await download.path()
    assert downloaded_path.read_bytes() == artifact_bytes
    assert Document(downloaded_path).paragraphs[0].text == "开源项目技术方案"
    assert await page.evaluate("document.documentElement.scrollWidth <= innerWidth")
    assert errors == []
