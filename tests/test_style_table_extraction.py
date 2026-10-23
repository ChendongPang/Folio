from copy import deepcopy

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

from folio.document.reader import read_docx, xpath
from folio.document.style_profile import prepare_style_profile
from folio.document.style_sampling import (
    measure_docx_layout,
    measure_docx_tables,
    read_style_source,
)
from folio.document.typography import prop, resolved_format
from folio.document.word import render
from folio.models import Block, DocumentPlan, Extraction, Section, SectionContent, SectionState


def marked_table(document, *, fill="DDEEFF", rows=3, columns=2, header=True):
    table = document.add_table(rows=rows, cols=columns)
    if header:
        prop(table.rows[0]._tr.get_or_add_trPr(), "tblHeader")
    prop(table._tbl.tblPr, "tblLook", firstRow=1 if header else 0)
    borders = prop(table._tbl.tblPr, "tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        prop(borders, edge, val="single", color="123456", sz=8)
    for row_index, row in enumerate(table.rows):
        for column, cell in enumerate(row.cells):
            cell.text = f"示例内容{row_index}-{column}"
            font = cell.paragraphs[0].runs[0].font
            font.name = "Arial"
            font.size = Pt(11 if row_index == 0 else 10)
            if row_index == 0:
                font.bold = True
                prop(cell._tc.get_or_add_tcPr(), "shd", val="clear", fill=fill)
    return table


async def test_explicit_table_style_is_extracted_and_reaches_real_render(tmp_path):
    source = tmp_path / "style.docx"
    document = Document()
    marked_table(document, fill="FF0000")  # Cover layout does not set global header colors.
    document.add_heading("示例主体", 1)
    document.add_paragraph("正式正文。")
    marked_table(document)
    document.save(source)
    profile = await prepare_style_profile([source], tmp_path / "profile")
    table = profile["layout"]["table"]
    assert table["header_fill"] == "DDEEFF"
    assert table["border_color"] == "123456"
    assert table["border_size_pt"] == 1
    assert table["header"]["font_size_pt"] == 11
    assert table["body"]["font_size_pt"] == 10
    coverage = profile["documents"][0]["layout_coverage"]["table"]
    assert coverage["tables_total"] == 2
    assert coverage["tables_measured"] == 1
    assert coverage["cells_measured"] == 6
    assert coverage["unconfirmed_tables"] == 1
    assert "表头底色：DDEEFF" in profile["markdown"]
    assert "网格边框宽度(pt)：1" in profile["markdown"]
    assert "状态：已完成程序统计" in profile["markdown"]
    assert "标题：1 级 1 个" in profile["markdown"]
    assert "{'1':" not in profile["markdown"]

    plan = DocumentPlan(sections=[Section(section_id="s1", node_id="n1", title="当前文档", level=1, display_number="1", instruction="描述当前材料", expected_table_headers=["项目", "说明"])])
    states = [SectionState(section_id="s1", content=SectionContent(blocks=[Block(block_id="t1", kind="table", rows=[["项目", "说明"], ["当前服务", "处理请求"]])]), extraction=Extraction(summary="说明"), evidence=[])]
    output = tmp_path / "output.docx"
    render(output, plan, states, title="当前文档", draft=False, defer_pagination=True, writing_profile=profile)
    result = Document(output)
    rendered = result.tables[0]
    assert xpath(rendered._tbl, "./w:tr[1]/w:tc/w:tcPr/w:shd/@w:fill") == ["DDEEFF", "DDEEFF"]
    assert xpath(rendered._tbl, "./w:tblPr/w:tblBorders/w:top/@w:color") == ["123456"]
    assert xpath(rendered._tbl, "./w:tblPr/w:tblBorders/w:top/@w:sz") == ["8"]
    assert xpath(resolved_format(result, rendered.cell(0, 0).paragraphs[0]._p)[1], "./w:sz/@w:val") == ["22"]
    assert xpath(resolved_format(result, rendered.cell(1, 0).paragraphs[0]._p)[1], "./w:sz/@w:val") == ["20"]


async def test_unmarked_layout_and_partial_borders_are_not_promoted(tmp_path):
    source = tmp_path / "layout.docx"
    document = Document()
    unmarked = marked_table(document, header=False)
    # Even a firstRow hint without a real body section cannot establish a layout table as data.
    prop(unmarked._tbl.tblPr, "tblLook", firstRow=1)
    document.add_heading("正文", 1)
    marked = marked_table(document)
    borders = marked._tbl.tblPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblBorders")
    borders.remove(borders[-1])
    document.save(source)
    profile = await prepare_style_profile([source], tmp_path / "profile")
    assert profile["layout"]["table"]["header_fill"] == "DDEEFF"
    assert "border_color" not in profile["layout"]["table"]
    assert profile["documents"][0]["layout_coverage"]["table"]["unconfirmed_tables"] == 1


def test_table_sampling_has_a_hard_cell_budget(tmp_path):
    source = tmp_path / "large-table.docx"
    document = Document()
    document.add_heading("正文", 1)
    marked_table(document, rows=40, columns=8)
    document.save(source)
    layout, coverage = measure_docx_tables(document, read_docx(source), max_cells=20)
    assert coverage["cells_measured"] == 20
    assert coverage["cells_available"] == 320
    assert coverage["complete"] is False
    assert layout["header_fill"] == "DDEEFF"


async def test_unknown_table_style_inheritance_does_not_become_measured_font(tmp_path):
    source = tmp_path / "inherited.docx"
    document = Document()
    document.add_heading("正文", 1)
    table = document.add_table(rows=2, cols=2)
    table.style = "Light Shading Accent 1"
    for row in table.rows:
        for cell in row.cells:
            cell.text = "内容"
    document.save(source)
    profile = await prepare_style_profile([source], tmp_path / "profile")
    assert profile["layout"]["table"]["header"] == {}
    assert profile["layout"]["table"]["body"] == {}
    assert "header_fill" not in profile["layout"]["table"]
    assert "未解析表格样式的条件继承" in profile["markdown"]


def test_native_table_measurement_includes_all_tables_rows_columns_and_cells(tmp_path):
    source = tmp_path / "all-tables.docx"
    document = Document()
    document.add_heading("正文", 1)
    small = marked_table(document, rows=2, columns=2)
    for _ in range(23):
        document.element.body.insert(-1, deepcopy(small._tbl))
    table = document.add_table(rows=40, cols=20)
    rows = table._tbl.findall(qn("w:tr"))
    prop(rows[0].get_or_add_trPr(), "tblHeader")
    for index, row in enumerate(rows):
        for cell in row.findall(qn("w:tc")):
            paragraph = cell.find(qn("w:p"))
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            prop(properties, "sz", val=24 if index >= 14 else 20)
            run.append(properties)
            text = OxmlElement("w:t")
            text.text = "后半部分也必须参与统计"
            run.append(text)
            paragraph.append(run)
    document.save(source)
    parsed = read_docx(source)
    layout, coverage = measure_docx_tables(document, parsed)
    assert coverage["tables_measured"] == coverage["tables_eligible"] == 25
    assert coverage["cells_measured"] == coverage["cells_available"] == 24 * 4 + 40 * 20
    assert coverage["complete"] is True
    assert layout["body"]["font_size_pt"] == 12
    _, limited = measure_docx_tables(document, parsed, max_tables=2)
    assert limited["tables_measured"] == 2
    assert limited["complete"] is False


def test_native_paragraph_measurement_has_no_implicit_2000_paragraph_cutoff(tmp_path):
    source = tmp_path / "all-paragraphs.docx"
    document = Document()
    document.add_heading("正文", 1)
    paragraph = document.add_paragraph("正文格式必须覆盖后半部分。")
    paragraph.runs[0].font.size = Pt(12)
    for _ in range(2002):
        document.element.body.insert(-1, deepcopy(paragraph._p))
    document.save(source)
    layout, coverage = measure_docx_layout(source)
    assert coverage["paragraphs_measured"] == coverage["paragraphs_total"] == 2004
    assert coverage["complete"] is True
    assert layout["body"]["font_size_pt"] == 12
    _, limited = measure_docx_layout(source, max_paragraphs=7)
    assert limited["paragraphs_measured"] == 7
    assert limited["complete"] is False


def native_regions_example(path):
    document = Document()
    document.styles["Normal"].font.size = Pt(12)
    document.add_paragraph("旧项目封面名称").runs[0].font.size = Pt(36)
    document.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    toc = document.add_paragraph("原始目录条目……1", style="TOC 1")
    toc.runs[0].font.size = Pt(14)
    document.add_heading("主体栏目", 1)
    paragraph = document.add_paragraph("正文内容用于观察一般的写作形式和段落组织。" * 8)
    drawing = OxmlElement("w:pict")
    shape = etree.SubElement(drawing, "{urn:schemas-microsoft-com:vml}shape")
    textbox = OxmlElement("w:txbxContent")
    shape.append(textbox)
    native = OxmlElement("w:p")
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    prop(properties, "sz", val=36)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = "文本框中的原生文字"
    run.append(text)
    native.append(run)
    textbox.append(native)
    paragraph._p.append(drawing)
    header = document.sections[0].header.paragraphs[0]
    header.text = "旧公司的页眉名称"
    header.runs[0].font.size = Pt(9)
    footer = document.sections[0].footer.paragraphs[0]
    footer.text = "示例页脚编号"
    footer.runs[0].font.size = Pt(8)
    document.save(path)
    return path


def test_style_reading_keeps_native_toc_header_footer_and_textbox_without_duplicates(tmp_path):
    path = native_regions_example(tmp_path / "native.docx")
    source = read_style_source(path)
    text = "\n".join(block.text for block in source.blocks)
    for expected in ("原始目录条目", "旧公司的页眉名称", "示例页脚编号", "文本框中的原生文字"):
        assert text.count(expected) == 1
    assert text.index("原始目录条目") < text.index("主体栏目")
    regions = {block.metadata.get("style_region") for block in source.blocks}
    assert {"toc", "header", "footer"} <= regions
    assert source.metadata["characters"] == sum(len(block.text) for block in source.blocks)
    business_text = "\n".join(block.text for block in read_docx(path).blocks)
    assert "原始目录条目" not in business_text
    assert "旧公司的页眉名称" not in business_text


async def test_native_regions_have_separate_format_observations_without_source_facts(tmp_path):
    path = native_regions_example(tmp_path / "native.docx")
    profile = await prepare_style_profile([path], tmp_path / "style")
    layout = profile["layout"]
    coverage = profile["documents"][0]["layout_coverage"]
    assert layout["body"]["font_size_pt"] == 12
    assert coverage["complete"] is True
    assert coverage["paragraphs_measured"] == coverage["paragraphs_total"] == 7
    observations = "\n".join(item["text"] for item in coverage["observations"])
    for region, size in (("封面/正文前", 36), ("目录", 14), ("页眉", 9), ("页脚", 8), ("文本框", 18)):
        assert any(region in line and f"字号(pt)：{size}" in line for line in observations.splitlines())
    for fact in ("旧项目封面名称", "原始目录条目", "旧公司的页眉名称", "示例页脚编号", "文本框中的原生文字"):
        assert fact not in profile["markdown"]


def test_native_extra_text_respects_revision_view(tmp_path):
    path = tmp_path / "revised.docx"
    document = Document()
    document.add_paragraph("正文。")
    header = document.sections[0].header.paragraphs[0]
    for marker, content in (("ins", "新版页眉"), ("del", "旧版页眉")):
        revision = OxmlElement("w:" + marker)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t" if marker == "ins" else "w:delText")
        text.text = content
        run.append(text)
        revision.append(run)
        header._p.append(revision)
    document.save(path)
    final = "\n".join(block.text for block in read_style_source(path).blocks)
    original = "\n".join(block.text for block in read_style_source(path, revision_view="original").blocks)
    assert "新版页眉" in final and "旧版页眉" not in final
    assert "旧版页眉" in original and "新版页眉" not in original


def test_native_region_measurement_failure_cannot_claim_complete(tmp_path, monkeypatch):
    from folio.document import style_sampling

    path = native_regions_example(tmp_path / "unreadable-format.docx")
    original = style_sampling.paragraph_format

    def fail_header(document, element):
        if any(ancestor.tag == qn("w:hdr") for ancestor in element.iterancestors()):
            raise ValueError("malformed header property")
        return original(document, element)

    monkeypatch.setattr(style_sampling, "paragraph_format", fail_header)
    _, coverage = measure_docx_layout(path)
    assert coverage["complete"] is False
    assert coverage["paragraphs_measured"] == coverage["paragraphs_total"] - 1
    assert "header1.xml" in coverage["limitations"][0]


def test_style_native_footnotes_and_endnotes_include_text_and_format_with_revision_view(tmp_path):
    path = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph("正文。")
    for kind, size in (("footnote", 9), ("endnote", 10)):
        root = OxmlElement("w:" + kind + "s")
        separator = OxmlElement("w:" + kind)
        separator.set(qn("w:id"), "-1")
        separator.set(qn("w:type"), "separator")
        paragraph = OxmlElement("w:p")
        paragraph.append(OxmlElement("w:separator"))
        separator.append(paragraph)
        root.append(separator)
        note = OxmlElement("w:" + kind)
        note.set(qn("w:id"), "1")
        paragraph = OxmlElement("w:p")
        for marker, content in (("ins", f"新版{kind}文字"), ("del", f"旧版{kind}文字")):
            revision = OxmlElement("w:" + marker)
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            prop(properties, "sz", val=size * 2)
            run.append(properties)
            text = OxmlElement("w:t" if marker == "ins" else "w:delText")
            text.text = content
            run.append(text)
            revision.append(run)
            paragraph.append(revision)
        note.append(paragraph)
        root.append(note)
        part = Part(
            PackURI(f"/word/{kind}s.xml"),
            f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml",
            etree.tostring(root), document.part.package,
        )
        document.part.relate_to(part, f"http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}s")
    document.save(path)
    final = read_style_source(path)
    original = read_style_source(path, revision_view="original")
    for kind, size, label in (("footnote", 9, "脚注"), ("endnote", 10, "尾注")):
        assert [block.text for block in final.blocks if block.metadata.get("style_region") == kind] == [f"新版{kind}文字"]
        assert [block.text for block in original.blocks if block.metadata.get("style_region") == kind] == [f"旧版{kind}文字"]
        _, coverage = measure_docx_layout(path)
        assert coverage["complete"] is True
        assert any(label in item["text"] and f"字号(pt)：{size}" in item["text"] for item in coverage["observations"])
        assert all(f"新版{kind}文字" not in item["text"] for item in coverage["observations"])
