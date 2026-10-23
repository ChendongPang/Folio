"""Document language follows user intent while source text and templates stay intact."""

import hashlib
import json
import re

import pytest
from docx import Document
from PIL import Image

from folio.decisions.routing import _explicit_style_targets
from folio.document.style_layout import resolve_style_policy
from folio.document.style_profile import prepare_style_profile, profile_for_stage
from folio.document.template_contract import read_template
from folio.document.word import render
from folio.i18n import translate_text
from folio.models import Block, DocumentPlan, Extraction, Section, SectionContent, SectionState
from folio.runtime.demo import DemoRuntime
from folio.runtime.image_inputs import image_messages
from folio.runtime.language import language_instruction, output_language


@pytest.mark.parametrize("instruction, expected", [
    ("Prepare a tender response", "en"),
    ("请写投标文件", "zh-CN"),
    ("请用英文撰写投标文件", "en"),
    ("Write this tender response in Chinese", "zh-CN"),
    ("Write in English. Actually, reply in Simplified Chinese.", "zh-CN"),
    ("Prepare an English-language report", "en"),
    ("Summarize the source.\n> 中文原始资料", "en"),
    ("Write a bid based on 招标文件.docx", "en"),
    ("Please draft a bid based on 《项目建设招标文件》", "en"),
    ("请根据英文报告.docx写一份总结。", "zh-CN"),
    ('Analyze the title "write in Chinese" and explain its wording.', "en"),
    ("请用英文根据文件.docx写报告", "en"),
    ("请用“英文”根据文件.docx写报告", "en"),
    ("请用`英文`根据文件.docx写报告", "en"),
    ('Write in "Chinese" using the English-language report.docx.', "zh-CN"),
    ("Please use Chinese.", "zh-CN"),
    ("请写报告。Please use English.", "en"),
    ("Please switch to Chinese.", "zh-CN"),
    ("请写报告。Switch to English.", "en"),
    ("Output language: Chinese", "zh-CN"),
    ("请写报告。Output language: English", "en"),
    ('Please use "Chinese".', "zh-CN"),
    ("Do not use Chinese; use English.", "en"),
    ("Do not switch to Chinese.", "en"),
    ("请用中文。Do not use English.", "zh-CN"),
    ("请用中文。Do not switch to English.", "zh-CN"),
    ('Explain the phrase "use Chinese" in this source.', "en"),
    ("请用中文。Use English.pdf as a reference.", "zh-CN"),
    ("Do not write in Chinese. Prepare a report.", "en"),
    ("请不要用英文，按资料写正文。", "zh-CN"),
    ("", "zh-CN"),
])
def test_document_language_uses_instruction(instruction, expected):
    assert output_language(instruction) == expected


def test_document_and_interface_languages_are_independent():
    instruction = language_instruction({"document_language": "zh-CN", "ui_language": "en"})
    assert "confirmed output language is Simplified Chinese" in instruction
    assert "status labels and temporary outline placeholders in English" in instruction
    assert "must not change this document language" in instruction


def test_word_generated_labels_use_document_language_without_translating_sources(tmp_path):
    section = Section(section_id="s1", node_id="n1", title="Scope", level=1, display_number="1", instruction="Write the scope")
    plan = DocumentPlan(sections=[section])
    states = [SectionState(section_id="s1", extraction=Extraction(summary="Scope"), evidence=[], content=SectionContent(
        blocks=[Block(block_id="b1", text="保留的中文原文")], pending_items=["Provide the bidder name."],
    ))]
    figures = {"s1": [{"figure_id": "f1", "caption": "Workflow", "origin": "model_proposal",
                        "status": "unresolved", "reason": "Missing business constraints", "source_block_ids": ["b1"]}]}
    path = tmp_path / "output.docx"
    render(path, plan, states, title="Tender", draft=True, language="en", figures=figures,
           include_evidence_appendix=True, defer_pagination=True)
    text = "\n".join(p.text for p in Document(path).paragraphs)
    assert "Pending confirmation: Provide the bidder name." in text
    assert "Figure 1  Workflow (proposed design)" in text
    assert "References" in text and "Contents" in text and "DRAFT" in text
    assert "保留的中文原文" in text
    assert not re.search(r"[\u3400-\u9fff]", text.replace("保留的中文原文", ""))
    render(path, plan, [], title="Tender", draft=True, language="en", defer_pagination=True)
    assert "Content pending." in "\n".join(p.text for p in Document(path).paragraphs)


def test_english_template_guidance_controls_optional_children(tmp_path):
    document = Document()
    document.add_heading("Requirements", 1)
    document.add_paragraph("You may add subsections to describe implementation.")
    document.add_heading("Pricing", 1)
    document.add_paragraph("Do not add subsections or change these headings.")
    path = tmp_path / "template.docx"
    document.save(path)
    contract = read_template(path)
    assert contract.sections[0].allow_children is True
    assert contract.sections[1].allow_children is False
    assert "Do not add subsections" in contract.sections[1].instruction


def test_english_formatting_requests_apply_and_negated_examples_do_not():
    profile = resolve_style_policy({}, "Body text font size 12 pt, Arial, line spacing 1.5; Heading 1 font size 16 pt bold; margins 2 cm; paper A4")
    assert profile["layout"]["body"] == {"font_name": "Arial", "font_size_pt": 12, "line_spacing": 1.5, "line_spacing_rule": "multiple"}
    assert profile["layout"]["headings"]["1"] == {"font_size_pt": 16, "bold": True}
    assert profile["layout"]["page"]["margin_left_mm"] == 20
    assert profile["layout"]["page"]["width_mm"] == 210
    assert "Explicit formatting requirements" in profile["markdown"]
    assert not re.search(r"[\u3400-\u9fff]", profile["markdown"])
    unchanged = resolve_style_policy({}, "Could the body font size be 12 pt? For example, margins 2 cm.")
    assert not unchanged["user_format_policy"]["overrides"]
    exact = resolve_style_policy({}, "Body text line spacing 18 pt")
    assert exact["layout"]["body"]["line_spacing_rule"] == "exactly"


async def test_english_style_report_has_english_program_text_and_preserves_user_rules(tmp_path):
    source = tmp_path / "example.md"
    source.write_text("# Example\n\nA clear and concise paragraph.\n\nAnother explanatory paragraph.")
    profile = await prepare_style_profile([source], tmp_path / "style", user_request="Analyze the writing style")
    assert profile["language"] == "en"
    for path in (tmp_path / "style" / "style.md", tmp_path / "style" / "D1" / "style.md"):
        text = path.read_text()
        assert "style" in text.lower()
        assert not re.search(r"[\u3400-\u9fff]", text), text
    profile["documents"][0]["rules"] = [{"dimension": "writing", "scope": "body", "description": "用户明确保留的规则", "confidence": "observed", "sample_ids": ["D1-S1"]}]
    assert "用户明确保留的规则" in profile_for_stage(profile, "writing")


def test_native_image_prompt_labels_are_english(tmp_path):
    image = tmp_path / "page.png"
    Image.new("RGB", (8, 8), "white").save(image)
    messages = image_messages([{"role": "system", "content": "Review the image"}, {"role": "user", "content": "Inspect this page"}], {
        "page_images": [{"path": str(image), "page": 2, "source_id": "D1", "sha256": hashlib.sha256(image.read_bytes()).hexdigest()}],
    }, supported=True)
    labels = [part["text"] for part in messages[1]["content"] if part["type"] == "text"]
    assert "Source D1, page 2; region: full page" in labels


def test_english_style_file_ordinals_are_resolved():
    files = [{"id": "a" * 32, "name": "first.pdf", "role": "style_reference"}, {"id": "b" * 32, "name": "second.pdf", "role": "style_reference"}]
    for message in ("Analyze the second example", "Analyze file 2", "Analyze the 2nd style reference"):
        ids, error = _explicit_style_targets({"latest_user_message": message, "workflow": {"files": files}})
        assert ids == ["b" * 32] and error is None


async def test_demo_messages_follow_requested_document_language():
    result = await DemoRuntime().complete("style_analysis", {"user_request": "Analyze the writing style"}, {}, "test")
    assert not re.search(r"[\u3400-\u9fff]", json.dumps(json.loads(result.content), ensure_ascii=False))


def test_dynamic_document_messages_translate_without_losing_details():
    assert translate_text("正在分析 D1 的表达风格，第 2 批，阅读 4 个连续分段；已读 17223/21555 字符。", "en") == (
        "Analyzing the expression style of D1, batch 2: reading 4 consecutive segments; 17223/21555 characters read."
    )
    assert translate_text("A required template table is missing: T1", "zh-CN") == "模板必填表格缺失：T1"
