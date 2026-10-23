"""Sources-only directory snapshots, historic previews and confirmed export integration."""

import hashlib
import json
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from folio.api.app import create_app
from folio.models import JobInput
from folio.orchestration.engine import Engine
from folio.previews import PreviewLibrary
from folio.research.formats import source_hash
from folio.research.source_index import SourceIndex
from folio.runtime.base import RuntimeResult
from folio.runtime.demo import DemoRuntime
from folio.source_documents import SourceDirectoryDrafts, source_catalog
from folio.store import Store


def test_catalog_retains_large_text_when_explicitly_requested(tmp_path):
    sources = tmp_path / "sources"
    sources.mkdir()
    text = "业务规则应保留全部原文，不得仅保留开头。" * 6500
    (sources / "rules.txt").write_text(text, encoding="utf-8")
    catalog = source_catalog(sources)
    assert len(catalog[0]["excerpt"]) > 100000
    assert catalog[0]["excerpt"] == text
    assert catalog[0]["excerpt_truncated"] is False
    # Discovery still starts with the manifest; a larger budget does not read
    # file bodies before the researcher asks for them.
    manifest = source_catalog(sources, include_content=False)
    assert "excerpt" not in manifest[0]


@pytest.fixture
def source_directory(tmp_path, monkeypatch):
    calls = []

    class Planner(DemoRuntime):
        def __init__(self, config):
            pass

        async def complete(self, task, context, schema, session_key):
            calls.append((task, context))
            if task == "revise_plan":
                return RuntimeResult(json.dumps({
                    "operations": [{"op": "update", "node_id": context["previous_directory"][0]["node_id"],
                                    "field": "title", "value_json": json.dumps("订单与退款")}],
                    "summary": "补充退款范围",
                }))
            return await super().complete(task, context, schema, session_key)

    monkeypatch.setattr("folio.directory_drafts.LLMRuntime", Planner)
    monkeypatch.setattr("folio.document.word.paginate", lambda *a, **k: {"status": "not_required"})
    monkeypatch.setattr("folio.document.pagination.paginate", lambda *a, **k: {"status": "not_required"})
    root = tmp_path / "state"
    drafts = SourceDirectoryDrafts(root, PreviewLibrary(root))
    path = tmp_path / "rules.md"
    path.write_text("# 订单规则\n订单付款前允许撤销。", encoding="utf-8")
    attachments = [{"id": "a" * 32, "name": "订单规则.md", "path": str(path)}]
    return root, drafts, attachments, calls


async def test_new_directory_uses_requirements_and_sources_without_legacy_inputs(source_directory):
    root, drafts, attachments, calls = source_directory
    result = await drafts.create_from_sources(None, "conversation", attachments, "# 订单\n说明撤销规则")
    loaded = drafts.load(result["revision"], "conversation")
    folder = loaded["folder"]
    assert result["source_count"] == 1
    assert result["preview"]["download_url"].startswith("/previews/")
    assert not (root / "templates").exists()
    assert not list(folder.glob("outline.*"))
    assert not (folder / "request.md").exists()
    assert not (folder / "template.docx").exists()
    assert not (folder / "layout.docx").exists()
    assert not list(drafts.root.glob("source-input-*"))
    assert calls[0][0] == "directory_research"
    plan_index, context = next((index, context) for index, (task, context) in enumerate(calls) if task == "plan")
    assert plan_index > 0
    assert context["requirements"] == "# 订单\n说明撤销规则"
    assert "excerpt" not in context["source_catalog"][0]
    assert "headings" not in context["source_catalog"][0]
    assert context["source_catalog"][0]["path"]
    assert "订单付款前允许撤销" in json.dumps(context["directory_research"]["evidence"], ensure_ascii=False)
    assert result["research"]["files"][0]["read"] is True
    assert result["confirmable"] is True
    assert result["task_usage"][0]["task_id"] == "directory_planning"
    assert result["task_usage"][0]["started_at"]
    assert not {"outline", "template", "directory_mode"}.intersection(context)
    assert set(json.loads((folder / "planning-proposal.json").read_text())["sections"][0]) == {
        "node_id", "title", "level", "instruction", "content_mode",
    }


async def test_source_directory_accepts_english_workspace_language(source_directory):
    _, drafts, attachments, _ = source_directory
    result = await drafts.create_from_sources(
        None,
        "conversation",
        attachments,
        "# Order Rules\nExplain the cancellation workflow.",
        language="en",
    )
    metadata = json.loads(
        (drafts.root / result["revision"] / "metadata.json").read_text(encoding="utf-8")
    )
    assert metadata["language"] == "en"
    assert result["title"] == "Order Rules"
    assert result["task_usage"][0]["title"] == "Outline planning"


async def test_source_identity_survives_directory_move_and_job_snapshot(source_directory):
    root, drafts, attachments, calls = source_directory
    result = await drafts.create_from_sources(None, "conversation", attachments, "# 订单\n说明撤销规则")
    loaded = drafts.load(result["revision"], "conversation")
    source_dir = loaded["knowledge_path"]
    filename = next(source_dir.iterdir()).name
    context = next(context for task, context in calls if task == "plan")
    original_version = context["source_catalog"][0]["source_version"]
    manifest = json.loads((source_dir.with_name("sources.index") / "manifest.json").read_text())
    assert manifest["files"][filename]["version"] == original_version
    index = SourceIndex(source_dir, cache_root=root / "source-cache")
    index.prepare()
    assert index.load(filename).source_version == index.structure(filename)["source_version"] == original_version
    Path(attachments[0]["path"]).write_text("上传原文件已改变，此内容不属于已确认目录。", encoding="utf-8")
    assert source_hash(source_dir / filename) == original_version
    store = Store(root)
    job_id = store.create(JobInput(
        title="订单需求", user_request="说明订单", knowledge_path=str(source_dir), runtime="demo",
    ), plan=loaded["plan"])
    job_sources = Path(store.get(job_id)["input"]["knowledge_path"])
    catalog = source_catalog(job_sources, cache_root=root / "source-cache")
    assert catalog[0]["source_version"] == original_version
    assert "原文件已改变" not in catalog[0]["excerpt"]


async def test_saved_legacy_directory_loads_and_revises_without_old_input_files(source_directory):
    _, drafts, attachments, calls = source_directory
    result = await drafts.create_from_sources(None, "conversation", attachments, "# 订单\n说明撤销规则")
    loaded = drafts.load(result["revision"], "conversation")
    folder = loaded["folder"]
    plan_path = folder / "document-plan.json"
    proposal_path = folder / "planning-proposal.json"
    old_plan = json.loads(plan_path.read_text())
    old_plan.update(planning_method="llm", unused_template_slots=[])
    for item in old_plan["sections"]:
        item.pop("node_id")
        item.update(outline_node_id="o1", outline_refs=["o1"], origin="outline", mapping_reason="用户要求")
    plan_path.write_text(json.dumps(old_plan), encoding="utf-8")
    old_proposal = json.loads(proposal_path.read_text())
    old_proposal["unused_template_slots"] = []
    for item in old_proposal["sections"]:
        item.update(outline_refs=["o1"], mapping_reason="用户要求", template_slot_id=None, template_table_ref=None)
    proposal_path.write_text(json.dumps(old_proposal), encoding="utf-8")
    previous = drafts.load(result["revision"], "conversation")
    revised = await drafts.create_from_sources(
        None, "conversation", attachments, "# 订单\n说明撤销规则及退款规则", previous=previous, advice="增加退款",
    )
    updated = drafts.load(revised["revision"], "conversation")
    assert updated["plan"].sections[0].section_id == previous["plan"].sections[0].section_id
    assert updated["plan"].sections[0].title == "订单与退款"
    revision_context = next(context for task, context in reversed(calls) if task == "revise_plan")
    assert not {"outline_refs", "mapping_reason", "template_slot_id"}.intersection(
        revision_context["previous_proposal"]["sections"][0]
    )
    assert json.loads(plan_path.read_text()) == old_plan
    assert json.loads(proposal_path.read_text()) == old_proposal


async def test_old_directory_without_source_snapshot_requires_regeneration(source_directory):
    _, drafts, attachments, _ = source_directory
    result = await drafts.create_from_sources(None, "conversation", attachments, "说明订单")
    folder = drafts.load(result["revision"], "conversation")["folder"]
    (folder / "sources").rename(folder / "preserved-old-inputs")
    with pytest.raises(ValueError, match="旧版目录未保存当前资料快照"):
        drafts.load(result["revision"], "conversation")
    assert (folder / "preserved-old-inputs").is_dir()


async def test_confirmed_source_directory_exports_without_replanning(source_directory):
    root, drafts, attachments, _ = source_directory
    result = await drafts.create_from_sources(None, "conversation", attachments, "# 订单\n说明订单撤销规则")
    loaded = drafts.load(result["revision"], "conversation")
    store = Store(root)
    job_id = store.create(JobInput(
        title="订单需求", user_request="说明订单撤销规则", knowledge_path=str(loaded["knowledge_path"]),
        runtime="demo",
    ), plan=loaded["plan"])

    class Writer(DemoRuntime):
        async def complete(self, task, *args):
            assert task != "plan", "A confirmed directory must not trigger another planning request"
            return await super().complete(task, *args)

    job = await Engine(store, Writer()).run(job_id)
    assert job["status"] == "draft"
    assert job["completed_sections"] == 1
    document = Document(job["artifact"])
    assert any("订单付款前允许撤销" in paragraph.text for paragraph in document.paragraphs)
    assert not {"template_path", "outline_path"}.intersection(job["input"])
    with TestClient(create_app(root)) as client:
        response = client.get(f"/jobs/{job_id}/artifact")
        assert response.status_code == 200
        assert response.content == Path(job["artifact"]).read_bytes()


def test_old_conversation_preview_links_resolve_without_template_api(tmp_path):
    root = tmp_path / "state"
    source = tmp_path / "saved-preview.docx"
    document = Document()
    document.add_paragraph("旧版目录预览")
    document.save(source)
    preview_id = hashlib.sha256(source.read_bytes()).hexdigest()[:32]
    legacy = root / "templates" / preview_id
    legacy.mkdir(parents=True)
    (legacy / "source.docx").write_bytes(source.read_bytes())
    (legacy / "metadata.json").write_text(json.dumps({"id": preview_id, "name": "已有目录", "listed": False}))
    preview = {"id": preview_id, "name": "已有目录", "download_url": f"/templates/{preview_id}/download",
               "preview_url": f"/templates/{preview_id}/preview"}
    old_state = {"messages": [{"role": "assistant", "content": "已有目录", "artifact": preview}],
                 "files": [], "active_job": None,
                 "directory_review": {"directory": {"preview": preview}}, "template_selection": {}}
    app = create_app(root)
    cid = "c" * 32
    with app.state.store.connect() as connection:
        connection.execute("INSERT INTO web_conversations(id,data) VALUES (?,?)", (cid, json.dumps(old_state)))
    with TestClient(app) as client:
        state = client.get(f"/conversations/{cid}").json()
        new_preview = state["directory_review"]["directory"]["preview"]
        assert new_preview["download_url"] == f"/previews/{preview_id}/download"
        assert state["messages"][0]["artifact"]["preview_url"] == f"/previews/{preview_id}/preview"
        assert "template_selection" not in state
        assert client.get(new_preview["download_url"]).content == source.read_bytes()
        assert client.get(f"/previews/{preview_id}").json()["blocks"][0]["text"] == "旧版目录预览"
        assert client.get("/templates").status_code == 404
    assert (legacy / "source.docx").read_bytes() == source.read_bytes()


def test_preview_library_accepts_plain_documents_without_structural_input_requirements(tmp_path):
    source = tmp_path / "preview.docx"
    document = Document()
    document.add_paragraph("仅含普通段落的生成产物，也可以查看。")
    document.save(source)
    library = PreviewLibrary(tmp_path / "state")
    first = library.add(source, "目录预览")
    second = library.add(source, "同一文档")
    assert first == second
    assert library.path(first["id"]).read_bytes() == source.read_bytes()
    assert library.structure(first["id"])[0]["text"] == "仅含普通段落的生成产物，也可以查看。"
