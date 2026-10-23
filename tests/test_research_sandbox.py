from __future__ import annotations

import asyncio
import hashlib
import json
import os
import shutil
import subprocess

import pytest

from folio.research.sandbox import SandboxError, SandboxManager, _safe


@pytest.fixture
def sandbox(tmp_path):
    sources = tmp_path / "sources"
    sources.mkdir()
    (sources / "requirements.txt").write_text("Order IDs must be unique.", encoding="utf-8")
    return SandboxManager(sources, tmp_path / "task", {})


def require_image(profile):
    if not shutil.which("docker"):
        pytest.skip("Docker is not installed")
    result = subprocess.run(
        ["docker", "image", "inspect", f"folio-research-{profile}:v2"],
        capture_output=True, timeout=10,
    )
    if result.returncode:
        pytest.skip(f"Research {profile} image is not installed")


@pytest.mark.parametrize("value", ["../outside", "/etc/passwd", "artifacts/../../outside"])
def test_reject_external_paths(tmp_path, value):
    with pytest.raises(SandboxError, match='[Pp]ath'):
        _safe(tmp_path, value, exists=False)


def test_reject_symlink_ancestor(tmp_path):
    root = tmp_path / "task"
    root.mkdir()
    outside = tmp_path / "outside"
    outside.mkdir()
    (root / ".research").symlink_to(outside, target_is_directory=True)
    with pytest.raises(SandboxError, match='Symbolic-link'):
        SandboxManager(tmp_path, root)


@pytest.mark.parametrize("package,version", [
    ("https://evil.example/x.whl", "1.0"), ("../package", "1.0"),
    ("networkx", ">=3"), ("networkx", "3.4.1"), ("unknown", "1.0"),
])
async def test_dependency_policy_rejects_unapproved_without_starting_docker(sandbox, monkeypatch, package, version):
    async def unexpected(*args, **kwargs):
        pytest.fail("Denied dependency must not start Docker")

    monkeypatch.setattr(sandbox, "_image", unexpected)
    with pytest.raises(SandboxError) as exc:
        await sandbox.request_dependency(package, version, "analysis")
    assert exc.value.code == "dependency_denied"


async def test_missing_image_explicitly_fails_without_host_fallback(sandbox, monkeypatch):
    async def missing(*args, **kwargs):
        return 1, "", "No such image"

    monkeypatch.setattr(sandbox, "_command", missing)
    with pytest.raises(SandboxError) as exc:
        await sandbox.execute("raise Exception('do not run')", execution_id="unavailable")
    assert exc.value.code == "sandbox_unavailable"
    assert sandbox.get_execution("unavailable") is None


def test_docker_command_enforces_isolation(sandbox):
    args = sandbox._base_args("test")
    assert args[args.index("--network") + 1] == "none"
    assert args[args.index("--cap-drop") + 1] == "ALL"
    assert "no-new-privileges" in args
    assert "--read-only" in args
    assert not args[args.index("--user") + 1].startswith("0:")
    assert {"--memory", "--memory-swap", "--cpus", "--pids-limit"}.issubset(args)


def test_dependency_hashes_and_transitive_versions_are_required(sandbox):
    sandbox.settings["dependency_allowlist"] = {"example": {"1.0": {"sha256": []}}}
    with pytest.raises(SandboxError, match="SHA-256"):
        sandbox._approved("example", "1.0")
    sandbox.settings["dependency_allowlist"]["example"]["1.0"] = {
        "sha256": ["1" * 64], "dependencies": {"unapproved": "2.0"}
    }
    with pytest.raises(SandboxError, match='not approved by operations'):
        sandbox._approved("example", "1.0")


def test_dependency_cache_is_scoped_to_image_and_exact_lock(sandbox):
    dependencies = sandbox._approved("networkx", "3.4.2")
    assert sandbox._cache_key("sha256:" + "1" * 64, dependencies) != sandbox._cache_key("sha256:" + "2" * 64, dependencies)
    changed = [dict(dependencies[0], version="3.4.1")]
    assert sandbox._cache_key("image", dependencies) != sandbox._cache_key("image", changed)


async def test_recovery_stops_old_container_without_reexecuting(sandbox, monkeypatch):
    code = "print('once')"
    record = {
        "execution_id": "interrupted", "profile": "base", "status": "running",
        "script_sha256": hashlib.sha256(code.encode()).hexdigest(), "container_name": "previous-container",
    }
    (sandbox.executions / "interrupted.json").write_text(json.dumps(record))
    removed = []

    async def remove(name):
        removed.append(name)

    async def unexpected(*args, **kwargs):
        pytest.fail("Interrupted execution must not run automatically")

    monkeypatch.setattr(sandbox, "_remove", remove)
    monkeypatch.setattr(sandbox, "_image", unexpected)
    result = await sandbox.execute(code, execution_id="interrupted")
    assert result["status"] == "interrupted"
    assert removed == ["previous-container"]


async def test_orphan_output_is_not_accepted_as_completed(sandbox, monkeypatch):
    async def image(profile):
        return "sha256:" + "1" * 64

    monkeypatch.setattr(sandbox, "_image", image)
    orphan = sandbox.output_root / "orphan"
    orphan.mkdir()
    (orphan / "partial.txt").write_text("unfinished")
    result = await sandbox.execute("print('new')", execution_id="orphan")
    assert result["status"] == "interrupted"
    assert result["artifacts"] == []


async def test_real_isolation_provenance_and_replay(sandbox):
    require_image("base")
    code = """
import json, os, pathlib, socket
source = pathlib.Path('/sources/requirements.txt')
try:
    source.write_text('changed')
    readonly = False
except OSError:
    readonly = True
try:
    socket.create_connection(('1.1.1.1', 443), timeout=0.3)
    network = True
except OSError:
    network = False
out = pathlib.Path(os.environ['FOLIO_OUTPUT_DIR']) / 'result.json'
out.write_text(json.dumps({'uid':os.getuid(),'readonly':readonly,'network':network,'text':source.read_text()}))
print('finished')
"""
    result = await sandbox.execute(code, execution_id="isolation")
    assert result["status"] == "completed", result
    assert result["stdout"].strip() == "finished"
    assert len(result["artifacts"]) == 1
    artifact = result["artifacts"][0]
    content = json.loads((sandbox.task_root / artifact["path"]).read_text())
    assert content == {"uid": content["uid"], "readonly": True, "network": False, "text": "Order IDs must be unique."}
    assert content["uid"] != 0
    assert result["image_digest"].startswith("sha256:")
    assert result["source_manifest"][0]["path"] == "requirements.txt"
    assert sandbox.verify_execution("isolation") == result
    assert await sandbox.execute(code, execution_id="isolation") == result
    with pytest.raises(SandboxError, match='another request'):
        await sandbox.execute("print(2)", execution_id="isolation")
    (sandbox.task_root / artifact["path"]).write_text("tampered")
    with pytest.raises(SandboxError, match='changed'):
        sandbox.verify_execution("isolation")


async def test_real_timeout_removes_container(sandbox):
    require_image("base")
    result = await sandbox.execute("import time; time.sleep(60)", execution_id="timeout", timeout_seconds=0.3)
    assert result["status"] == "timed_out"
    code, _, _ = await sandbox._command("inspect", result["container_name"])
    assert code != 0


async def test_real_cancel_removes_container_and_records_state(sandbox):
    require_image("base")
    task = asyncio.create_task(sandbox.execute("import time; time.sleep(60)", execution_id="cancel"))
    for _ in range(100):
        await asyncio.sleep(0.05)
        record = sandbox.get_execution("cancel")
        if record and record["status"] == "running":
            break
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    record = sandbox.get_execution("cancel")
    assert record["status"] == "cancelled"
    assert (await sandbox._command("inspect", record["container_name"]))[0] != 0


async def test_real_stdout_limit(sandbox):
    require_image("base")
    sandbox.settings["max_output_bytes"] = 1000
    result = await sandbox.execute("print('x' * 1000000)", execution_id="output")
    assert result["status"] == "output_limit"
    assert len(result["stdout"].encode()) + len(result["stderr"].encode()) <= 1000


async def test_real_symlink_output_rejected(sandbox):
    require_image("base")
    code = "import os; os.symlink('/etc/passwd', os.environ['FOLIO_OUTPUT_DIR'] + '/bad')"
    result = await sandbox.execute(code, execution_id="symlink")
    assert result["status"] in {"invalid_artifacts", "artifact_limit"}
    assert result["artifacts"] == []


async def test_real_artifact_byte_limit(sandbox):
    require_image("base")
    sandbox.settings["max_artifact_bytes"] = 10
    code = "import os,pathlib; pathlib.Path(os.environ['FOLIO_OUTPUT_DIR'],'big').write_bytes(b'x'*100)"
    result = await sandbox.execute(code, execution_id="bytes")
    assert result["status"] == "invalid_artifacts"
    assert result["error_code"] == "artifact_limit"


async def test_real_duplicate_inflight_request_cannot_cancel_execution(sandbox):
    require_image("base")
    code = "import time; time.sleep(1); print('ok')"
    task = asyncio.create_task(sandbox.execute(code, execution_id="concurrent"))
    for _ in range(100):
        await asyncio.sleep(0.02)
        if sandbox.get_execution("concurrent"):
            break
    with pytest.raises(SandboxError) as exc:
        await sandbox.execute(code, execution_id="concurrent")
    assert exc.value.code == "execution_busy"
    assert (await task)["status"] == "completed"


async def test_real_office_conversion_and_recalculation(sandbox):
    require_image("office")
    from docx import Document
    from openpyxl import Workbook, load_workbook

    document = Document()
    document.add_paragraph("Research conversion acceptance")
    document.save(sandbox.source_root / "input.docx")
    result = await sandbox.convert("office_pdf", "input.docx", execution_id="office-pdf")
    assert result["status"] == "completed", result
    assert any(item["path"].endswith(".pdf") for item in result["artifacts"])
    workbook = Workbook()
    workbook.active.append([2, 3, "=A1+B1"])
    workbook.save(sandbox.source_root / "input.xlsx")
    result = await sandbox.convert("office_recalculate", "input.xlsx", execution_id="office-calc")
    assert result["status"] == "completed", result
    calculated = next(item for item in result["artifacts"] if item["path"].endswith(".xlsx"))
    assert load_workbook(sandbox.task_root / calculated["path"], data_only=True).active["C1"].value == 5


async def test_real_ocr_image(sandbox):
    require_image("ocr")
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (900, 140), "white")
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 50)
    ImageDraw.Draw(image).text((20, 30), "ORDER TOTAL 1250", font=font, fill="black")
    image.save(sandbox.source_root / "scan.png")
    result = await sandbox.convert("ocr", "scan.png", execution_id="ocr")
    assert result["status"] == "completed", result
    extraction = next(item for item in result["artifacts"] if item["path"].endswith("extraction.json"))
    value = json.loads((sandbox.task_root / extraction["path"]).read_text())
    assert "1250" in value["text"]
    assert value["blocks"][0]["locator"] == "page:1"
    assert value["warnings"]


@pytest.mark.parametrize("suffix", ["docx", "pptx"])
async def test_real_office_ocr_keeps_original_source_and_page_count(sandbox, suffix):
    require_image("ocr")
    from docx import Document
    from docx.shared import Pt

    filename = f"scan-office.{suffix}"
    if suffix == "docx":
        document = Document()
        document.styles["Normal"].font.size = Pt(24)
        document.add_paragraph("INVOICE TOTAL 1250")
        document.add_page_break()
        document.add_paragraph("ORDER COUNT 42")
        document.save(sandbox.source_root / filename)
    else:
        from pptx import Presentation

        presentation = Presentation()
        for title in ("INVOICE TOTAL 1250", "ORDER COUNT 42"):
            slide = presentation.slides.add_slide(presentation.slide_layouts[0])
            slide.shapes.title.text = title
        presentation.save(sandbox.source_root / filename)
    result = await sandbox.convert("office_ocr", filename, execution_id="office-ocr")
    assert result["status"] == "completed", result
    extraction = next(item for item in result["artifacts"] if item["path"].endswith("extraction.json"))
    value = json.loads((sandbox.task_root / extraction["path"]).read_text())
    assert value["operation"] == "office_ocr"
    assert value["page_count"] == 2
    assert "1250" in value["blocks"][0]["text"]
    assert "42" in value["blocks"][1]["text"]
    assert {item["path"] for item in result["source_manifest"]} == {"requirements.txt", filename}


async def test_real_previous_artifact_is_readonly_and_provenance_is_chained(sandbox):
    require_image("base")
    initial = "import os,pathlib; pathlib.Path(os.environ['FOLIO_OUTPUT_DIR'],'count.txt').write_text('42')"
    first = await sandbox.execute(initial, execution_id="first")
    assert first["status"] == "completed", first
    second = await sandbox.execute("""
import os,pathlib
source=pathlib.Path(os.environ['FOLIO_PREVIOUS_ARTIFACTS'],'first','count.txt')
try:
    source.write_text('99')
    raise RuntimeError('Previous output was writable')
except OSError:
    pass
pathlib.Path(os.environ['FOLIO_OUTPUT_DIR'],'doubled.txt').write_text(str(int(source.read_text())*2))
""", execution_id="second")
    assert second["status"] == "completed", second
    assert (sandbox.output_root / "second" / "doubled.txt").read_text() == "84"
    assert second["input_artifacts"] == [dict(first["artifacts"][0], execution_id="first")]
    assert sandbox.verify_execution("second") == second
    (sandbox.output_root / "first" / "count.txt").write_text("tampered")
    with pytest.raises(SandboxError, match='Preceding outputs.*changed'):
        sandbox.verify_execution("second")


async def test_real_data_profile(sandbox):
    require_image("data")
    result = await sandbox.execute("import pandas as pd; print(pd.Series([1,2,2]).nunique())", execution_id="data", profile="data")
    assert result["status"] == "completed", result
    assert result["stdout"].strip() == "2"


async def test_real_approved_dependency_install_and_reuse(sandbox):
    if os.environ.get("FOLIO_TEST_DEPENDENCY_INSTALL") != "1":
        pytest.skip("Set FOLIO_TEST_DEPENDENCY_INSTALL=1 to test the trusted package index")
    require_image("base")
    result = await sandbox.request_dependency("networkx", "3.4.2", "Compare dependency graphs")
    assert result["status"] == "approved"
    environment = await sandbox.environment()
    assert environment["dependencies"] == result["dependencies"]
    output = await sandbox.execute("import networkx; print(networkx.__version__)", execution_id="dependency")
    assert output["status"] == "completed", output
    assert output["stdout"].strip() == "3.4.2"
    repeated = await sandbox.request_dependency("networkx", "3.4.2", "Reuse the locked graph library")
    assert repeated["cache_key"] == result["cache_key"]
