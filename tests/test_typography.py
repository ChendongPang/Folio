from copy import deepcopy
from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree

from folio.document.pagination import cache_page_references, office_runtime, paginate
from folio.document.presentation import prose_paragraphs
from folio.document.reader import read_docx, xpath
from folio.document.typography import clean_heading, install_styles
from folio.document.word import render
from folio.models import Block, DocumentPlan, Extraction, Section, SectionContent, SectionState


@pytest.fixture
def formatted_document(tmp_path):
    plan = DocumentPlan(
        sections=[
            Section(
                section_id="sec_input",
                node_id="a",
                title="输入内容",
                level=1,
                display_number="1",
                instruction="输入",
                expected_table_headers=["输入项", "说明"],
            )
        ]
    )
    state = SectionState(
        section_id="sec_input",
        content=SectionContent(
            blocks=[
                Block(block_id="heading", kind="heading", text="1.1 业务背景", heading_level=4),
                Block(block_id="text", text="唯一正文：ISO20022，金额精度0.01。"),
                Block(block_id="heading2", kind="heading", text="1.2 文件参数", heading_level=4),
                Block(
                    block_id="table",
                    kind="table",
                    rows=[["交易号", "保留原始单元格内容"], ["交易金额", "保留精度说明"]],
                ),
            ],
            pending_items=["保留原始待确认事项"],
        ),
        extraction=Extraction(summary="测试"),
        evidence=[],
    )
    original = deepcopy(state.model_dump())
    target = tmp_path / "result.docx"
    render(target, plan, [state], title="测试需求", draft=True)
    assert state.model_dump() == original
    return target, plan


def test_layout_corrects_structure_without_changing_prose_or_rows(formatted_document):
    path, _ = formatted_document
    doc = Document(path)
    headings = [b for b in read_docx(path).blocks if b.kind == "heading"]
    assert [(b.text, b.level) for b in headings] == [
        ("输入内容", 1),
        ("业务背景", 2),
        ("文件参数", 2),
    ]
    assert [p.text for p in doc.paragraphs].count("唯一正文：ISO20022，金额精度0.01。") == 1
    assert "待确认：保留原始待确认事项" in [p.text for p in doc.paragraphs]
    assert [[c.text for c in r.cells] for r in doc.tables[0].rows] == [
        ["输入项", "说明"],
        ["交易号", "保留原始单元格内容"],
        ["交易金额", "保留精度说明"],
    ]
    assert xpath(doc.tables[0]._tbl, "./w:tr[1]/w:trPr/w:tblHeader")
    assert len(xpath(doc.tables[0]._tbl, "./w:tr/w:trPr/w:cantSplit")) == 3
    toc = "".join(xpath(doc.element, ".//w:sdt//w:t/text()"))
    assert "1.1  业务背景" in toc and "1.2  文件参数" in toc
    assert "1.1 1.1" not in toc
    for name, family, size in [
        ("Folio Body", "Calibri", "22"),
        ("Folio Heading 1", "Calibri", "36"),
    ]:
        assert xpath(doc.styles[name].element, "./w:rPr/w:rFonts/@w:ascii") == [family]
        assert xpath(doc.styles[name].element, "./w:rPr/w:sz/@w:val") == [size]
    assert xpath(doc.styles["Folio TOC 1"].element, "./w:pPr/w:ind/@w:firstLineChars") == ["0"]
    body = doc.styles["Folio Body"].element
    assert xpath(body, "./w:pPr/w:ind/@w:firstLineChars") == ["200"]
    assert xpath(body, "./w:pPr/w:ind/@w:firstLine") == []
    assert xpath(body, "./w:pPr/w:jc/@w:val") == ["both"]
    assert xpath(body, "./w:pPr/w:spacing/@w:line") == ["360"]


def test_english_body_uses_half_inch_indent_and_compact_paragraph_spacing():
    document = Document()
    install_styles(document, language="en")
    body = document.styles["Folio Body"].element

    assert xpath(body, "./w:pPr/w:ind/@w:firstLine") == ["720"]
    assert xpath(body, "./w:pPr/w:ind/@w:firstLineChars") == []
    assert xpath(body, "./w:pPr/w:jc/@w:val") == ["left"]
    assert xpath(body, "./w:pPr/w:spacing/@w:line") == ["276"]
    assert xpath(body, "./w:pPr/w:spacing/@w:after") == ["120"]


def test_blank_lines_are_not_rendered_as_empty_body_paragraphs():
    assert prose_paragraphs("First paragraph.\n\nSecond paragraph.\n \nThird paragraph.") == [
        "First paragraph.",
        "Second paragraph.",
        "Third paragraph.",
    ]


def test_cached_toc_pages_are_idempotent_and_preserve_other_parts(formatted_document):
    path, _ = formatted_document
    with ZipFile(path) as z:
        before = {n: z.read(n) for n in z.namelist()}
    pages = {"sec_input": 4, "sec_input_b0": 4, "sec_input_b2": 5}
    assert cache_page_references(path, pages) == 3
    assert cache_page_references(path, pages) == 3
    with ZipFile(path) as z:
        for name, value in before.items():
            if name not in {"word/document.xml", "word/settings.xml"}:
                assert z.read(name) == value
        root = etree.fromstring(z.read("word/document.xml"))
    entries = xpath(root, ".//w:sdtContent/w:p")
    assert [xpath(p, ".//w:t/text()")[-1] for p in entries] == ["4", "4", "5"]
    assert len(xpath(root, './/w:instrText[contains(text(), "PAGEREF")]')) == 3


def test_missing_office_does_not_claim_pagination(formatted_document, monkeypatch):
    path, _ = formatted_document
    monkeypatch.setattr("folio.document.pagination.office_runtime", lambda: None)
    original = path.read_bytes()
    assert paginate(path)["status"] == "requires_office"
    with pytest.raises(RuntimeError, match="LibreOffice"):
        paginate(path, required=True)
    assert path.read_bytes() == original


def test_real_pagination_keeps_text_and_fills_all_toc_pages(formatted_document):
    if office_runtime() is None:
        pytest.skip("LibreOffice/UNO is not installed")
    path, _ = formatted_document
    before = read_docx(path)
    pdf = path.with_suffix(".pdf")
    result = paginate(path, pdf=pdf, required=True)
    assert result["status"] == "paginated" and result["toc_entries"] == 3
    assert result["pages"]["sec_input"] == 2
    assert pdf.read_bytes().startswith(b"%PDF-")
    after = read_docx(path)
    assert [(b.text, b.rows) for b in before.blocks] == [(b.text, b.rows) for b in after.blocks]
    assert result["pages"]["sec_input_b0"] == result["pages"]["sec_input"]


@pytest.mark.parametrize(
    "text,expected",
    [
        ("3.1.3.1 处理步骤", "处理步骤"),
        ("3.1.3.1. 处理步骤", "处理步骤"),
        ("ISO20022接口", "ISO20022接口"),
        ("3D 图形", "3D 图形"),
        ("1.2GHz clock", "1.2GHz clock"),
    ],
)
def test_heading_number_cleanup_does_not_strip_business_codes(text, expected):
    assert clean_heading(text) == expected
