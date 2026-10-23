"""User templates constrain output without becoming project evidence."""

import hashlib
import json
from pathlib import Path

import pytest
from docx import Document

from folio.document.plan_repair import DirectoryValidationError
from folio.document.planner import build_plan
from folio.document.template_contract import read_template, validate_template_plan
from folio.models import PlanItem, PlanProposal
from folio.previews import PreviewLibrary
from folio.research.source_index import SourceIndex
from folio.runtime.base import RuntimeResult
from folio.runtime.demo import DemoRuntime
from folio.source_documents import SourceDirectoryDrafts, _prepare_inputs


def template_file(tmp_path):
    path = tmp_path / "需求规格模板.docx"
    doc = Document()
    doc.add_paragraph("企业模板封面")
    doc.add_heading("项目概述", 1)
    doc.add_paragraph("请描述本项目目标和范围。")
    doc.add_paragraph("旧项目为北桥工程，总投资 999999 万元。")
    doc.add_heading("业务需求", 1)
    doc.add_paragraph("可按业务需要增加子章节。")
    doc.add_heading("接口清单", 1)
    for _ in range(2):
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "接口名称"
        table.cell(0, 1).text = "用途"
        table.cell(1, 0).text = "旧项目秘密接口"
        table.cell(1, 1).text = "仅为格式示例"
    doc.save(path)
    return path


def proposal():
    return PlanProposal(sections=[
        PlanItem(node_id="overview", title="项目概述", level=1, instruction="说明范围"),
        PlanItem(node_id="business", title="业务需求", level=1, content_mode="container"),
        PlanItem(node_id="order", title="订单处理", level=2, instruction="说明订单规则"),
        PlanItem(node_id="interfaces", title="接口清单", level=1, instruction="填入当前项目接口"),
    ])


def test_template_metadata_keeps_constraints_but_excludes_example_facts(tmp_path):
    path = template_file(tmp_path)
    template = read_template(path)
    context = json.dumps(template.planning_context(), ensure_ascii=False)
    assert template.source_hash == hashlib.sha256(path.read_bytes()).hexdigest()
    assert "北桥工程" not in context and "999999" not in context and "秘密接口" not in context
    assert str(tmp_path) not in context
    assert template.sections[0].instruction == "请描述本项目目标和范围。"
    assert [section.allow_children for section in template.sections] == [False, True, False]
    assert len(template.sections[-1].tables) == 2
    plan = build_plan("project", proposal(), template=template)
    assert plan.sections[0].allow_subheadings is False
    assert plan.sections[2].template_section_id is None
    assert len(plan.sections[-1].required_tables) == 2
    assert len({table.table_id for table in plan.sections[-1].required_tables}) == 2
    validate_template_plan(plan)


@pytest.mark.parametrize("change", ["missing", "rename", "reorder", "extra_root", "closed_child", "table_container"])
def test_fixed_template_structure_cannot_be_waived(tmp_path, change):
    template = read_template(template_file(tmp_path))
    candidate = proposal()
    if change == "missing":
        candidate.sections.pop(0)
    elif change == "rename":
        candidate.sections[0].title = "项目背景"
    elif change == "reorder":
        candidate.sections[0], candidate.sections[-1] = candidate.sections[-1], candidate.sections[0]
    elif change == "extra_root":
        candidate.sections.append(PlanItem(node_id="extra", title="自由新增", level=1, instruction="新章节"))
    elif change == "closed_child":
        candidate.sections.insert(1, PlanItem(node_id="extra", title="擅自细分", level=2, instruction="新章节"))
    else:
        candidate.sections[-1].content_mode = "container"
        candidate.sections.append(PlanItem(node_id="extra", title="接口说明", level=2, instruction="写接口"))
    with pytest.raises(DirectoryValidationError):
        build_plan("project", candidate, template=template)


def test_changed_template_and_weakened_saved_constraints_are_rejected(tmp_path):
    path = template_file(tmp_path)
    plan = build_plan("project", proposal(), template=read_template(path))
    changed = plan.model_copy(deep=True)
    changed.sections[-1].required_tables.pop()
    with pytest.raises(ValueError, match="constraints.*inconsistent"):
        validate_template_plan(changed)
    path.write_bytes(b"replacement")
    with pytest.raises(ValueError, match="original template is missing or changed"):
        validate_template_plan(plan)


def test_plain_text_document_is_not_silently_treated_as_a_template(tmp_path):
    path = tmp_path / "not-structured.docx"
    doc = Document()
    doc.add_paragraph("1 项目概述")
    doc.save(path)
    with pytest.raises(ValueError, match="Heading 1"):
        read_template(path)


def test_negative_expansion_instruction_stays_closed(tmp_path):
    path = tmp_path / "fixed.docx"
    doc = Document()
    doc.add_heading("固定章节", 1)
    doc.add_paragraph("不允许增加子章节。")
    doc.save(path)
    assert read_template(path).sections[0].allow_children is False


@pytest.mark.parametrize("instruction", ["可以按业务需要增加子章节。", "允许在本节增加子章节。"])
def test_common_expansion_wording_and_explicit_denial(tmp_path, instruction):
    path = tmp_path / "expansion.docx"
    doc = Document()
    doc.add_heading("业务需求", 1)
    doc.add_paragraph(instruction)
    doc.save(path)
    assert read_template(path).sections[0].allow_children is True
    doc.add_paragraph("不能新增子章节。")
    doc.save(path)
    assert read_template(path).sections[0].allow_children is False


def test_fixed_title_protection_does_not_disable_explicit_child_expansion(tmp_path):
    path = tmp_path / "fixed-titles.docx"
    doc = Document()
    doc.add_heading("业务需求", 1)
    doc.add_paragraph("填写说明：可增加子章节，但不得修改原有章节标题。")
    doc.save(path)
    assert read_template(path).sections[0].allow_children is True


def test_duplicate_fixed_sibling_titles_fail_before_model_planning(tmp_path):
    path = tmp_path / "duplicate.docx"
    doc = Document()
    doc.add_heading("业务范围", 1)
    doc.add_heading("业务范围", 1)
    doc.save(path)
    with pytest.raises(ValueError, match="duplicate headings"):
        read_template(path)


def test_template_is_absent_from_source_index_even_with_no_business_sources(tmp_path):
    path = template_file(tmp_path)
    folder = tmp_path / "inputs"
    catalog = _prepare_inputs(folder, [{"id": "t" * 32, "name": path.name, "path": str(path), "role": "template"}])
    assert catalog == []
    assert not list((folder / "sources").iterdir())
    assert (folder / "template.docx").read_bytes() == path.read_bytes()
    assert json.loads((folder / "sources.index" / "manifest.json").read_text())["files"] == {}


def test_duplicate_template_cannot_be_read_as_business_source_during_planning(tmp_path):
    path = template_file(tmp_path)
    duplicate = tmp_path / "另一个名字.docx"
    duplicate.write_bytes(path.read_bytes())
    with pytest.raises(ValueError, match="与当前模板内容相同"):
        _prepare_inputs(tmp_path / "inputs", [
            {"id": "a" * 32, "name": path.name, "path": str(path), "role": "template"},
            {"id": "b" * 32, "name": duplicate.name, "path": str(duplicate), "role": "source"},
        ])


@pytest.mark.parametrize("tag", ["ins", "del", "moveTo", "moveFrom", "pPrChange"])
def test_body_revisions_are_rejected_including_top_level_wrappers(tmp_path, tag):
    from docx.oxml import OxmlElement

    path = tmp_path / "revisions.docx"
    doc = Document()
    doc.add_heading("固定章节", 1)
    paragraph = doc.add_paragraph("未决修订文字")
    wrapper = OxmlElement("w:" + tag)
    if tag.endswith("PrChange"):
        paragraph._p.get_or_add_pPr().append(wrapper)
    else:
        paragraph._p.addprevious(wrapper)
        wrapper.append(paragraph._p)
    doc.save(path)
    with pytest.raises(ValueError, match="unresolved revisions"):
        read_template(path)


async def test_directory_retries_template_violation_and_freezes_separate_inputs(tmp_path, monkeypatch):
    path = template_file(tmp_path)
    source = tmp_path / "rules.txt"
    source.write_text("当前项目订单付款前允许取消。", encoding="utf-8")
    calls = []

    class Planner(DemoRuntime):
        def __init__(self, config):
            pass

        async def complete(self, task, context, schema, session_key):
            calls.append((task, context))
            if task != "plan":
                return await super().complete(task, context, schema, session_key)
            candidate = proposal()
            if sum(task == "plan" for task, _ in calls) == 1:
                candidate.sections[0].title = "擅自改名"
            return RuntimeResult(candidate.model_dump_json())

        async def aclose(self):
            pass

    def preview(target, plan, states, **kwargs):
        doc = Document()
        for section in plan.sections:
            doc.add_heading(section.title, section.level)
        doc.save(target)
        return {"status": "not_required"}

    monkeypatch.setattr("folio.directory_drafts.LLMRuntime", Planner)
    monkeypatch.setattr("folio.document.prepare.render", preview)
    root = tmp_path / "state"
    drafts = SourceDirectoryDrafts(root, PreviewLibrary(root))
    attachments = [
        {"id": "a" * 32, "name": path.name, "path": str(path), "role": "template"},
        {"id": "b" * 32, "name": source.name, "path": str(source), "role": "source"},
    ]
    result = await drafts.create_from_sources(None, "session", attachments, "按模板写需求规格说明书")
    saved = drafts.load(result["revision"], "session")
    plan_calls = [context for task, context in calls if task == "plan"]
    assert len(plan_calls) == 2
    assert "模板" in plan_calls[1]["format_feedback"]
    assert len(plan_calls[0]["source_catalog"]) == 1
    assert calls[0][0] == "directory_research"
    assert "当前项目订单付款前允许取消" in json.dumps(plan_calls[0]["directory_research"]["evidence"], ensure_ascii=False)
    assert "999999" not in json.dumps(calls, ensure_ascii=False)
    assert result["source_count"] == 1
    assert saved["source_roles"] == [{"id": item["id"], "role": item["role"]} for item in attachments]
    assert Path(saved["plan"].template.path) == saved["folder"] / "template.docx"
    assert (saved["folder"] / "template-contract.json").is_file()
    assert len(SourceIndex(saved["knowledge_path"]).prepare()["files"]) == 1
    path.write_bytes(b"upload changed after planning")
    validate_template_plan(saved["plan"])
