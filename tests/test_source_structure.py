"""Structure indexes navigate canonical body blocks without copying their text."""

from copy import deepcopy
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from folio.research.formats import READER_VERSION, SourceBlock, SourceDocument, read_source
from folio.research.structure import build_structure
from folio.research.workspace import evidence_for_block


def only(structure, kind):
    return [node for node in structure["nodes"] if node["kind"] == kind]


def assert_ranges(document, structure):
    by_id = {node["node_id"]: node for node in structure["nodes"]}
    assert len(by_id) == len(structure["nodes"])
    for node in structure["nodes"]:
        start, end = node["range"]["start_block"], node["range"]["end_block"]
        assert 0 <= start <= end <= len(document.blocks)
        selected = node.get("block_indices", list(range(start, end)))
        assert selected == sorted(set(selected))
        assert all(start <= index < end for index in selected)
        if "block_indices" in node:
            assert len(selected) < end - start
        if node["parent_id"] is not None:
            parent = by_id[node["parent_id"]]
            assert parent["range"]["start_block"] <= start <= end <= parent["range"]["end_block"]
            parent_selected = parent.get("block_indices", range(parent["range"]["start_block"], parent["range"]["end_block"]))
            assert set(selected) <= set(parent_selected)


def test_word_nested_sections_include_tables_and_duplicate_titles_have_stable_ids(tmp_path):
    path = tmp_path / "source.docx"
    word = Document()
    word.add_paragraph("前言正文不应复制到结构文件")
    word.add_heading("业务规则", 1)
    word.add_paragraph("第一节正文")
    word.add_heading("输入", 2)
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "字段", "要求"
    table.cell(1, 0).text, table.cell(1, 1).text = "id", "必填"
    word.add_heading("业务规则", 1)
    word.add_paragraph("第二节正文")
    word.save(path)
    document = read_source(path)
    before = deepcopy(document.model_dump())
    evidence = [evidence_for_block(path.name, document, block) for block in document.blocks]
    structure = build_structure(document)
    sections = only(structure, "section")
    assert [node["title"] for node in sections] == ["业务规则", "输入", "业务规则"]
    assert [node["range"] for node in sections] == [
        {"start_block": 1, "end_block": 6}, {"start_block": 3, "end_block": 6}, {"start_block": 6, "end_block": 8},
    ]
    assert sections[1]["parent_id"] == sections[0]["node_id"]
    assert sections[0]["node_id"] != sections[2]["node_id"]
    assert build_structure(read_source(path)) == structure
    assert document.model_dump() == before
    assert evidence == [evidence_for_block(path.name, document, block) for block in document.blocks]
    assert "前言正文" not in str(structure)
    assert structure["reader_version"] == READER_VERSION
    assert_ranges(document, structure)


def test_markdown_ranges_preserve_lines_and_ignore_code_fence_headings(tmp_path):
    path = tmp_path / "source.md"
    path.write_text("概述\n====\n描述\n## 子节\n代码示例\n```python\n# 不是标题\n```\n# 后一章\n结束", encoding="utf-8")
    document = read_source(path)
    structure = build_structure(document)
    sections = only(structure, "section")
    assert [node["title"] for node in sections] == ["概述", "子节", "后一章"]
    assert sections[0]["range"] == {"start_block": 0, "end_block": 8}
    assert sections[1]["range"] == {"start_block": 3, "end_block": 8}
    assert sections[1]["parent_id"] == sections[0]["node_id"]
    assert document.blocks[6].locator == "line 7"
    assert document.blocks[6].kind == "paragraph"
    assert_ranges(document, structure)


def test_plain_text_without_real_headings_has_only_full_document_node(tmp_path):
    path = tmp_path / "source.txt"
    path.write_text("第一章 看起来像标题\n# 仍是普通文本\n正文", encoding="utf-8")
    document = read_source(path)
    structure = build_structure(document)
    assert structure["structure_status"] == "no_outline"
    assert len(structure["nodes"]) == 1
    assert structure["nodes"][0]["range"] == {"start_block": 0, "end_block": 3}


def test_excel_sheet_ranges_headers_and_merge_metadata_are_bounded(tmp_path, monkeypatch):
    from folio.research import formats

    path = tmp_path / "source.xlsx"
    book = Workbook()
    first = book.active
    first.title = "Fields"
    first.append(["字段分组", None, "说明"])
    first.append(["字段", "要求", "备注"])
    first.append(["id", "必填", "主键"])
    first.merge_cells("A1:B1")
    first.merge_cells("A4:B4")
    book.create_sheet("Empty")
    last = book.create_sheet("Hidden")
    last.sheet_state = "hidden"
    last.append(["other", "value"])
    book.save(path)
    monkeypatch.setattr(formats, "MAX_MERGE_RANGES", 1)
    document = read_source(path)
    structure = build_structure(document)
    sheets = only(structure, "sheet")
    assert [node["title"] for node in sheets] == ["Fields", "Empty", "Hidden"]
    assert sheets[0]["range"] == {"start_block": 0, "end_block": 3}
    assert sheets[1]["range"] == {"start_block": 3, "end_block": 3}
    assert sheets[2]["range"] == {"start_block": 3, "end_block": 4}
    metadata = sheets[0]["metadata"]
    assert metadata["headers"] == ["字段分组", "", "说明"]
    assert metadata["header_basis"] == "first_nonempty_row"
    assert metadata["header_row"] == 1
    assert metadata["merged_cells"] == ["A1:B1"]
    assert metadata["merged_cells_total"] == 2 and metadata["merged_cells_truncated"]
    assert sheets[2]["metadata"]["hidden"]
    assert any('limited' in warning for warning in structure["warnings"])
    assert "cells" not in str(structure["metadata"])
    assert_ranges(document, structure)


def test_excel_inflated_dimension_uses_actual_data_bounds(tmp_path):
    path = tmp_path / "inflated.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Actual"
    sheet.append(["field", "required"])
    sheet.append(["id", "yes"])
    book.save(path)

    members = {}
    with ZipFile(path) as archive:
        for info in archive.infolist():
            members[info.filename] = archive.read(info.filename)
    sheet_path = "xl/worksheets/sheet1.xml"
    members[sheet_path] = members[sheet_path].replace(
        b'<dimension ref="A1:B2"/>',
        b'<dimension ref="A1:XFD1048576"/>',
    )
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)

    document = read_source(path)
    summary = document.metadata["sheets"][0]
    assert summary["rows"] == 2
    assert summary["columns"] == 2
    assert summary["data_cells"] == 4
    assert summary["range_cells"] == 4
    assert [block.metadata["cells"] for block in document.blocks] == [
        [
            {"cell": "A1", "value": "field"},
            {"cell": "B1", "value": "required"},
        ],
        [
            {"cell": "A2", "value": "id"},
            {"cell": "B2", "value": "yes"},
        ],
    ]


def pdf(path, bookmarks=(), pages=4):
    writer = PdfWriter()
    for number in range(pages):
        page = writer.add_blank_page(width=600, height=800)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 20 760 Td (Page {number + 1} body) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    parents = {}
    for identifier, title, page_number, parent in bookmarks:
        parents[identifier] = writer.add_outline_item(title, page_number, parent=parents.get(parent))
    writer.write(path)


def test_pdf_real_bookmark_hierarchy_and_chapter_ranges(tmp_path):
    path = tmp_path / "source.pdf"
    pdf(path, [("one", "Rules", 0, None), ("child", "Inputs", 1, "one"), ("two", "Rules", 2, None)])
    document = read_source(path)
    structure = build_structure(document)
    sections = only(structure, "section")
    assert structure["structure_status"] == "outline"
    assert [node["title"] for node in sections] == ["Rules", "Inputs", "Rules"]
    assert [node["range"] for node in sections] == [
        {"start_block": 0, "end_block": 3}, {"start_block": 1, "end_block": 3}, {"start_block": 2, "end_block": 4},
    ]
    assert sections[1]["parent_id"] == sections[0]["node_id"]
    assert all(node["metadata"]["range_precision"] == "page" for node in sections)
    assert [node["metadata"]["boundary_page_overlap"] for node in sections] == [True, True, False]
    assert len(only(structure, "page")) == 4
    assert "Page 1 body" not in str(structure)
    assert_ranges(document, structure)


def test_pdf_without_bookmarks_does_not_guess_headings(tmp_path):
    path = tmp_path / "source.pdf"
    pdf(path, pages=2)
    document = read_source(path)
    structure = build_structure(document)
    assert structure["structure_status"] == "no_outline"
    assert not only(structure, "section")
    assert [node["range"] for node in only(structure, "page")] == [
        {"start_block": 0, "end_block": 1}, {"start_block": 1, "end_block": 2},
    ]
    assert any('without guessing sections from body text or font size' in warning for warning in structure["warnings"])
    assert_ranges(document, structure)


def test_pdf_same_page_bookmarks_preserve_text_with_explicit_page_precision(tmp_path):
    path = tmp_path / "source.pdf"
    pdf(path, [("one", "First", 0, None), ("two", "Second", 0, None)], pages=2)
    document = read_source(path)
    structure = build_structure(document)
    sections = only(structure, "section")
    assert sections[0]["range"] == {"start_block": 0, "end_block": 1}
    assert sections[1]["range"] == {"start_block": 0, "end_block": 2}
    assert any('adjacent sections' in warning for warning in structure["warnings"])


def test_pdf_same_page_parent_and_last_page_child_preserve_descendant_boundaries(tmp_path):
    path = tmp_path / "source.pdf"
    pdf(path, [("one", "Parent", 0, None), ("child", "Child", 0, "one"),
               ("two", "Second", 1, None), ("last", "Final", 3, "two")])
    document = read_source(path)
    structure = build_structure(document)
    sections = only(structure, "section")
    assert [node["range"] for node in sections] == [
        {"start_block": 0, "end_block": 2}, {"start_block": 0, "end_block": 2},
        {"start_block": 1, "end_block": 4}, {"start_block": 3, "end_block": 4},
    ]
    assert sections[-1]["metadata"]["end_page"] == 4
    assert_ranges(document, structure)


def test_pdf_ocr_appended_after_native_text_uses_exact_page_and_section_membership(tmp_path):
    path = tmp_path / "source.pdf"
    pdf(path, [("one", "First", 0, None), ("two", "Second", 2, None)])
    document = read_source(path)
    before = build_structure(document)
    document.blocks.extend([
        SourceBlock("page 1/ocr exe/block 1", "First OCR", metadata={"page": 1, "extraction": {"execution_id": "exe"}}),
        SourceBlock("page 3/ocr exe/block 2", "Second OCR", metadata={"page": 3, "extraction": {"execution_id": "exe"}}),
        SourceBlock("page 1/ocr exe/block 3", "More first OCR", metadata={"page": 1, "extraction": {"execution_id": "exe"}}),
    ])
    structure = build_structure(document)
    sections = only(structure, "section")
    pages = only(structure, "page")
    assert sections[0]["block_indices"] == [0, 1, 2, 4, 5, 6]
    assert sections[1]["block_indices"] == [2, 3, 5]
    assert pages[0]["block_indices"] == [0, 4, 6]
    assert pages[2]["block_indices"] == [2, 5]
    assert "block_indices" not in pages[1]
    assert [node["node_id"] for node in only(before, "section")] == [node["node_id"] for node in sections]
    assert [node["node_id"] for node in only(before, "page")] == [node["node_id"] for node in pages]
    assert_ranges(document, structure)


def test_office_ocr_pages_do_not_attach_to_last_original_section(tmp_path):
    path = tmp_path / "source.docx"
    word = Document()
    word.add_heading("Section", 1)
    word.add_paragraph("Original text")
    word.save(path)
    document = read_source(path)
    document.blocks.extend([
        SourceBlock("page 1/ocr exe/block 1", "Page 1 diagram", metadata={"page": 1, "extraction": {"execution_id": "exe"}}),
        SourceBlock("page 2/ocr exe/block 2", "Page 2 diagram", metadata={"page": 2, "extraction": {"execution_id": "exe"}}),
        SourceBlock("page 1/ocr exe/block 3", "Page 1 note", metadata={"page": 1, "extraction": {"execution_id": "exe"}}),
    ])
    structure = build_structure(document)
    section = only(structure, "section")[0]
    assert section["range"] == {"start_block": 0, "end_block": 2}
    pages = only(structure, "page")
    assert pages[0]["block_indices"] == [2, 4]
    assert pages[1]["range"] == {"start_block": 3, "end_block": 4}
    assert all(page["metadata"]["requires_extraction_review"] for page in pages)
    assert any('not been reliably mapped to original sections' in warning for warning in structure["warnings"])
    assert_ranges(document, structure)


def test_pdf_out_of_order_or_invalid_bookmarks_fall_back_to_navigation(tmp_path):
    path = tmp_path / "source.pdf"
    pdf(path, [("one", "Later", 2, None), ("child", "Earlier", 0, "one"), ("invalid", "No page", None, None)])
    document = read_source(path)
    structure = build_structure(document)
    assert structure["structure_status"] == "outline_navigation_only"
    assert not only(structure, "section")
    assert len(only(structure, "bookmark")) == 2
    assert len(only(structure, "page")) == 4
    assert any("lacking valid page numbers" in warning for warning in structure["warnings"])
    assert_ranges(document, structure)


def test_empty_scanned_pdf_page_keeps_empty_range_and_unread_status(tmp_path):
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(path)
    document = read_source(path)
    structure = build_structure(document)
    assert not structure["complete"]
    page = only(structure, "page")[0]
    assert page["range"] == {"start_block": 0, "end_block": 0}
    assert page["metadata"]["needs_ocr"]


def test_slides_and_image_frames_have_explicit_navigation(tmp_path):
    from pptx import Presentation

    path = tmp_path / "slides.pptx"
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "Repeated"
    first.placeholders[1].text = "Body"
    second = presentation.slides.add_slide(presentation.slide_layouts[1])
    second.shapes.title.text = "Repeated"
    presentation.save(path)
    document = read_source(path)
    structure = build_structure(document)
    slides = only(structure, "slide")
    assert len(slides) == 2 and slides[0]["node_id"] != slides[1]["node_id"]
    assert slides[0]["range"]["end_block"] == slides[1]["range"]["start_block"]
    assert_ranges(document, structure)
    image_path = tmp_path / "frames.tiff"
    image = Image.new("RGB", (12, 20))
    image.save(image_path, save_all=True, append_images=[image.copy()])
    structure = build_structure(read_source(image_path))
    assert len(only(structure, "image")) == 2
    assert not structure["complete"]
    assert structure["structure_status"] == "images"


def test_title_change_does_not_change_position_identity_or_embed_body_metadata():
    document = SourceDocument("v1", {"format": "md", "complete": True, "comments": [{"text": "private body"}]},
                              blocks=[SourceBlock("line 1", "# Before", "heading", {"level": 1}), SourceBlock("line 2", "body", metadata={"text": "body"})])
    first = build_structure(document)
    document.blocks[0].text = "# After"
    second = build_structure(document)
    assert only(first, "section")[0]["node_id"] == only(second, "section")[0]["node_id"]
    assert "private body" not in str(first)
