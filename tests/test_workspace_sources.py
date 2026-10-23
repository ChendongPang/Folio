"""Many files bind to their original conversation and freeze before generation."""

import asyncio
import io
import json
from pathlib import Path

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from filelock import FileLock

from folio.api.app import create_app
from folio.model_settings import ModelSettingsInput, ModelSettingsStore
from folio.research.source_index import SourceIndex
from folio.runtime.demo import DemoRuntime
from folio.source_documents import SourceDirectoryDrafts, _prepare_inputs


def save_state(app, cid, state):
    with app.state.store.connect() as conn:
        conn.execute("UPDATE web_conversations SET data=? WHERE id=?", (json.dumps(state), cid))


class TitledDemoRuntime(DemoRuntime):
    def __init__(self, title):
        self.title = title

    async def complete(self, task, context, schema, session_key):
        result = await super().complete(task, context, schema, session_key)
        if task == "plan":
            result.content = json.dumps(json.loads(result.content) | {
                "document_title": self.title,
            }, ensure_ascii=False)
        return result


@pytest.fixture
def source_app(tmp_path, monkeypatch):
    ModelSettingsStore(tmp_path).save(
        ModelSettingsInput(base_url="http://unused.test/v1", model="test", no_auth=True)
    )
    monkeypatch.setattr("folio.directory_drafts.LLMRuntime", lambda _: TitledDemoRuntime("交易审批需求说明书"))
    monkeypatch.setattr(
        "folio.document.word.paginate", lambda *a, **k: {"status": "requires_office"}
    )
    return create_app(tmp_path)


def plan(client, cid):
    response = client.post(
        f"/conversations/{cid}/messages", json={"text": "交易审批需求说明书", "action": "generate"}
    )
    assert "directory.review" in response.text and "message.failed" not in response.text
    return client.get(f"/conversations/{cid}").json()["directory_review"]


def test_sources_and_text_draft_persist_independently_and_delete_only_reference(source_app):
    app = source_app
    with TestClient(app) as client:
        cid = client.post("/conversations").json()["id"]
        other = client.post("/conversations").json()["id"]
        source = client.post(
            f"/conversations/{cid}/sources?filename=facts.md", content="交易审批需要两人复核"
        ).json()
        assert source["role"] == "source" and "path" not in source
        assert (
            client.put(f"/conversations/{cid}/draft", json={"text": "编写交易规则"}).status_code
            == 200
        )
        state = client.get(f"/conversations/{cid}").json()
        assert state["draft"] == "编写交易规则" and state["files"][0]["id"] == source["id"]
        assert client.delete(f"/conversations/{other}/sources/{source['id']}").status_code == 404
        # A legacy client may reference the same blob. Removing one binding must
        # not break another conversation, completed job, or its frozen snapshot.
        assert (
            client.put(
                f"/conversations/{other}/draft", json={"files": [{"id": source["id"]}]}
            ).status_code
            == 200
        )
        selected = plan(client, cid)
        snapshot = SourceDirectoryDrafts(app.state.store.root, app.state.preview_library).load(
            selected["directory"]["revision"], cid
        )
        assert client.delete(f"/conversations/{cid}/sources/{source['id']}").status_code == 200
        state = client.get(f"/conversations/{cid}").json()
        assert state["files"] == [] and "directory_review" not in state
        assert state["draft"] == ""  # Sending generation consumed the draft.
        assert client.get(f"/conversations/{other}").json()["files"][0]["id"] == source["id"]
        with app.state.store.connect() as conn:
            blob = conn.execute(
                "SELECT path FROM web_uploads WHERE id=?", (source["id"],)
            ).fetchone()[0]
        assert Path(blob).read_text() == "交易审批需要两人复核"
        assert next(snapshot["knowledge_path"].iterdir()).read_text() == "交易审批需要两人复核"
        stale = client.post(
            f"/conversations/{cid}/messages",
            json={
                "text": "确认",
                "action": "confirm_directory",
                "selection_revision": selected["revision"],
            },
        )
        assert stale.status_code == 409
    with TestClient(create_app(app.state.store.root)) as client:
        assert client.get(f"/conversations/{other}").json()["files"][0]["id"] == source["id"]


def test_source_addition_revokes_directory_confirmation(source_app):
    with TestClient(source_app) as client:
        cid = client.post("/conversations").json()["id"]
        client.post(f"/conversations/{cid}/sources?filename=one.md", content="审批规则")
        selected = plan(client, cid)
        added = client.post(f"/conversations/{cid}/sources?filename=two.txt", content="补充规则")
        assert added.json()["directory_invalidated"] is True
        assert "directory_review" not in client.get(f"/conversations/{cid}").json()
        assert (
            client.post(
                f"/conversations/{cid}/messages",
                json={
                    "text": "确认",
                    "action": "confirm_directory",
                    "selection_revision": selected["revision"],
                },
            ).status_code
            == 409
        )
        with source_app.state.store.connect() as conn:
            assert conn.execute("SELECT count(*) FROM jobs").fetchone()[0] == 0


def upload_docx(client, cid, name="template.docx"):
    document = Document()
    document.add_heading("业务范围", level=1)
    document.add_paragraph("请填写本项目的业务范围。")
    buffer = io.BytesIO()
    document.save(buffer)
    response = client.post(
        f"/conversations/{cid}/sources?filename={name}", content=buffer.getvalue()
    )
    assert response.status_code == 201
    return response.json()


def test_template_chat_context_excludes_example_business_values(tmp_path):
    from folio.api.workspace import material_context

    document = Document()
    document.add_paragraph("旧项目名称：老商城；旧预算：9876543 元。")
    document.add_heading("业务范围", level=1)
    document.add_paragraph("请填写本项目的业务范围。")
    document.add_paragraph("旧项目审批金额为 543210 元。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "要求"
    table.cell(1, 0).text = "旧业务名称"
    table.cell(1, 1).text = "旧项目示例值"
    path = tmp_path / "template.docx"
    document.save(path)
    context = material_context(
        [{"id": "a" * 32, "name": path.name, "path": str(path), "role": "template"}]
    )
    encoded = json.dumps(context, ensure_ascii=False)
    assert all(
        value not in encoded
        for value in ["老商城", "9876543", "543210", "旧业务名称", "旧项目示例值"]
    )
    assert "业务范围" in encoded and "本项目的业务范围" in encoded
    assert context[0]["template"]["sections"][0]["tables"][0]["headers"] == ["字段", "要求"]
    assert str(tmp_path) not in encoded


def test_template_chat_returns_structure_problem_without_exposing_example_prose(tmp_path):
    from folio.api.workspace import material_context

    document = Document()
    document.add_paragraph("旧项目绝密样例金额 9876543")
    path = tmp_path / "unstructured.docx"
    document.save(path)
    context = material_context(
        [{"id": "b" * 32, "name": path.name, "path": str(path), "role": "template"}]
    )
    assert "No template headings were recognized" in context[0]["template_error"]
    assert "9876543" not in json.dumps(context, ensure_ascii=False)


def test_template_duplicate_as_source_does_not_leak_samples_into_chat(tmp_path):
    from folio.api.workspace import material_context

    document = Document()
    document.add_heading("业务范围", level=1)
    document.add_paragraph("旧项目示例预算 9876543 元")
    template = tmp_path / "template.docx"
    document.save(template)
    duplicate = tmp_path / "ordinary.docx"
    duplicate.write_bytes(template.read_bytes())
    facts = tmp_path / "facts.docx"
    document = Document()
    document.add_paragraph("当前项目的业务依据")
    document.save(facts)
    context = material_context(
        [
            {"id": "a" * 32, "name": template.name, "path": str(template), "role": "template"},
            {"id": "b" * 32, "name": duplicate.name, "path": str(duplicate), "role": "source"},
            {"id": "c" * 32, "name": facts.name, "path": str(facts), "role": "source"},
        ]
    )
    assert context[1]["excluded_as_template_duplicate"] is True
    assert "parsed" not in context[1] and "was not read as business evidence" in context[1]["note"]
    encoded = json.dumps(context, ensure_ascii=False)
    assert "9876543" not in encoded and "当前项目的业务依据" in encoded


def test_template_role_persists_and_replacement_requires_explicit_source_reclassification(
    source_app,
):
    with TestClient(source_app) as client:
        cid = client.post("/conversations").json()["id"]
        first = upload_docx(client, cid, "first.docx")
        second = upload_docx(client, cid, "second.DOCX")
        url = f"/conversations/{cid}/sources/{first['id']}"
        response = client.patch(url, json={"role": "template"})
        assert response.status_code == 200 and response.json()["directory_invalidated"]
        assert [item["role"] for item in response.json()["files"]] == ["template", "source"]
        client.put(f"/conversations/{cid}/draft", json={"text": "待补充要求"})
        restored = client.get(f"/conversations/{cid}").json()
        assert restored["files"][0]["role"] == "template" and restored["draft"] == "待补充要求"
        # Older callers omitting the role must retain the saved purpose.
        client.put(
            f"/conversations/{cid}/draft",
            json={"files": [{"id": first["id"]}, {"id": second["id"]}]},
        )
        assert client.get(f"/conversations/{cid}").json()["files"][0]["role"] == "template"
        response = client.patch(
            f"/conversations/{cid}/sources/{second['id']}", json={"role": "template"}
        )
        assert response.status_code == 422 and "先移除原模板" in response.json()["detail"]
        assert [item["role"] for item in client.get(f"/conversations/{cid}").json()["files"]] == [
            "template",
            "source",
        ]
        assert client.patch(url, json={"role": "source"}).status_code == 200
        response = client.patch(
            f"/conversations/{cid}/sources/{second['id']}", json={"role": "template"}
        )
        assert response.status_code == 200
        assert [item["role"] for item in response.json()["files"]] == ["source", "template"]
        assert all("path" not in item for item in response.json()["files"])
    with TestClient(create_app(source_app.state.store.root)) as client:
        assert [item["role"] for item in client.get(f"/conversations/{cid}").json()["files"]] == [
            "source",
            "template",
        ]


def test_template_role_changes_revoke_confirmation_but_noop_retains_directory(source_app):
    with TestClient(source_app) as client:
        cid = client.post("/conversations").json()["id"]
        source = upload_docx(client, cid)
        selected = plan(client, cid)
        url = f"/conversations/{cid}/sources/{source['id']}"
        unchanged = client.patch(url, json={"role": "source"}).json()
        assert unchanged["directory_invalidated"] is False
        assert (
            client.get(f"/conversations/{cid}").json()["directory_review"]["revision"]
            == selected["revision"]
        )
        assert client.patch(url, json={"role": "template"}).json()["directory_invalidated"]
        assert "directory_review" not in client.get(f"/conversations/{cid}").json()
        assert (
            client.post(
                f"/conversations/{cid}/messages",
                json={
                    "text": "确认",
                    "action": "confirm_directory",
                    "selection_revision": selected["revision"],
                },
            ).status_code
            == 409
        )
        with source_app.state.store.connect() as conn:
            assert conn.execute("SELECT count(*) FROM jobs").fetchone()[0] == 0


def test_template_role_validates_ownership_format_lock_and_legacy_inputs(source_app):
    with TestClient(source_app) as client:
        cid = client.post("/conversations").json()["id"]
        other = client.post("/conversations").json()["id"]
        text = client.post(
            f"/conversations/{cid}/sources?filename=facts.md", content=b"rules"
        ).json()
        first = upload_docx(client, cid)
        second = upload_docx(client, cid, "other.docx")
        url = f"/conversations/{cid}/sources/{first['id']}"
        assert (
            client.patch(
                f"/conversations/{other}/sources/{first['id']}", json={"role": "template"}
            ).status_code
            == 404
        )
        assert (
            client.patch(
                f"/conversations/{cid}/sources/{text['id']}", json={"role": "template"}
            ).status_code
            == 422
        )
        assert client.patch(url, json={"role": "example"}).status_code == 422
        assert (
            client.patch(
                url, json={"role": "template"}, headers={"Origin": "https://foreign.test"}
            ).status_code
            == 403
        )
        with FileLock(source_app.state.store.root / f"conversation-{cid}.lock"):
            assert client.patch(url, json={"role": "template"}).status_code == 409
        assert (
            client.put(
                f"/conversations/{cid}/draft",
                json={"files": [{"id": text["id"], "role": "template"}]},
            ).status_code
            == 422
        )
        multiple = [{"id": item["id"], "role": "template"} for item in (first, second)]
        assert (
            client.put(f"/conversations/{cid}/draft", json={"files": multiple}).status_code == 422
        )
        assert len(client.get(f"/conversations/{cid}").json()["files"]) == 3


def test_confirmation_rechecks_frozen_file_roles(source_app):
    with TestClient(source_app) as client:
        cid = client.post("/conversations").json()["id"]
        source = upload_docx(client, cid)
        selected = plan(client, cid)
        metadata_path = (
            source_app.state.store.root
            / "directory-drafts"
            / selected["directory"]["revision"]
            / "metadata.json"
        )
        metadata = json.loads(metadata_path.read_text())
        metadata["source_roles"] = [{"id": source["id"], "role": "template"}]
        metadata_path.write_text(json.dumps(metadata))
        rejected = client.post(
            f"/conversations/{cid}/messages",
            json={
                "text": "确认",
                "action": "confirm_directory",
                "selection_revision": selected["revision"],
            },
        )
        assert rejected.status_code == 409 and "文件用途已改变" in rejected.json()["detail"]
        with source_app.state.store.connect() as conn:
            assert conn.execute("SELECT count(*) FROM jobs").fetchone()[0] == 0


async def test_parallel_uploads_allow_draft_saves_and_bind_only_original_conversation(tmp_path):
    app = create_app(tmp_path)
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://testserver") as client:
        cid = (await client.post("/conversations")).json()["id"]
        other = (await client.post("/conversations")).json()["id"]
        arrived = [asyncio.Event() for _ in range(3)]
        release = asyncio.Event()

        async def body(i):
            yield f"# file {i}\n".encode()
            arrived[i].set()
            await release.wait()
            yield b"complete factual content"

        pending = [
            asyncio.create_task(
                client.post(f"/conversations/{cid}/sources?filename=file-{i}.md", content=body(i))
            )
            for i in range(3)
        ]
        await asyncio.wait_for(asyncio.gather(*(event.wait() for event in arrived)), timeout=3)
        # Each upload is still streaming; no one may monopolize the session lock.
        assert (
            await client.put(f"/conversations/{cid}/draft", json={"text": "original draft"})
        ).status_code == 200
        assert (
            await client.put(f"/conversations/{other}/draft", json={"text": "other draft"})
        ).status_code == 200
        release.set()
        responses = await asyncio.gather(*pending)
        assert all(response.status_code == 201 for response in responses)
        restored = (await client.get(f"/conversations/{cid}")).json()
        assert len(restored["files"]) == 3 and restored["draft"] == "original draft"
        assert {f["id"] for f in restored["files"]} == {r.json()["id"] for r in responses}
        assert "knowledge_engineering" not in restored
        assert (await client.get(f"/conversations/{other}")).json()["files"] == []


async def test_job_started_during_upload_rejects_binding_and_cleans_only_new_blob(
    tmp_path, monkeypatch
):
    app = create_app(tmp_path)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://testserver"
    ) as client:
        cid = (await client.post("/conversations")).json()["id"]
        first = (
            await client.post(f"/conversations/{cid}/sources?filename=first.md", content=b"keep")
        ).json()

        async def body():
            yield b"partial"
            state = (await client.get(f"/conversations/{cid}")).json()
            state["active_job"] = "running-job"
            save_state(app, cid, state)
            monkeypatch.setattr(app.state.store, "get", lambda _: {"status": "running"})
            yield b"last"

        rejected = await client.post(
            f"/conversations/{cid}/sources?filename=late.md", content=body()
        )
        assert rejected.status_code == 409
        with app.state.store.connect() as conn:
            saved = conn.execute("SELECT id,path FROM web_uploads").fetchall()
        assert len(saved) == 1 and saved[0]["id"] == first["id"]
        assert Path(saved[0]["path"]).read_bytes() == b"keep"
        assert len(list((tmp_path / "uploads").iterdir())) == 1


@pytest.mark.parametrize("status", ["queued", "running", "pausing", "paused"])
def test_source_mutations_lock_during_unfinished_jobs(tmp_path, monkeypatch, status):
    app = create_app(tmp_path)
    with TestClient(app) as client:
        cid = client.post("/conversations").json()["id"]
        added = client.post(
            f"/conversations/{cid}/sources?filename=first.md", content=b"keep"
        ).json()
        state = client.get(f"/conversations/{cid}").json()
        state["active_job"] = "job"
        save_state(app, cid, state)
        monkeypatch.setattr(app.state.store, "get", lambda _: {"status": status})
        assert (
            client.post(
                f"/conversations/{cid}/sources?filename=extra.md", content=b"no"
            ).status_code
            == 409
        )
        assert client.delete(f"/conversations/{cid}/sources/{added['id']}").status_code == 409
        assert (
            client.patch(
                f"/conversations/{cid}/sources/{added['id']}", json={"role": "source"}
            ).status_code
            == 409
        )
        assert len(client.get(f"/conversations/{cid}").json()["files"]) == 1


def test_source_limit_origin_and_conversation_lock(tmp_path):
    app = create_app(tmp_path)
    with TestClient(app) as client:
        cid = client.post("/conversations").json()["id"]
        url = f"/conversations/{cid}/sources?filename=file.md"
        assert (
            client.post(url, content=b"x", headers={"Origin": "https://foreign.test"}).status_code
            == 403
        )
        with FileLock(tmp_path / f"conversation-{cid}.lock"):
            assert client.post(url, content=b"x").status_code == 409
        for i in range(200):
            assert client.post(url, content=f"file {i}".encode()).status_code == 201
        assert client.post(url, content=b"201st").status_code == 422
        state = client.get(f"/conversations/{cid}").json()
        assert len(state["files"]) == 200
        assert "knowledge_engineering" not in state
        with app.state.store.connect() as conn:
            assert conn.execute("SELECT count(*) FROM web_uploads").fetchone()[0] == 200


def test_many_sources_catalog_budget_includes_headings_names_and_json(tmp_path):
    source = tmp_path / "input.md"
    content = ("# " + 'head\\"' * 80 + "\n") * 12 + "细节事实" * 5000
    source.write_text(content)
    attachments = [
        {"id": f"{i:032x}", "name": "很长的资料名称" * 25 + ".md", "path": str(source)}
        for i in range(200)
    ]
    catalog = _prepare_inputs(tmp_path / "prepared", attachments)
    assert len(catalog) == 200
    assert len(json.dumps(catalog, ensure_ascii=False)) <= 100000
    assert {entry["source_id"] for entry in catalog} == {item["id"] for item in attachments}
    assert all(not {"headings", "excerpt", "headings_truncated", "excerpt_truncated"}.intersection(entry)
               for entry in catalog)
    assert all(entry["name_truncated"] and entry["path"] for entry in catalog)
    snapshots = list((tmp_path / "prepared" / "sources").iterdir())
    assert len(snapshots) == 200 and all(p.read_text() == content for p in snapshots)
    assert all(len(p.name.encode()) < 255 for p in snapshots)
    index = SourceIndex(tmp_path / "prepared" / "sources")
    index.prepare()
    structure = index.structure(catalog[0]["path"])
    assert [node["title"] for node in structure["nodes"] if node["kind"] == "section"] == [
        content.splitlines()[0].removeprefix("# ")
    ] * 12
    assert "细节事实" * 5000 in "\n".join(block.text for block in index.load(catalog[0]["path"]).blocks)


def test_uploaded_docx_is_factual_source_and_internal_layout_is_independent(source_app, inputs):
    app = source_app
    with TestClient(app) as client:
        cid = client.post("/conversations").json()["id"]
        import io

        from docx import Document

        source = Document()
        source.add_paragraph("这是一份业务事实资料")
        buffer = io.BytesIO()
        source.save(buffer)
        original = buffer.getvalue()
        client.post(f"/conversations/{cid}/sources?filename=notes.docx", content=original)
        selected = plan(client, cid)
        snapshot = SourceDirectoryDrafts(app.state.store.root, app.state.preview_library).load(
            selected["directory"]["revision"], cid
        )
        assert next(snapshot["knowledge_path"].iterdir()).read_bytes() == original
        assert "template_path" not in snapshot and "outline_path" not in snapshot
        assert snapshot["title"] == "交易审批需求说明书"
        assert "cover_values" not in selected and "template" not in selected
        with pytest.raises(ValueError, match="不属于"):
            SourceDirectoryDrafts(app.state.store.root, app.state.preview_library).load(
                selected["directory"]["revision"], "another"
            )


@pytest.mark.parametrize(
    "action",
    [
        "confirm_template",
        "confirm_cover",
        "upload_template",
        "manual_directory",
        "upload_directory",
    ],
)
def test_removed_gui_prerequisites_are_not_message_actions(tmp_path, action):
    with TestClient(create_app(tmp_path)) as client:
        cid = client.post("/conversations").json()["id"]
        assert (
            client.post(
                f"/conversations/{cid}/messages", json={"text": "test", "action": action}
            ).status_code
            == 422
        )


@pytest.mark.parametrize("generate_text", ["部署方案", "根据已上传资料和我们的讨论规划目录"])
def test_document_title_uses_latest_writing_request_not_greeting_or_button_label(
    source_app, generate_text, monkeypatch
):
    monkeypatch.setattr("folio.directory_drafts.LLMRuntime", lambda _: TitledDemoRuntime("部署方案"))
    app = source_app
    with TestClient(app) as client:
        cid = client.post("/conversations").json()["id"]
        client.post(
            f"/conversations/{cid}/sources?filename=部署.md", content="部署与回滚采用版本发布流程。"
        )
        state = client.get(f"/conversations/{cid}").json()
        state["messages"] = [
            {"role": "user", "content": "你好", "action": "chat"},
            {"role": "assistant", "content": "你好，请说明文档要求。"},
            {"role": "user", "content": "部署方案", "action": "chat"},
        ]
        save_state(app, cid, state)
        response = client.post(
            f"/conversations/{cid}/messages", json={"text": generate_text, "action": "generate"}
        )
        assert "directory.review" in response.text
        review = client.get(f"/conversations/{cid}").json()["directory_review"]
        assert review["title"] == "部署方案"
        assert "你好" not in review["request"]  # Greetings and button labels are not requirements.
        assert "部署方案" in review["request"]
        snapshot = SourceDirectoryDrafts(app.state.store.root, app.state.preview_library).load(
            review["directory"]["revision"], cid
        )
        assert snapshot["title"] == "部署方案"
        assert "部署方案" in json.dumps(snapshot["requirements"], ensure_ascii=False)
