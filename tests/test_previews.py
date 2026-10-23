import pytest
from docx import Document
from fastapi.testclient import TestClient

from folio.api.app import create_app


@pytest.fixture
def app(tmp_path):
    application = create_app(tmp_path / "web")
    source = tmp_path / "generated.docx"
    doc = Document()
    doc.add_heading("业务需求", 1)
    doc.add_paragraph("已生成的目录预览")
    doc.save(source)
    application.state.fixture_preview = application.state.preview_library.add(source, "目录预览")
    return application


def test_preview_cache_preserves_generated_source(app, monkeypatch):
    item = app.state.fixture_preview
    source = app.state.preview_library.path(item["id"])
    before = source.read_bytes()

    def paginate(copy, pdf, required):
        copy.write_bytes(b"changed copy")
        pdf.write_bytes(b"%PDF-test")

    monkeypatch.setattr("folio.previews.paginate", paginate)
    with TestClient(app) as client:
        result = client.get(item["preview_url"])
        assert result.status_code == 200 and result.content == b"%PDF-test"
        assert source.read_bytes() == before
        monkeypatch.setattr(
            "folio.previews.paginate", lambda *a, **k: pytest.fail("cache unused")
        )
        assert client.get(item["preview_url"]).status_code == 200
        assert client.get("/previews/not-a-preview/preview").status_code == 404
        assert client.get("/templates").status_code == 404


def test_preview_page_routes_validate_bounds(app, monkeypatch):
    item = app.state.fixture_preview
    folder = app.state.preview_library.path(item["id"]).parent / "pages"
    folder.mkdir()
    (folder / "page-1.png").write_bytes(b"PNG-ONE")
    (folder / "page-2.png").write_bytes(b"PNG-TWO")
    monkeypatch.setattr(
        app.state.preview_library, "preview", lambda _: folder.parent / "preview.pdf"
    )
    with TestClient(app) as client:
        pages = client.get(f"/previews/{item['id']}/pages").json()["pages"]
        assert len(pages) == 2
        assert client.get(pages[0]).content == b"PNG-ONE"
        assert client.get(pages[1]).headers["content-type"] == "image/png"
        assert client.get(f"/previews/{item['id']}/pages/0").status_code == 404
        assert client.get(f"/previews/{item['id']}/pages/3").status_code == 404


def test_page_failure_does_not_leak_renderer_details_or_break_structure(app, monkeypatch):
    def failed(*args):
        raise RuntimeError("private renderer details")

    monkeypatch.setattr(app.state.preview_library, "pages", failed)
    with TestClient(app) as client:
        item = app.state.fixture_preview
        response = client.get(f"/previews/{item['id']}/pages")
        assert response.status_code == 503 and "private renderer" not in response.text
        assert client.get(f"/previews/{item['id']}").json()["blocks"]
