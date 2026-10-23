from docx import Document
from docx.oxml import OxmlElement

from folio.document.reader import read_docx


def test_tracked_text_view_preserves_insertions_and_excludes_deletions(tmp_path):
    doc = Document()
    doc.add_heading("功能", 1)
    p = doc.add_paragraph("当前：")
    inserted = OxmlElement("w:ins")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "新要求"
    r.append(t)
    inserted.append(r)
    p._p.append(inserted)
    deleted = OxmlElement("w:del")
    r = OxmlElement("w:r")
    t = OxmlElement("w:delText")
    t.text = "旧要求"
    r.append(t)
    deleted.append(r)
    p._p.append(deleted)
    path = tmp_path / "source.docx"
    doc.save(path)
    assert read_docx(path).blocks[-1].text == "当前：新要求"
    assert read_docx(path, "original").blocks[-1].text == "当前：旧要求"


def test_word_outline_uses_style_levels_and_paths_not_numeric_style_ids(tmp_path):
    doc = Document()
    for text, level in [("功能", 1), ("甲", 2), ("输入", 3), ("乙", 2), ("输入", 3)]:
        doc.add_heading(text, level)
    path = tmp_path / "outline.docx"
    doc.save(path)
    nodes = read_docx(path).blocks
    assert [n.level for n in nodes] == [1, 2, 3, 2, 3]
    assert nodes[2].locator != nodes[4].locator
    assert nodes[4].path == ["功能", "乙", "输入"]


def test_table_stays_between_its_surrounding_paragraphs(tmp_path):
    doc = Document()
    doc.add_heading("输入", 1)
    doc.add_paragraph("前文")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "金额"
    table.cell(1, 1).text = "交易金额"
    doc.add_paragraph("后文")
    path = tmp_path / "source.docx"
    doc.save(path)
    blocks = read_docx(path).blocks
    assert [b.kind for b in blocks] == ["heading", "paragraph", "table", "paragraph"]
    assert blocks[2].rows[1] == ["金额", "交易金额"]
