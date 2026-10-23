"""Only generated directory presentations are refreshed; confirmed inputs stay frozen."""

import copy
import json
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from folio.api.app import create_app
from folio.directory_drafts import DirectoryDrafts
from folio.document.planner import build_plan
from folio.document.prepare import DIRECTORY_PREVIEW_VERSION, prepare
from folio.document.word import validate_render
from folio.models import PlanItem, PlanProposal
from folio.previews import PreviewLibrary
from folio.source_documents import SourceDirectoryDrafts


@pytest.fixture
def saved_directory(tmp_path, monkeypatch):
    def unexpected(*args, **kwargs):
        pytest.fail("Refreshing a presentation must not call a model or the office worker")

    monkeypatch.setattr("folio.directory_drafts.LLMRuntime", unexpected)
    monkeypatch.setattr("folio.document.word.paginate", unexpected)
    root = tmp_path / "state"
    library = PreviewLibrary(root)
    drafts = SourceDirectoryDrafts(root, library)
    revision = "a" * 32
    folder = drafts.root / revision
    folder.mkdir(parents=True)
    plan = build_plan(
        "confirmed-document",
        PlanProposal(sections=[
            PlanItem(node_id="orders", title="订单", level=1, content_mode="container"),
            PlanItem(node_id="cancel", title="取消", level=2, instruction="说明订单取消"),
        ]),
    )
    (folder / "document-plan.json").write_text(plan.model_dump_json(), encoding="utf-8")
    metadata = {
        "conversation_id": "conversation",
        "revision": revision,
        "title": "订单需求",
        "requirements": {"revision": "frozen-requirements", "text": "说明取消规则"},
        "source_ids": ["frozen-source"],
    }
    (folder / "metadata.json").write_text(json.dumps(metadata), encoding="utf-8")
    (folder / "sources").mkdir()
    (folder / "sources" / "rules.txt").write_text("付款前可取消", encoding="utf-8")
    (folder / "sources.index").mkdir()
    (folder / "sources.index" / "manifest.json").write_text('{"version":"frozen"}')
    old_source = folder / "文档-目录预览.docx"
    old_doc = Document()
    old_doc.add_paragraph("待基于知识库生成。")
    old_doc.add_paragraph("待确认：尚未检索知识库和生成正文")
    old_doc.save(old_source)
    old_preview = library.add(old_source, "已有目录")
    old_cache = library.path(old_preview["id"]).parent
    (old_cache / "preview.pdf").write_bytes(b"old PDF")
    (old_cache / "pages").mkdir()
    (old_cache / "pages" / "page-1.png").write_bytes(b"old page")
    directory = {
        "revision": revision,
        "document_id": "confirmed-document",
        "requirements_revision": "frozen-requirements",
        "sections": [s.model_dump(mode="json") for s in plan.sections],
        "preview": old_preview,
    }
    return drafts, directory, plan, folder


def snapshot(folder):
    return {str(path.relative_to(folder)): path.read_bytes() for path in folder.rglob("*")
            if path.is_file()}


def test_directory_refresh_restores_specific_failure_without_mutating_saved_inputs(saved_directory):
    drafts, directory, _, folder = saved_directory
    metadata_file = folder / "metadata.json"
    metadata = json.loads(metadata_file.read_text())
    issue = {"code": "tool_error", "message": "研究工具执行失败，已保存研究笔记。"}
    metadata.update(confirmable=False, research={"status": "partial", "termination_reason": "tool_error"},
                    confirmation_issues=[issue])
    metadata_file.write_text(json.dumps(metadata))
    before = snapshot(folder)
    refreshed = drafts.refresh_preview(directory | {"confirmable": False}, "conversation")
    assert refreshed["confirmation_issues"] == [issue]
    assert refreshed["confirmable"] is False
    assert snapshot(folder) == before
    assert drafts.refresh_preview(refreshed, "conversation")["confirmation_issues"] == [issue]


def test_directory_refresh_preserves_authoritative_plan_and_all_original_files(saved_directory):
    drafts, directory, plan, folder = saved_directory
    original_directory = copy.deepcopy(directory)
    original_inputs = snapshot(folder)
    old_preview_folder = drafts.library.path(directory["preview"]["id"]).parent
    original_preview = snapshot(old_preview_folder)

    refreshed = drafts.refresh_preview(directory, "conversation")

    assert directory == original_directory
    presentation_fields = {"preview", "confirmable", "research"}
    assert {k: v for k, v in refreshed.items() if k not in presentation_fields} == {
        k: v for k, v in directory.items() if k not in presentation_fields
    }
    assert set(refreshed) - set(directory) == {"confirmable", "research"}
    assert refreshed["confirmable"] is False
    assert refreshed["research"]["status"] == "partial"
    assert "旧目录" in refreshed["research"]["summary"]
    assert refreshed["research"]["unresolved"]
    assert refreshed["preview"]["presentation_version"] == DIRECTORY_PREVIEW_VERSION
    assert refreshed["preview"]["id"] != directory["preview"]["id"]
    assert snapshot(folder) == original_inputs
    assert snapshot(old_preview_folder) == original_preview
    new_source = drafts.library.path(refreshed["preview"]["id"])
    text = "\n".join(p.text for p in Document(new_source).paragraphs)
    assert "订单需求" in text and "目录预览" in text
    assert "确认目录后，将研究上传资料" in text
    assert "知识库" not in text and "待确认：" not in text
    validate_render(new_source, plan)
    assert DirectoryDrafts.load(drafts, directory["revision"], "conversation")["plan"] == plan


def test_repeated_old_directory_reads_reuse_versioned_cache(saved_directory, monkeypatch):
    drafts, directory, _, _ = saved_directory
    refreshed = drafts.refresh_preview(directory, "conversation")

    def unexpected(*args, **kwargs):
        pytest.fail("A cached directory preview must not render again")

    monkeypatch.setattr("folio.directory_drafts.prepare", unexpected)
    assert drafts.refresh_preview(directory, "conversation") == refreshed
    assert drafts.refresh_preview(refreshed, "conversation") == refreshed


def test_new_presentation_version_regenerates_only_derived_preview(saved_directory, monkeypatch):
    drafts, directory, _, folder = saved_directory
    original = snapshot(folder)
    first = drafts.refresh_preview(directory, "conversation")
    calls = []

    def tracked(*args, **kwargs):
        calls.append(kwargs)
        return prepare(*args, **kwargs)

    monkeypatch.setattr("folio.directory_drafts.prepare", tracked)
    monkeypatch.setattr("folio.directory_drafts.DIRECTORY_PREVIEW_VERSION", 6)
    next_version = drafts.refresh_preview(first, "conversation")
    assert next_version["preview"]["presentation_version"] == 6
    assert len(calls) == 1 and calls[0]["defer_pagination"]
    assert snapshot(folder) == original
    assert len(list((drafts.root.parent / "directory-preview-cache").glob("*.json"))) == 2


def test_refresh_failure_never_publishes_or_changes_the_saved_directory(saved_directory, monkeypatch):
    drafts, directory, _, folder = saved_directory
    original = snapshot(folder)
    original_directory = copy.deepcopy(directory)

    def failed(*args, **kwargs):
        raise RuntimeError("render unavailable")

    monkeypatch.setattr("folio.directory_drafts.prepare", failed)
    with pytest.raises(RuntimeError, match="render unavailable"):
        drafts.refresh_preview(directory, "conversation")
    assert directory == original_directory
    assert snapshot(folder) == original
    assert not list((drafts.root.parent / "directory-preview-cache").glob("*.json"))
    assert not list((drafts.root.parent / "directory-preview-cache").glob("render-*"))


def test_refresh_only_applies_to_saved_generated_directories(saved_directory):
    drafts, directory, _, folder = saved_directory
    normal_document = {"preview": directory["preview"]}
    assert drafts.refresh_preview(normal_document, "conversation") is normal_document
    with pytest.raises(ValueError, match="不属于当前会话"):
        drafts.refresh_preview(directory, "someone-else")
    assert not (drafts.root.parent / "directory-preview-cache").exists()
    # Viewing a saved plan remains possible even if its old source snapshot is unavailable.
    (folder / "sources").rename(folder / "old-sources")
    assert drafts.refresh_preview(directory, "conversation")["preview"]["presentation_version"]


def test_concurrent_refreshes_publish_one_cached_presentation(saved_directory, monkeypatch):
    drafts, directory, _, _ = saved_directory
    calls = []

    def tracked(*args, **kwargs):
        calls.append(kwargs)
        return prepare(*args, **kwargs)

    monkeypatch.setattr("folio.directory_drafts.prepare", tracked)
    with ThreadPoolExecutor(max_workers=2) as workers:
        results = list(workers.map(lambda _: drafts.refresh_preview(directory, "conversation"), range(2)))
    assert results[0] == results[1]
    assert len(calls) == 1


def test_unwritten_preview_content_is_not_a_business_confirmation_issue(saved_directory, tmp_path):
    _, _, plan, _ = saved_directory
    result = prepare(tmp_path / "new-preview", plan=plan, defer_pagination=True)
    assert not result["missing_sections"]
    assert all(not item["pending_items"] for item in result["directory"])
    assert result["content_sections"] == 1
    trace = json.loads((Path(result["artifact"]).parent / "trace.json").read_text())
    assert trace[0]["content"]["blocks"] == []  # Container headings have no body placeholder.
    assert trace[1]["content"]["blocks"][0]["text"].startswith("尚未生成正文。")


def test_english_directory_preview_localizes_generated_labels(saved_directory, tmp_path):
    _, _, plan, _ = saved_directory
    result = prepare(
        tmp_path / "english-preview",
        plan=plan,
        title="Order requirements",
        defer_pagination=True,
        language="en",
    )
    artifact = Path(result["artifact"])
    document = Document(artifact)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]

    assert artifact.name == "Document-outline-preview.docx"
    assert "Order requirements" in paragraphs and "OUTLINE PREVIEW" in paragraphs
    assert "Contents" in paragraphs
    trace = json.loads((artifact.parent / "trace.json").read_text())
    assert trace[1]["content"]["blocks"][0]["text"].startswith("Body not generated yet.")
    assert (artifact.parent / "inspection-report.md").is_file()


def test_restored_conversation_serves_refreshed_preview_without_rewriting_saved_state(saved_directory):
    drafts, directory, _, folder = saved_directory
    original_files = snapshot(folder)
    state = {
        "messages": [{"role": "assistant", "content": "已生成目录"}],
        "files": [],
        "active_job": None,
        "directory_review": {"phase": "directory_review", "directory": directory},
    }
    original_state = json.dumps(state)
    app = create_app(drafts.root.parent)
    with app.state.store.connect() as connection:
        connection.execute(
            "INSERT INTO web_conversations(id,data) VALUES (?,?)", ("conversation", original_state)
        )
    with TestClient(app) as client:
        restored = client.get("/conversations/conversation").json()
        refreshed = restored["directory_review"]["directory"]["preview"]
        assert refreshed["id"] != directory["preview"]["id"]
        response = client.get(f"/previews/{refreshed['id']}")
        assert response.status_code == 200
        assert "知识库" not in response.text and "待确认：" not in response.text
        assert "确认目录后，将研究上传资料" in response.text
        assert client.get(refreshed["download_url"]).status_code == 200
    with app.state.store.connect() as connection:
        saved = connection.execute(
            "SELECT data FROM web_conversations WHERE id=?", ("conversation",)
        ).fetchone()[0]
    assert saved == original_state
    assert snapshot(folder) == original_files
