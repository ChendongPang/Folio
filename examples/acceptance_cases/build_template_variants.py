"""Build optional-template variants without changing the original ten case packages.

Run ``uv run python examples/acceptance_cases/build_template_variants.py``.
Only local DOCX construction and the production template parser are used.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from folio.document.template_contract import read_template

ROOT = Path(__file__).resolve().parent
VARIANTS = (
    ("TV01", "01_expense_urd", "固定章节需求说明书"),
    ("TV02", "02_order_srs", "同列名的两张独立必填表"),
    ("TV05", "05_inventory_report", "两行合并表头的分析报告"),
    ("TV06", "06_operations_manual", "资料不足的完整运维模板"),
    ("TV08", "08_data_platform_srs", "允许子章节扩展的长文规格书"),
)
TEMPLATE_NAME = "文档模板.docx"
SAMPLE_PROJECT = "北辰旧项目示例"


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def file_record(path, folder, role):
    return {
        "path": os.path.relpath(path, folder),
        "role": role,
        "format": path.suffix.removeprefix(".").lower(),
        "bytes": path.stat().st_size,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    }


def document(variant_id):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "仿宋")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.5
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体")
    doc.styles["Heading 1"].font.size = Pt(17)
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin, section.bottom_margin = Cm(2.4), Cm(2.2)
        section.left_margin, section.right_margin = Cm(2.8), Cm(2.1)
        section.header.paragraphs[0].text = f"Folio 合成模板 · {variant_id} · 文档标准"
        footer = section.footer.paragraphs[0]
        footer.text = "仅供合成验收使用 · 第 "
        page = OxmlElement("w:fldSimple")
        page.set(qn("w:instr"), "PAGE")
        run, text = OxmlElement("w:r"), OxmlElement("w:t")
        text.text = "1"
        run.append(text)
        page.append(run)
        footer._p.append(page)
        footer.add_run(" 页")
    doc.add_paragraph("{{document_title}}", "Title")
    doc.add_paragraph("合成验收文档 · 按模板填写")
    doc.add_page_break()
    return doc


def mark_header(row):
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(marker)


def required_table(doc, headers):
    table = doc.add_table(rows=2, cols=len(headers))
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, headers, strict=True):
        cell.text = label
    mark_header(table.rows[0])
    # The blank data row carries formatting, not an answer or business example.
    return table


def merged_header_table(doc):
    table = doc.add_table(rows=3, cols=4)
    table.style = "Table Grid"
    table.cell(0, 0).merge(table.cell(0, 1)).text = "商品资料"
    table.cell(0, 2).merge(table.cell(0, 3)).text = "补货结论"
    for cell, label in zip(table.rows[1].cells, ["SKU", "计量单位", "建议数量", "计算依据"], strict=True):
        cell.text = label
    for row in table.rows[:2]:
        mark_header(row)
    return table


def section_instruction(variant_id, index):
    expansion = variant_id == "TV08" and 2 <= index <= 11
    policy = (
        "允许根据业务需要增加子章节；保留本章标题、位置和已有章节顺序。"
        if expansion
        else "不得新增子章节或修改目录；用段落和表格展开本章内容。"
    )
    return "填写说明：" + policy + "从业务资料研究本章主题；缺少依据的参数明确待确认。"


def build_variant(variant_id, slug, name, output_root):
    recipe = json.loads((ROOT / "recipes" / f"{slug}.json").read_text(encoding="utf-8"))
    base = ROOT / "cases" / slug
    folder = Path(output_root) / slug
    folder.mkdir(parents=True, exist_ok=True)
    old_amount = str(987654320 + int(variant_id[2:]))
    doc = document(variant_id)
    for index, chapter in enumerate(recipe["expected"]["chapters"]):
        title = chapter["title"]
        doc.add_heading(title, 1)
        doc.add_paragraph(section_instruction(variant_id, index))
        if index == 0:
            doc.add_paragraph(
                f"旧项目示例（不得沿用）：{SAMPLE_PROJECT}的预算是 {old_amount} 元，"
                "审批人为“示例负责人”。这些只是历史示例，不能作为本项目事实。"
            )
        if variant_id == "TV02" and title == "支付窗口与并发竞争":
            doc.add_paragraph(
                "填写说明：下列两张同列名表必须分别保留。第一张整理支付回调与重复消息；"
                "第二张整理超时关闭与支付并发。各表独立填写至少一条资料支持的场景，不能合并。"
            )
            headers = ["场景编号", "触发条件", "预期处理", "验收依据"]
            required_table(doc, headers)
            required_table(doc, headers)
        if variant_id == "TV05" and title == "SKU补货建议与执行清单":
            doc.add_paragraph(
                "填写说明：保留两行分组表头；按真实工作簿计算各 SKU 建议数量及依据，"
                "不得把空白样式行当成已完成的数据记录。"
            )
            merged_header_table(doc)
        if variant_id == "TV06" and title == "数据库、备份与恢复验证":
            doc.add_paragraph(
                "填写说明：分别列出 RTO、RPO、恢复验证状态及其项目依据；"
                "未获批准的目标和值保留待确认，不能用旧项目示例或通用经验填成已批准政策。"
            )
            required_table(doc, ["参数或验证事项", "已确认值或待确认", "依据与核实方式"])
    template_path = folder / TEMPLATE_NAME
    doc.save(template_path)
    contract = read_template(template_path)
    base_prompt = (base / "prompt.txt").read_text(encoding="utf-8").strip()
    additions = [
        f"本次按另外上传并标记为“模板”的《{TEMPLATE_NAME}》生成文档，其他文件均为业务资料。",
        "保留模板固定章节的标题、层级及顺序；只在模板明确允许的章节内扩展子章节。"
        "这项模板约束用于明确本次的目录与格式要求。",
        "模板必填表格逐张保留并填写，不能只留表头，也不能用同一张表满足两张表的要求。",
        "保留模板封面格式、页眉页脚、页面边距、字体和表格表头结构；封面 document_title 填写当前文档名称。"
        "模板没有规定的排版可以采用系统默认值。",
        "模板中的填写说明、空白行和旧项目示例不作为业务事实，成品中删除填写说明和旧项目示例；"
        "资料不足的项目参数明确写待确认，模型建议保持可辨认，不伪造批准记录。",
    ]
    if variant_id == "TV06":
        additions.append(
            "GUI 人工验收时关闭联网；若使用隔离的测试驱动，仅允许原 TC06 的离线 Mock 响应。"
            "必需章节与表格不得因缺资料而删除，特别检查 RTO/RPO 待确认项。"
        )
    if variant_id == "TV08":
        additions.append("文档仍须达到原 5 万至 6.5 万有效字符目标，长内容在允许扩展的章节下合理拆分。")
    (folder / "prompt.txt").write_text(
        base_prompt + "\n\n本次模板要求：\n" + "\n".join(additions) + "\n", encoding="utf-8"
    )
    expected = {
        "variant_id": variant_id,
        "base_case_id": recipe["id"],
        "name": name,
        "document_type": recipe["document_type"],
        "output_filename": recipe["output_filename"],
        "target_chars": recipe["target_chars"],
        "base_expected_path": os.path.relpath(base / "expected" / "expected.json", folder),
        "template_path": TEMPLATE_NAME,
        "sources_path": os.path.relpath(base / "sources", folder),
        "fixed_sections": [s.model_dump(exclude={"body_index"}) for s in contract.sections],
        "format_expectations": {
            "normal_latin_font": "Arial", "normal_east_asian_font": "仿宋",
            "normal_font_size_pt": 10.5, "normal_line_spacing": 1.5,
            "normal_space_after_pt": 8,
            "margins_cm": {"top": 2.4, "bottom": 2.2, "left": 2.8, "right": 2.1},
            "page_size_cm": [21, 29.7], "preserve_header_footer": True,
            "cover_document_title": recipe["output_title"],
        },
        "forbidden_template_examples": [SAMPLE_PROJECT, old_amount, "示例负责人"],
        "acceptance": [
            "继承原案例的事实、计算、缺口、图表、覆盖范围和有效字数验收。",
            "固定章节标题、层级及顺序与模板一致，新增子章节仅位于允许范围。",
            "每张必填表独立存在，逻辑列名完整匹配，至少有一行资料支持或明确待确认的数据。",
            "模板原件及其重复上传副本不得进入业务证据索引或全局事实。",
            "成品删除旧项目示例与填写提示；不把待确认参数伪装成确定事实。",
            "以真实 Word/PDF 成品检查模板字体、边距、页眉页脚、表头与分页可读性。",
        ],
        "variant_checks": {
            "TV01": ["全部章节固定，不接受任意同义改名或增加正文子标题。"],
            "TV02": ["支付并发章节的两张同列表头表不能合并；分别覆盖回调和超时关闭。"],
            "TV05": ["原 Word 保留两行合并表头；数据行不合并；不以空白样式行充数。"],
            "TV06": ["RTO/RPO 等未批准参数保留待确认；通用联网信息不能证明项目内部批准。"],
            "TV08": ["第 3 至 12 章允许新增子章节，其他固定章节不能扩展；不靠重复内容达到字数。"],
        }[variant_id],
        "validation_status": "fixture_structure_checked_only",
        "full_long_document_generation_tested": False,
        "requires_manual_review": True,
    }
    write_json(folder / "expected.json", expected)
    sources = sorted(path for path in (base / "sources").rglob("*") if path.is_file())
    write_json(folder / "manifest.json", {
        "variant_id": variant_id, "base_case_id": recipe["id"], "synthetic": True,
        "files": [file_record(template_path, folder, "template"),
                  *[file_record(path, folder, "source") for path in sources]],
        "do_not_upload": ["expected.json", "manifest.json", expected["base_expected_path"]],
        "user_prompt": "prompt.txt",
        "notes": "仅检查夹具可解析及约束结构；未进行真实模型长文生成验收。",
    })
    return folder


def build_all(output_root=None):
    output = Path(output_root) if output_root else ROOT / "template_variants"
    return [build_variant(*variant, output) for variant in VARIANTS]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "template_variants")
    args = parser.parse_args()
    folders = build_all(args.output)
    print(f"已离线生成并解析 {len(folders)} 套模板变体：{args.output}")


if __name__ == "__main__":
    main()
