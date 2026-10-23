"""Build synthetic, uploadable document acceptance cases without calling a model.

Run with ``uv run python examples/acceptance_cases/build_cases.py``.
Only PDF/PNG rendering needs the project's installed Playwright Chromium.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import io
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches
from pptx.util import Pt as SlidePt
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parent
COUNTING = (
    "正文有效字符数：汉字、英文字母、数字逐字符计数，计入表格，"
    "不计封面、目录、标题、页眉页脚、空白和标点。页数只作观察，不作字数替代。"
)
CSS = """
@page { size: A4; margin: 20mm; }
* { box-sizing: border-box; }
body { font-family: 'Noto Sans CJK SC', 'Microsoft YaHei', sans-serif;
       color: #17293c; margin: 0; font-size: 15px; line-height: 1.7; }
h1 { font-size: 27px; } h2 { font-size: 21px; color: #174a76; }
p { white-space: pre-wrap; overflow-wrap: anywhere; }
table { border-collapse: collapse; width: 100%; font-size: 13px; }
th, td { border: 1px solid #91a3b5; padding: 9px; text-align: left; overflow-wrap: anywhere; }
th { background: #e9f2fa; }
.page { padding: 32px; } .badge { color: #62788d; font-size: 12px; }
.prototype { padding: 20px; border: 2px solid #9cb0c5; border-radius: 8px; }
.prototype h2 { background: #edf4fb; padding: 10px; }
.appbar { background: #174a76; color: white; padding: 18px; font-size: 22px; }
.fields { display: grid; grid-template-columns: 1fr 1fr; gap: 18px; padding: 24px; }
.field label { display: block; color: #536d84; font-size: 14px; }
.field .value { border: 1px solid #9fb4c8; border-radius: 4px; padding: 10px; background: white; }
.tabs { display: flex; gap: 24px; padding: 16px; border-bottom: 2px solid #9fb4c8; }
.actions { display: flex; flex-wrap: wrap; gap: 12px; padding: 20px; }
.button { border: 1px solid #4178a1; border-radius: 4px; padding: 8px 16px; background: #edf4fb; }
.node { border: 2px solid #4178a1; border-radius: 10px; padding: 16px;
        margin: 12px auto; width: 80%; text-align: center; }
"""


def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def safe_path(root, relative):
    path = root / relative
    if Path(relative).is_absolute() or ".." in Path(relative).parts:
        raise ValueError(f"Not a relative fixture path: {relative}")
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def paragraphs(value):
    return [value] if isinstance(value, str) else list(value or [])


def table_rows(table):
    return [table["headers"], *table.get("rows", [])]


def document():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.space_after = Pt(6)
    for style in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        doc.styles[style].font.name = "黑体"
        doc.styles[style].element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2.2)
        section.footer.paragraphs[0].text = "Folio 合成测试资料 · 不对应真实企业"
    return doc


def doc_table(doc, data):
    rows = table_rows(data)
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.style = "Table Grid"
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = str(value)
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)


def sections_doc(doc, sections):
    for section in sections:
        doc.add_heading(section["heading"], level=section.get("level", 1))
        for para in paragraphs(section.get("text")):
            doc.add_paragraph(para)
        if section.get("table"):
            doc_table(doc, section["table"])


def source_html(source):
    chunks = [f'<h1>{html.escape(source.get("title", source.get("heading", "资料")))}</h1>']
    if source.get("ui"):
        ui = source["ui"]
        chunks.extend(['<div class="prototype">', f'<div class="appbar">{html.escape(ui["page_title"])}</div>', '<div class="fields">'])
        for field in ui["fields"]:
            chunks.append(f'<div class="field"><label>{html.escape(field["label"])}</label><div class="value">{html.escape(str(field["value"]))}</div></div>')
        chunks.append('</div><div class="tabs">' + "".join(f'<span>{html.escape(tab)}</span>' for tab in ui.get("tabs", [])) + "</div>")
        chunks.append('<div class="actions">' + "".join(f'<span class="button">{html.escape(button)}</span>' for button in ui.get("buttons", [])) + "</div></div>")
    sections = source.get("sections", [])
    if "text" in source:
        sections = [{"heading": "", "text": source["text"]}, *sections]
    if source.get("table"):
        sections = [*sections, {"heading": "", "table": source["table"]}]
    for section in sections:
        chunks.append(f'<h2>{html.escape(section.get("heading", ""))}</h2>')
        chunks.extend(f"<p>{html.escape(p)}</p>" for p in paragraphs(section.get("text")))
        if section.get("table"):
            chunks.append("<table>")
            for index, row in enumerate(table_rows(section["table"])):
                tag = "th" if index == 0 else "td"
                chunks.append("<tr>" + "".join(
                    f"<{tag}>{html.escape(str(cell))}</{tag}>" for cell in row
                ) + "</tr>")
            chunks.append("</table>")
    if source.get("nodes"):
        labels = {node["id"]: node["label"] for node in source["nodes"]}
        for node in source["nodes"]:
            chunks.append(f'<div class="node">{html.escape(node["label"])}</div>')
        for edge in source.get("edges", []):
            chunks.append("<p>" + html.escape(
                f'{labels[edge["from"]]} → {labels[edge["to"]]}：{edge.get("label", "")}'
            ) + "</p>")
    cls = "prototype" if source.get("layout") == "prototype" else ""
    return f'<html lang="zh"><meta charset="utf-8"><style>{CSS}</style><body><main class="page {cls}">' + "".join(chunks) + "</main></body></html>"


class Renderer:
    def __init__(self):
        self.playwright = self.browser = None

    def page(self):
        if self.browser is None:
            from playwright.sync_api import sync_playwright

            self.playwright = sync_playwright().start()
            self.browser = self.playwright.chromium.launch()
        page = self.browser.new_page(viewport={"width": 1100, "height": 1500}, device_scale_factor=1)
        page.route("**/*", lambda route: route.abort())
        return page

    def render(self, source, path):
        page = self.page()
        try:
            if source["kind"] == "png":
                page.set_content(source_html(source))
                page.screenshot(path=str(path), full_page=True)
                return
            writer = PdfWriter()
            pages = source.get("pages") or [{"heading": source["title"], "sections": source.get("sections", [])}]
            for content in pages:
                page.set_content(source_html(content))
                if source["kind"] == "scan_pdf" or content.get("scan"):
                    # One image-only physical page: no duplicate hidden text layer.
                    raw = page.screenshot(full_page=True)
                    with Image.open(io.BytesIO(raw)) as im:
                        stream = io.BytesIO()
                        im.convert("RGB").save(stream, format="PDF", resolution=140)
                        writer.append(PdfReader(io.BytesIO(stream.getvalue())))
                else:
                    raw = page.pdf(format="A4", print_background=True)
                    native = PdfReader(io.BytesIO(raw))
                    if len(native.pages) != 1:
                        raise ValueError(f"{path}: one page recipe overflowed to {len(native.pages)} pages")
                    writer.append(native)
            with path.open("wb") as stream:
                writer.write(stream)
        finally:
            page.close()

    def close(self):
        if self.browser:
            self.browser.close()
        if self.playwright:
            self.playwright.stop()


def write_source(source, path, renderer):
    kind = source["kind"]
    if kind in {"md", "txt"}:
        lines = ["# " + source["title"], ""]
        for section in source.get("sections", []):
            lines.extend(["## " + section["heading"], "", *paragraphs(section.get("text")), ""])
            if section.get("table"):
                rows = table_rows(section["table"])
                lines.append("| " + " | ".join(map(str, rows[0])) + " |")
                lines.append("| " + " | ".join("---" for _ in rows[0]) + " |")
                lines.extend("| " + " | ".join(map(str, row)) + " |" for row in rows[1:])
                lines.append("")
        path.write_text("\n".join(lines), encoding="utf-8")
    elif kind == "docx":
        doc = document()
        doc.add_heading(source["title"], 0)
        sections_doc(doc, source.get("sections", []))
        doc.save(path)
    elif kind == "xlsx":
        book = Workbook()
        book.remove(book.active)
        for sheet in source["sheets"]:
            ws = book.create_sheet(sheet["name"])
            ws.append(sheet["headers"])
            for row in sheet["rows"]:
                ws.append(row)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="245A81")
            for row in ws:
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            for index in range(1, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(index)].width = 25
            if sheet.get("hidden"):
                ws.sheet_state = "hidden"
        book.save(path)
    elif kind == "pptx":
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.33), Inches(7.5)
        for content in source["slides"]:
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = content["title"]
            frame = slide.placeholders[1].text_frame
            frame.clear()
            frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            for index, bullet in enumerate(content["bullets"]):
                p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                p.text = bullet
                p.font.name = "Noto Sans CJK SC"
                p.font.size = SlidePt(20)
        deck.save(path)
    elif kind in {"pdf", "scan_pdf", "png"}:
        renderer.render(source, path)
    else:
        raise ValueError(f"Unsupported fixture kind {kind}")


def validate_recipe(case):
    paths = [source["path"] for source in case["sources"]]
    if len(paths) != len(set(paths)) or len(paths) > 200:
        raise ValueError(f"{case['id']}: duplicate paths or more than 200 sources")
    expected = case["expected"]
    minimum = sum(chapter["min_chars"] for chapter in expected["chapters"])
    if not case["target_chars"]["min"] <= minimum <= case["target_chars"]["max"]:
        raise ValueError(f"{case['id']}: chapter length allocation {minimum} outside total target")
    for chapter in expected["chapters"]:
        for source in chapter["source_paths"]:
            if source not in paths:
                raise ValueError(f"{case['id']}: missing source {source}")
    for fact in expected["facts"]:
        if fact["source"] not in paths:
            raise ValueError(f"{case['id']}: missing fact source {fact['source']}")
    for source in case["sources"]:
        for table in [*(section["table"] for section in source.get("sections", []) if section.get("table")), *source.get("sheets", [])]:
            if any(len(row) != len(table["headers"]) for row in table["rows"]):
                raise ValueError(f"{case['id']} / {source['path']}: ragged table")


def write_expectations(case, root):
    expected = {key: case[key] for key in ("id", "name", "document_type", "output_title", "output_filename", "target_chars", "audience", "focus", "expected", "execution")}
    expected["counting_method"] = COUNTING
    write_json(root / "expected" / "expected.json", expected)
    doc = document()
    doc.add_heading(case["name"] + " · 交付预期与验收", 0)
    doc.add_paragraph("本文件供测试人员核对成品，不上传资料库。它是验收蓝图，不是已经生成的完整业务文档。")
    doc.add_paragraph(f"目标文件：{case['output_filename']}；类型：{case['document_type']}；正文目标：{case['target_chars']['min']:,}—{case['target_chars']['max']:,} 字符。")
    doc.add_paragraph(COUNTING)
    doc.add_heading("一、预期目录与各章展开程度", 1)
    for index, chapter in enumerate(case["expected"]["chapters"], 1):
        doc.add_heading(f"{index}. {chapter['title']}（至少 {chapter['min_chars']} 字符）", 2)
        for topic in chapter["topics"]:
            doc.add_paragraph(topic, style="List Bullet")
        doc.add_paragraph("依据：" + "、".join(chapter["source_paths"]))
    doc.add_heading("二、必须准确表达的事实", 1)
    for fact in case["expected"]["facts"]:
        doc.add_paragraph(f"{fact['id']} [{fact['severity']}] {fact['statement']}")
        doc.add_paragraph(f"依据：{fact['source']} / {fact['locator']}")
    if case["expected"].get("coverage_ids"):
        doc.add_paragraph("完整覆盖编号：" + "、".join(case["expected"]["coverage_ids"]))
    doc.add_heading("三、计算与统计口径", 1)
    for calc in case["expected"].get("calculations", []):
        doc.add_paragraph(f"{calc['id']} {calc['description']}：{calc['formula']} = {calc['result']} {calc['unit']}")
    doc.add_heading("四、图表和原型", 1)
    for visual in case["expected"].get("visuals", []):
        doc.add_heading(visual["title"] + " / " + visual["kind"], 2)
        for text in visual["requirements"]:
            doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("五、缺口与不得写成事实的结论", 1)
    for gap in case["expected"].get("gaps", []):
        doc.add_paragraph(f"{gap['topic']}：{gap['reason']}；预期：{gap['expected_handling']}")
    for text in case["expected"].get("must_not_claim", []):
        doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("六、人工语义与版式验收", 1)
    for text in case["expected"]["acceptance"]:
        doc.add_paragraph(text, style="List Bullet")
    for text in ["章节层级与目录一致，标题不孤悬；页码在 Word/WPS 更新后核对。", "表头跨页重复，表格和图片不超出页边距；正文、标题、表格使用一致的中文字体。", "流程图分支、回路、例外条件与正文一致；缩放到实际页面后图中文字清晰。", "模型建议与已证实事实明确区分，重复段落和空泛扩写不能用于补足字数。", "达到字数或命中关键词不代表语义验收通过；逐项确认事实、关系、缺口与来源。"]:
        doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("七、执行与过程验收", 1)
    doc.add_paragraph(json.dumps(case["execution"], ensure_ascii=False, indent=2))
    doc.save(root / "expected" / "交付预期.docx")


def build_case(case, output, renderer):
    validate_recipe(case)
    root = output / case["slug"]
    root.mkdir(parents=True, exist_ok=True)
    manifest = []
    for source in case["sources"]:
        path = safe_path(root / "sources", source["path"])
        write_source(source, path, renderer)
        size = path.stat().st_size
        manifest.append({"path": source["path"], "kind": source["kind"], "title": source["title"], "bytes": size, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
    current = {item["path"] for item in manifest}
    stale = [str(path) for path in (root / "sources").rglob("*") if path.is_file() and path.relative_to(root / "sources").as_posix() not in current]
    if stale:
        raise ValueError(f"Stale sources; rebuild in a new output directory: {stale}")
    write_json(root / "manifest.json", {"case_id": case["id"], "synthetic": True, "files": manifest})
    prompt = case["prompt"] + "\n\n" + COUNTING + "\n所有资料均为合成项目资料。请先生成完整目录供确认，再生成 Word；有信息缺口时明确标注，不能虚构批准记录或外部依据。\n"
    (root / "prompt.txt").write_text(prompt, encoding="utf-8")
    write_expectations(case, root)
    if case.get("mock_web"):
        write_json(root / "mock_web.json", case["mock_web"])
    with ZipFile(root / "upload_sources.zip", "w", ZIP_DEFLATED) as archive:
        for item in manifest:
            archive.write(root / "sources" / item["path"], arcname=item["path"])
    return {"id": case["id"], "name": case["name"], "path": case["slug"], "document_type": case["document_type"], "output_filename": case["output_filename"], "target_chars": case["target_chars"], "source_count": len(manifest), "source_bytes": sum(item["bytes"] for item in manifest), "formats": sorted({item["kind"] for item in manifest}), "chapter_count": len(case["expected"]["chapters"]), "fact_count": len(case["expected"]["facts"]), "focus": case["focus"]}


def write_catalog(summary, path):
    doc = document()
    doc.add_heading("Folio · 10 套智能体长文档测试案例", 0)
    doc.add_paragraph("合成资料包与预期交付蓝图 · 2026-10-15")
    doc.add_paragraph("每套包含可上传原件、用户生成指令、文件Manifest、预期Word结构与事实/计算/图表验收规则。所有企业、人物、规则和数字均为虚构，供测试使用。")
    doc.add_paragraph("先上传sources中的文件，发送prompt.txt并确认目录，再把生成的Word与expected中的验收蓝图逐项比对。预期蓝图不上传、不作为实际长文成品。")
    doc.add_paragraph(COUNTING)
    doc_table(doc, {"headers": ["案例", "文档", "目标字符", "资料数"], "rows": [
        [item["id"] + " " + item["name"], item["document_type"], f'{item["target_chars"]["min"]:,}—{item["target_chars"]["max"]:,}', str(item["source_count"])]
        for item in summary
    ]})
    for item in summary:
        doc.add_page_break()
        doc.add_heading(item["id"] + " · " + item["name"], 1)
        doc.add_paragraph("预期输出：" + item["output_filename"])
        doc.add_paragraph(f'正文目标 {item["target_chars"]["min"]:,}—{item["target_chars"]["max"]:,} 有效字符；{item["chapter_count"]} 个主章；{item["fact_count"]} 个关键事实锚点。')
        doc.add_paragraph(f'{item["source_count"]} 份资料，格式：' + "、".join(item["formats"]))
        for focus in item["focus"]:
            doc.add_paragraph(focus, style="List Bullet")
        base = "cases/" + item["path"]
        doc.add_paragraph("上传资料：" + base + "/sources/")
        doc.add_paragraph("生成指令：" + base + "/prompt.txt")
        doc.add_paragraph("验收蓝图：" + base + "/expected/交付预期.docx")
        doc.add_paragraph("判定需同时检查业务正确性、来源与缺口、完整覆盖、字数和实际Word版式。关键词或编号命中仅作遗漏初筛，不能证明事实语义正确。")
    doc.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "cases")
    parser.add_argument("--case", help="Build a single case ID, e.g. TC01")
    args = parser.parse_args()
    recipes = [json.loads(path.read_text(encoding="utf-8")) for path in sorted((ROOT / "recipes").glob("*.json"))]
    if args.case:
        recipes = [case for case in recipes if case["id"] == args.case]
    if not recipes:
        parser.error("No matching recipes")
    renderer = Renderer()
    summary = []
    try:
        for case in recipes:
            summary.append(build_case(case, args.output, renderer))
            print(f"{case['id']}: {len(case['sources'])} sources → {args.output / case['slug']}")
    finally:
        renderer.close()
    if not args.case:
        write_json(args.output / "suite.json", {"version": 1, "synthetic": True, "counting_method": COUNTING, "cases": summary})
        write_catalog(summary, args.output.parent / "测试案例总览.docx")
        with ZipFile(args.output.parent / "folio-10-cases.zip", "w", ZIP_DEFLATED) as archive:
            for path in [ROOT / name for name in ["README.md", "README.zh-CN.md", "build_cases.py", "check_output.py", "mock_web.py"]]:
                if path.exists():
                    archive.write(path, arcname=path.name)
            for path in sorted((ROOT / "recipes").glob("*.json")):
                archive.write(path, arcname="recipes/" + path.name)
            archive.write(args.output.parent / "测试案例总览.docx", arcname="测试案例总览.docx")
            for path in sorted(args.output.rglob("*")):
                if path.is_file():
                    archive.write(path, arcname="cases/" + path.relative_to(args.output).as_posix())
    print("Fixture generation only: no LLM or Tavily calls made.")


if __name__ == "__main__":
    main()
